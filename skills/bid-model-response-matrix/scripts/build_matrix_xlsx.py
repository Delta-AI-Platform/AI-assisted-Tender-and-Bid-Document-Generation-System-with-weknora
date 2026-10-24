"""Build the 8-column model standard response matrix (车型标准响应参考表) as XLSX.

The agent does the judgement (what the answer is, whether it deviates); this
script does the mechanical assembly and enforces the house rules so they can
never silently drift:

* numbers are embedded in the 招标要求 / 投标人响应 columns, never a separate
  column, and are always ``<section>.<original number>`` (see numbering.py)
* no duplicated prefix when a body already starts with its number
* ``无偏离`` rows carry the answer only — any ``补充说明`` / ``证明材料`` block is
  stripped (说明只用于正/负偏离)
* ``已满足`` phrasing is removed
* module columns (车型/子系统/部件) fill down and merge vertically
* 正偏离 without an 优于补充说明 block is reported as a rule violation

Usage
-----
    python build_matrix_xlsx.py matrix-spec.json -o out.xlsx
    python build_matrix_xlsx.py matrix-spec.json -o out.xlsx --strict

Spec (see references/data-schema.md)::

    {
      "sheet": "车型标准响应对照",
      "note": "说明：标准车型答案列取自……",
      "rows": [
        {"sec": "5.1.1", "kind": "table", "orig_no": "15",
         "model": "矿用防爆指挥车 WLR-5", "sub": "动力与行驶系统",
         "part": "整车基本参数", "func": "续航里程（km）",
         "std": "矿井综合运行里程 ≥120km……",
         "req_body": "续航里程（km）：≥120",
         "resp_body": "续航里程（km）：≥120（矿井综合运行工况）",
         "dev": "无偏离"},
        {"sec": "6.2.6", "kind": "para", "orig_no": "",
         "func": "电驱系统", "req_body": "驱动电机为防爆永磁同步电机……",
         "resp_body": "驱动电机为……", "dev": "无偏离", "sep": "、"}
      ]
    }
"""
import argparse
import json
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from numbering import derive_para_numbers, derive_table_numbers, embed

HEAD = ["车型", "子系统模块", "部件模块", "功能模块",
        "标准车型答案（历史投标响应）", "本次招标要求", "投标人响应", "偏离说明"]
WIDTHS = [22, 20, 22, 24, 46, 46, 46, 12]
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="D9E2F3")
MOD_FILL = PatternFill("solid", fgColor="FCE4D6")
HEAD_FONT = Font(name="微软雅黑", size=10, bold=True)
BODY_FONT = Font(name="微软雅黑", size=9)
NOTE_FONT = Font(name="微软雅黑", size=9, italic=True, color="C00000")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def tidy(text):
    """Drop the '已满足' phrasing the business rejects, and collapse punctuation."""
    s = str(text or "")
    if not s:
        return s
    s = re.sub(r"[，,；;、]?\s*已满足(?=\s*(?:[。；;]|\n|$))", "", s)
    s = s.replace("已满足", "满足")
    s = re.sub(r"[，,]\s*\n", "\n", s)
    s = re.sub(r"。\s*。", "。", s).replace("；；", "；")
    return s.strip()


def strip_notes(resp):
    """Remove the 补充说明 / 证明材料 layers (used when 偏离 == 无偏离)."""
    return re.sub(r"\n\s*\n\s*(?:优于)?补充说明：[\s\S]*$", "", str(resp or "")).rstrip()


def assign_numbers(rows):
    """Fill ``row['_no']`` for every row, grouping consecutive rows by section."""
    groups = []
    for i, r in enumerate(rows):
        sec = str(r.get("sec") or "")
        if groups and groups[-1][0] == sec:
            groups[-1][1].append(i)
        else:
            groups.append((sec, [i]))

    for sec, idxs in groups:
        items = [rows[i] for i in idxs]
        kinds = {str(r.get("kind") or "para") for r in items}
        derived = [("", "")] * len(items)
        if sec and kinds == {"table"}:
            derived = derive_table_numbers(sec, items)
        elif sec and kinds == {"para"}:
            derived = derive_para_numbers(sec, [
                {"orig_no": r.get("orig_no"), "text": r.get("text") or r.get("req_body")}
                for r in items])
        for i, (no, sep) in zip(idxs, derived):
            explicit = str(rows[i].get("no") or "")
            rows[i]["_no"] = explicit or no
            # explicit numbers keep the plain space form unless a row says otherwise
            rows[i]["_sep"] = str(rows[i].get("sep") or (sep if not explicit else " ") or " ")
    return rows


def build(spec, out_path, strict=False):
    rows = assign_numbers([dict(r) for r in spec.get("rows") or []])
    if not rows:
        raise SystemExit("spec has no rows")

    problems, warnings = [], []
    for n, r in enumerate(rows, 1):
        dev = str(r.get("dev") or "无偏离").strip()
        resp = tidy(r.get("resp_body", r.get("resp", "")))
        if dev == "无偏离" and re.search(r"(?:优于)?补充说明：", resp):
            warnings.append("row %d「%s」判为无偏离却带补充说明，已剥离" % (n, r.get("func", "")))
            resp = strip_notes(resp)
        if dev == "正偏离" and "优于补充说明" not in resp:
            problems.append("row %d「%s」判为正偏离但缺优于补充说明" % (n, r.get("func", "")))
        if dev == "负偏离" and not re.search(r"(?:补充说明|差异说明)：", resp):
            problems.append("row %d「%s」判为负偏离但缺差异说明" % (n, r.get("func", "")))
        if dev not in ("正偏离", "无偏离", "负偏离"):
            problems.append("row %d 偏离说明取值非法：%s" % (n, dev))
        if "已满足" in resp:
            problems.append("row %d「%s」响应仍含“已满足”" % (n, r.get("func", "")))
        r["_resp"], r["_dev"] = resp, dev

    wb = Workbook()
    ws = wb.active
    ws.title = str(spec.get("sheet") or "车型标准响应对照")
    ws.append(HEAD)

    # row 2 = single merged note line
    note = str(spec.get("note") or "")
    ws.append([note] + [""] * 7)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=8)
    c = ws.cell(row=2, column=1)
    c.font = NOTE_FONT
    c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 42 if note else 8

    last = ["", "", ""]
    for r in rows:
        trio = [str(r.get(k) or "").strip() for k in ("model", "sub", "part")]
        trio = [v or last[i] for i, v in enumerate(trio)]      # fill down
        last = trio
        no, sep = r["_no"], r["_sep"]
        req = r.get("req") or embed(no, r.get("req_body", r.get("req", "")), sep)
        resp = r.get("resp") or (embed(no, r["_resp"], sep) if r["_resp"] else "")
        ws.append(trio + [str(r.get("func") or ""), str(r.get("std") or ""),
                          req, resp, r["_dev"]])

    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill, cell.font, cell.alignment, cell.border = HEAD_FILL, HEAD_FONT, CENTER, BORDER
    ws.row_dimensions[1].height = 32
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, max_col=8):
        for i, cell in enumerate(row):
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = CENTER if i < 4 else WRAP
            if i < 4:
                cell.fill = MOD_FILL
    ws.freeze_panes = "E2"

    # vertical merge of the module columns, constrained to one 车型 run
    runs, start = [], 3
    for r in range(3, ws.max_row + 2):
        v = ws.cell(row=r, column=1).value if r <= ws.max_row else None
        if r == ws.max_row + 1 or v != ws.cell(row=start, column=1).value:
            runs.append((start, r - 1))
            start = r

    def merge_same(col, r0, r1):
        r = r0
        while r <= r1:
            v = ws.cell(row=r, column=col).value
            end = r
            while end + 1 <= r1 and ws.cell(row=end + 1, column=col).value == v:
                end += 1
            if end > r and v not in (None, ""):
                ws.merge_cells(start_row=r, start_column=col, end_row=end, end_column=col)
            ws.cell(row=r, column=col).alignment = CENTER
            r = end + 1

    for r0, r1 in runs:
        if r1 > r0:
            ws.merge_cells(start_row=r0, start_column=1, end_row=r1, end_column=1)
        ws.cell(row=r0, column=1).alignment = CENTER
        merge_same(2, r0, r1)
        merge_same(3, r0, r1)

    ws.auto_filter.ref = "A1:H%d" % ws.max_row
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)

    print("saved: %s" % out_path)
    print("rows: %d | 正偏离 %d / 无偏离 %d / 负偏离 %d" % (
        len(rows), sum(1 for r in rows if r["_dev"] == "正偏离"),
        sum(1 for r in rows if r["_dev"] == "无偏离"),
        sum(1 for r in rows if r["_dev"] == "负偏离")))
    print("no-number rows: %d" % sum(1 for r in rows if not r["_no"]))
    for w in warnings:
        print("WARN  " + w)
    for p in problems:
        print("ERROR " + p)
    if problems and strict:
        raise SystemExit(1)
    return problems


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("spec")
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--strict", action="store_true", help="exit 1 on rule violations")
    a = ap.parse_args()
    spec = json.load(open(a.spec, encoding="utf-8"))
    problems = build(spec, a.out, a.strict)
    sys.exit(1 if (a.strict and problems) else 0)


if __name__ == "__main__":
    main()
