"""Pre-delivery verification for the model standard response matrix.

Four gates, each one maps to a defect that actually reached a business review:

A 结构   : 8 columns, row count, empty responses, rows without a number
B 编号   : level format, per-section gaps/duplicates, requirement/response
           alignment, double prefix (``5.1.1.1 1.``)
C 偏离   : 正偏离 really better (numeric direction) and on-topic
           (核心词命中), 无偏离 carries no 说明, no "已满足"
D 原文   : every requirement body is found verbatim in the source text dump
           (optional; pass --source)

Usage
-----
    python verify_matrix.py matrix.xlsx
    python verify_matrix.py matrix.xlsx --source tender_text.txt --report chk/report.txt
    python verify_matrix.py matrix.xlsx --allow-dup 5.3.5,6.2.5

Exit code 0 only when no hard error is found.
"""
import argparse
import json
import os
import re
import sys
from collections import Counter, defaultdict

from openpyxl import load_workbook

from numbering import head_number, strip_lead

SUBJECT = {
    "电机": ("电机", "电驱", "驱动"),
    "电池": ("电池", "电源", "蓄电", "电芯", "续航", "电压", "容量", "能量"),
    "电芯": ("电芯", "电池"),
    "充电": ("充电",),
    "制动": ("制动", "刹车", "驻车", "应急"),
    "转向": ("转向",),
    "轮胎": ("轮胎", "行走"),
    "环视": ("环视", "全景", "影像"),
    "监控": ("监控", "报警", "传感", "瓦斯"),
    "车速": ("车速", "速度", "限速"),
    "续航": ("续航", "里程"),
    "灭火": ("灭火", "消防"),
}


def norm(s):
    return re.sub(r"[\s\u3000（）()：:，,。；;、/·\-—]", "", str(s or ""))


def metric(text):
    m = re.search(r"[：:]\s*([≥≤<>]=?)\s*(\d+(?:\.\d+)?)", str(text or ""))
    return (m.group(1), float(m.group(2))) if m else (None, None)


def first_number(resp):
    body = re.split(r"\n\s*\n|优于补充说明|补充说明", str(resp or ""))[0]
    body = re.sub(r"^\s*\d+(?:\.\d+)+\s*", "", body)
    m = re.search(r"[：:]\s*[≥≤<>]?=?\s*(\d+(?:\.\d+)?)", body)
    return float(m.group(1)) if m else None


def bigrams(text, n=16):
    t = re.sub(r"^[\d.．、\s]+", "", str(text or ""))
    t = re.sub(r"[（(].*?[)）]|\s|[／/]", "", t)[:n]
    if not t:
        return set()
    return {t[i:i + 2] for i in range(len(t) - 1)} | {t}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--source", help="text dump (or .txt extracted from PDF/DOCX) for verbatim check")
    ap.add_argument("--allow-dup", default="", help="comma list of level numbers the tender itself repeats")
    ap.add_argument("--report")
    ap.add_argument("--json", dest="json_out")
    a = ap.parse_args()

    allowed = {x.strip() for x in a.allow_dup.split(",") if x.strip()}
    wb = load_workbook(a.xlsx)
    ws = wb[wb.sheetnames[0]]
    out, errors, warns = [], [], []

    def say(line=""):
        out.append(line)
        print(line)

    say("=" * 78)
    say("A. 结构")
    say("=" * 78)
    head = [(ws.cell(1, c).value or "") for c in range(1, ws.max_column + 1)]
    say("工作表：%s" % wb.sheetnames)
    say("列数：%d%s" % (ws.max_column, "" if ws.max_column == 8 else "  ← 应为 8"))
    if ws.max_column != 8:
        errors.append("列数不是 8")
    say("表头：%s" % " | ".join(str(h) for h in head[:8]))

    rows, empty_resp, no_number = [], [], []
    header_row = 1
    for r in range(1, min(ws.max_row, 6) + 1):
        if "偏离说明" in "".join(str(ws.cell(r, c).value or "") for c in (3, 4, 8)):
            header_row = r
            break
    for r in range(header_row + 1, ws.max_row + 1):
        v = [ws.cell(r, c).value or "" for c in range(1, 9)]
        # skip empty rows and the single merged note row (only column A filled)
        if not any(str(x).strip() for x in v) or not any(str(x).strip() for x in v[1:]):
            continue
        num = head_number(v[5])
        if not num and not str(v[5]).strip().startswith(("（", "(")):
            no_number.append((r, str(v[5])[:40]))
        if not str(v[6]).strip():
            empty_resp.append((r, str(v[3])[:24]))
        body = re.sub(r"^\s*" + re.escape(num) + r"[\s\u3000、]*", "", str(v[5])) if num else str(v[5])
        rows.append({"row": r, "v": v, "num": num, "req": body, "resp": str(v[6])})
    say("数据行：%d" % len(rows))
    say("无编号行：%d" % len(no_number))
    for x in no_number[:8]:
        say("   R%d %s" % x)
    say("响应为空：%d" % len(empty_resp))
    for x in empty_resp[:8]:
        say("   R%d %s" % x)

    say("")
    say("=" * 78)
    say("B. 编号")
    say("=" * 78)
    per_sec, bad_fmt = defaultdict(list), []
    for x in rows:
        if not x["num"]:
            continue
        if not re.fullmatch(r"\d+(?:\.\d+)+", x["num"]):
            bad_fmt.append((x["row"], x["num"]))
            continue
        sec, n = x["num"].rsplit(".", 1)
        per_sec[sec].append((int(n), x["row"]))
    say("层级编号行：%d，覆盖 %d 节；格式非法：%d" % (
        sum(len(v) for v in per_sec.values()), len(per_sec), len(bad_fmt)))
    if bad_fmt:
        errors.append("编号格式非法 %d 处" % len(bad_fmt))
        for x in bad_fmt[:6]:
            say("   格式非法 R%d %s" % x)

    for sec in sorted(per_sec, key=lambda s: [int(p) for p in s.split(".")]):
        lst = per_sec[sec]
        nums = [n for n, _ in lst]
        dup = {n: c for n, c in Counter(nums).items() if c > 1}
        # A single-row "section" is usually a clause number that happens to look
        # like a section (e.g. 6.2.6 with no sub-numbers) — do not demand 1..n.
        gap = [i for i in range(1, max(nums) + 1) if i not in nums] if len(lst) >= 2 else []
        flags = []
        if dup:
            flags.append("重复 %s" % dup)
        if gap:
            flags.append("断号 %s" % gap)
        hard_dup = {k: v for k, v in dup.items() if "%s.%d" % (sec, k) not in allowed}
        if gap or hard_dup:
            errors.append("编号 %s：%s" % (sec, "；".join(flags)))
        if flags:
            say("   %-9s %2d 行 | %s" % (sec, len(lst), "；".join(flags)))
    if not allowed:
        say("   提示：招标原文若本身重复编号，请用 --allow-dup 列出，避免误判")

    say("")
    say("=" * 78)
    say("C. 偏离与语义")
    say("=" * 78)
    pos = 0
    for x in rows:
        req, resp, dev, func = str(x["v"][5]), str(x["v"][6]), str(x["v"][7]), str(x["v"][3])
        rnum = head_number(resp)
        if x["num"] and rnum and rnum != x["num"]:
            errors.append("R%d 要求/响应编号不一致：%s vs %s" % (x["row"], x["num"], rnum))
        if re.match(r"^\d+(?:\.\d+)+\s+\d+[.．、]", resp):
            errors.append("R%d 双前缀：%s" % (x["row"], resp[:24]))
        if "已满足" in resp:
            errors.append("R%d 含“已满足”" % x["row"])
        if dev == "无偏离" and re.search(r"(?:优于)?补充说明：", resp):
            errors.append("R%d 无偏离却带补充说明" % x["row"])
        if dev == "正偏离":
            pos += 1
            if "优于补充说明" not in resp:
                errors.append("R%d 正偏离缺优于补充说明" % x["row"])
            seg = resp.split("优于补充说明", 1)[1].split("证明材料", 1)[0]
            allowed_subj = {k for k, owners in SUBJECT.items()
                            if any(o in func or o in str(x["v"][2]) for o in owners)}
            stray = [k for k in SUBJECT if k in seg and k not in allowed_subj]
            if stray:
                warns.append("R%d「%s」说明含其他族词 %s" % (x["row"], func, "/".join(stray)))
            if not any(b in seg for b in (bigrams(func) | bigrams(x["req"]))):
                errors.append("R%d「%s」优于说明未出现本条核心词（疑似串条）" % (x["row"], func))
            d, rq = metric(req)
            rv = first_number(resp)
            if d and rq is not None and rv is not None:
                better = rv > rq if d in (">", "≥", ">=") else rv < rq
                if not better:
                    errors.append("R%d「%s」%s%g，响应 %g，不构成正偏离" % (x["row"], func, d, rq, rv))
    say("正偏离：%d 条" % pos)
    say("硬错误：%d 条" % len(errors))
    for e in errors:
        say("   [改] " + e)
    say("提示：%d 条" % len(warns))
    for w in warns:
        say("   [看] " + w)

    if a.source:
        say("")
        say("=" * 78)
        say("D. 原文逐字核对（%s）" % os.path.basename(a.source))
        say("=" * 78)
        blob = norm(open(a.source, encoding="utf-8", errors="ignore").read())
        miss = soft = 0
        for x in rows:
            if "：" in x["req"]:
                # composed table row (名称：指标（备注）): compare the segments,
                # because the sheet legitimately re-expresses units/values
                segs = [p for p in re.split(r"[：（）()]", x["req"]) if len(norm(p)) >= 4]
                bad = [p for p in segs if norm(p) not in blob]
                if bad:
                    soft += 1
                    warns.append("R%d %s 原文未全命中：%s" % (x["row"], x["num"], "/".join(bad)[:40]))
            elif norm(x["req"])[:40] not in blob:
                miss += 1
                say("   未命中 R%d %s | %s" % (x["row"], x["num"], x["req"][:34]))
        say("段落条未命中：%d / %d" % (miss, sum(1 for x in rows if "：" not in x["req"])))
        say("表格条分段提示：%d" % soft)
        if miss:
            errors.append("段落原文未命中 %d 条" % miss)

    say("")
    say("=" * 78)
    say("结论：%s（硬错误 %d / 提示 %d）" % ("通过" if not errors else "未通过", len(errors), len(warns)))
    say("=" * 78)

    if a.report:
        os.makedirs(os.path.dirname(os.path.abspath(a.report)), exist_ok=True)
        open(a.report, "w", encoding="utf-8").write("\n".join(out) + "\n")
        print("report:", a.report)
    if a.json_out:
        json.dump({"errors": errors, "warnings": warns,
                   "rows": len(rows), "positive": pos},
                  open(a.json_out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
