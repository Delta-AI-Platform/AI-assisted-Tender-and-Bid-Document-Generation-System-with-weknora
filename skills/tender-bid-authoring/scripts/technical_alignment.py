#!/usr/bin/env python3
"""Build and validate a source-locked internal technical alignment workbook.

This is a loss-detection tool, not a semantic classifier. It groups the goods
table and working environment as whole points, while retaining their source
blocks and exact text in the manifest. It never copies historical responses.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.text.paragraph import Paragraph
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


MODULE_HEADERS = ["子系统模块", "部件模块", "功能模块"]
MODEL_HEADER = "标准车型及报告级数据"
TAIL_HEADERS = ["本次招标要求", "投标人响应", "偏离说明"]
HEADERS = MODULE_HEADERS + [MODEL_HEADER] + TAIL_HEADERS
NOTICE = "内部技术审查稿｜偏离说明为拟应标结论，不代表现有报告已实测满足；招标原文列不得修改"
TAXONOMY_PATH = Path(__file__).resolve().parent.parent / "assets" / "technical-module-taxonomy-v138.json"


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def valid_module_triples() -> set[tuple[str, str, str]]:
    """Use only actual, non-placeholder A:C mappings from the source taxonomy."""
    taxonomy = json.loads(TAXONOMY_PATH.read_text(encoding="utf-8"))
    return {
        (entry["subsystem"], entry["component"], entry["function"])
        for entry in taxonomy["entries"]
        if all(str(entry[key]).strip() not in {"", "/", "模块"}
               for key in ("subsystem", "component", "function"))
    }


def selected_blocks(path: Path, start: str | None, end: str | None) -> list[dict]:
    doc = Document(path)
    active = start is None
    saw_start = start is None
    saw_end = end is None
    blocks: list[dict] = []
    title_path: list[str] = []
    for body_index, element in enumerate(doc.element.body.iterchildren()):
        unsupported = [
            tag for tag in ("w:drawing", "w:pict", "w:txbxContent", "w:footnoteReference", "w:endnoteReference")
            if any(node.tag == qn(tag) for node in element.iter())
        ]
        if isinstance(element, CT_P):
            paragraph = Paragraph(element, doc)
            value = paragraph.text
            if not active and start and start in value:
                active, saw_start = True, True
            if active and end and end in value and (start is None or saw_start):
                saw_end = True
                break
            if not active or not value.strip():
                continue
            if unsupported:
                raise ValueError(f"Unsupported object {unsupported} at body block {body_index}; manual Word/page review required")
            style_num = paragraph.style.element.pPr.numPr if paragraph.style and paragraph.style.element.pPr is not None else None
            ppr = element.find(qn("w:pPr"))
            direct_num = ppr is not None and ppr.find(qn("w:numPr")) is not None
            if direct_num or style_num is not None:
                raise ValueError(f"Word automatic numbering at body block {body_index}; visible number needs manual verification")
            style = paragraph.style.name if paragraph.style else ""
            kind = "heading" if style.lower().startswith("heading") or style.startswith("标题") else "paragraph"
            if kind == "heading":
                title_path.append(value)
            blocks.append({"kind": kind, "text": value, "body_index": body_index, "title_path": list(title_path)})
        elif isinstance(element, CT_Tbl):
            if not active:
                continue
            if unsupported:
                raise ValueError(f"Unsupported table object {unsupported} at body block {body_index}; manual Word/page review required")
            rows: list[str] = []
            for tr in element.findall(qn("w:tr")):
                cells = []
                for tc in tr.findall(qn("w:tc")):
                    parts = [node.text or "" for node in tc.iter(qn("w:t"))]
                    cells.append("".join(parts))
                rows.append("\t".join(cells))
            value = "\n".join(rows)
            if value.strip():
                blocks.append({"kind": "table", "text": value, "body_index": body_index, "title_path": list(title_path), "source_rows": len(rows)})
    if not saw_start or not saw_end:
        raise ValueError("Start/end heading not found; no partial workbook was written")
    if not blocks:
        raise ValueError("No nonempty Word blocks in selected technical range")
    grouped: list[dict] = []
    pending: list[dict] = []
    pending_kind: str | None = None
    headings = {
        "一、货物需求一览表": "goods_table",
        "二、工作环境": "environment",
        "三、技术参数及要求": "technical_parameters",
    }

    def flush() -> None:
        nonlocal pending, pending_kind
        if pending:
            grouped.append({
                "kind": pending_kind,
                "text": "\n".join(part["text"] for part in pending),
                "body_index": pending[0]["body_index"],
                "title_path": pending[0]["title_path"],
                "source_parts": pending,
            })
            pending = []
            pending_kind = None

    for block in blocks:
        stripped = re.sub(r"\s+", "", block["text"].splitlines()[0])
        current_heading = next((kind for label, kind in headings.items() if stripped.startswith(label)), None)
        if current_heading:
            flush()
            if current_heading == "technical_parameters":
                grouped.append({**block, "kind": "heading"})
            else:
                pending_kind = current_heading
                pending = [block]
            continue
        if pending and block["kind"] == "heading":
            flush()
        if pending:
            pending.append(block)
        else:
            if block["kind"] == "paragraph" and len(re.findall(r"(?:^|\n)\s*\d+(?:\.\d+)+[\s、.]", block["text"])) > 1:
                raise ValueError(f"Multiple numbered clauses in one Word paragraph at block {block['body_index']}; manual split required")
            grouped.append(block)
    flush()
    return grouped


def checked_scope(scope: dict, items: list[dict]) -> dict:
    """Require an explicit project-only model list and source-order section map."""
    if scope.get("layout") not in {"combined", "separate"}:
        raise ValueError("Model scope layout must be combined or separate")
    models = scope.get("models")
    if not isinstance(models, list) or not models:
        raise ValueError("Project model scope is empty")
    ids = []
    for model in models:
        if not all(str(model.get(key, "")).strip() for key in ("model_id", "tender_model")):
            raise ValueError("Each project model needs model_id and tender_model")
        ids.append(model["model_id"])
    if len(ids) != len(set(ids)):
        raise ValueError("Duplicate project model_id")
    sections = scope.get("sections")
    if not isinstance(sections, list) or not sections:
        raise ValueError("Source section scope is empty")
    next_order = 1
    seen_models = set()
    for section in sections:
        first, last = section.get("first_source_order"), section.get("last_source_order")
        active = section.get("model_ids")
        if (not isinstance(first, int) or not isinstance(last, int)
                or first != next_order or last < first or last > len(items)):
            raise ValueError("Source sections must partition original rows in order without gaps or overlap")
        if not isinstance(active, list) or not active or not set(active).issubset(ids):
            raise ValueError("Section must reference current-project models only")
        first_line = items[first - 1]["text"].splitlines()[0].strip()
        if section.get("source_heading") != first_line:
            raise ValueError(f"Section anchor differs from source row {first}")
        seen_models.update(active)
        next_order = last + 1
    if next_order != len(items) + 1 or seen_models != set(ids):
        raise ValueError("Source sections do not cover all rows and project models")
    if scope["layout"] == "combined" and (len(sections) != 1 or set(sections[0]["model_ids"]) != set(ids)):
        raise ValueError("Combined tender must remain one combined source section")
    if scope["layout"] == "separate" and len(models) > 1 and len(sections) < 2:
        raise ValueError("Separately written models need separate source sections")
    return scope


def section_for_order(scope: dict, order: int) -> dict:
    return next(s for s in scope["sections"] if s["first_source_order"] <= order <= s["last_source_order"])


def model_headers(scope: dict) -> list[str]:
    return [f"标准车型答案｜{model['tender_model']}｜{model['model_id']}" for model in scope["models"]]


def build(args: argparse.Namespace) -> None:
    source = Path(args.source).resolve()
    output = Path(args.output).resolve()
    manifest_path = Path(args.manifest).resolve()
    blocks = selected_blocks(source, args.start, args.end)
    source_hash = digest(source.read_bytes())
    items = []
    for index, block in enumerate(blocks, 1):
        text_hash = digest(block["text"].encode("utf-8"))
        items.append({"requirement_id": f"src-{index:05d}-{text_hash[:10]}", "source_order": index, "text_sha256": text_hash, **block})
    scope_path = Path(args.scope).resolve()
    scope = checked_scope(json.loads(scope_path.read_text(encoding="utf-8")), items)
    headers = MODULE_HEADERS + model_headers(scope) + TAIL_HEADERS
    requirement_col = len(headers) - 2

    wb = Workbook()
    ws = wb.active
    ws.title = "技术对标"
    last_col = get_column_letter(len(headers))
    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = NOTICE
    ws["A1"].fill = PatternFill("solid", fgColor="FFF2CC")
    ws["A1"].font = Font(bold=True, color="9C5700")
    for col, header in enumerate(headers, 1):
        cell = ws.cell(2, col, header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for row, item in enumerate(items, 3):
        cell = ws.cell(row, requirement_col, item["text"])
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if item["kind"] == "heading" or item["source_order"] in {s["first_source_order"] for s in scope["sections"]}:
            cell.font = Font(bold=True)
    for col, width in enumerate([20, 20, 22] + [50] * len(scope["models"]) + [65, 65, 20], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = f"{get_column_letter(requirement_col)}3"
    ws.auto_filter.ref = f"A2:{last_col}{ws.max_row}"
    evidence = wb.create_sheet("证据索引")
    evidence.append(["requirement_id", "本项目车型ID", "标准型号", "子型号", "报告唯一ID", "材料原名", "版本/哈希", "页码或表格位置", "本项目附件位置", "审查状态"])
    wb.create_sheet("核对记录").append(["requirement_id", "确认人", "时间", "依据", "修改前", "修改后"])

    manifest = {"version": "v1.38", "source_path": str(source), "source_sha256": source_hash, "start": args.start, "end": args.end, "row_count": len(items), "project_scope": scope, "items": items}
    output.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    validate(source, output, manifest_path, require_responses=False)
    print(json.dumps({"status": "built", "rows": len(items), "workbook": str(output)}, ensure_ascii=False))


def validate_response_format(value: str, row: int) -> None:
    """Check the two user-approved cell layouts, not the factual conclusion."""
    response = value.replace("\r\n", "\n").replace("\r", "\n").strip()
    if response.startswith(("*", "＊", "★", "☆")):
        raise ValueError(f"Response starts with requirement star at workbook row {row}")
    if "优于补充说明：" in response:
        parts = response.split("\n\n")
        if (len(parts) != 3 or not parts[0].strip()
                or not parts[1].startswith("优于补充说明：")
                or not parts[1][len("优于补充说明："):].strip()
                or not parts[2].startswith("证明材料详见：")
                or not parts[2][len("证明材料详见："):].strip()):
            raise ValueError(f"Positive response spacing/evidence format invalid at workbook row {row}")
    elif "补充说明：" in response:
        parts = response.split("\n\n")
        if (len(parts) != 2 or not parts[0].strip()
                or not parts[1].startswith("补充说明：")
                or not parts[1][len("补充说明："):].strip()):
            raise ValueError(f"Adjustment response spacing format invalid at workbook row {row}")


def validate(source: Path, workbook: Path, manifest_path: Path, require_responses: bool = False) -> None:
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if manifest["source_sha256"] != digest(source.read_bytes()):
        raise ValueError("Word source hash changed")
    fresh = selected_blocks(source, manifest.get("start"), manifest.get("end"))
    items = manifest["items"]
    if len(fresh) != len(items) or manifest["row_count"] != len(items):
        raise ValueError("Source block count differs from manifest")
    wb = load_workbook(workbook, read_only=True, data_only=True)
    try:
        if "技术对标" not in wb:
            raise ValueError("Technical alignment sheet missing")
        ws = wb["技术对标"]
        headers = [ws.cell(2, col).value for col in range(1, ws.max_column + 1)]
        scope = manifest.get("project_scope")
        if scope:
            checked_scope(scope, items)
            expected_headers = MODULE_HEADERS + model_headers(scope) + TAIL_HEADERS
        else:
            if require_responses:
                raise ValueError("Project model scope missing; cannot finalize alignment")
            expected_headers = HEADERS  # read-only compatibility for earlier blank v1.38 workbooks
        if ws["A1"].value != NOTICE or headers != expected_headers:
            raise ValueError("Workbook notice or column order changed")
        requirement_col = len(headers) - 2
        response_col = len(headers) - 1
        if ws.max_row != len(items) + 2:
            raise ValueError("Workbook source row count changed")
        evidence_pairs: set[tuple[str, str]] = set()
        valid_modules = valid_module_triples()
        if require_responses and scope:
            if "证据索引" not in wb:
                raise ValueError("Report-level evidence index missing")
            evidence = wb["证据索引"]
            source_orders = {item["requirement_id"]: item["source_order"] for item in items}
            for evidence_row in evidence.iter_rows(min_row=2, values_only=True):
                requirement_id, model_id = evidence_row[0], evidence_row[1]
                if not requirement_id and not model_id:
                    continue
                if model_id not in {model["model_id"] for model in scope["models"]}:
                    raise ValueError(f"Evidence index contains unrelated project model: {model_id}")
                if requirement_id not in source_orders or model_id not in section_for_order(scope, source_orders[requirement_id])["model_ids"]:
                    raise ValueError("Evidence index points to unrelated requirement/model section")
                if not all(evidence_row[position] for position in (0, 2, 4, 5, 6, 7)):
                    raise ValueError("Evidence index lacks standard model, report ID or source location")
                evidence_pairs.add((str(requirement_id), str(model_id)))
        for index, (original, item) in enumerate(zip(fresh, items), 3):
            value = ws.cell(index, requirement_col).value
            if original["text"] != item["text"] or value != item["text"]:
                raise ValueError(f"Original requirement differs at workbook row {index}")
            if digest(str(value).encode("utf-8")) != item["text_sha256"]:
                raise ValueError(f"Requirement hash differs at workbook row {index}")
            module = tuple(str(ws.cell(index, col).value or "").strip() for col in range(1, 4))
            if any(module) and module not in valid_modules:
                raise ValueError(f"Module classification is incomplete or not in taxonomy at workbook row {index}")
            if original["kind"] == "heading" and any(module):
                raise ValueError(f"Heading is not a business module at workbook row {index}")
            if scope:
                active = set(section_for_order(scope, item["source_order"])["model_ids"])
                for model_col, model in enumerate(scope["models"], 4):
                    model_value = ws.cell(index, model_col).value
                    if original["kind"] == "heading" and model_value not in (None, ""):
                        raise ValueError(f"Standard answer on heading at workbook row {index}")
                    if (require_responses and original["kind"] not in {"heading", "goods_table", "environment"}
                            and model_value not in (None, "") and not any(module)):
                        raise ValueError(f"Standard answer lacks same-row module mapping at workbook row {index}")
                    if model["model_id"] not in active and model_value not in (None, ""):
                        raise ValueError(f"Unrelated model data at workbook row {index}")
                    if require_responses and model_value not in (None, "") and (item["requirement_id"], model["model_id"]) not in evidence_pairs:
                        raise ValueError(f"Report-level evidence missing at workbook row {index}")
            if require_responses and original["kind"] != "heading":
                response = str(ws.cell(index, response_col).value or "")
                if not response.strip():
                    raise ValueError(f"Response missing at workbook row {index}")
                validate_response_format(response, index)
    finally:
        wb.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    make = commands.add_parser("build", help="Build source-locked blank workbook from a Word technical range")
    make.add_argument("--source", required=True)
    make.add_argument("--output", required=True)
    make.add_argument("--manifest", required=True)
    make.add_argument("--scope", required=True, help="JSON with current tender models and source-order sections")
    make.add_argument("--start", help="Exact start-heading substring, included")
    make.add_argument("--end", help="Exact end-heading substring, excluded")
    check = commands.add_parser("validate", help="Compare workbook requirements with original Word")
    check.add_argument("--source", required=True)
    check.add_argument("--workbook", required=True)
    check.add_argument("--manifest", required=True)
    check.add_argument("--require-responses", action="store_true")
    args = parser.parse_args()
    if args.command == "build":
        build(args)
    else:
        validate(Path(args.source), Path(args.workbook), Path(args.manifest), args.require_responses)
        print(json.dumps({"status": "valid"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
