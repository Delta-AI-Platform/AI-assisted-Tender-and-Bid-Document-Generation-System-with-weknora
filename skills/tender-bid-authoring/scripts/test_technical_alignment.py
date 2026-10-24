import argparse
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from technical_alignment import build, selected_blocks, validate, validate_response_format


class TechnicalAlignmentTest(unittest.TestCase):
    def test_positive_and_adjustment_spacing(self):
        validate_response_format("符合原文条件。\n\n优于补充说明：经核对优于要求。\n\n证明材料详见：R2 第3页。", 3)
        validate_response_format("符合原文条件。\n\n补充说明：按项目要求调整，待技术确认。", 4)
        with self.assertRaisesRegex(ValueError, "spacing"):
            validate_response_format("符合原文条件。\n优于补充说明：优于。\n证明材料详见：R2。", 5)
        with self.assertRaisesRegex(ValueError, "spacing"):
            validate_response_format("符合原文条件。\n补充说明：待确认。", 6)

    def test_automatic_numbering_requires_manual_review(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "numbered.docx"
            doc = Document()
            doc.add_paragraph("技术要求")
            doc.add_paragraph("电机功率要求", style="List Number")
            doc.save(source)
            with self.assertRaisesRegex(ValueError, "automatic numbering"):
                selected_blocks(source, "技术要求", None)

    def test_build_preserves_order_and_detects_changed_requirement(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source, workbook, manifest, scope_file = root / "source.docx", root / "alignment.xlsx", root / "manifest.json", root / "scope.json"
            doc = Document()
            doc.add_paragraph("前言")
            doc.add_heading("一、货物需求一览表", 1)
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "名称"
            table.cell(0, 1).text = "数量"
            table.cell(1, 0).text = "车辆"
            table.cell(1, 1).text = "2"
            doc.add_paragraph("二、工作环境：温度不低于-20℃。")
            doc.add_paragraph("三、技术参数及要求")
            doc.add_paragraph("1.1 电机功率须不低于100kW。")
            doc.add_paragraph("附件")
            doc.save(source)
            scope_file.write_text(json.dumps({
                "layout": "combined",
                "models": [{"model_id": "A", "tender_model": "19座人车"}, {"model_id": "B", "tender_model": "3吨料车"}],
                "sections": [{"source_heading": "一、货物需求一览表", "first_source_order": 1, "last_source_order": 4, "model_ids": ["A", "B"]}],
            }, ensure_ascii=False), encoding="utf-8")
            build(argparse.Namespace(source=str(source), output=str(workbook), manifest=str(manifest), scope=str(scope_file), start="一、货物需求一览表", end="附件"))
            validate(source, workbook, manifest)
            items = json.loads(manifest.read_text(encoding="utf-8"))["items"]
            self.assertEqual(4, len(items))
            self.assertEqual("goods_table", items[0]["kind"])
            self.assertEqual("table", items[0]["source_parts"][1]["kind"])
            self.assertEqual("environment", items[1]["kind"])
            self.assertEqual("1.1 电机功率须不低于100kW。", items[-1]["text"])
            wb = load_workbook(workbook)
            ws = wb["技术对标"]
            ws["F6"] = "1.1 电机功率须不低于100kW。"
            for row in (3, 4, 6):
                ws.cell(row, 7, "满足本项要求。")
            ws["A6"] = "永磁电驱动子系统"
            ws["B6"] = "驱动电机"
            ws["C6"] = "驱动电机功率"
            ws["D6"] = "100kW；报告R1第3页"
            wb.save(workbook)
            wb.close()
            validate(source, workbook, manifest)
            with self.assertRaisesRegex(ValueError, "Report-level evidence missing"):
                validate(source, workbook, manifest, require_responses=True)
            wb = load_workbook(workbook)
            wb["证据索引"].append([items[-1]["requirement_id"], "A", "WLR-19", "V1", "R1", "检验报告.docx", "sha256:test", "第3页", "技术附件一", "待审核"])
            wb.save(workbook)
            wb.close()
            validate(source, workbook, manifest, require_responses=True)
            wb = load_workbook(workbook)
            ws = wb["技术对标"]
            ws["A6"] = None
            ws["B6"] = None
            ws["C6"] = None
            wb.save(workbook)
            wb.close()
            with self.assertRaisesRegex(ValueError, "same-row module mapping"):
                validate(source, workbook, manifest, require_responses=True)
            wb = load_workbook(workbook)
            ws = wb["技术对标"]
            ws["A6"] = "永磁电驱动子系统"
            ws["B6"] = "驱动电机"
            ws["C6"] = "驱动电机功率"
            wb.save(workbook)
            wb.close()
            validate(source, workbook, manifest, require_responses=True)
            wb = load_workbook(workbook)
            ws = wb["技术对标"]
            ws["F6"] = "1.1 电机功率须不低于90kW。"
            wb.save(workbook)
            wb.close()
            with self.assertRaisesRegex(ValueError, "differs"):
                validate(source, workbook, manifest)

    def test_unrequested_modules_do_not_create_rows_or_standard_answers(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source, workbook, manifest, scope_file = (root / name for name in ("source.docx", "alignment.xlsx", "manifest.json", "scope.json"))
            doc = Document()
            doc.add_paragraph("1. 驱动电机功率不低于100kW。")
            doc.save(source)
            scope_file.write_text(json.dumps({
                "layout": "combined",
                "models": [{"model_id": "A", "tender_model": "19座人车"}],
                "sections": [{"source_heading": "1. 驱动电机功率不低于100kW。", "first_source_order": 1, "last_source_order": 1, "model_ids": ["A"]}],
            }, ensure_ascii=False), encoding="utf-8")
            build(argparse.Namespace(source=str(source), output=str(workbook), manifest=str(manifest), scope=str(scope_file), start=None, end=None))
            wb = load_workbook(workbook)
            ws = wb["技术对标"]
            self.assertEqual(3, ws.max_row)  # taxonomy entries are not extra tender rows
            self.assertIsNone(ws["D3"].value)
            ws["A3"] = "永磁电驱动子系统"
            ws["B3"] = "驱动电机"
            ws["C3"] = "不存在的功能"
            wb.save(workbook)
            wb.close()
            with self.assertRaisesRegex(ValueError, "not in taxonomy"):
                validate(source, workbook, manifest)

    def test_separate_model_sections_reject_cross_model_data(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source, workbook, manifest, scope_file = (root / name for name in ("source.docx", "alignment.xlsx", "manifest.json", "scope.json"))
            doc = Document()
            doc.add_heading("车型A技术要求", 1)
            doc.add_paragraph("1. 电机功率不低于100kW。")
            doc.add_heading("车型B技术要求", 1)
            doc.add_paragraph("1. 电机功率不低于80kW。")
            doc.save(source)
            scope_file.write_text(json.dumps({
                "layout": "separate",
                "models": [{"model_id": "A", "tender_model": "车型A"}, {"model_id": "B", "tender_model": "车型B"}],
                "sections": [
                    {"source_heading": "车型A技术要求", "first_source_order": 1, "last_source_order": 2, "model_ids": ["A"]},
                    {"source_heading": "车型B技术要求", "first_source_order": 3, "last_source_order": 4, "model_ids": ["B"]},
                ],
            }, ensure_ascii=False), encoding="utf-8")
            build(argparse.Namespace(source=str(source), output=str(workbook), manifest=str(manifest), scope=str(scope_file), start=None, end=None))
            validate(source, workbook, manifest)
            wb = load_workbook(workbook)
            wb["技术对标"]["E4"] = "来自车型B的报告参数"
            wb.save(workbook)
            wb.close()
            with self.assertRaisesRegex(ValueError, "Unrelated model"):
                validate(source, workbook, manifest)


if __name__ == "__main__":
    unittest.main()
