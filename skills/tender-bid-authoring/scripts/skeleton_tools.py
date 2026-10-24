#!/usr/bin/env python3
"""Build, freeze and verify bid-document skeletons and controlled fill slots.

The visible skeleton contains only tender-controlled headings, fixed text,
tables and material captions.  Mutable content lives inside tagged Word
content controls (w:sdt).  Structure comparison removes the mutable content
before hashing, so text, tables and images may grow naturally inside a slot
without authorising changes to headings, sections or fixed table cells.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
import shutil
import sys
from pathlib import Path
from typing import Any, Iterable
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree


SLOT_PREFIX = "BID_SLOT:"
TOOL_VERSION = "1.3.0"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}
SKILL_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_STYLE_PROFILE = SKILL_ROOT / "assets" / "skeleton-format-standard.json"
DEFAULT_STYLE_TEMPLATE = SKILL_ROOT / "assets" / "skeleton-style-template.docx"

STYLE_NAMES = {
    "body": "Bid Body",
    "table": "Bid Table",
    "technical_table": "Bid Technical Table",
    "material_caption": "Bid Material Caption",
    "cover_title": "Bid Cover Title",
    "cover_volume": "Bid Cover Volume",
    "cover_info": "Bid Cover Info",
    "cover_project_start": "Bid Cover Project Start",
    "cover_bidder": "Bid Cover Bidder",
    "cover_date": "Bid Cover Date",
    "cover_footer": "Bid Cover Footer",
    "toc_title": "Bid TOC Title",
    "signature": "Bid Signature",
}


class SkeletonError(RuntimeError):
    pass


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def canonical_hash(value: Any) -> str:
    payload = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def read_json(path: Path) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise SkeletonError(f"Missing JSON file: {path}") from exc
    except json.JSONDecodeError as exc:
        raise SkeletonError(f"Invalid JSON: {path}: {exc}") from exc


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def resolve(path_value: str, base: Path) -> Path:
    path = Path(path_value)
    return path if path.is_absolute() else (base / path).resolve()


def clean_id(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9_.-]+", "_", str(value)).strip("_")
    if not value:
        raise SkeletonError("Stable ids must contain a letter, digit, dot, underscore or dash")
    return value


def _rgb(value: str | None) -> RGBColor | None:
    if not value:
        return None
    value = str(value).lstrip("#").upper()
    if not re.fullmatch(r"[0-9A-F]{6}", value):
        raise SkeletonError(f"Invalid RGB colour: {value}")
    return RGBColor.from_string(value)


def _set_font(font: Any, element: Any, spec: dict[str, Any]) -> None:
    east_asia = spec.get("font_east_asia", "宋体")
    latin = spec.get("font_latin", "Times New Roman")
    font.name = latin
    r_pr = element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for key, value in (("w:ascii", latin), ("w:hAnsi", latin), ("w:eastAsia", east_asia), ("w:cs", latin)):
        r_fonts.set(qn(key), str(value))
    if spec.get("size_pt") is not None:
        font.size = Pt(float(spec["size_pt"]))
    if spec.get("bold") is not None:
        font.bold = bool(spec["bold"])
    if _rgb(spec.get("color")) is not None:
        font.color.rgb = _rgb(spec.get("color"))


def _set_paragraph_format(paragraph_format: Any, spec: dict[str, Any]) -> None:
    if spec.get("space_before_pt") is not None:
        paragraph_format.space_before = Pt(float(spec["space_before_pt"]))
    if spec.get("space_after_pt") is not None:
        paragraph_format.space_after = Pt(float(spec["space_after_pt"]))
    if spec.get("line_spacing") is not None:
        paragraph_format.line_spacing = float(spec["line_spacing"])
    if spec.get("first_line_indent_pt") is not None:
        paragraph_format.first_line_indent = Pt(float(spec["first_line_indent_pt"]))
    if spec.get("left_indent_cm") is not None:
        paragraph_format.left_indent = Cm(float(spec["left_indent_cm"]))
    if spec.get("right_indent_cm") is not None:
        paragraph_format.right_indent = Cm(float(spec["right_indent_cm"]))
    if spec.get("keep_with_next") is not None:
        paragraph_format.keep_with_next = bool(spec["keep_with_next"])


def _alignment(value: str | None) -> Any:
    mapping = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    return mapping.get(str(value).lower()) if value is not None else None


def _ensure_paragraph_style(document: Any, name: str, spec: dict[str, Any], base: str = "Normal") -> Any:
    styles = document.styles
    style = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    if base in styles and style.base_style is None and name != base:
        style.base_style = styles[base]
    _set_font(style.font, style._element, spec)
    _set_paragraph_format(style.paragraph_format, spec)
    align = _alignment(spec.get("alignment"))
    if align is not None:
        style.paragraph_format.alignment = align
    return style


def _load_style_profile(spec: dict[str, Any], spec_path: Path) -> tuple[dict[str, Any], Path]:
    reference = spec.get("style_profile")
    path = resolve(str(reference), spec_path.parent) if reference else DEFAULT_STYLE_PROFILE
    profile = read_json(path)
    if profile.get("schema_version") != "1.0" or not profile.get("profile_id"):
        raise SkeletonError(f"Invalid skeleton style profile: {path}")
    return profile, path


def _apply_style_profile(document: Any, profile: dict[str, Any]) -> None:
    styles = profile.get("styles", {})
    normal_spec = styles.get("normal", {})
    normal = document.styles["Normal"]
    _set_font(normal.font, normal._element, normal_spec)
    _set_paragraph_format(normal.paragraph_format, normal_spec)
    for level in range(1, 5):
        name = f"Heading {level}"
        if name in document.styles:
            spec = styles.get(f"heading_{level}", styles.get("heading_3", {}))
            _set_font(document.styles[name].font, document.styles[name]._element, spec)
            _set_paragraph_format(document.styles[name].paragraph_format, spec)
            align = _alignment(spec.get("alignment"))
            if align is not None:
                document.styles[name].paragraph_format.alignment = align
    for key, name in STYLE_NAMES.items():
        _ensure_paragraph_style(document, name, styles.get(key, normal_spec))


def _style_id(document: Any, role: str) -> str:
    return document.styles[STYLE_NAMES[role]].style_id


def _style_paragraph(paragraph: Any, document: Any, role: str) -> None:
    paragraph.style = document.styles[STYLE_NAMES[role]]


def _slot_sdt(slot_id: str, *, inline: bool = False, paragraph_style_id: str | None = None, placeholder: str | None = None) -> OxmlElement:
    slot_id = clean_id(slot_id)
    sdt = OxmlElement("w:sdt")
    props = OxmlElement("w:sdtPr")
    control_id = OxmlElement("w:id")
    control_id.set(qn("w:val"), str(int(hashlib.sha256(slot_id.encode("utf-8")).hexdigest()[:7], 16)))
    tag = OxmlElement("w:tag")
    tag.set(qn("w:val"), SLOT_PREFIX + slot_id)
    alias = OxmlElement("w:alias")
    alias.set(qn("w:val"), slot_id)
    lock = OxmlElement("w:lock")
    lock.set(qn("w:val"), "sdtLocked")
    props.extend([control_id, tag, alias, lock])
    if placeholder is not None:
        props.append(OxmlElement("w:showingPlcHdr"))
    content = OxmlElement("w:sdtContent")
    child = OxmlElement("w:r" if inline else "w:p")
    if not inline and paragraph_style_id:
        p_pr = OxmlElement("w:pPr")
        p_style = OxmlElement("w:pStyle")
        p_style.set(qn("w:val"), paragraph_style_id)
        p_pr.append(p_style)
        child.append(p_pr)
    if placeholder is not None:
        if inline:
            text = OxmlElement("w:t")
            text.text = str(placeholder)
            child.append(text)
        else:
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = str(placeholder)
            run.append(text)
            child.append(run)
    content.append(child)
    sdt.extend([props, content])
    return sdt


def _append_slot(container: Any, slot_id: str, paragraph_style_id: str | None = None) -> None:
    body = getattr(container, "_body", None)
    target = body._element if body is not None else container._element
    control = _slot_sdt(slot_id, paragraph_style_id=paragraph_style_id)
    section_properties = target.find(qn("w:sectPr"))
    if section_properties is not None:
        section_properties.addprevious(control)
    else:
        target.append(control)


def _replace_cell_with_slot(cell: Any, slot_id: str, paragraph_style_id: str | None = None) -> None:
    tc = cell._tc
    for child in list(tc):
        if child.tag != qn("w:tcPr"):
            tc.remove(child)
    tc.append(_slot_sdt(slot_id, paragraph_style_id=paragraph_style_id))


def _repeat_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _set_cell_text(
    cell: Any,
    value: str,
    profile: dict[str, Any],
    *,
    role: str = "table_body",
    bold: bool | None = None,
    alignment: str | None = None,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    run = paragraph.add_run(str(value))
    style = copy.deepcopy(profile.get("styles", {}).get(role, {}))
    if bold is not None:
        style["bold"] = bold
    _set_font(run.font, run._element, style)
    paragraph.alignment = _alignment(alignment or style.get("alignment"))
    _set_paragraph_format(paragraph.paragraph_format, style)
    if role in {"table_header", "technical_table_header"}:
        fill = profile.get("palette", {}).get("table_header_fill")
        if fill:
            tc_pr = cell._tc.get_or_add_tcPr()
            shd = tc_pr.find(qn("w:shd"))
            if shd is None:
                shd = OxmlElement("w:shd")
                tc_pr.append(shd)
            shd.set(qn("w:fill"), str(fill).lstrip("#").upper())


def _append_inline_parts(
    paragraph: Any,
    parts: list[Any],
    slots_out: list[dict[str, Any]],
    slot_defs: dict[str, dict[str, Any]],
    volume_id: str,
    section_id: str,
    title: str,
) -> None:
    """Append fixed text and invisible, tagged inline fill positions."""
    for part in parts:
        if isinstance(part, str):
            paragraph.add_run(part)
        elif "text" in part:
            paragraph.add_run(str(part.get("text", "")))
        else:
            slot_id = clean_id(part["slot_id"])
            paragraph._p.append(_slot_sdt(slot_id, inline=True, placeholder=part.get("placeholder")))
            _register_slot(
                slots_out, slot_defs, slot_id, volume_id, section_id, title,
                {"type": "inline_content_control", "tag": SLOT_PREFIX + slot_id}, part,
            )


def _append_toc(document: Any, levels: str = "1-3", entries: list[str] | None = None) -> None:
    """Insert an updatable Word TOC field."""
    paragraph = document.add_paragraph()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText"); instruction.set(qn("xml:space"), "preserve")
    instruction.text = f' TOC \\o "{levels}" \\h \\z \\u '
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    result_nodes = []
    for index, value in enumerate(entries or ["目录将在打开或最终渲染时更新"]):
        run = OxmlElement("w:r"); value_node = OxmlElement("w:t")
        value_node.text = str(value); run.append(value_node); result_nodes.append(run)
        if index < len(entries or []) - 1:
            break_run = OxmlElement("w:r"); break_run.append(OxmlElement("w:br")); result_nodes.append(break_run)
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    paragraph._p.extend([begin, instruction, separate, *result_nodes, end])
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields"); settings.append(update)
    update.set(qn("w:val"), "true")


def _apply_section(section: Any, config: dict[str, Any]) -> None:
    orientation = str(config.get("orientation", "portrait")).lower()
    if str(config.get("size", "A4")).upper() == "A4":
        portrait_width, portrait_height = Cm(21.0), Cm(29.7)
        section.page_width = portrait_width
        section.page_height = portrait_height
    if orientation == "landscape":
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = portrait_height
        section.page_height = portrait_width
    elif orientation == "portrait":
        section.orientation = WD_ORIENT.PORTRAIT
    else:
        raise SkeletonError(f"Unknown orientation: {orientation}")
    margins = config.get("margins_cm", {})
    for key, attr in (("top", "top_margin"), ("bottom", "bottom_margin"), ("left", "left_margin"), ("right", "right_margin")):
        if key in margins:
            setattr(section, attr, Cm(float(margins[key])))


def _page_config(profile: dict[str, Any], override: dict[str, Any] | None = None) -> dict[str, Any]:
    base = copy.deepcopy(profile.get("page", {}))
    base.pop("landscape_table_margins_cm", None)
    override = copy.deepcopy(override or {})
    base_margins = base.get("margins_cm", {})
    base_margins.update(override.pop("margins_cm", {}))
    base.update(override)
    base["margins_cm"] = base_margins
    return base


def create_style_template(profile_path: Path, output: Path) -> dict[str, Any]:
    profile = read_json(profile_path)
    document = Document()
    _apply_style_profile(document, profile)
    _apply_section(document.sections[0], _page_config(profile))
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return {
        "output": str(output.resolve()),
        "sha256": sha256(output),
        "profile_id": profile["profile_id"],
        "profile_sha256": sha256(profile_path),
    }


def _register_slot(
    slots_out: list[dict[str, Any]],
    slot_defs: dict[str, dict[str, Any]],
    slot_id: str,
    volume_id: str,
    section_id: str,
    title: str,
    anchor: dict[str, Any],
    overrides: dict[str, Any] | None = None,
) -> None:
    slot_id = clean_id(slot_id)
    if any(item["slot_id"] == slot_id for item in slots_out):
        raise SkeletonError(f"Duplicate slot id: {slot_id}")
    supplied = copy.deepcopy(slot_defs.get(slot_id, {}))
    supplied.update(copy.deepcopy(overrides or {}))
    item = {
        "slot_id": slot_id,
        "volume_id": volume_id,
        "section_id": section_id,
        "display_title": supplied.pop("display_title", title),
        "slot_type": supplied.pop("slot_type", "narrative"),
        "requirement_ids": supplied.pop("requirement_ids", []),
        "scoring_ids": supplied.pop("scoring_ids", []),
        "allowed_material_types": supplied.pop("allowed_material_types", []),
        "applicability": supplied.pop("applicability", {}),
        "retrieval": supplied.pop("retrieval", {"keywords": [title] if title else []}),
        "allow_blank": bool(supplied.pop("allow_blank", False)),
        "allow_repeat": bool(supplied.pop("allow_repeat", False)),
        "natural_overflow": bool(supplied.pop("natural_overflow", True)),
        "allowed_stages": supplied.pop("allowed_stages", []),
        "required_by_stage": supplied.pop("required_by_stage", None),
        "anchor": anchor,
        **supplied,
    }
    slots_out.append(item)


def build_skeleton(spec_path: Path, output_dir: Path) -> dict[str, Path]:
    spec = read_json(spec_path)
    style_profile, style_profile_path = _load_style_profile(spec, spec_path)
    project_id = str(spec.get("project_id", "")).strip()
    volumes = spec.get("volumes")
    if not project_id or not isinstance(volumes, list) or not volumes:
        raise SkeletonError("Skeleton spec needs project_id and non-empty volumes")
    output_dir.mkdir(parents=True, exist_ok=True)
    frozen_style_profile = output_dir / "skeleton_style_profile.json"
    if style_profile_path.resolve() != frozen_style_profile.resolve():
        shutil.copy2(style_profile_path, frozen_style_profile)
    slot_defs = {str(item["slot_id"]): item for item in spec.get("slots", [])}
    slots_out: list[dict[str, Any]] = []
    requirements_out: list[dict[str, Any]] = []
    bundle_items: list[dict[str, Any]] = []

    for volume in volumes:
        volume_id = clean_id(volume["id"])
        role = volume.get("role", volume_id)
        output_name = volume.get("output", f"{project_id}_{volume_id}_blank.docx")
        output_path = (output_dir / output_name).resolve()
        template_ref = spec.get("style_template")
        template_path = resolve(str(template_ref), spec_path.parent) if template_ref else DEFAULT_STYLE_TEMPLATE
        document = Document(str(template_path)) if template_path.is_file() else Document()
        volume_page = _page_config(style_profile, volume.get("page", {}))
        _apply_section(document.sections[0], volume_page)
        _apply_style_profile(document, style_profile)

        cover = volume.get("cover")
        if cover:
            title = document.add_paragraph()
            _style_paragraph(title, document, "cover_title")
            run = title.add_run(str(cover.get("title", "投标文件")))
            title_spec = copy.deepcopy(style_profile.get("styles", {}).get("cover_title", {}))
            if cover.get("title_size_pt") is not None:
                title_spec["size_pt"] = float(cover["title_size_pt"])
            _set_font(run.font, run._element, title_spec)
            for line_spec in cover.get("lines", []):
                paragraph = document.add_paragraph()
                line_role = "cover_info"
                if isinstance(line_spec, dict):
                    line_role = str(line_spec.get("style_role", line_role))
                    label_probe = str(line_spec.get("label", ""))
                    if line_role == "cover_info" and "投标人" in label_probe:
                        line_role = "cover_bidder"
                    elif line_role == "cover_info" and "日期" in label_probe:
                        line_role = "cover_date"
                _style_paragraph(paragraph, document, line_role if line_role in STYLE_NAMES else "cover_info")
                if isinstance(line_spec, str):
                    run = paragraph.add_run(line_spec)
                    _set_font(run.font, run._element, style_profile.get("styles", {}).get("cover_info", {}))
                    continue
                label = str(line_spec.get("label", ""))
                if label:
                    run = paragraph.add_run(label)
                    _set_font(run.font, run._element, style_profile.get("styles", {}).get(line_role, {}))
                if "text" in line_spec:
                    run = paragraph.add_run(str(line_spec.get("text", "")))
                    _set_font(run.font, run._element, style_profile.get("styles", {}).get(line_role, {}))
                if "slot_id" in line_spec:
                    slot_id = clean_id(line_spec["slot_id"])
                    paragraph._p.append(_slot_sdt(slot_id, inline=True, placeholder=line_spec.get("placeholder")))
                    _register_slot(
                        slots_out,
                        slot_defs,
                        slot_id,
                        volume_id,
                        "cover",
                        str(cover.get("title", "投标文件")),
                        {"type": "content_control", "tag": SLOT_PREFIX + slot_id},
                        line_spec,
                    )
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

        first_section = True
        table_index = 0
        for section_spec in volume.get("sections", []):
            section_id = clean_id(section_spec["id"])
            if not first_section and section_spec.get("section_break"):
                new_section = document.add_section(WD_SECTION.NEW_PAGE)
                _apply_section(new_section, _page_config(style_profile, section_spec.get("page", volume_page)))
            title = str(section_spec.get("title", "")).strip()
            if not title:
                raise SkeletonError(f"Section {section_id} has no title")
            level = max(1, min(9, int(section_spec.get("level", 1))))
            heading = document.add_heading(title, level=level)
            heading.paragraph_format.page_break_before = bool(
                section_spec.get("page_break_before", not first_section)
            )
            heading.paragraph_format.keep_with_next = True
            first_section = False

            for block in section_spec.get("blocks", []):
                kind = block.get("kind")
                if kind == "fixed_paragraph":
                    paragraph = document.add_paragraph()
                    _style_paragraph(paragraph, document, "body")
                    if isinstance(block.get("parts"), list):
                        _append_inline_parts(paragraph, block["parts"], slots_out, slot_defs, volume_id, section_id, title)
                    else:
                        paragraph.add_run(str(block.get("text", "")))
                elif kind == "toc":
                    _append_toc(document, str(block.get("levels", "1-3")), [str(x) for x in block.get("entries", [])] or None)
                elif kind == "slot":
                    slot_id = clean_id(block["slot_id"])
                    _append_slot(document, slot_id, _style_id(document, "body"))
                    _register_slot(
                        slots_out, slot_defs, slot_id, volume_id, section_id, title,
                        {"type": "content_control", "tag": SLOT_PREFIX + slot_id}, block,
                    )
                elif kind == "material_slot":
                    material_title = str(block["title"])
                    paragraph = document.add_paragraph(material_title)
                    _style_paragraph(paragraph, document, "material_caption")
                    slot_id = clean_id(block["slot_id"])
                    _append_slot(document, slot_id, _style_id(document, "body"))
                    _register_slot(
                        slots_out, slot_defs, slot_id, volume_id, section_id, material_title,
                        {"type": "content_control", "tag": SLOT_PREFIX + slot_id}, block,
                    )
                elif kind == "table":
                    headers = [str(x) for x in block.get("headers", [])]
                    rows = block.get("rows", [])
                    has_header = bool(block.get("has_header", True))
                    columns = len(headers) if headers else int(block.get("columns", 0))
                    if columns < 1 or (has_header and not headers):
                        raise SkeletonError(f"Table in {section_id} needs headers or an explicit column count")
                    table = document.add_table(rows=1 if has_header else 0, cols=columns)
                    table.style = block.get("style", "Table Grid")
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    if has_header:
                        for index, header in enumerate(headers):
                            _set_cell_text(table.rows[0].cells[index], header, style_profile, role="table_header", bold=True)
                        if block.get("repeat_header", True):
                            _repeat_header(table.rows[0])
                    row_offset = 1 if has_header else 0
                    merge_specs = []
                    for logical_row_index, row_spec in enumerate(rows):
                        row_index = logical_row_index + row_offset
                        cells = row_spec.get("cells", [])
                        if len(cells) != columns:
                            raise SkeletonError(f"Table row width mismatch in {section_id}")
                        row = table.add_row()
                        for column, cell_spec in enumerate(cells):
                            if "slot_id" in cell_spec:
                                slot_id = clean_id(cell_spec["slot_id"])
                                _replace_cell_with_slot(row.cells[column], slot_id, _style_id(document, "table"))
                                _register_slot(
                                    slots_out, slot_defs, slot_id, volume_id, section_id, title,
                                    {"type": "table_cell", "table_index": table_index, "row": row_index, "column": column, "tag": SLOT_PREFIX + slot_id},
                                    cell_spec,
                                )
                            elif isinstance(cell_spec.get("parts"), list):
                                cell = row.cells[column]
                                cell.text = ""
                                paragraph = cell.paragraphs[0]
                                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                                paragraph.alignment = _alignment(cell_spec.get("alignment", "left"))
                                for part in cell_spec["parts"]:
                                    if isinstance(part, str) or "text" in part:
                                        value = part if isinstance(part, str) else str(part.get("text", ""))
                                        run = paragraph.add_run(value)
                                        _set_font(run.font, run._element, style_profile.get("styles", {}).get("table_body", {}))
                                    else:
                                        slot_id = clean_id(part["slot_id"])
                                        paragraph._p.append(_slot_sdt(slot_id, inline=True, placeholder=part.get("placeholder")))
                                        _register_slot(
                                            slots_out, slot_defs, slot_id, volume_id, section_id, title,
                                            {"type": "table_cell_inline", "table_index": table_index, "row": row_index, "column": column, "tag": SLOT_PREFIX + slot_id},
                                            part,
                                        )
                            else:
                                _set_cell_text(row.cells[column], str(cell_spec.get("fixed", "")), style_profile)
                        for column, cell_spec in enumerate(cells):
                            merge_right = int(cell_spec.get("merge_right", 0))
                            merge_down = int(cell_spec.get("merge_down", 0))
                            if merge_right or merge_down:
                                if column + merge_right >= columns or logical_row_index + merge_down >= len(rows):
                                    raise SkeletonError(f"Table merge exceeds table bounds in {section_id}")
                                merge_specs.append((row_index, column, merge_down, merge_right))
                    for row_index, column, merge_down, merge_right in merge_specs:
                        table.cell(row_index, column).merge(table.cell(row_index + merge_down, column + merge_right))
                    table_index += 1
                elif kind == "technical_response_table":
                    headers = [str(x) for x in block.get("headers", ["序号", "招标人要求", "投标人响应", "偏离说明"])]
                    if len(headers) < 4:
                        raise SkeletonError("Technical response table needs at least four columns")
                    table_id = clean_id(block["id"])
                    table = document.add_table(rows=1, cols=len(headers))
                    table.style = block.get("style", "Table Grid")
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    for index, header in enumerate(headers):
                        _set_cell_text(table.rows[0].cells[index], header, style_profile, role="technical_table_header", bold=True)
                    _repeat_header(table.rows[0])
                    for row_index, requirement in enumerate(block.get("requirements", []), 1):
                        requirement_id = clean_id(requirement["requirement_id"])
                        response_slot = clean_id(requirement.get("response_slot_id", f"{requirement_id}.response"))
                        deviation_slot = clean_id(requirement.get("deviation_slot_id", f"{requirement_id}.deviation"))
                        row = table.add_row()
                        _set_cell_text(row.cells[0], str(requirement.get("number", row_index)), style_profile, role="technical_table_body", alignment="center")
                        _set_cell_text(row.cells[1], str(requirement["requirement_text"]), style_profile, role="technical_table_body", alignment="left")
                        _replace_cell_with_slot(row.cells[2], response_slot, _style_id(document, "technical_table"))
                        _replace_cell_with_slot(row.cells[3], deviation_slot, _style_id(document, "technical_table"))
                        _register_slot(
                            slots_out, slot_defs, response_slot, volume_id, section_id, title,
                            {"type": "table_cell", "table_index": table_index, "row": row_index, "column": 2, "tag": SLOT_PREFIX + response_slot},
                            {"slot_type": "technical_response", "requirement_ids": [requirement_id], "allowed_stages": [7, 8, 9], "allow_blank": True, "blank_reason": "business_directed_blank"},
                        )
                        _register_slot(
                            slots_out, slot_defs, deviation_slot, volume_id, section_id, title,
                            {"type": "table_cell", "table_index": table_index, "row": row_index, "column": 3, "tag": SLOT_PREFIX + deviation_slot},
                            {"slot_type": "technical_deviation", "requirement_ids": [requirement_id], "allowed_stages": [7, 8, 9], "allow_blank": True, "blank_reason": "business_directed_blank"},
                        )
                        requirements_out.append({
                            "requirement_id": requirement_id,
                            "volume_id": volume_id,
                            "table_id": table_id,
                            "table_index": table_index,
                            "row_index": row_index,
                            "number": str(requirement.get("number", row_index)),
                            "requirement_text": str(requirement["requirement_text"]),
                            "source_location": requirement.get("source_location"),
                            "composite_parent": requirement.get("composite_parent"),
                            "requirement_column": 1,
                            "response_column": 2,
                            "deviation_column": 3,
                            "response_slot_id": response_slot,
                            "deviation_slot_id": deviation_slot,
                        })
                    table_index += 1
                elif kind == "page_break":
                    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                else:
                    raise SkeletonError(f"Unsupported skeleton block kind: {kind}")

        document.save(output_path)
        bundle_items.append({
            "id": volume_id,
            "role": role,
            "path": str(output_path),
            "sha256": sha256(output_path),
        })

    slot_manifest = {
        "schema_version": "1.0",
        "project_id": project_id,
        "skill_version": spec.get("skill_version", "v1.37"),
        "slots": slots_out,
    }
    technical_manifest = {
        "schema_version": "1.0",
        "project_id": project_id,
        "requirements": requirements_out,
    }
    bundle = {
        "schema_version": "1.0",
        "project_id": project_id,
        "style_profile": {
            "profile_id": style_profile["profile_id"],
            "path": str(frozen_style_profile.resolve()),
            "sha256": sha256(frozen_style_profile),
        },
        "volumes": bundle_items,
    }
    bundle_path = output_dir / "skeleton_bundle_manifest.json"
    slot_path = output_dir / "slot_manifest.json"
    technical_path = output_dir / "technical_requirement_manifest.json"
    write_json(bundle_path, bundle)
    write_json(slot_path, slot_manifest)
    write_json(technical_path, technical_manifest)
    return {"bundle": bundle_path, "slots": slot_path, "technical": technical_path}


def _slot_tag(element: etree._Element) -> str | None:
    tags = element.xpath("./w:sdtPr/w:tag/@w:val", namespaces=NS)
    return tags[0] if tags else None


def _slot_content_has_value(sdt: etree._Element) -> bool:
    if sdt.find("./" + qn("w:sdtPr") + "/" + qn("w:showingPlcHdr")) is not None:
        return False
    content = sdt.find(qn("w:sdtContent"))
    if content is None:
        return False
    text = "".join(content.xpath(".//w:t/text()", namespaces=NS)).strip()
    return bool(text or content.xpath(".//w:drawing|.//w:pict|.//w:tbl", namespaces=NS))


def _normalise_tree(xml_bytes: bytes) -> bytes:
    parser = etree.XMLParser(remove_blank_text=False, resolve_entities=False)
    root = etree.fromstring(xml_bytes, parser)
    for element in root.xpath("//w:sdt", namespaces=NS):
        tag = _slot_tag(element)
        if tag and tag.startswith(SLOT_PREFIX):
            content = element.find(qn("w:sdtContent"))
            if content is not None:
                for child in list(content):
                    content.remove(child)
                content.append(etree.Element(qn("w:r" if element.getparent().tag == qn("w:p") else "w:p")))
        marker = element.find("./" + qn("w:sdtPr") + "/" + qn("w:showingPlcHdr"))
        if marker is not None:
            marker.getparent().remove(marker)
    for element in root.xpath("//w:lastRenderedPageBreak|//w:proofErr", namespaces=NS):
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)
    for element in root.iter():
        for attr in list(element.attrib):
            if etree.QName(attr).localname.startswith("rsid"):
                del element.attrib[attr]
    for paragraph in root.xpath("//w:p[.//w:instrText]", namespaces=NS):
        for text in paragraph.xpath(".//w:t", namespaces=NS):
            text.text = "{FIELD_RESULT}"
    return etree.tostring(root, method="c14n", exclusive=True, with_comments=False)


def structure_snapshot(path: Path) -> dict[str, Any]:
    try:
        with ZipFile(path) as archive:
            names = set(archive.namelist())
            if "word/document.xml" not in names:
                raise SkeletonError(f"Not a DOCX document: {path}")
            document_xml = archive.read("word/document.xml")
            root = etree.fromstring(document_xml)
            slots = []
            for element in root.xpath("//w:sdt", namespaces=NS):
                tag = _slot_tag(element)
                if tag and tag.startswith(SLOT_PREFIX):
                    slots.append({"slot_id": tag[len(SLOT_PREFIX):], "filled": _slot_content_has_value(element)})
            locked_parts = {}
            for name in sorted(names):
                if name in {"word/styles.xml", "word/numbering.xml", "word/settings.xml", "word/fontTable.xml", "word/theme/theme1.xml"} or re.fullmatch(r"word/(header|footer)\d+\.xml", name):
                    locked_parts[name] = hashlib.sha256(_normalise_tree(archive.read(name))).hexdigest()
            normalised = _normalise_tree(document_xml)
    except BadZipFile as exc:
        raise SkeletonError(f"Invalid DOCX ZIP: {path}") from exc
    doc = Document(path)
    headings = []
    for index, paragraph in enumerate(doc.paragraphs):
        style = paragraph.style.name if paragraph.style else ""
        if style.lower().startswith("heading") or paragraph._p.xpath("./w:pPr/w:outlineLvl"):
            headings.append({"index": index, "text": paragraph.text, "style": style, "page_break_before": bool(paragraph.paragraph_format.page_break_before)})
    tables = [{"index": i, "rows": len(table.rows), "columns": len(table.columns)} for i, table in enumerate(doc.tables)]
    sections = []
    for index, section in enumerate(doc.sections):
        sections.append({
            "index": index,
            "orientation": int(section.orientation) if section.orientation is not None else None,
            "page_width": int(section.page_width or 0),
            "page_height": int(section.page_height or 0),
            "top_margin": int(section.top_margin or 0),
            "bottom_margin": int(section.bottom_margin or 0),
            "left_margin": int(section.left_margin or 0),
            "right_margin": int(section.right_margin or 0),
        })
    return {
        "path": str(path.resolve()),
        "sha256": sha256(path),
        "structure_sha256": hashlib.sha256(normalised).hexdigest(),
        "locked_parts": locked_parts,
        "headings": headings,
        "tables": tables,
        "sections": sections,
        "slots": slots,
    }


def _bundle_volumes(bundle_path: Path) -> list[dict[str, Any]]:
    bundle = read_json(bundle_path)
    volumes = bundle.get("volumes")
    if not isinstance(volumes, list) or not volumes:
        raise SkeletonError("Skeleton bundle needs non-empty volumes")
    resolved = []
    ids = set()
    for item in volumes:
        volume_id = clean_id(item["id"])
        if volume_id in ids:
            raise SkeletonError(f"Duplicate volume id: {volume_id}")
        ids.add(volume_id)
        path = resolve(item["path"], bundle_path.parent)
        if not path.is_file():
            raise SkeletonError(f"Missing skeleton volume: {path}")
        if item.get("sha256") and item["sha256"].lower() != sha256(path):
            raise SkeletonError(f"Skeleton volume hash mismatch: {volume_id}")
        resolved.append({**item, "id": volume_id, "path": path})
    return resolved


def _slots_by_id(slot_manifest_path: Path) -> dict[str, dict[str, Any]]:
    manifest = read_json(slot_manifest_path)
    raw = manifest.get("slots")
    if not isinstance(raw, list) or not raw:
        raise SkeletonError("Slot manifest needs non-empty slots")
    result = {}
    required = {"slot_id", "volume_id", "section_id", "slot_type", "anchor", "allow_blank", "natural_overflow"}
    for item in raw:
        missing = required - set(item)
        if missing:
            raise SkeletonError(f"Slot is missing fields: {sorted(missing)}")
        slot_id = clean_id(item["slot_id"])
        if slot_id in result:
            raise SkeletonError(f"Duplicate slot id: {slot_id}")
        result[slot_id] = item
    return result


def check_technical_blank_table(docx_path: Path, technical_manifest_path: Path, volume_id: str | None = None) -> list[str]:
    errors: list[str] = []
    data = read_json(technical_manifest_path)
    rows = [item for item in data.get("requirements", []) if volume_id is None or item.get("volume_id") == volume_id]
    if not rows:
        return ["technical requirement manifest has no rows for the selected volume"]
    document = Document(docx_path)
    seen = set()
    for expected_index, item in enumerate(rows, 1):
        requirement_id = item.get("requirement_id")
        if requirement_id in seen:
            errors.append(f"duplicate requirement id: {requirement_id}")
            continue
        seen.add(requirement_id)
        table_index = int(item["table_index"])
        row_index = int(item["row_index"])
        if table_index >= len(document.tables) or row_index >= len(document.tables[table_index].rows):
            errors.append(f"missing technical row: {requirement_id}")
            continue
        row = document.tables[table_index].rows[row_index]
        req_col = int(item["requirement_column"])
        if row.cells[req_col].text.strip() != str(item["requirement_text"]).strip():
            errors.append(f"requirement text changed: {requirement_id}")
        for key in ("response_column", "deviation_column"):
            column = int(item[key])
            tc = row.cells[column]._tc
            values = [node for node in tc.xpath(".//w:sdt") if (_slot_tag(node) or "").startswith(SLOT_PREFIX)]
            if len(values) != 1 or _slot_content_has_value(values[0]):
                errors.append(f"technical blank cell is not empty/tagged: {requirement_id}:{key}")
    ordered = sorted(rows, key=lambda x: (int(x["table_index"]), int(x["row_index"])))
    if [x.get("requirement_id") for x in rows] != [x.get("requirement_id") for x in ordered]:
        errors.append("technical requirements are not in physical table order")
    return errors


def _style_font_snapshot(style: Any) -> dict[str, Any]:
    east = style._element.xpath("./w:rPr/w:rFonts/@w:eastAsia")
    colour = style.font.color.rgb
    return {
        "font_east_asia": east[0] if east else None,
        "font_latin": style.font.name,
        "size_pt": round(style.font.size.pt, 2) if style.font.size is not None else None,
        "bold": style.font.bold,
        "color": str(colour) if colour is not None else None,
    }


def _iter_visible_paragraphs(document: Any) -> Iterable[Any]:
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def check_format_standard(docx_path: Path, profile_path: Path) -> list[str]:
    profile = read_json(profile_path)
    document = Document(docx_path)
    errors: list[str] = []
    style_map = {
        "normal": "Normal",
        "heading_1": "Heading 1",
        "heading_2": "Heading 2",
        "heading_3": "Heading 3",
        "heading_4": "Heading 4",
        **{key: value for key, value in STYLE_NAMES.items()},
    }
    for profile_key, style_name in style_map.items():
        expected = profile.get("styles", {}).get(profile_key)
        if not expected:
            continue
        if style_name not in document.styles:
            errors.append(f"missing required style: {style_name}")
            continue
        actual = _style_font_snapshot(document.styles[style_name])
        for field in ("font_east_asia", "font_latin", "size_pt", "bold", "color"):
            if field not in expected:
                continue
            wanted = expected[field]
            got = actual[field]
            if field == "color":
                wanted = str(wanted).lstrip("#").upper()
            if field == "size_pt":
                wanted = round(float(wanted), 2)
            if got != wanted:
                errors.append(f"style mismatch {style_name}.{field}: expected={wanted}, actual={got}")
    palette = profile.get("palette", {})
    if not palette.get("external_accent_allowed", False):
        allowed = {str(palette.get("text", "000000")).lstrip("#").upper(), "000000"}
        for paragraph in _iter_visible_paragraphs(document):
            for run in paragraph.runs:
                if not run.text.strip() or run.font.color.rgb is None:
                    continue
                value = str(run.font.color.rgb).upper()
                if value not in allowed:
                    errors.append(f"unauthorised visible text colour {value}: {run.text[:30]}")
                    if len(errors) >= 30:
                        return errors
    return errors


def check_skeleton_contract(bundle_path: Path, slot_manifest_path: Path, technical_manifest_path: Path) -> list[str]:
    errors: list[str] = []
    slots = _slots_by_id(slot_manifest_path)
    volumes = _bundle_volumes(bundle_path)
    bundle = read_json(bundle_path)
    style_info = bundle.get("style_profile", {})
    style_path = resolve(style_info.get("path", ""), bundle_path.parent) if style_info.get("path") else None
    if style_path is None or not style_path.is_file():
        errors.append("skeleton bundle has no readable style profile")
    elif style_info.get("sha256") != sha256(style_path):
        errors.append("skeleton style profile hash mismatch")
    volume_ids = {item["id"] for item in volumes}
    found: dict[str, str] = {}
    for volume in volumes:
        snap = structure_snapshot(volume["path"])
        if style_path is not None and style_path.is_file():
            errors.extend(f"{volume['id']}: {item}" for item in check_format_standard(volume["path"], style_path))
        for slot in snap["slots"]:
            slot_id = slot["slot_id"]
            if slot_id in found:
                errors.append(f"slot appears more than once: {slot_id}")
            found[slot_id] = volume["id"]
            if slot["filled"]:
                errors.append(f"skeleton slot is not blank: {slot_id}")
    if set(found) != set(slots):
        errors.extend(f"slot missing from skeleton: {slot_id}" for slot_id in sorted(set(slots) - set(found)))
        errors.extend(f"unregistered slot in skeleton: {slot_id}" for slot_id in sorted(set(found) - set(slots)))
    for slot_id, item in slots.items():
        if item["volume_id"] not in volume_ids:
            errors.append(f"slot references unknown volume: {slot_id}")
        elif found.get(slot_id) and found[slot_id] != item["volume_id"]:
            errors.append(f"slot volume mismatch: {slot_id}")
    technical = read_json(technical_manifest_path).get("requirements", [])
    for volume in volumes:
        if any(item.get("volume_id") == volume["id"] for item in technical):
            errors.extend(check_technical_blank_table(volume["path"], technical_manifest_path, volume["id"]))
    return errors


def freeze_bundle(bundle_path: Path, slot_manifest_path: Path, technical_manifest_path: Path, output: Path) -> dict[str, Any]:
    errors = check_skeleton_contract(bundle_path, slot_manifest_path, technical_manifest_path)
    if errors:
        raise SkeletonError("Skeleton validation failed: " + "; ".join(errors))
    volumes = _bundle_volumes(bundle_path)
    snapshots = []
    for volume in volumes:
        snapshot = structure_snapshot(volume["path"])
        snapshots.append({"id": volume["id"], "role": volume.get("role"), **snapshot})
    material = {
        "tool_version": TOOL_VERSION,
        "bundle_sha256": sha256(bundle_path),
        "slot_manifest_sha256": sha256(slot_manifest_path),
        "technical_manifest_sha256": sha256(technical_manifest_path),
        "style_profile": read_json(bundle_path).get("style_profile"),
        "volumes": [{"id": item["id"], "structure_sha256": item["structure_sha256"], "locked_parts": item["locked_parts"]} for item in snapshots],
    }
    result = {
        "schema_version": "1.0",
        "tool_version": TOOL_VERSION,
        "skeleton_id": canonical_hash(material),
        **material,
        "bundle_path": str(bundle_path.resolve()),
        "slot_manifest_path": str(slot_manifest_path.resolve()),
        "technical_manifest_path": str(technical_manifest_path.resolve()),
        "volumes": snapshots,
    }
    write_json(output, result)
    return result


def _lock_volume(lock: dict[str, Any], volume_id: str) -> dict[str, Any]:
    matches = [item for item in lock.get("volumes", []) if item.get("id") == volume_id]
    if len(matches) != 1:
        raise SkeletonError(f"Skeleton lock has no unique volume: {volume_id}")
    return matches[0]


def check_docx_against_lock(docx_path: Path, lock_path: Path, volume_id: str) -> list[str]:
    lock = read_json(lock_path)
    expected = _lock_volume(lock, volume_id)
    actual = structure_snapshot(docx_path)
    errors = []
    if actual["structure_sha256"] != expected["structure_sha256"]:
        errors.append(f"document structure differs from frozen skeleton: {volume_id}")
    if actual["locked_parts"] != expected.get("locked_parts", {}):
        errors.append(f"styles, numbering, settings, header or footer changed: {volume_id}")
    expected_slots = {item["slot_id"] for item in expected.get("slots", [])}
    actual_slots = {item["slot_id"] for item in actual.get("slots", [])}
    if actual_slots != expected_slots:
        errors.append(f"slot set changed: {volume_id}")
    return errors


def check_structure_frozen(candidate_path: Path, lock_path: Path, volume_id: str | None = None) -> list[str]:
    lock = read_json(lock_path)
    if candidate_path.suffix.lower() == ".docx":
        if not volume_id:
            raise SkeletonError("volume_id is required when checking one DOCX")
        return check_docx_against_lock(candidate_path, lock_path, volume_id)
    errors = []
    for volume in _bundle_volumes(candidate_path):
        errors.extend(check_docx_against_lock(volume["path"], lock_path, volume["id"]))
    expected = {item["id"] for item in lock.get("volumes", [])}
    actual = {item["id"] for item in _bundle_volumes(candidate_path)}
    if expected != actual:
        errors.append(f"bundle volume set changed; missing={sorted(expected-actual)}, extra={sorted(actual-expected)}")
    return errors


def _docx_slot_values(path: Path) -> dict[str, bool]:
    with ZipFile(path) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    values = {}
    for element in root.xpath("//w:sdt", namespaces=NS):
        tag = _slot_tag(element)
        if tag and tag.startswith(SLOT_PREFIX):
            values[tag[len(SLOT_PREFIX):]] = _slot_content_has_value(element)
    return values


def check_slot_fill(docx_path: Path, slot_manifest_path: Path, fill_manifest_path: Path, lock_path: Path) -> list[str]:
    errors = check_docx_against_lock(docx_path, lock_path, str(read_json(fill_manifest_path).get("volume_id", "")))
    slots = _slots_by_id(slot_manifest_path)
    fill = read_json(fill_manifest_path)
    lock = read_json(lock_path)
    if fill.get("skeleton_id") != lock.get("skeleton_id"):
        errors.append("fill manifest is not bound to the current skeleton")
    stage = int(fill.get("stage", -1))
    if not 2 <= stage <= 9:
        errors.append("fill manifest stage must be 2..9")
    items = fill.get("items")
    if not isinstance(items, list):
        return errors + ["fill manifest items must be a list"]
    actual = _docx_slot_values(docx_path)
    seen = set()
    valid_status = {"filled", "blank", "unresolved", "not_applicable"}
    for item in items:
        slot_id = item.get("slot_id")
        if slot_id in seen:
            errors.append(f"duplicate fill item: {slot_id}")
            continue
        seen.add(slot_id)
        slot = slots.get(slot_id)
        if not slot:
            errors.append(f"unknown fill slot: {slot_id}")
            continue
        if slot.get("volume_id") != fill.get("volume_id"):
            errors.append(f"fill item belongs to another volume: {slot_id}")
        status = item.get("status")
        if status not in valid_status:
            errors.append(f"invalid fill status: {slot_id}:{status}")
            continue
        allowed = [int(x) for x in slot.get("allowed_stages", [])]
        if status == "filled" and allowed and stage not in allowed:
            errors.append(f"stage {stage} may not fill slot: {slot_id}")
        has_value = actual.get(slot_id, False)
        if status == "filled":
            if not has_value:
                errors.append(f"slot marked filled but is empty: {slot_id}")
            if not item.get("sources") or not item.get("content_sha256"):
                errors.append(f"filled slot lacks sources or content hash: {slot_id}")
        elif has_value:
            errors.append(f"slot has content but status is {status}: {slot_id}")
        if status in {"blank", "unresolved", "not_applicable"} and not str(item.get("reason", "")).strip():
            errors.append(f"non-filled slot lacks a specific reason: {slot_id}")
        if status in {"blank", "unresolved"} and not slot.get("allow_blank") and slot.get("required_by_stage") is not None and stage >= int(slot["required_by_stage"]):
            errors.append(f"required slot remains blank at stage {stage}: {slot_id}")
    for slot_id, has_value in actual.items():
        slot = slots.get(slot_id)
        if slot and slot.get("volume_id") == fill.get("volume_id") and has_value and slot_id not in seen:
            errors.append(f"filled slot is missing from fill manifest: {slot_id}")
    if stage == 9:
        expected = {slot_id for slot_id, slot in slots.items() if slot.get("volume_id") == fill.get("volume_id")}
        missing = expected - seen
        extra = seen - expected
        if missing:
            errors.append(f"final fill manifest does not cover all volume slots: {sorted(missing)}")
        if extra:
            errors.append(f"final fill manifest contains slots from another volume: {sorted(extra)}")
    return errors


def check_requirement_coverage(section_tree_path: Path, coverage_path: Path, slot_manifest_path: Path) -> list[str]:
    section_tree = read_json(section_tree_path)
    coverage = read_json(coverage_path)
    slots = _slots_by_id(slot_manifest_path)
    required_ids = {str(x) for x in section_tree.get("requirement_ids", [])}
    mappings = coverage.get("mappings")
    if not isinstance(mappings, list):
        return ["coverage matrix needs mappings"]
    mapped: dict[str, int] = {}
    errors = []
    for item in mappings:
        requirement_id = str(item.get("requirement_id", ""))
        slot_ids = item.get("slot_ids", [])
        if not requirement_id or not slot_ids:
            errors.append("coverage mapping needs requirement_id and slot_ids")
            continue
        mapped[requirement_id] = mapped.get(requirement_id, 0) + 1
        for slot_id in slot_ids:
            if slot_id not in slots:
                errors.append(f"coverage references unknown slot: {requirement_id}:{slot_id}")
    errors.extend(f"requirement is not mapped to a slot: {value}" for value in sorted(required_ids - set(mapped)))
    errors.extend(f"coverage has unknown requirement: {value}" for value in sorted(set(mapped) - required_ids))
    return errors


def validate_binding(binding_path: Path, lock_path: Path, stage: int, expected_volume: str | None = None) -> list[str]:
    binding = read_json(binding_path)
    lock = read_json(lock_path)
    errors = []
    if binding.get("skeleton_id") != lock.get("skeleton_id"):
        errors.append("binding skeleton_id differs from current skeleton")
    if int(binding.get("stage", -1)) != stage:
        errors.append("binding stage differs from workflow stage")
    if expected_volume and binding.get("volume_id") != expected_volume:
        errors.append(f"binding volume must be {expected_volume}")
    return errors


def fill_text_slots(source: Path, values_path: Path, output: Path) -> dict[str, Any]:
    values = read_json(values_path)
    shutil.copy2(source, output)
    with ZipFile(output, "r") as source_zip:
        files = {name: source_zip.read(name) for name in source_zip.namelist()}
    root = etree.fromstring(files["word/document.xml"])
    updated = []
    for element in root.xpath("//w:sdt", namespaces=NS):
        tag = _slot_tag(element)
        if not tag or not tag.startswith(SLOT_PREFIX):
            continue
        slot_id = tag[len(SLOT_PREFIX):]
        if slot_id not in values:
            continue
        content = element.find(qn("w:sdtContent"))
        props = element.find(qn("w:sdtPr"))
        if props is not None:
            marker = props.find(qn("w:showingPlcHdr"))
            if marker is not None:
                props.remove(marker)
        prototype_ppr = None
        if element.getparent().tag != qn("w:p"):
            first_paragraph = content.find(qn("w:p"))
            if first_paragraph is not None:
                source_ppr = first_paragraph.find(qn("w:pPr"))
                if source_ppr is not None:
                    prototype_ppr = copy.deepcopy(source_ppr)
        for child in list(content):
            content.remove(child)
        lines = str(values[slot_id]).splitlines() or [""]
        if element.getparent().tag == qn("w:p"):
            for index, line in enumerate(lines):
                run = etree.Element(qn("w:r"))
                text = etree.SubElement(run, qn("w:t"))
                text.text = line
                if index < len(lines) - 1:
                    etree.SubElement(run, qn("w:br"))
                content.append(run)
        else:
            for line in lines:
                paragraph = etree.Element(qn("w:p"))
                if prototype_ppr is not None:
                    paragraph.append(copy.deepcopy(prototype_ppr))
                run = etree.SubElement(paragraph, qn("w:r"))
                text = etree.SubElement(run, qn("w:t"))
                text.text = line
                content.append(paragraph)
        updated.append(slot_id)
    files["word/document.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
    from zipfile import ZIP_DEFLATED
    temp = output.with_suffix(output.suffix + ".tmp")
    with ZipFile(temp, "w", ZIP_DEFLATED) as target:
        for name, payload in files.items():
            target.writestr(name, payload)
    temp.replace(output)
    return {"output": str(output.resolve()), "sha256": sha256(output), "updated_slots": sorted(updated)}


def emit(value: Any) -> None:
    print(json.dumps(value, ensure_ascii=False, indent=2))


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    build = sub.add_parser("build", help="Build blank DOCX skeletons from a versioned spec")
    build.add_argument("--spec", required=True, type=Path)
    build.add_argument("--output-dir", required=True, type=Path)
    inspect = sub.add_parser("inspect", help="Inspect one DOCX skeleton or filled copy")
    inspect.add_argument("docx", type=Path)
    freeze = sub.add_parser("freeze", help="Validate and freeze a skeleton bundle")
    freeze.add_argument("--bundle", required=True, type=Path)
    freeze.add_argument("--slots", required=True, type=Path)
    freeze.add_argument("--technical", required=True, type=Path)
    freeze.add_argument("--output", required=True, type=Path)
    validate = sub.add_parser("validate", help="Validate a blank skeleton bundle")
    validate.add_argument("--bundle", required=True, type=Path)
    validate.add_argument("--slots", required=True, type=Path)
    validate.add_argument("--technical", required=True, type=Path)
    validate.add_argument("--output", type=Path)
    compare = sub.add_parser("compare", help="Compare a filled copy or bundle with a skeleton lock")
    compare.add_argument("--candidate", required=True, type=Path)
    compare.add_argument("--lock", required=True, type=Path)
    compare.add_argument("--volume-id")
    fill = sub.add_parser("fill", help="Fill plain-text values into tagged slots without changing structure")
    fill.add_argument("--source", required=True, type=Path)
    fill.add_argument("--values", required=True, type=Path)
    fill.add_argument("--output", required=True, type=Path)
    style_template = sub.add_parser("create-style-template", help="Create a blank DOCX style template from a skeleton format profile")
    style_template.add_argument("--profile", type=Path, default=DEFAULT_STYLE_PROFILE)
    style_template.add_argument("--output", type=Path, default=DEFAULT_STYLE_TEMPLATE)
    args = parser.parse_args()
    if args.command == "build":
        emit({key: str(value.resolve()) for key, value in build_skeleton(args.spec, args.output_dir).items()})
        return 0
    if args.command == "inspect":
        emit(structure_snapshot(args.docx))
        return 0
    if args.command == "freeze":
        emit(freeze_bundle(args.bundle, args.slots, args.technical, args.output))
        return 0
    if args.command == "validate":
        errors = check_skeleton_contract(args.bundle, args.slots, args.technical)
        result = {"status": "pass" if not errors else "blocked", "errors": errors}
        if args.output:
            write_json(args.output, result)
        emit(result)
        return 0 if not errors else 2
    if args.command == "compare":
        errors = check_structure_frozen(args.candidate, args.lock, args.volume_id)
        emit({"status": "pass" if not errors else "blocked", "errors": errors})
        return 0 if not errors else 2
    if args.command == "create-style-template":
        emit(create_style_template(args.profile, args.output))
        return 0
    emit(fill_text_slots(args.source, args.values, args.output))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (SkeletonError, ValueError, KeyError, IndexError, OSError) as exc:
        print(json.dumps({"status": "error", "error": str(exc)}, ensure_ascii=False), file=sys.stderr)
        raise SystemExit(2)
