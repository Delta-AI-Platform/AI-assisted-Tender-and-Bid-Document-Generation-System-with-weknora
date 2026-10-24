"""Read report-locked vehicle parameters from the maintained technical workbook.

Model lookup only enumerates candidate report identities. Facts are returned
only after an exact report_id is selected; no same-model columns are merged.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError as exc:
    raise SystemExit("读取技术参数工作簿需要 openpyxl；请使用工作区提供的 Python 依赖环境") from exc


DEFAULT_WORKBOOK = Path(
    r"C:\Users\DELL\Documents\ChatGPT\技术部分整理\outputs\车型技术参数\国能18个响应表_按检验报告补齐.xlsx"
)


def comparable_model(value: str) -> str:
    text = "".join(value.split())
    return text.removesuffix("（大电池报告）")


def load_reports(path: Path) -> tuple[str, dict[str, dict]]:
    if not path.is_file():
        raise ValueError(f"技术参数工作簿不存在: {path}")
    workbook_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    workbook = load_workbook(path, read_only=True, data_only=True)
    required = {"18个技术响应表", "检验报告来源"}
    if not required.issubset(workbook.sheetnames):
        raise ValueError(f"工作簿缺少必要工作表: {sorted(required - set(workbook.sheetnames))}")
    table = workbook["18个技术响应表"]
    sources = workbook["检验报告来源"]
    reports: dict[str, dict] = {}
    seen_cells = set()
    for row_number, row in enumerate(sources.iter_rows(min_row=2, values_only=True), 2):
        values = list(row) + [None] * max(0, 14 - len(row))
        model, number, title, target = values[:4]
        report_id, fingerprint = values[12:14]
        if not report_id or not fingerprint:
            raise ValueError(f"来源表第{row_number}行缺少报告唯一ID或内容指纹")
        if target in seen_cells:
            raise ValueError(f"来源表重复目标单元格: {target}")
        seen_cells.add(target)
        actual = table[str(target)].value
        if actual != values[7]:
            raise ValueError(f"来源值与响应表不一致: {target}")
        identity = (str(model), str(number), str(title), str(fingerprint))
        item = reports.setdefault(str(report_id), {
            "model": identity[0], "report_number": identity[1],
            "report_title": identity[2], "report_fingerprint": identity[3],
            "fields": [],
        })
        if tuple(item[key] for key in ("model", "report_number", "report_title", "report_fingerprint")) != identity:
            raise ValueError(f"同一报告ID出现冲突身份: {report_id}")
        item["fields"].append({
            "cell": target,
            "subsystem": values[4], "component": values[5], "module": values[6],
            "value": values[7], "printed_page": values[8],
            "image_sequence": values[9], "image_sha256": values[10],
            "source_path": values[11],
        })
    if not reports:
        raise ValueError("没有可追溯的检验报告字段")
    return workbook_hash, reports


def main() -> None:
    parser = argparse.ArgumentParser(description="按报告唯一ID查询技术参数和原件页位；车型查询仅列候选报告")
    parser.add_argument("--workbook", type=Path, default=Path(os.environ.get("BID_TECH_REPORT_WORKBOOK", DEFAULT_WORKBOOK)))
    choice = parser.add_mutually_exclusive_group(required=True)
    choice.add_argument("--model", help="仅枚举相同车型下的独立报告，不输出参数")
    choice.add_argument("--report-id", help="选定原件后用唯一报告ID查询参数")
    parser.add_argument("--field", help="仅显示模块名包含该字词的字段；须与 --report-id 配用")
    args = parser.parse_args()
    if args.field and not args.report_id:
        parser.error("--field 须与 --report-id 一起使用")
    try:
        workbook_hash, reports = load_reports(args.workbook)
        if args.model:
            requested = comparable_model(args.model)
            candidates = [
                {"report_id": report_id, "model": item["model"], "report_number": item["report_number"],
                 "report_fingerprint": item["report_fingerprint"], "field_count": len(item["fields"])}
                for report_id, item in reports.items()
                if comparable_model(item["model"]) == requested
            ]
            result = {"mode": "candidate_reports_only", "requested_model": args.model, "reports": candidates}
        else:
            item = reports.get(args.report_id)
            if item is None:
                raise ValueError(f"未找到报告唯一ID: {args.report_id}")
            fields = item["fields"]
            if args.field:
                fields = [field for field in fields if args.field in str(field["module"])]
            result = {"mode": "report_locked_fields", "report_id": args.report_id,
                      **{key: item[key] for key in ("model", "report_number", "report_title", "report_fingerprint")},
                      "fields": fields}
        print(json.dumps({"workbook": str(args.workbook.resolve()), "workbook_sha256": workbook_hash,
                          **result}, ensure_ascii=False, indent=2))
    except ValueError as exc:
        parser.exit(2, f"错误：{exc}\n")


if __name__ == "__main__":
    main()
