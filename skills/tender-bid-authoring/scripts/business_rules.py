"""Deterministic checks for the confirmed business assembly rules."""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

FORM_TYPES = {"performance", "commercial_deviation", "technical_deviation"}


def read_json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"Expected JSON object: {path}")
    return value


def compact(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or ""))


def _document_text(doc: Any) -> str:
    parts = [p.text for p in _all_paragraphs(doc)]
    parts.extend(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    return compact("".join(parts))


def _all_paragraphs(doc: Any) -> list[Any]:
    """Return body paragraphs including those nested in Word content controls.

    ``python-docx`` deliberately omits paragraphs inside ``w:sdt`` elements
    from ``Document.paragraphs``.  Skeleton-first documents store every
    mutable slot in an SDT, so using only ``Document.paragraphs`` makes visible
    commitments and certificate captions look absent to the Rulebase.
    """
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    return [Paragraph(element, doc) for element in doc.element.body.iter(qn("w:p"))]


def _cell_set(raw: Any, rows: int, columns: int, label: str) -> set[tuple[int, int]]:
    if not isinstance(raw, list):
        raise ValueError(f"{label} must be a list")
    result: set[tuple[int, int]] = set()
    for coordinate in raw:
        if not isinstance(coordinate, list) or len(coordinate) != 2:
            raise ValueError(f"{label} needs [row, column] coordinates")
        row, column = coordinate
        if type(row) is not int or type(column) is not int or not 0 <= row < rows or not 0 <= column < columns:
            raise ValueError(f"{label} has an out-of-range coordinate: {coordinate}")
        if (row, column) in result:
            raise ValueError(f"{label} has a duplicate coordinate: {coordinate}")
        result.add((row, column))
    return result


def check_blank_forms(candidate_path: Path, manifest_path: Path) -> list[str]:
    """Require every original form cell to be either fixed and unchanged or blank."""
    from docx import Document

    manifest = read_json(manifest_path)
    forms = manifest.get("forms")
    if not isinstance(forms, list) or {form.get("type") for form in forms if isinstance(form, dict)} != FORM_TYPES or len(forms) != 3:
        return ["Blank-form manifest must contain exactly the three confirmed form types"]
    candidate = Document(candidate_path)
    candidate_text = _document_text(candidate)
    errors: list[str] = []
    for form in forms:
        kind = form["type"]
        try:
            source_file = Path(form["source_file"])
            source_file = source_file if source_file.is_absolute() else (manifest_path.parent / source_file).resolve()
            source = Document(source_file)
            source_table = source.tables[int(form["source_table"])]
            target_table = candidate.tables[int(form["target_table"])]
            rows, columns = len(source_table.rows), len(source_table.columns)
            if (rows, columns) != (len(target_table.rows), len(target_table.columns)):
                raise ValueError("table dimensions differ from the original")
            title = compact(form.get("title"))
            if not title or title not in _document_text(source) or title not in candidate_text:
                raise ValueError("original form title is absent")
            fixed = _cell_set(form.get("fixed_cells"), rows, columns, "fixed_cells")
            response = _cell_set(form.get("response_cells"), rows, columns, "response_cells")
            all_cells = {(r, c) for r in range(rows) for c in range(columns)}
            if fixed & response or fixed | response != all_cells or not response:
                raise ValueError("fixed and response cells must partition the whole form")
            for row, column in fixed:
                if compact(source_table.cell(row, column).text) != compact(target_table.cell(row, column).text):
                    errors.append(f"{kind} fixed cell {row},{column} changed")
            for row, column in response:
                if compact(target_table.cell(row, column).text):
                    errors.append(f"{kind} response cell {row},{column} is not blank")
        except (KeyError, IndexError, TypeError, ValueError, OSError) as exc:
            errors.append(f"{kind}: {exc}")
    return errors


def check_blank_policy(path: Path) -> list[str]:
    data = read_json(path)
    if data.get("decision") != "leave_blank" or set(data.get("forms", [])) != FORM_TYPES:
        return ["Blank policy must name all three forms and the leave_blank decision"]
    if not str(data.get("authority", "")).strip():
        return ["Blank policy needs the current business decision as authority"]
    return []


def check_qualification(path: Path) -> list[str]:
    data = read_json(path)
    errors: list[str] = []
    if data.get("performance_submitted") is not False:
        errors.append("Performance must be recorded as not submitted")
    if type(data.get("performance_required")) is not bool:
        errors.append("performance_required must be a boolean")
    if data.get("overall_status") not in {"met", "not_met", "unresolved"}:
        errors.append("qualification overall_status must be met, not_met or unresolved")
    if data.get("performance_required") is True and data.get("overall_status") == "met":
        errors.append("Qualification cannot be met when required performance is not submitted")
    return errors


def check_scoring_and_qualification(qualification_path: Path, scoring_path: Path, certification_path: Path) -> list[str]:
    errors = check_qualification(qualification_path)
    scoring = read_json(scoring_path)
    max_points = scoring.get("performance_points_max")
    if type(max_points) not in {int, float} or max_points < 0:
        errors.append("performance_points_max must be nonnegative")
    if scoring.get("performance_status") != "not_submitted" or scoring.get("claimed_performance_points") != 0:
        errors.append("Performance scoring must show not_submitted and zero claimed points")
    certs = read_json(certification_path)
    items = certs.get("items")
    if not isinstance(items, list):
        errors.append("Certification placement needs an items list")
    else:
        ids: set[str] = set()
        for item in items:
            ident = str(item.get("id", ""))
            if not ident or ident in ids:
                errors.append("Certification IDs must be present and unique")
            ids.add(ident)
            scored = item.get("scoring_related")
            section = item.get("section")
            if type(scored) is not bool:
                errors.append(f"Certification {ident} scoring_related must be boolean")
            elif scored and (not item.get("scoring_section") or section != item["scoring_section"]):
                errors.append(f"Certification {ident} must be in its scoring section")
            elif not scored and section != "其他商务资料":
                errors.append(f"Certification {ident} must be in 其他商务资料")
            source = item.get("source_file")
            if not source:
                errors.append(f"Certification {ident} needs a source file")
            else:
                source_path = Path(source)
                source_path = source_path if source_path.is_absolute() else (certification_path.parent / source_path).resolve()
                if not source_path.is_file():
                    errors.append(f"Certification {ident} source file is missing")
    return errors


def check_certification_placement(docx_path: Path, certification_path: Path) -> list[str]:
    """Check the visible certificate title sits inside its declared section."""
    from docx import Document

    items = read_json(certification_path).get("items", [])
    doc = Document(docx_path)
    paragraphs = [(compact(p.text), p.style.name) for p in _all_paragraphs(doc) if compact(p.text)]
    errors: list[str] = []
    for item in items:
        ident = str(item.get("id", ""))
        section, caption = compact(item.get("section")), compact(item.get("caption"))
        if not section or not caption:
            errors.append(f"Certification {ident} needs section and visible caption")
            continue
        starts = [i for i, (text, _) in enumerate(paragraphs) if text == section]
        if len(starts) != 1:
            errors.append(f"Certification {ident} section is missing or ambiguous")
            continue
        start = starts[0]
        style = paragraphs[start][1]
        level = int(style.split(" ")[1]) if style.startswith("Heading ") and style.split(" ")[1].isdigit() else 1
        end = len(paragraphs)
        for i in range(start + 1, len(paragraphs)):
            later_style = paragraphs[i][1]
            if later_style.startswith("Heading ") and later_style.split(" ")[1].isdigit() and int(later_style.split(" ")[1]) <= level:
                end = i
                break
        if not any(caption == text for text, _ in paragraphs[start + 1:end]):
            errors.append(f"Certification {ident} caption is not under {item.get('section')}")
    return errors


def check_commitment_coverage(
    prohibition_path: Path, responsiveness_path: Path, inventory_path: Path, candidate_path: Path
) -> list[str]:
    from docx import Document

    prohibitions = read_json(prohibition_path)
    responsiveness = read_json(responsiveness_path)
    inventory = read_json(inventory_path)
    document = Document(candidate_path)
    paragraphs = [compact(paragraph.text) for paragraph in _all_paragraphs(document)]
    errors: list[str] = []
    letters = inventory.get("letters")
    if not isinstance(letters, list):
        return ["Commitment inventory needs letters"]
    letter_map = {str(item.get("id")): item for item in letters if item.get("id")}
    if len(letter_map) != len(letters):
        errors.append("Commitment letter IDs must be unique and nonempty")
    response_letter = str(inventory.get("responsiveness_letter_id", ""))
    if response_letter not in letter_map:
        errors.append("A dedicated responsiveness commitment letter is required")
    starts: dict[str, int] = {}
    for letter_id, letter in letter_map.items():
        title = compact(letter.get("title"))
        matches = [index for index, value in enumerate(paragraphs) if value == title]
        if not matches:
            errors.append(f"Commitment letter {letter_id} title is absent from the DOCX")
        else:
            starts[letter_id] = matches[-1]
    letter_bodies: dict[str, str] = {}
    for letter_id, start in starts.items():
        end = min((value for other, value in starts.items() if other != letter_id and value > start), default=len(paragraphs))
        letter_bodies[letter_id] = "".join(paragraphs[start + 1:end])
    if not isinstance(prohibitions.get("items"), list):
        errors.append("Prohibition ledger needs items")
    else:
        seen: set[str] = set()
        dedicated: dict[str, list[str]] = {}
        for item in prohibitions["items"]:
            ident = str(item.get("id", ""))
            if not ident or ident in seen:
                errors.append("Prohibition IDs must be present and unique")
            seen.add(ident)
            if not item.get("source") or not item.get("clause"):
                errors.append(f"Prohibition {ident} needs source and clause")
            if type(item.get("applicable")) is not bool:
                errors.append(f"Prohibition {ident} needs applicability")
            if item.get("applicable") is True:
                letter_id = str(item.get("letter_id", ""))
                response_text = compact(item.get("response_text"))
                if letter_id not in letter_map or not response_text or response_text not in letter_bodies.get(letter_id, ""):
                    errors.append(f"Prohibition {ident} has no matching commitment text")
                if item.get("requires_separate_letter"):
                    dedicated.setdefault(letter_id, []).append(ident)
        for letter_id, ids in dedicated.items():
            uses = [str(item.get("id")) for item in prohibitions["items"] if item.get("applicable") and str(item.get("letter_id")) == letter_id]
            if len(uses) != 1 or uses != ids:
                errors.append(f"Independent prohibition letter {letter_id} must be dedicated")
    if not isinstance(responsiveness.get("items"), list):
        errors.append("Responsiveness ledger needs items")
    else:
        seen_response: set[str] = set()
        for item in responsiveness["items"]:
            ident = str(item.get("id", ""))
            if not ident or ident in seen_response:
                errors.append("Responsiveness IDs must be present and unique")
            seen_response.add(ident)
            if not item.get("source") or not item.get("clause"):
                errors.append(f"Responsiveness item {ident} needs source and clause")
            if type(item.get("requires_response")) is not bool:
                errors.append(f"Responsiveness item {ident} needs requires_response")
            if item.get("requires_response") is True:
                response_text = compact(item.get("response_text"))
                if item.get("letter_id") != response_letter or not response_text or response_text not in letter_bodies.get(response_letter, ""):
                    errors.append(f"Responsiveness item {ident} is not covered by the dedicated letter")
    return errors
