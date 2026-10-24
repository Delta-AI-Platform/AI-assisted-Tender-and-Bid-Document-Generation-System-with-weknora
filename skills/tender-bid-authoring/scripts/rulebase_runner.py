#!/usr/bin/env python3
"""Deterministic bid-production checks with a versioned, hash-based run record."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile
from business_rules import (
    FORM_TYPES,
    check_blank_forms,
    check_certification_placement,
    check_commitment_coverage,
    check_scoring_and_qualification,
)
from skeleton_tools import (
    check_format_standard,
    check_requirement_coverage,
    check_skeleton_contract,
    check_slot_fill,
    check_structure_frozen,
    check_technical_blank_table,
)

RULEBASE_VERSION = "1.4.0"
SCRIPT_SHA256 = hashlib.sha256(Path(__file__).read_bytes()).hexdigest()


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def docx_snapshot(path: Path) -> dict[str, Any]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required for DOCX checks") from exc
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    tables = []
    for i, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])
        tables.append({
            "index": i,
            "rows": len(table.rows),
            "columns": len(table.columns),
            "text_sha256": hashlib.sha256(
                json.dumps(rows, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
            ).hexdigest(),
            "column_sha256": [
                hashlib.sha256(
                    json.dumps([row[c] for row in rows], ensure_ascii=False, separators=(",", ":")).encode("utf-8")
                ).hexdigest()
                for c in range(len(table.columns))
            ],
        })
    sections = []
    for section in doc.sections:
        sections.append({
            "page_width_emu": section.page_width,
            "page_height_emu": section.page_height,
            "top_margin_emu": section.top_margin,
            "bottom_margin_emu": section.bottom_margin,
            "left_margin_emu": section.left_margin,
            "right_margin_emu": section.right_margin,
        })
    return {
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "image_count": len(doc.inline_shapes),
        "text_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
        "tables": tables,
        "sections": sections,
        "text": text,
    }


def file_record(item: dict[str, Any], base: Path) -> dict[str, Any]:
    path = Path(item["path"])
    if not path.is_absolute():
        path = (base / path).resolve()
    record = {"id": item.get("id"), "path": str(path), "exists": path.is_file()}
    if record["exists"]:
        record["size"] = path.stat().st_size
        record["sha256"] = sha256(path)
        if path.suffix.lower() == ".docx":
            try:
                with ZipFile(path) as zf:
                    record["docx_zip_valid"] = "word/document.xml" in zf.namelist()
            except BadZipFile:
                record["docx_zip_valid"] = False
    return record


def resolve_file(file_id: str, files: dict[str, dict[str, Any]]) -> Path:
    if file_id not in files:
        raise ValueError(f"Unknown file id: {file_id}")
    return Path(files[file_id]["path"])


def check_fill_manifest(path: Path, manifest: dict[str, Any]) -> list[dict[str, Any]]:
    """Compare final cells with the independently prepared field inventory.

    Table labels must select exactly one table. Never use the last matching
    table or silently ignore a missing field after document assembly.
    """
    from docx import Document
    doc = Document(path)
    norm = lambda value: re.sub(r"\s+", "", str(value))
    outcomes = []
    forms = manifest.get("forms", [])
    if not forms:
        raise ValueError("Empty form inventory")
    ids = set()
    for form in forms:
        labels = form.get("table_labels", [])
        if not labels or not form.get("fields"):
            raise ValueError("Each form needs labels and inventoried fields")
        matches = [t for t in doc.tables if all(
            any(norm(label) == norm(c.text) for r in t.rows for c in r.cells)
            for label in labels)]
        for field in form["fields"]:
            key = form["id"] + ":" + field["id"]
            if key in ids:
                raise ValueError("Duplicate field id: " + key)
            ids.add(key)
            item = {"id": key, "passed": False}
            try:
                if len(matches) != 1:
                    raise ValueError(f"Expected one table, found {len(matches)}")
                table = matches[0]
                row = field.get("row")
                if "row_label" in field:
                    rows = [i for i, r in enumerate(table.rows) if any(
                        norm(c.text) == norm(field["row_label"]) for c in r.cells)]
                    if len(rows) != 1:
                        raise ValueError("Row label missing or ambiguous")
                    row = rows[0]
                col = int(field["column"])
                actual = table.cell(int(row), col).text
                item.update(row=int(row), column=col, actual=actual)
                status = field.get("status")
                if status in {"allowed_blank", "not_applicable", "business_directed_blank"}:
                    if not field.get("reason") or not field.get("source"):
                        raise ValueError("Exemption requires reason and authority")
                elif status != "required":
                    raise ValueError("Unresolved field blocks completion")
                if "expected" not in field or not field.get("source"):
                    raise ValueError("Expected value and source are required")
                expected = str(field["expected"])
                if status == "business_directed_blank" and norm(expected):
                    raise ValueError("Business-directed blank field must have an empty expected value")
                if status == "business_directed_blank" and form.get("type") not in FORM_TYPES:
                    raise ValueError("Business-directed blank is limited to the three confirmed form types")
                if status == "required" and (not norm(expected) or re.fullmatch(
                    r"[/／_＿—\-\.…\s]*", expected) or any(x in expected for x in
                    ["待填", "待确认", "待补充", "详见附件"])):
                    raise ValueError("Placeholder cannot satisfy a required field")
                item["passed"] = norm(actual) == norm(expected)
                if not item["passed"]:
                    item["error"] = "Value differs from approved field inventory"
            except Exception as exc:
                item["error"] = str(exc)
            outcomes.append(item)
    return outcomes


def check_rule(rule: dict[str, Any], files: dict[str, dict[str, Any]]) -> dict[str, Any]:
    kind = rule.get("type", rule.get("kind"))
    if not kind:
        raise ValueError("Rule must define a type")
    result: dict[str, Any] = {"id": rule.get("id"), "type": kind, "passed": False}
    try:
        if kind == "docx_skeleton_contract":
            errors = check_skeleton_contract(
                resolve_file(rule["bundle_file"], files),
                resolve_file(rule["slot_manifest_file"], files),
                resolve_file(rule["technical_manifest_file"], files),
            )
            result.update(errors=errors)
            result["passed"] = not errors
        elif kind == "docx_skeleton_format_standard":
            errors = check_format_standard(
                resolve_file(rule["file"], files),
                resolve_file(rule["profile_file"], files),
            )
            result.update(errors=errors)
            result["passed"] = not errors
        elif kind == "docx_technical_blank_table":
            errors = check_technical_blank_table(
                resolve_file(rule["file"], files),
                resolve_file(rule["manifest_file"], files),
                rule.get("volume_id"),
            )
            result.update(errors=errors)
            result["passed"] = not errors
        elif kind == "docx_slot_fill":
            errors = check_slot_fill(
                resolve_file(rule["file"], files),
                resolve_file(rule["slot_manifest_file"], files),
                resolve_file(rule["fill_manifest_file"], files),
                resolve_file(rule["lock_file"], files),
            )
            result.update(errors=errors)
            result["passed"] = not errors
        elif kind == "docx_structure_frozen":
            errors = check_structure_frozen(
                resolve_file(rule["candidate_file"], files),
                resolve_file(rule["lock_file"], files),
                rule.get("volume_id"),
            )
            result.update(errors=errors)
            result["passed"] = not errors
        elif kind == "docx_requirement_coverage":
            errors = check_requirement_coverage(
                resolve_file(rule["section_tree_file"], files),
                resolve_file(rule["coverage_file"], files),
                resolve_file(rule["slot_manifest_file"], files),
            )
            result.update(errors=errors)
            result["passed"] = not errors
        elif kind == "docx_fill_manifest":
            manifest = read_json(resolve_file(rule["manifest_file"], files))
            checks = check_fill_manifest(resolve_file(rule["file"], files), manifest)
            result.update(fields=checks, checked_count=len(checks))
            result["passed"] = bool(checks) and all(x["passed"] for x in checks)
        elif kind == "docx_intentional_blank_forms":
            errors = check_blank_forms(
                resolve_file(rule["file"], files),
                resolve_file(rule["manifest_file"], files),
            )
            result.update(errors=errors)
            result["passed"] = not errors
        elif kind == "business_outcomes":
            errors = check_scoring_and_qualification(
                resolve_file(rule["qualification_file"], files),
                resolve_file(rule["scoring_file"], files),
                resolve_file(rule["certification_file"], files),
            )
            result.update(errors=errors)
            result["passed"] = not errors
        elif kind == "docx_certification_placement":
            errors = check_certification_placement(
                resolve_file(rule["file"], files),
                resolve_file(rule["certification_file"], files),
            )
            result.update(errors=errors)
            result["passed"] = not errors
        elif kind == "commitment_coverage":
            errors = check_commitment_coverage(
                resolve_file(rule["prohibition_file"], files),
                resolve_file(rule["responsiveness_file"], files),
                resolve_file(rule["inventory_file"], files),
                resolve_file(rule["file"], files),
            )
            result.update(errors=errors)
            result["passed"] = not errors
        elif kind == "file_exists":
            rec = files[rule["file"]]
            result["actual"] = rec["exists"]
            result["passed"] = bool(rec["exists"])
        elif kind == "sha256_equals":
            rec = files[rule["file"]]
            expected = rule["expected"].lower()
            result.update(expected=expected, actual=rec.get("sha256"))
            result["passed"] = rec.get("sha256") == expected
        elif kind == "docx_zip_valid":
            rec = files[rule["file"]]
            result["actual"] = rec.get("docx_zip_valid")
            result["passed"] = rec.get("docx_zip_valid") is True
        elif kind in {"docx_required_terms", "docx_forbidden_terms", "docx_table_shape", "docx_locked_columns", "docx_sections_equal"}:
            path = resolve_file(rule["file"], files)
            snap = docx_snapshot(path)
            if kind == "docx_required_terms":
                missing = [x for x in rule["terms"] if x not in snap["text"]]
                result.update(missing=missing)
                result["passed"] = not missing
            elif kind == "docx_forbidden_terms":
                found = [x for x in rule["terms"] if x in snap["text"]]
                result.update(found=found)
                result["passed"] = not found
            elif kind == "docx_table_shape":
                idx = int(rule["table"])
                actual = snap["tables"][idx]
                expected = {"rows": int(rule["rows"]), "columns": int(rule["columns"])}
                result.update(expected=expected, actual={k: actual[k] for k in expected})
                result["passed"] = all(actual[k] == v for k, v in expected.items())
            elif kind == "docx_locked_columns":
                before = docx_snapshot(resolve_file(rule["source_file"], files))["tables"][int(rule["source_table"])]
                after = snap["tables"][int(rule["table"])]
                columns = [int(c) for c in rule["columns"]]
                mismatches = [c for c in columns if before["column_sha256"][c] != after["column_sha256"][c]]
                result.update(columns=columns, mismatches=mismatches)
                result["passed"] = not mismatches
            else:
                other = docx_snapshot(resolve_file(rule["source_file"], files))
                result.update(expected=other["sections"], actual=snap["sections"])
                result["passed"] = snap["sections"] == other["sections"]
        elif kind == "json_set_equal":
            left = set(read_json(resolve_file(rule["left_file"], files))[rule["left_key"]])
            right = set(read_json(resolve_file(rule["right_file"], files))[rule["right_key"]])
            result.update(left_only=sorted(left-right), right_only=sorted(right-left))
            result["passed"] = left == right
        elif kind == "attachments_exist":
            manifest_path = resolve_file(rule["file"], files)
            data = read_json(manifest_path)
            missing = []
            for item in data[rule.get("key", "attachments")]:
                p = Path(item["path"])
                if not p.is_absolute():
                    p = (manifest_path.parent / p).resolve()
                if not p.is_file():
                    missing.append(str(p))
            result.update(missing=missing)
            result["passed"] = not missing
        else:
            raise ValueError(f"Unsupported rule type: {kind}")
    except Exception as exc:
        result["error"] = f"{type(exc).__name__}: {exc}"
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("spec", type=Path, help="UTF-8 JSON rulebase task specification")
    parser.add_argument("--out", required=True, type=Path, help="JSON run record")
    args = parser.parse_args()
    spec = read_json(args.spec)
    base = args.spec.resolve().parent
    all_items = spec.get("inputs", []) + spec.get("outputs", [])
    records = [file_record(item, base) for item in all_items]
    files = {str(r["id"]): r for r in records if r.get("id") is not None}
    results = [check_rule(rule, files) for rule in spec.get("rules", [])]
    required_failures = [r for r, src in zip(results, spec.get("rules", [])) if src.get("severity", "error") == "error" and not r["passed"]]
    run = {
        "schema_version": "1.0",
        "rulebase_version": RULEBASE_VERSION,
        "script_sha256": SCRIPT_SHA256,
        "run_at_utc": datetime.now(timezone.utc).isoformat(),
        "project_id": spec.get("project_id"),
        "stage": spec.get("stage"),
        "spec_path": str(args.spec.resolve()),
        "spec_sha256": sha256(args.spec),
        "files": records,
        "results": results,
        "status": "pass" if not required_failures else "blocked",
        "required_failure_count": len(required_failures),
    }
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(run, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"status": run["status"], "result": str(args.out)}, ensure_ascii=False))
    return 0 if run["status"] == "pass" else 2


if __name__ == "__main__":
    raise SystemExit(main())
