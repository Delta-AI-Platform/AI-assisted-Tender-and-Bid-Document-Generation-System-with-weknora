"""Small local helpers. No model calls, cloud access or inferred response facts."""
import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def emit(value):
    print(json.dumps(value, ensure_ascii=False, indent=2))


def snapshot(path, table=None):
    doc = Document(path)
    with ZipFile(path) as archive:
        bad = archive.testzip()
        if bad:
            raise ValueError('Corrupt ZIP member: ' + bad)
        media = {n: hashlib.sha256(archive.read(n)).hexdigest()
                 for n in archive.namelist() if n.startswith('word/media/')}
    result = {
        'sha256': sha(path),
        'tables': [{'index': i, 'rows': len(t.rows), 'columns': len(t.columns),
                    'header': [c.text for c in t.rows[0].cells] if t.rows else []}
                   for i, t in enumerate(doc.tables)],
        'media': media,
        'body_text': '\n'.join(doc.element.xpath('.//w:t/text()')),
        'sections': [s._sectPr.xml for s in doc.sections],
        'image_occurrences': [hashlib.sha256(doc.part.related_parts[r].blob).hexdigest()
                              for r in doc.element.xpath('.//a:blip/@r:embed')],
    }
    if table is not None:
        result['rows'] = [[c.text for c in row.cells] for row in doc.tables[table].rows]
    return result


def compare(before, after, table, allowed_columns):
    a, b = snapshot(before, table), snapshot(after, table)
    errors = []
    for key in ('media', 'sections', 'image_occurrences', 'tables'):
        if a[key] != b[key]:
            errors.append(key + ' changed')
    if allowed_columns and table is None:
        raise ValueError('Allowed text columns require --table')
    if not allowed_columns and a['body_text'] != b['body_text']:
        errors.append('Document text changed during style-only revision')
    if table is not None:
        if len(a['rows']) != len(b['rows']):
            errors.append('Row count changed')
        for i, (old, new) in enumerate(zip(a['rows'], b['rows'])):
            if len(old) != len(new):
                errors.append(f'Row {i}: column count changed')
            for j, (left, right) in enumerate(zip(old, new)):
                if j not in allowed_columns and left != right:
                    errors.append(f'Row {i}, column {j}: protected text changed')
    return {'passed': not errors, 'errors': errors,
            'before_sha256': a['sha256'], 'after_sha256': b['sha256'],
            'scope': 'Selected columns/text, section settings and images; not factual or visual approval'}


def apply_style(source, spec_path, output):
    source, output = Path(source), Path(output)
    if output.exists() or source.resolve() == output.resolve():
        raise ValueError('Output must be a new file, never overwrite an existing artifact')
    spec = json.loads(Path(spec_path).read_text(encoding='utf-8'))
    if spec['source_sha256'] != sha(source):
        raise ValueError('Source SHA-256 mismatch; rebind the current source')
    doc = Document(source)
    targets = set()
    for item in spec['paragraphs']:
        indices = [item[k] for k in ('table', 'row', 'column', 'paragraph')]
        if any(type(v) is not int or v < 0 for v in indices):
            raise ValueError('All physical indices must be nonnegative integers')
        t, r, c, p = indices
        paragraph = doc.tables[t].rows[r].cells[c].paragraphs[p]
        if paragraph._p in targets:
            raise ValueError('Duplicate physical paragraph, possibly aliased merged cells')
        targets.add(paragraph._p)
        if paragraph.text != item['expected_text']:
            raise ValueError('Text binding mismatch; do not apply stale offsets')
        # Reject fields, hyperlinks, drawings, tabs, breaks or anchors rather than flatten them.
        if any(ch.tag not in {qn('w:pPr'), qn('w:r')} for ch in paragraph._p):
            raise ValueError('Complex paragraph requires manual OOXML handling')
        for run in paragraph.runs:
            if any(ch.tag not in {qn('w:rPr'), qn('w:t')} for ch in run._r):
                raise ValueError('Complex run requires manual OOXML handling')
        spans = sorted(item['spans'], key=lambda x: x['start'])
        end = 0
        for span in spans:
            start, stop = span['start'], span['end']
            if (type(start) is not int or type(stop) is not int or
                    not end <= start < stop <= len(paragraph.text)):
                raise ValueError('Invalid or overlapping spans; resolve layered priorities first')
            if not re.fullmatch('[0-9A-Fa-f]{6}', span['color']) or type(span['bold']) is not bool:
                raise ValueError('Expected hex RGB and boolean bold')
            end = stop
        offset = 0
        for run in list(paragraph.runs):
            text = run.text
            finish = offset + len(text)
            cuts = {offset, finish}
            for span in spans:
                cuts.update(v for v in (span['start'], span['end']) if offset < v < finish)
            cursor = run._r
            for start, stop in zip(sorted(cuts), sorted(cuts)[1:]):
                clone = deepcopy(run._r)
                for child in list(clone):
                    if child.tag != qn('w:rPr'):
                        clone.remove(child)
                from docx.text.run import Run
                piece = Run(clone, paragraph)
                piece.text = text[start-offset:stop-offset]
                for span in spans:
                    if span['start'] <= start and stop <= span['end']:
                        piece.font.color.rgb = RGBColor.from_string(span['color'])
                        piece.bold = span['bold']
                cursor.addnext(clone)
                cursor = clone
            paragraph._p.remove(run._r)
            offset = finish
    doc.save(output)
    result = compare(source, output, None, [])
    if not result['passed']:
        raise ValueError('Diagnostic candidate only: ' + json.dumps(result['errors']))
    return {'output': str(output.resolve()), 'sha256': sha(output),
            'paragraphs_styled': len(targets), 'preservation': result}


def pdf_pages(path):
    from pypdf import PdfReader
    reader = PdfReader(path)
    entries = []

    def walk(items, level=0):
        for item in items:
            if isinstance(item, list):
                walk(item, level+1)
            else:
                page = reader.get_destination_page_number(item)
                entries.append({'title': str(item.title), 'level': level,
                                'physical_page': page+1 if page is not None and page >= 0 else None})
    walk(reader.outline)
    return {'sha256': sha(path), 'pages': len(reader.pages), 'outline': entries,
            'warning': 'Physical PDF pages, not printed numbering; no outline does not mean no chapters'}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest='command', required=True)
    inspect = sub.add_parser('inspect')
    inspect.add_argument('docx')
    inspect.add_argument('--table', type=int)
    compare_parser = sub.add_parser('compare')
    compare_parser.add_argument('before')
    compare_parser.add_argument('after')
    compare_parser.add_argument('--table', type=int)
    compare_parser.add_argument('--allow-text-columns', type=int, nargs='*', default=[])
    style = sub.add_parser('style')
    style.add_argument('source')
    style.add_argument('bindings')
    style.add_argument('output')
    pages = sub.add_parser('pages')
    pages.add_argument('pdf')
    args = parser.parse_args()
    if args.command == 'inspect':
        emit(snapshot(args.docx, args.table))
    elif args.command == 'compare':
        result = compare(args.before, args.after, args.table, args.allow_text_columns)
        emit(result)
        return 0 if result['passed'] else 1
    elif args.command == 'style':
        emit(apply_style(args.source, args.bindings, args.output))
    else:
        emit(pdf_pages(args.pdf))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
