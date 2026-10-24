"""Session-bound project lifecycle. Scope is signed by the native server, never the model."""
import hashlib,hmac,io,json,mimetypes,os,re,sqlite3,subprocess,tarfile,tempfile,time,uuid,zipfile
from pathlib import Path, PurePosixPath
from contextlib import contextmanager
APP=Path(__file__).resolve().parents[2]
DB=APP.parent/'weknora-build-tools/data/weknora.db'
STATE=APP/'runs/native-projects.sqlite'
ROOT=APP/'runs/native-projects'
AGENT='5546cce7-872d-483a-a260-893339fbe05d'

def sha(p):
 with Path(p).open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()

def verify_scope(scope):
 key=os.environ.get('DETA_PROJECT_SCOPE_KEY','')
 raw='\n'.join(str(scope.get(k,'')) for k in ('tenant_id','session_id','user_id'))
 if not key or not hmac.compare_digest(hmac.new(key.encode(),raw.encode(),hashlib.sha256).hexdigest(),str(scope.get('signature',''))):raise ValueError('Trusted session scope required')
 with sqlite3.connect(DB) as c:
  columns={x[1] for x in c.execute('pragma table_info(sessions)')}
  select='tenant_id,user_id,agent_id,deleted_at'+(',agent_config' if 'agent_config' in columns else '')
  row=c.execute('select '+select+' from sessions where id=?',(scope['session_id'],)).fetchone()
 owner=str(scope.get('user_id',''))
 if owner.startswith('web_user:'):owner=owner[len('web_user:'):]
 fallback=json.loads(row[4] or '{}').get('agent_id') if row and len(row)>4 else None
 if not row or row[3] or str(row[0])!=str(scope['tenant_id']) or row[1]!=owner or (row[2] or fallback)!=AGENT:raise ValueError('Session ownership or agent mismatch')
 # Libraries are currently configured for this workspace only; fail closed elsewhere.
 if int(scope['tenant_id'])!=1:raise ValueError('Project libraries are not configured for this workspace')
 return scope

@contextmanager
def db():
 STATE.parent.mkdir(exist_ok=True,parents=True)
 c=sqlite3.connect(STATE,timeout=30);c.row_factory=sqlite3.Row
 c.executescript('''CREATE TABLE IF NOT EXISTS bindings(tenant INTEGER,session TEXT,project TEXT,run TEXT,title TEXT,source_hash TEXT,source_path TEXT,checkpoint TEXT,updated REAL,owner TEXT,PRIMARY KEY(tenant,session)); CREATE TABLE IF NOT EXISTS audits(id TEXT PRIMARY KEY,tenant INTEGER,session TEXT,project TEXT,created REAL,payload TEXT); CREATE TABLE IF NOT EXISTS continuation_tasks(project TEXT,stage INTEGER,task_key TEXT,status TEXT,attempts INTEGER DEFAULT 0,heartbeat REAL,last_error TEXT,PRIMARY KEY(project,stage,task_key)); CREATE TABLE IF NOT EXISTS continuation_runs(project TEXT PRIMARY KEY,tenant INTEGER,session TEXT,owner TEXT,status TEXT,requested REAL,heartbeat REAL,pid INTEGER,last_error TEXT);''')
 cols={r[1] for r in c.execute('pragma table_info(bindings)')}
 if 'owner' not in cols:c.execute('alter table bindings add column owner TEXT')
 for row in c.execute("select session from bindings where owner is null or owner='' ").fetchall():
  with sqlite3.connect(DB) as auth:
   found=auth.execute('select user_id from sessions where id=?',(row[0],)).fetchone()
  if found:c.execute('update bindings set owner=? where session=?',(found[0],row[0]))
 try:
  with c:yield c
 finally:c.close()

def binding(s):
 with db() as c:r=c.execute('select * from bindings where tenant=? and session=?',(s['tenant_id'],s['session_id'])).fetchone()
 return dict(r) if r else None

def container(s):
 r=subprocess.run(['docker','ps','-q','--filter','label=weknora_session_id='+s['session_id'],'--filter','label=weknora_tenant_id='+str(s['tenant_id'])],capture_output=True,text=True,timeout=15,check=True)
 ids=r.stdout.split()
 if len(ids)!=1:raise ValueError('No unique live session sandbox; restore session first')
 return ids[0]

def sandbox_read(s,path,max_bytes=80*1024*1024):
 p=PurePosixPath(path)
 if '..' in p.parts or not str(p).startswith(('/workspace/input/','/workspace/output/')):raise ValueError('Only session input/output files allowed')
 proc=subprocess.Popen(['docker','cp',container(s)+':'+str(p),'-'],stdout=subprocess.PIPE,stderr=subprocess.PIPE)
 try:
  with tarfile.open(fileobj=proc.stdout,mode='r|*') as archive:
   member=next(iter(archive))
   if not member.isfile() or member.size>max_bytes:raise ValueError('Regular file within size limit required')
   data=archive.extractfile(member).read(max_bytes+1)
   if len(data)!=member.size:raise ValueError('Incomplete sandbox file')
  proc.stdout.close();proc.wait(timeout=20)
  if proc.returncode:raise ValueError('Sandbox copy failed')
  return data
 finally:
  if proc.poll() is None:proc.kill();proc.wait()

def safe(root,path):
 p=PurePosixPath(path)
 if p.is_absolute() or '..' in p.parts or not p.parts:raise ValueError('Relative project path required')
 dest=Path(root)/str(p)
 for part in [dest,*dest.parents]:
  if part.is_symlink():raise ValueError('Symlink forbidden')
 if not dest.resolve().is_relative_to(Path(root).resolve()):raise ValueError('Path outside project')
 return dest

def atomic(path,data):
 path.parent.mkdir(parents=True,exist_ok=True)
 tmp=path.with_name(path.name+'.'+uuid.uuid4().hex+'.tmp');tmp.write_bytes(data);tmp.replace(path)

def inspect_input(name,data):
 """Deterministic container-level readability check. It does not certify semantics."""
 ext=Path(name).suffix.lower();result={'detected_type':ext.lstrip('.') or 'unknown','mime':mimetypes.guess_type(name)[0] or 'application/octet-stream','readable':False,'encrypted':False,'page_count':None,'text_extractable':None,'notes':[]}
 try:
  if ext=='.pdf':
   if not data.startswith(b'%PDF-'):raise ValueError('PDF signature missing')
   from pypdf import PdfReader
   reader=PdfReader(io.BytesIO(data),strict=False);result['encrypted']=bool(reader.is_encrypted)
   if result['encrypted']:
    try:
     if not reader.decrypt(''):raise ValueError('encrypted PDF requires a password')
    except Exception:raise ValueError('encrypted PDF requires a password')
   result['page_count']=len(reader.pages)
   if not result['page_count']:raise ValueError('PDF contains no pages')
   chars=0;failed=[];page_records=[];text_pages=0;fingerprints={}
   for number,page in enumerate(reader.pages,1):
    try:
     text=(page.extract_text() or '').strip();chars+=len(text)
     if text:text_pages+=1
     normalized=''.join(text.split())
     fingerprint=hashlib.sha256(normalized.encode()).hexdigest() if normalized else None
     if fingerprint:fingerprints.setdefault(fingerprint,[]).append(number)
     page_records.append({'physical_page':number,'rotation':int(page.get('/Rotate',0) or 0)%360,'text_characters':len(text),'text_sha256':fingerprint})
    except Exception:failed.append(number)
   result['text_extractable']=chars>0
   result['text_page_count']=text_pages
   result['scanned_page_count']=len(reader.pages)-text_pages
   result['scanned_page_ratio']=round((len(reader.pages)-text_pages)/len(reader.pages),6)
   result['rotated_pages']=[x['physical_page'] for x in page_records if x['rotation'] not in (0,)]
   result['duplicate_text_page_groups']=[pages for pages in fingerprints.values() if len(pages)>1]
   result['page_inspection']=page_records
   result['declared_page_consistency']='not_determinable_without_directory_or_printed_page_analysis'
   if failed:result['notes'].append('text extraction failed on pages '+','.join(map(str,failed[:20])))
   if not chars:result['notes'].append('image-only PDF; OCR/visual reading is required in the parsing step')
   result['readable']=True
  elif ext=='.docx':
   with zipfile.ZipFile(io.BytesIO(data)) as z:
    if 'word/document.xml' not in z.namelist():raise ValueError('DOCX document.xml missing')
    xml=z.read('word/document.xml')
   result['text_extractable']=bool(re.search(br'<w:t(?:\s[^>]*)?>',xml));result['readable']=True
  elif ext=='.xlsx':
   with zipfile.ZipFile(io.BytesIO(data)) as z:
    if 'xl/workbook.xml' not in z.namelist():raise ValueError('XLSX workbook.xml missing')
   result['text_extractable']=True;result['readable']=True
  elif ext in ('.png','.jpg','.jpeg','.tif','.tiff','.bmp','.webp'):
   from PIL import Image
   with Image.open(io.BytesIO(data)) as image:image.verify()
   result['text_extractable']=False;result['readable']=True;result['notes'].append('image input; visual reading/OCR is required in the parsing step')
  else:result['notes'].append('unsupported project input type')
 except Exception as exc:result['notes'].append(str(exc)[:240])
 return result

def input_manifest(root,project_id,title):
 meta_path=Path(root)/'project.json';meta=json.loads(meta_path.read_text()) if meta_path.exists() else {}
 roles=meta.get('input_roles',{});relations=meta.get('input_relations',{})
 entries=[]
 folder=Path(root)/'input'
 for path in sorted(folder.iterdir()) if folder.exists() else []:
  if not path.is_file() or path.is_symlink():continue
  data=path.read_bytes();digest=hashlib.sha256(data).hexdigest();role=roles.get(digest,'项目补充原件')
  priority={'招标主文件':100,'补遗':110,'澄清':110,'报价清单':90,'技术附件':80,'图纸':70,'投标格式模板':80}.get(role,50)
  entries.append({'name':path.name,'role':role,'priority':priority,'relation':relations.get(digest,{'type':'supplement','related_sha256':''}),'project_path':'input/'+path.name,'bytes':len(data),'sha256':digest,**inspect_input(path.name,data)})
 duplicate_files=[]
 by_hash={}
 for item in entries:by_hash.setdefault(item['sha256'],[]).append(item['name'])
 duplicate_files=[names for names in by_hash.values() if len(names)>1]
 return {'version':'source-lock-v2','project_id':project_id,'title':title,'generated_by':'project_gateway','lock_updated_at':meta.get('input_lock_updated',meta.get('created')),'inputs':entries,'input_count':len(entries),'duplicate_file_groups':duplicate_files,'all_readable':bool(entries) and all(x['readable'] for x in entries),'scope_note':'锁定文件、角色、关系及可检测的页级异常；缺页仅能由目录、印刷页码或声明页数的后续解析确认，不代表内容事实或业务批准'}

def write_manifest(root,project_id,title):
 manifest=input_manifest(root,project_id,title)
 atomic(Path(root)/'analysis/input-manifest.json',json.dumps(manifest,ensure_ascii=False,indent=2).encode())
 return manifest

def start(s,args,template_state,template_run):
 from bid_agent.store import Store
 source=args['source_path'];data=sandbox_read(s,source)
 if not source.startswith('/workspace/input/'):raise ValueError('Project baseline must be a session input attachment')
 digest=hashlib.sha256(data).hexdigest()
 if digest!=args['sha256']:raise ValueError('Attachment hash mismatch')
 inspection=inspect_input(PurePosixPath(source).name,data)
 if not inspection['readable']:raise ValueError('Attachment is not readable: '+'; '.join(inspection['notes']))
 if not args['project_title'].strip():raise ValueError('Project title required')
 owner=str(s['user_id']);ident=hashlib.sha256((str(s['tenant_id'])+'\n'+owner+'\n'+digest).encode()).hexdigest()[:24]
 root=ROOT/ident
 with db() as c:
  c.execute('BEGIN IMMEDIATE')
  old=c.execute('select * from bindings where tenant=? and session=?',(s['tenant_id'],s['session_id'])).fetchone()
  if old and old['source_hash']==digest:
   c.commit()
   return snapshot(s)
  prior=c.execute('select * from bindings where tenant=? and owner=? and source_hash=? order by updated desc limit 1',(s['tenant_id'],owner,digest)).fetchone()
  if prior and (ROOT/prior['project']).is_dir():
   c.execute('insert or replace into bindings(tenant,session,project,run,title,source_hash,source_path,checkpoint,updated,owner) values(?,?,?,?,?,?,?,?,?,?)',(s['tenant_id'],s['session_id'],prior['project'],prior['run'],prior['title'],prior['source_hash'],prior['source_path'],prior['checkpoint'],time.time(),owner))
   c.commit()
   return snapshot(s)
  cached=root/'binding.json'
  if cached.exists():
   saved=json.loads(cached.read_text())
   c.execute('insert or replace into bindings(tenant,session,project,run,title,source_hash,source_path,checkpoint,updated,owner) values(?,?,?,?,?,?,?,?,?,?)',tuple(saved.get(k) for k in ['tenant','session','project','run','title','source_hash','source_path','checkpoint','updated'])+(owner,))
   return { 'status':'restored','project_id':ident,'next_action':'调用resume_project读取已有检查点；未新建任务'}
  store=Store(Path(template_state));template=store.get(template_run)['scope']
  root.mkdir(parents=True,exist_ok=True)
  target=safe(root,'input/'+PurePosixPath(source).name)
  atomic(target,data)
  run=store.create('制作 '+args['project_title'],{'project':str(root),'skills':template.get('skills',{}),'libraries':template.get('libraries',[]),'tenant_id':s['tenant_id'],'session_id':s['session_id'],'source_sha256':digest,'historical_replay':False})
  checkpoint={'phase':'source_registered','next_action':'读取主技能和本项目招标全文，建立条款与证据矩阵','open_items':[],'status':'in_progress'}
  c.execute('insert or replace into bindings(tenant,session,project,run,title,source_hash,source_path,checkpoint,updated,owner) values(?,?,?,?,?,?,?,?,?,?)',(s['tenant_id'],s['session_id'],ident,run,args['project_title'],digest,str(target),json.dumps(checkpoint,ensure_ascii=False),time.time(),owner))
  row=c.execute('select * from bindings where tenant=? and session=?',(s['tenant_id'],s['session_id'])).fetchone()
  atomic(root/'binding.json',json.dumps(dict(row),ensure_ascii=False).encode())
  created=time.time()
  atomic(root/'project.json',json.dumps({'project_id':ident,'title':args['project_title'],'source_sha256':digest,'historical_replay':False,'created':created,'input_lock_updated':created,'input_roles':{digest:'招标主文件'},'input_relations':{digest:{'type':'primary','related_sha256':''}}},ensure_ascii=False,indent=2).encode())
  write_manifest(root,ident,args['project_title'])
 return snapshot(s)

def snapshot(s):
 b=binding(s)
 if not b:return {'status':'unbound','next_action':'调用 begin_project，以当前附件原件SHA256独立建档；不得访问旧项目','project_id':None}
 root=ROOT/b['project'];files=[{'path':str(p.relative_to(root)),'bytes':p.stat().st_size} for p in root.rglob('*') if p.is_file() and not p.is_symlink()]
 source=Path(b['source_path']);actual=sha(source) if source.is_file() else None
 manifest=write_manifest(root,b['project'],b['title'])
 from project_stages import REQUIRED,status
 from workflow_policy import load_policy
 policy=load_policy()
 stage_rows=status(root);fact_files=sorted((root/'analysis/step4').glob('project-facts*.json')) if (root/'analysis/step4').exists() else []
 next_row=next((row for row in stage_rows if row.get('status')!='passed'),None)
 facts=fact_files[-1] if fact_files else None;locks=sorted((root/'analysis/step6').glob('configuration-lock-*.json')) if (root/'analysis/step6').exists() else []
 payload_dir=root/'analysis/step2/tender-analysis-payloads'
 pending_payloads=[{'path':str(x.relative_to(root)),'sha256':sha(x),'bytes':x.stat().st_size} for x in sorted(payload_dir.glob('pending*.json'))] if payload_dir.exists() else []
 quality='agent_scope_complete' if len(stage_rows)==10 and all(x.get('status')=='passed' for x in stage_rows) else 'in_progress'
 return {'policy':policy,'stages':stage_rows,'next_stage':next_row.get('stage') if next_row else None,
         'next_stage_status':next_row.get('status') if next_row else 'complete',
         'required_artifacts':REQUIRED.get(next_row.get('stage'),[]) if next_row else [],
         'project_id':b['project'],'title':b['title'],'source_sha256':b['source_hash'],'source_path':'input/'+Path(b['source_path']).name,'source_integrity':{'exists':source.is_file(),'expected_sha256':b['source_hash'],'actual_sha256':actual,'matched':actual==b['source_hash']},'input_manifest':manifest,'project_facts':{'path':str(facts.relative_to(root)) if facts else None,'exists':bool(facts),'sha256':sha(facts) if facts else None},'configuration_locks':[{'path':str(x.relative_to(root)),'sha256':sha(x)} for x in locks],'pending_tender_analysis_payloads':pending_payloads,'sandbox_output_dir':'/workspace/output/'+b['project'],'checkpoint':json.loads(b['checkpoint']),'files':files[:100],'historical_replay':False,'quality_status':quality}

def add_input(s,args):
 b=binding(s)
 if not b:raise ValueError('Call begin_project first')
 source=args['source_path']
 if not args['role'].strip():raise ValueError('Input role required')
 if not source.startswith('/workspace/input/'):raise ValueError('Project input must be a session input attachment')
 data=sandbox_read(s,source);digest=hashlib.sha256(data).hexdigest()
 if digest!=args['sha256']:raise ValueError('Attachment hash mismatch')
 inspection=inspect_input(PurePosixPath(source).name,data)
 if not inspection['readable']:raise ValueError('Attachment is not readable: '+'; '.join(inspection['notes']))
 root=ROOT/b['project'];current=input_manifest(root,b['project'],b['title'])
 duplicate=next((x for x in current['inputs'] if x['sha256']==digest),None)
 if duplicate:return {'status':'duplicate','input':duplicate,'manifest':current}
 name=PurePosixPath(source).name
 if not name or name in ('.','..'):raise ValueError('Input file name required')
 dest=safe(root,'input/'+name)
 if dest.exists():
  stem,suffix=Path(name).stem,Path(name).suffix;dest=safe(root,'input/'+stem+'-'+digest[:8]+suffix)
 atomic(dest,data)
 project_file=root/'project.json';project=json.loads(project_file.read_text());project.setdefault('input_roles',{})[digest]=args['role'].strip();project['input_lock_updated']=time.time();atomic(project_file,json.dumps(project,ensure_ascii=False,indent=2).encode())
 manifest=write_manifest(root,b['project'],b['title'])
 return {'status':'registered','role':args['role'],'input':next(x for x in manifest['inputs'] if x['sha256']==digest),'manifest':manifest}

def relate_input(s,args):
 b=binding(s)
 if not b:raise ValueError('Call begin_project first')
 root=ROOT/b['project'];project_file=root/'project.json';project=json.loads(project_file.read_text())
 known=set(project.get('input_roles',{}));digest=args['sha256'];related=args['related_sha256']
 if digest not in known:raise ValueError('Input SHA256 is not registered in this project')
 allowed={'primary','supplement','clarifies','replaces','deletes_requirements_from','price_schedule_for','technical_attachment_for','drawing_for','template_for'}
 if args['relation_type'] not in allowed:raise ValueError('Invalid input relation type')
 if args['relation_type']!='primary' and related not in known:raise ValueError('Related SHA256 is not registered in this project')
 if digest==related:raise ValueError('Input cannot relate to itself')
 project.setdefault('input_relations',{})[digest]={'type':args['relation_type'],'related_sha256':related}
 project['input_lock_updated']=time.time();atomic(project_file,json.dumps(project,ensure_ascii=False,indent=2).encode())
 return {'status':'related','manifest':write_manifest(root,b['project'],b['title'])}

def index_technical_source(s,args):
 """Index every table row and non-table text block in the selected tender pages."""
 b=binding(s)
 if not b:raise ValueError('Call begin_project first')
 if 'physical_pages' in args:
  selected=[int(x) for x in args['physical_pages']]
  if not selected or selected!=sorted(set(selected)) or selected[0]<1:raise ValueError('Physical pages must be unique, ascending, and positive')
 else:
  start_page,end_page=int(args['start_page']),int(args['end_page'])
  if start_page<1 or end_page<start_page:raise ValueError('Invalid physical page range')
  selected=list(range(start_page,end_page+1))
 import pymupdf
 source=Path(b['source_path'])
 if not source.is_file() or sha(source)!=b['source_hash']:raise ValueError('Locked tender source is missing or changed')
 pages=[];table_row_count=0;text_block_count=0
 with pymupdf.open(source) as document:
  if selected[-1]>len(document):raise ValueError('Technical page selection exceeds tender page count')
  for number in selected:
   page=document[number-1];found=page.find_tables().tables;table_boxes=[pymupdf.Rect(t.bbox) for t in found]
   tables=[]
   for table_index,table in enumerate(found,1):
    rows=[]
    for row_index,cells in enumerate(table.extract(),1):
     values=[None if cell is None else str(cell) for cell in cells]
     text_value=' | '.join('' if cell is None else str(cell).replace('\n',' ').strip() for cell in cells)
     rows.append({'id':f'TECH-P{number:03}-T{table_index:02}-R{row_index:03}','row':row_index,'cells':values,'text':text_value,'text_sha256':hashlib.sha256(text_value.encode()).hexdigest()})
    table_row_count+=len(rows);tables.append({'table':table_index,'bbox':[round(x,3) for x in table.bbox],'rows':rows})
   blocks=[]
   for block_index,block in enumerate(page.get_text('blocks'),1):
    rect=pymupdf.Rect(block[:4]);text_value=str(block[4]).strip()
    if not text_value:continue
    overlap=False
    for box in table_boxes:
     intersection=rect & box
     if not intersection.is_empty and rect.get_area()>0 and intersection.get_area()/rect.get_area()>=0.5:overlap=True;break
    if overlap:continue
    blocks.append({'id':f'TECH-P{number:03}-B{block_index:03}','block':block_index,'bbox':[round(x,3) for x in rect],'text':text_value,'text_sha256':hashlib.sha256(text_value.encode()).hexdigest()})
   text_block_count+=len(blocks);pages.append({'physical_page':number,'printed_page':number-2 if number>=3 else None,'tables':tables,'text_blocks':blocks})
 result={'version':'technical-source-index-v1','project_id':b['project'],'source_path':'input/'+source.name,'source_sha256':b['source_hash'],'physical_pages':selected,'physical_page_range':[selected[0],selected[-1]],'pages':pages,'stats':{'page_count':len(pages),'table_count':sum(len(p['tables']) for p in pages),'table_row_count':table_row_count,'text_block_count':text_block_count},'semantic_split':'not_performed','scope_note':'只索引逐页分析明确分类为technical的主招标文件页；逐页保存表格行及表外文本块，保持来源顺序，不生成响应、偏离或企业事实'}
 target=ROOT/b['project']/'analysis/step2/technical-source-index.json';atomic(target,json.dumps(result,ensure_ascii=False,indent=2).encode())
 return {'path':'analysis/step2/technical-source-index.json','sha256':sha(target),'stats':result['stats']}

def extract_project_source(s):
 """Extract every physical PDF page from the locked project inputs."""
 b=binding(s)
 if not b:raise ValueError('Call begin_project first')
 import pymupdf
 root=ROOT/b['project'];manifest=input_manifest(root,b['project'],b['title']);documents=[]
 for item in manifest['inputs']:
  path=safe(root,item['project_path'])
  if path.suffix.lower()!='.pdf':
   documents.append({'name':item['name'],'path':item['project_path'],'sha256':item['sha256'],'role':item['role'],'status':'non_pdf_requires_native_document_reader','pages':[]})
   continue
  pages=[]
  with pymupdf.open(path) as document:
   for number,page in enumerate(document,1):
    text_value=page.get_text('text')
    pages.append({'physical_page':number,'text':text_value,'text_sha256':hashlib.sha256(text_value.encode()).hexdigest()})
  documents.append({'name':item['name'],'path':item['project_path'],'sha256':item['sha256'],'role':item['role'],'status':'extracted','page_count':len(pages),'pages':pages})
 folder=root/'analysis/step2';folder.mkdir(parents=True,exist_ok=True);sequence=1
 while (folder/('source-extraction-v%d.json'%sequence)).exists():sequence+=1
 target=folder/('source-extraction-v%d.json'%sequence)
 payload={'version':'project-source-extraction-v1','project_id':b['project'],'source_sha256':b['source_hash'],'documents':documents,'generated_by':'project_gateway','created_at':time.time()}
 atomic(target,json.dumps(payload,ensure_ascii=False,indent=2).encode())
 return {'path':str(target.relative_to(root)),'sha256':sha(target),'documents':len(documents),'pages':sum(len(x['pages']) for x in documents)}

def build_tender_analysis(s,args):
 """Persist model classifications only after exact page-quote verification."""
 b=binding(s)
 if not b:raise ValueError('Call begin_project first')
 root=ROOT/b['project'];source=safe(root,args['source_extraction_path'])
 if not source.is_file():raise ValueError('Source extraction file is missing')
 extraction=json.loads(source.read_text());allowed={'qualification','rejection','scoring','business','technical','pricing','format','signature'}
 page_map={(d['sha256'],p['physical_page']):p['text'] for d in extraction.get('documents',[]) for p in d.get('pages',[])}
 document_hashes=list(dict.fromkeys(d['sha256'] for d in extraction.get('documents',[]) if d.get('pages')))
 default_hash=document_hashes[0] if len(document_hashes)==1 else None
 items=[];ids=[]
 for raw in args['items']:
  item=dict(raw)
  item['source_sha256']=item.get('source_sha256') or default_hash
  items.append(item)
 if not items:raise ValueError('Tender analysis items cannot be empty')
 for position,item in enumerate(items,1):
  if item.get('category') not in allowed:raise ValueError('Invalid tender analysis category at item '+str(position))
  if not item.get('requirement_id') or item['requirement_id'] in ids:raise ValueError('Tender analysis IDs are empty or duplicated')
  ids.append(item['requirement_id']);text_value=page_map.get((item.get('source_sha256'),item.get('physical_page')))
  if text_value is None:raise ValueError('Tender analysis item points outside extracted pages: '+item['requirement_id'])
  quote=item.get('source_quote','')
  if not quote or quote not in text_value:raise ValueError('Tender analysis quote is not verbatim on its source page: '+item['requirement_id'])
 page_reviews=[];review_keys=[];item_pairs={(x['source_sha256'],x['physical_page'],x['category']) for x in items}
 for compact in args['page_reviews']:
  digest=compact.get('source_sha256') or default_hash
  start=int(compact.get('physical_page') or 0);end=int(compact.get('end_physical_page') or start)
  if start<1 or end<start:raise ValueError('Tender analysis page review range is invalid')
  if compact.get('disposition')=='requirements_extracted' and end!=start:raise ValueError('A requirement page review cannot span multiple pages')
  for number in range(start,end+1):
   review={k:v for k,v in compact.items() if k!='end_physical_page'}
   review['source_sha256']=digest;review['physical_page']=number
   key=(digest,number)
   if key not in page_map or key in review_keys:raise ValueError('Tender analysis page reviews are duplicated or outside extracted pages')
   review_keys.append(key);categories=review.get('categories') or [];disposition=review.get('disposition');reason=str(review.get('reason') or '').strip()
   if disposition=='requirements_extracted':
    if not categories or any(x not in allowed for x in categories):raise ValueError('Reviewed requirement page needs valid categories')
    if any((key[0],key[1],category) not in item_pairs for category in categories):raise ValueError('Every declared page category needs at least one same-page verbatim item')
   elif disposition=='no_actionable_requirements':
    if categories or not reason:raise ValueError('No-actionable page review needs an empty category list and a reason')
   else:raise ValueError('Invalid page review disposition')
   page_reviews.append(review)
 if review_keys!=list(page_map):raise ValueError('Tender analysis must review every extracted physical page exactly once in source order')
 reviewed_pairs={(x['source_sha256'],x['physical_page']) for x in page_reviews if x['disposition']=='requirements_extracted'}
 if any((x['source_sha256'],x['physical_page']) not in reviewed_pairs for x in items):raise ValueError('Tender analysis item is not covered by its page review')
 amendments=args['amendment_effects'];project_data=json.loads((root/'project.json').read_text());known=set(project_data.get('input_roles',{}));relations=project_data.get('input_relations',{})
 for effect in amendments:
  digest=effect.get('source_sha256')
  if digest not in known or relations.get(digest,{}).get('type') not in {'supplement','clarifies','replaces','deletes_requirements_from'}:raise ValueError('Amendment effect is not bound to a registered amendment input')
  if effect.get('effect') not in {'supplements','clarifies','replaces','deletes'} or not effect.get('target_requirement_id'):raise ValueError('Amendment effect is incomplete')
 folder=root/'analysis/step2';folder.mkdir(parents=True,exist_ok=True);sequence=1
 while (folder/('tender-analysis-v%d.json'%sequence)).exists():sequence+=1
 target=folder/('tender-analysis-v%d.json'%sequence)
 payload={'version':'tender-analysis-v1','project_id':b['project'],'source_sha256':b['source_hash'],'source_extraction_path':args['source_extraction_path'],'source_extraction_sha256':sha(source),'items':items,'page_reviews':page_reviews,'reviewed_page_count':len(page_reviews),'amendment_effects':amendments,'category_counts':{x:sum(i['category']==x for i in items) for x in sorted(allowed)},'created_at':time.time(),'generated_by':'project_gateway'}
 atomic(target,json.dumps(payload,ensure_ascii=False,indent=2).encode())
 return {'path':str(target.relative_to(root)),'sha256':sha(target),'items':len(items),'category_counts':payload['category_counts']}

def save_tender_analysis_batch(s,args):
 """Validate and persist one contiguous page batch without claiming full coverage."""
 b=binding(s)
 if not b:raise ValueError('Call begin_project first')
 root=ROOT/b['project'];source=safe(root,args['source_extraction_path'])
 if not source.is_file():raise ValueError('Source extraction file is missing')
 extraction=json.loads(source.read_text());digest=args['source_sha256'];start=int(args['start_page']);end=int(args['end_page'])
 if start<1 or end<start:raise ValueError('Tender analysis batch page range is invalid')
 if end-start+1>20:raise ValueError('Tender analysis batch cannot exceed 20 physical pages')
 page_map={(d['sha256'],p['physical_page']):p['text'] for d in extraction.get('documents',[]) for p in d.get('pages',[])}
 expected=[(digest,n) for n in range(start,end+1)]
 if any(key not in page_map for key in expected):raise ValueError('Tender analysis batch points outside one extracted document')
 allowed={'qualification','rejection','scoring','business','technical','pricing','format','signature'};items=[];ids=set()
 for position,raw in enumerate(args['items'],1):
  item=dict(raw);item['source_sha256']=item.get('source_sha256') or digest
  if item['source_sha256']!=digest or not start<=int(item.get('physical_page') or 0)<=end:raise ValueError('Tender analysis batch item points outside its batch')
  if item.get('category') not in allowed or not item.get('requirement_id') or item['requirement_id'] in ids:raise ValueError('Tender analysis batch category/ID is invalid at item '+str(position))
  ids.add(item['requirement_id']);quote=item.get('source_quote','')
  if not quote or quote not in page_map[(digest,item['physical_page'])]:raise ValueError('Tender analysis batch quote is not verbatim: '+item['requirement_id'])
  items.append(item)
 item_pairs={(x['source_sha256'],x['physical_page'],x['category']) for x in items};reviews=[];keys=[]
 for compact in args['page_reviews']:
  review_digest=compact.get('source_sha256') or digest;review_start=int(compact.get('physical_page') or 0);review_end=int(compact.get('end_physical_page') or review_start)
  if review_digest!=digest or review_start<start or review_end> end or review_end<review_start:raise ValueError('Tender analysis page review points outside its batch')
  if compact.get('disposition')=='requirements_extracted' and review_end!=review_start:raise ValueError('A requirement page review cannot span multiple pages')
  for number in range(review_start,review_end+1):
   review={k:v for k,v in compact.items() if k!='end_physical_page'};review['source_sha256']=digest;review['physical_page']=number;key=(digest,number)
   if key in keys:raise ValueError('Tender analysis batch page reviews are duplicated')
   keys.append(key);categories=review.get('categories') or [];disposition=review.get('disposition');reason=str(review.get('reason') or '').strip()
   if disposition=='requirements_extracted':
    if not categories or any(x not in allowed for x in categories) or any((digest,number,x) not in item_pairs for x in categories):raise ValueError('Tender analysis batch page category lacks a same-page verbatim item')
   elif disposition=='no_actionable_requirements':
    if categories or not reason:raise ValueError('No-actionable batch page needs an empty category list and a reason')
   else:raise ValueError('Invalid tender analysis batch disposition')
   reviews.append(review)
 if keys!=expected:raise ValueError('Tender analysis batch must review every page in its range exactly once')
 folder=root/'analysis/step2/tender-analysis-batches';folder.mkdir(parents=True,exist_ok=True);target=folder/(digest[:12]+'-p%04d-p%04d.json'%(start,end))
 payload={'version':'tender-analysis-batch-v1','project_id':b['project'],'source_sha256':b['source_hash'],'document_sha256':digest,'source_extraction_path':args['source_extraction_path'],'source_extraction_sha256':sha(source),'start_page':start,'end_page':end,'items':items,'page_reviews':reviews,'created_at':time.time(),'generated_by':'project_gateway'}
 atomic(target,json.dumps(payload,ensure_ascii=False,indent=2).encode())
 return {'path':str(target.relative_to(root)),'sha256':sha(target),'pages':len(reviews),'items':len(items)}

def finalize_tender_analysis_batches(s,args):
 b=binding(s)
 if not b:raise ValueError('Call begin_project first')
 root=ROOT/b['project'];items=[];reviews=[]
 if not args['batch_paths']:raise ValueError('Tender analysis batch paths cannot be empty')
 for rel in args['batch_paths']:
  path=safe(root,rel)
  if not path.is_file():raise ValueError('Tender analysis batch file is missing: '+rel)
  data=json.loads(path.read_text())
  if data.get('version')!='tender-analysis-batch-v1' or data.get('project_id')!=b['project'] or data.get('source_extraction_path')!=args['source_extraction_path']:raise ValueError('Tender analysis batch belongs to another project or extraction')
  source=safe(root,args['source_extraction_path'])
  if not source.is_file() or data.get('source_extraction_sha256')!=sha(source):raise ValueError('Tender analysis batch source extraction changed')
  items.extend(data.get('items',[]));reviews.extend(data.get('page_reviews',[]))
 return build_tender_analysis(s,{'source_extraction_path':args['source_extraction_path'],'items':items,'page_reviews':reviews,'amendment_effects':args['amendment_effects']})

def import_tender_analysis_payload(s,args):
 """Import a large sandbox or resumable project JSON, then validate/persist it in page batches."""
 b=binding(s)
 if not b:raise ValueError('Call begin_project first')
 root=ROOT/b['project'];payload_path=args['payload_path']
 prefix='/workspace/output/'+b['project']+'/'
 if payload_path.startswith(prefix):raw=sandbox_read(s,payload_path,20*1024*1024);persisted=None
 else:
  persisted=safe(root,payload_path);allowed=root/'analysis/step2/tender-analysis-payloads'
  if persisted.parent!=allowed or not persisted.name.startswith('pending') or persisted.suffix!='.json' or not persisted.is_file():raise ValueError('Persistent tender analysis payload must be a pending JSON in the project payload queue')
  raw=persisted.read_bytes()
  if len(raw)>20*1024*1024:raise ValueError('Tender analysis payload exceeds size limit')
 digest=hashlib.sha256(raw).hexdigest()
 if digest!=args['payload_sha256']:raise ValueError('Tender analysis payload hash mismatch')
 try:payload=json.loads(raw)
 except Exception as exc:raise ValueError('Tender analysis payload is not valid JSON') from exc
 if payload.get('source_extraction_path')!=args['source_extraction_path']:raise ValueError('Tender analysis payload source extraction path differs')
 source=safe(root,args['source_extraction_path'])
 if not source.is_file():raise ValueError('Source extraction file is missing')
 extraction=json.loads(source.read_text());documents=[d for d in extraction.get('documents',[]) if d.get('pages')]
 hashes=[d['sha256'] for d in documents];default_hash=hashes[0] if len(hashes)==1 else None
 items=[]
 for raw_item in payload.get('items',[]):
  item=dict(raw_item);item['source_sha256']=item.get('source_sha256') or default_hash;items.append(item)
 expanded=[]
 for compact in payload.get('page_reviews',[]):
  review_digest=compact.get('source_sha256') or default_hash;start=int(compact.get('physical_page') or 0);end=int(compact.get('end_physical_page') or start)
  if start<1 or end<start:raise ValueError('Tender analysis payload page review range is invalid')
  for number in range(start,end+1):
   review={k:v for k,v in compact.items() if k!='end_physical_page'};review['source_sha256']=review_digest;review['physical_page']=number;expanded.append(review)
 batch_paths=[]
 for document in documents:
  doc_hash=document['sha256'];page_numbers=[p['physical_page'] for p in document['pages']]
  for offset in range(0,len(page_numbers),20):
   selected=page_numbers[offset:offset+20];start,end=selected[0],selected[-1]
   batch_items=[x for x in items if x.get('source_sha256')==doc_hash and start<=int(x.get('physical_page') or 0)<=end]
   batch_reviews=[x for x in expanded if x.get('source_sha256')==doc_hash and start<=int(x.get('physical_page') or 0)<=end]
   saved=save_tender_analysis_batch(s,{'source_extraction_path':args['source_extraction_path'],'source_sha256':doc_hash,'start_page':start,'end_page':end,'items':batch_items,'page_reviews':batch_reviews})
   batch_paths.append(saved['path'])
 result=finalize_tender_analysis_batches(s,{'source_extraction_path':args['source_extraction_path'],'batch_paths':batch_paths,'amendment_effects':payload.get('amendment_effects',[])})
 if persisted is not None:persisted.replace(persisted.with_name('imported-'+str(int(time.time()))+'.json'))
 return {**result,'payload_sha256':digest,'batch_paths':batch_paths,'batch_count':len(batch_paths)}

def checkpoint(s,args):
 b=binding(s)
 if not b:raise ValueError('Call begin_project first')
 # A model checkpoint is a plan, never approval or a passed quality gate.
 record={'phase':args['phase'],'next_action':args['next_action'],'open_items':args['open_items'],'status':'blocked' if args.get('blocked') else 'in_progress','assertions_verified':False}
 with db() as c:c.execute('update bindings set checkpoint=?,updated=? where tenant=? and session=?',(json.dumps(record,ensure_ascii=False),time.time(),s['tenant_id'],s['session_id']))
 with db() as c:
  row=c.execute('select * from bindings where tenant=? and session=?',(s['tenant_id'],s['session_id'])).fetchone()
  atomic(ROOT/b['project']/'binding.json',json.dumps(dict(row),ensure_ascii=False).encode())
 return record

def import_file(s,args):
 b=binding(s)
 if not b:raise ValueError('Call begin_project first')
 source=args['source_path'];prefix='/workspace/output/'+b['project']+'/'
 if not source.startswith(prefix):raise ValueError('Output must belong to active project output directory')
 data=sandbox_read(s,source);dest=safe(ROOT/b['project'],args['target'])
 if not str(PurePosixPath(args['target'])).startswith('analysis/'):
  raise ValueError('capture_project_file is for internal analysis files only; the three approved DOCX/DOCX/XLSX website outputs enter outputs through stage 9')
 if dest.exists() and dest.read_bytes()!=data:raise ValueError('Use a new version; existing file differs')
 atomic(dest,data)
 return {'path':args['target'],'sha256':sha(dest),'bytes':len(data),'quality_status':'not_verified'}

def export_file(s,args):
 b=binding(s)
 if not b:raise ValueError('Call begin_project first')
 src=safe(ROOT/b['project'],args['path'])
 if not src.is_file() or src.stat().st_size>80*1024*1024:raise ValueError('Project file missing or too large')
 target='/workspace/output/'+b['project']+'/'+str(PurePosixPath(args['path']))
 cid=container(s)
 subprocess.run(['docker','exec',cid,'mkdir','-p',str(PurePosixPath(target).parent)],check=True,capture_output=True,timeout=20)
 # Docker follows a pre-existing destination symlink; do not overwrite any destination.
 check=subprocess.run(['docker','exec',cid,'test','-e',target],timeout=10)
 if check.returncode==0:raise ValueError('Sandbox target exists; choose a new project version')
 # Validate each parent against symlinks inside the isolated container.
 code='import pathlib,sys; p=pathlib.Path(sys.argv[1]); assert not any(x.is_symlink() for x in [p,*p.parents])'
 subprocess.run(['docker','exec',cid,'python3','-c',code,target],check=True,timeout=10,capture_output=True)
 subprocess.run(['docker','cp',str(src),cid+':'+target],check=True,timeout=60,capture_output=True)
 return {'sandbox_path':target,'sha256':sha(src)}

def audit_docx(path):
 from docx import Document
 issues=[]
 with zipfile.ZipFile(path) as z:
  xml=z.read('word/document.xml');media=[n for n in z.namelist() if n.startswith('word/media/')];footers=[n for n in z.namelist() if n.startswith('word/footer')]
 d=Document(path);text='\n'.join([p.text for p in d.paragraphs]+[c.text for t in d.tables for r in t.rows for c in r.cells])
 if not media:issues.append('没有内嵌证明原页；附件清单不能算装配完成')
 if '目录' not in text and b'TOC ' not in xml:issues.append('缺少目录')
 if not footers:issues.append('缺少页脚页码部件')
 if any(x in text for x in ['待补','待财务','待核','留空']):issues.append('正文仍有待补或待核字段')
 if any(x in text for x in ['完全响应','无偏离']):issues.append('无偏离/完全响应声明须逐条事实核验，文本本身不构成证明')
 if '非递交' not in text:issues.append('草稿正文缺少非递交标识')
 return {'file':path.name,'sha256':sha(path),'embedded_media':len(media),'issues':issues,'quality_status':'draft_not_verified'}

def audit(s):
 b=binding(s)
 if not b:return {'status':'unbound','issues':['未绑定独立项目；不能沿用旧项目状态或宣称完整完成'],'files':[]}
 root=ROOT/b['project'];files=[];issues=[]
 # Stage recording already captures files with hashes. Never sweep the whole
 # sandbox into outputs/: doing so used to publish skeletons, fragments, JSON
 # and duplicate outputs/outputs trees as if they were deliverables.
 for path in sorted((root/'outputs').rglob('*.docx')) if (root/'outputs').exists() else []:
  try:
   item=audit_docx(path)
   from bid_agent.response_logic import check_logic
   item['logic']=check_logic(b['source_path'],path)
   files.append(item)
  except Exception as exc:issues.append(path.name+'：文档检查失败 '+str(exc)[:100])
 if not files:issues.append('当前项目尚无已核验的DOCX产物')
 # No arbitrary model-written JSON can confer semantic/visual approval.
 issues+=['条款全覆盖、原件适用性、证据装配和最终逐页视觉检查尚无可信完整验收回执']
 result={'status':'draft_not_verified','project_id':b['project'],'files':files,'issues':issues,'submission_ready':False,'checkpoint':json.loads(b['checkpoint'])}
 with db() as c:c.execute('insert into audits values(?,?,?,?,?,?)',(uuid.uuid4().hex,s['tenant_id'],s['session_id'],b['project'],time.time(),json.dumps(result,ensure_ascii=False)))
 atomic(root/'analysis/latest-delivery-audit.json',json.dumps(result,ensure_ascii=False,indent=2).encode())
 return result

def request_continuation(s):
 b=binding(s)
 if not b:raise ValueError('Call begin_project first')
 with db() as c:
  current=c.execute('select * from continuation_runs where project=?',(b['project'],)).fetchone()
  if current and current['status'] in ('requested','launching','running'):
   return {'project_id':b['project'],'status':current['status'],'requested_at':current['requested'],'heartbeat':current['heartbeat']}
  now=time.time();c.execute('insert or replace into continuation_runs(project,tenant,session,owner,status,requested,heartbeat,pid,last_error) values(?,?,?,?,?,?,?,?,?)',(b['project'],s['tenant_id'],s['session_id'],s['user_id'],'requested',now,now,None,''))
 return {'project_id':b['project'],'status':'requested','requested_at':now,'next_action':'通用续跑调度器将从当前最早未通过阶段继续'}

def continuation_status(s):
 b=binding(s)
 if not b:return {'status':'unbound'}
 with db() as c:
  run=c.execute('select * from continuation_runs where project=?',(b['project'],)).fetchone()
  tasks=[dict(x) for x in c.execute('select * from continuation_tasks where project=? order by stage,task_key',(b['project'],))]
 return {'project_id':b['project'],'run':dict(run) if run else {'status':'not_requested'},'tasks':tasks}

def create_stage_validation(s,args):
 """Create the hash-bound stage receipt; the model cannot author this receipt."""
 b=binding(s)
 if not b:raise ValueError('Call begin_project first')
 from project_stages import REQUIRED
 stage=int(args['stage']);artifacts=dict(args['artifacts'])
 if stage not in REQUIRED:raise ValueError('Stage must be 1..10')
 expected=set(REQUIRED[stage])-{'stage_validation'}
 if set(artifacts)!=expected:
  missing=sorted(expected-set(artifacts));extra=sorted(set(artifacts)-expected)
  details=[]
  if missing:details.append('missing: '+', '.join(missing))
  if extra:details.append('unknown: '+', '.join(extra))
  raise ValueError('Stage validation requires the exact non-validation artifact set ('+'; '.join(details)+')')
 root=ROOT/b['project'];hashes={}
 for kind,rel in artifacts.items():
  prefix='outputs/' if stage==9 and kind in {'final_bid_docx','eight_column_docx','eight_column_xlsx'} else 'analysis/'
  if not str(PurePosixPath(rel)).startswith(prefix):raise ValueError(kind+' must be stored under '+prefix)
  dest=safe(root,rel)
  if not dest.is_file() or not dest.stat().st_size:
   data=sandbox_read(s,'/workspace/output/'+b['project']+'/'+rel)
   if not data:raise ValueError('Missing/empty artifact: '+rel)
   atomic(dest,data)
  hashes[kind]=sha(dest)
 folder=root/'analysis'/('step%d'%stage);folder.mkdir(parents=True,exist_ok=True)
 sequence=1
 while (folder/('stage-validation-v%d.json'%sequence)).exists():sequence+=1
 target=folder/('stage-validation-v%d.json'%sequence)
 payload={'version':'stage-validation-v1','project_id':b['project'],'source_sha256':b['source_hash'],
          'stage':stage,'status':'passed','artifact_hashes':hashes,
          'checks':[{'id':'server_artifact_integrity','status':'pass'}],
          'awarded_quality_state':'agent_scope_complete' if stage==10 else 'in_progress',
          'created_at':time.time(),'generated_by':'project_gateway'}
 atomic(target,json.dumps(payload,ensure_ascii=False,indent=2).encode())
 return {'path':str(target.relative_to(root)),'sha256':sha(target),'stage':stage,'artifact_hashes':hashes}

SPECS=[
 ('analyze_tender_payload','长招标文件的首屏恢复入口；功能等同import_project_tender_analysis_payload。直接传入resume_project返回的pending_tender_analysis_payloads项目相对路径和SHA-256，服务器读取载荷、按每批不超过20页验证并生成正式tender_analysis。禁止先用read_json_value重读载荷。',{'source_extraction_path':{'type':'string'},'payload_path':{'type':'string'},'payload_sha256':{'type':'string'}}),
 ('extract_project_source','抽取当前项目全部已锁定PDF原件的每个物理页文字并记录页级SHA256，生成第2步source_extraction；不分类、不生成响应。',{}),
 ('build_project_tender_analysis','将资格、否决、评分、商务、技术、报价、格式、签章八类动态分析写入当前项目；每条source_quote必须逐字存在于source_extraction指定物理页。单原件项目由服务器补全重复的source_sha256。page_reviews必须按来源顺序覆盖每个物理页；连续无可执行要求页可用end_physical_page压缩，但含要求页必须逐页提交并绑定同页逐字条目。禁止抽样分析及响应结论。',{'source_extraction_path':{'type':'string'},'items':{'type':'array','items':{'type':'object','properties':{'requirement_id':{'type':'string'},'category':{'type':'string','enum':['qualification','rejection','scoring','business','technical','pricing','format','signature']},'source_quote':{'type':'string'},'source_sha256':{'type':'string'},'physical_page':{'type':'integer','minimum':1},'source_location':{'type':'string'},'scope':{'type':'string'}},'required':['requirement_id','category','source_quote','physical_page','source_location','scope'],'additionalProperties':False}},'page_reviews':{'type':'array','items':{'type':'object','properties':{'source_sha256':{'type':'string'},'physical_page':{'type':'integer','minimum':1},'end_physical_page':{'type':'integer','minimum':1},'disposition':{'type':'string','enum':['requirements_extracted','no_actionable_requirements']},'categories':{'type':'array','items':{'type':'string','enum':['qualification','rejection','scoring','business','technical','pricing','format','signature']}},'reason':{'type':'string'}},'required':['physical_page','disposition','categories','reason'],'additionalProperties':False}},'amendment_effects':{'type':'array','items':{'type':'object','properties':{'source_sha256':{'type':'string'},'effect':{'type':'string','enum':['supplements','clarifies','replaces','deletes']},'target_requirement_id':{'type':'string'},'source_quote':{'type':'string'}},'required':['source_sha256','effect','target_requirement_id','source_quote'],'additionalProperties':False}}}),
 ('save_project_tender_analysis_batch','长招标文件按同一原件连续物理页分批保存逐页分析；每批仍须覆盖范围内每页、逐字绑定条目，无要求连续页可用end_physical_page压缩。该工具只保存批次，不授予完整分析。',{'source_extraction_path':{'type':'string'},'source_sha256':{'type':'string'},'start_page':{'type':'integer','minimum':1},'end_page':{'type':'integer','minimum':1},'items':{'type':'array','items':{'type':'object'}},'page_reviews':{'type':'array','items':{'type':'object'}}}),
 ('finalize_project_tender_analysis_batches','按来源顺序汇总已由服务器验证的招标分析批次，再次核对全部原件物理页恰好覆盖一次、全部引文逐字绑定后生成正式tender_analysis。',{'source_extraction_path':{'type':'string'},'batch_paths':{'type':'array','items':{'type':'string'}},'amendment_effects':{'type':'array','items':{'type':'object'}}}),
 ('import_project_tender_analysis_payload','长文件分析先在当前项目sandbox output目录写一个JSON（source_extraction_path、items、page_reviews、amendment_effects），再以真实SHA256调用本工具。中断恢复时可直接使用resume_project返回的pending_tender_analysis_payloads项目相对路径和哈希，不要重读或重抄载荷。服务器按每20页拆成持久批次、逐条核对逐字引文和逐页覆盖，最后合并生成正式tender_analysis。',{'source_extraction_path':{'type':'string'},'payload_path':{'type':'string'},'payload_sha256':{'type':'string'}}),
 ('create_stage_validation','在本阶段其他真实产物全部完成后调用。服务器从当前项目或会话沙箱读取文件、核对精确产物种类并按真实SHA256生成不可伪造的stage_validation；将返回的path加入record_project_stage的artifacts。',{'stage':{'type':'integer','minimum':1,'maximum':10},'artifacts':{'type':'object','additionalProperties':{'type':'string'}}}),
 ('record_project_stage','保存本项目十步阶段的真实产物哈希及缺项；验证文件可读性，下游自动失效；只记待审核，不授予人工批准。先resume_project查看各步required_artifacts。',{'stage':{'type':'integer','minimum':1,'maximum':10},'artifacts':{'type':'object','additionalProperties':{'type':'string'}},'open_items':{'type':'array','items':{'type':'string'}}}),
 ('resume_project','开始或恢复时首先调用：一次返回当前项目、原件哈希、最新文件、检查点和下一步；不搜索旧事件。',{}),
 ('begin_project','以当前会话招标原件建立独立项目；同原件幂等，新原件新项目，不继承历史报价授权。',{'source_path':{'type':'string'},'sha256':{'type':'string'},'project_title':{'type':'string'}}),
 ('add_project_input','向当前项目登记补充原件；系统按真实哈希去重、检查容器可读性并更新唯一输入清单。',{'source_path':{'type':'string'},'sha256':{'type':'string'},'role':{'type':'string'}}),
 ('relate_project_input','记录补遗、澄清、替换、报价清单、技术附件、图纸或模板与已登记项目原件的关系。关系只绑定真实SHA256；具体覆盖条款在第2步解析。',{'sha256':{'type':'string'},'relation_type':{'type':'string'},'related_sha256':{'type':'string'}}),
 ('index_project_technical_source','按物理页索引当前招标技术章节的全部表格行和表外文本块；只建立原文基线，不生成响应或偏离。',{'start_page':{'type':'integer','minimum':1},'end_page':{'type':'integer','minimum':1}}),
 ('index_project_technical_pages','按升序物理页清单索引逐页分析已分类为technical的全部主招标文件页，适用于非连续技术页；不得包含联系方式、响应格式等非技术页。只建立原文基线，不生成响应或偏离。',{'physical_pages':{'type':'array','items':{'type':'integer','minimum':1}}}),
 ('save_project_checkpoint','保存当前阶段、明确下一步及缺项。此操作不能授予检查通过或完整完成。',{'phase':{'type':'string'},'next_action':{'type':'string'},'open_items':{'type':'array','items':{'type':'string'}},'blocked':{'type':'boolean'}}),
 ('capture_project_file','将本项目沙箱内部工作文件按真实哈希保存到analysis/；不覆盖旧版本。网站三项DOCX/DOCX/XLSX最终产物只能通过第9步登记进入outputs/，PDF始终留在analysis。',{'source_path':{'type':'string'},'target':{'type':'string'}}),
 ('stage_project_file','将持久项目原件/产物复制到本会话沙箱的活动项目目录；不覆盖。',{'path':{'type':'string'}}),
 ('audit_project_delivery','检查真实DOCX结构、附件装配和交付状态并保存审计；未核证不可称完整完成。',{}),
 ('start_project_auto_continue','为当前会话和当前真实项目提交持久化自动续跑请求。调度器从最早未通过阶段继续；不会使用固定旧session、项目ID或轮数。',{}),
 ('get_project_auto_continue_status','读取当前项目自动续跑、阶段批次、心跳、重试次数和错误状态。',{}),
]
def catalog():
 return [{'name':n,'description':d,'inputSchema':{'type':'object','properties':p,'required':list(p),'additionalProperties':False},'annotations':{'readOnlyHint':n=='resume_project'}} for n,d,p in SPECS]

def route(state,run,request):
 args=dict(request.get('arguments') or {});scope=verify_scope(args.pop('__deta_scope',{}));name=request['name']
 from bid_agent.tools import Tool,validate_args
 for item in catalog():
  if item['name']==name:validate_args(Tool(name,item['description'],item['inputSchema'],None),args);break
 if name=='record_project_stage':
  b=binding(scope)
  if not b:raise ValueError('Call begin_project first')
  from project_stages import record,REQUIRED
  root=ROOT/b['project']
  if args['stage'] not in REQUIRED:raise ValueError('Stage must be 1..10')
  missing=[kind for kind in REQUIRED[args['stage']] if kind not in args['artifacts']]
  if missing:raise ValueError('Missing stage artifact kinds: '+', '.join(missing))
  # A stage call may refer to files still in this session's sandbox. Capture
  # them in one transaction-sized operation before validating the receipt.
  for path in args['artifacts'].values():
   dest=safe(root,path)
   if not str(PurePosixPath(path)).startswith(('outputs/','analysis/')):raise ValueError('Stage artifact must be in outputs/ or analysis/')
   if dest.is_file() and dest.stat().st_size:continue
   source='/workspace/output/'+b['project']+'/'+path
   data=sandbox_read(scope,source)
   if not data:raise ValueError('Missing/empty artifact: '+path)
   atomic(dest,data)
  result=record(root,**args)
  remaining=[n for n in range(1,11) if n>args['stage']]
  next_stage=remaining[0] if remaining else None
  next_action=('执行第%d步并登记真实产物；对缺失证据标注未决，继续可独立完成的工作'%next_stage) if next_stage else '十步候选产物已登记；等待人工审核和业务确认，不得称可递交'
  progress={'phase':'stage_%d_recorded'%args['stage'],'next_action':next_action,'open_items':args['open_items'],'status':'in_progress' if next_stage else 'review_pending','assertions_verified':False}
  with db() as c:
   c.execute('update bindings set checkpoint=?,updated=? where tenant=? and session=?',(json.dumps(progress,ensure_ascii=False),time.time(),scope['tenant_id'],scope['session_id']))
   row=c.execute('select * from bindings where tenant=? and session=?',(scope['tenant_id'],scope['session_id'])).fetchone()
   atomic(root/'binding.json',json.dumps(dict(row),ensure_ascii=False).encode())
  return {'ok':True,'data':result,'checkpoint':progress}
 if name=='extract_project_source':return {'ok':True,'data':extract_project_source(scope)}
 if name=='build_project_tender_analysis':return {'ok':True,'data':build_tender_analysis(scope,args)}
 if name=='save_project_tender_analysis_batch':return {'ok':True,'data':save_tender_analysis_batch(scope,args)}
 if name=='finalize_project_tender_analysis_batches':return {'ok':True,'data':finalize_tender_analysis_batches(scope,args)}
 if name in {'analyze_tender_payload','import_project_tender_analysis_payload'}:return {'ok':True,'data':import_tender_analysis_payload(scope,args)}
 if name=='create_stage_validation':return {'ok':True,'data':create_stage_validation(scope,args)}
 if name=='resume_project':return {'ok':True,'data':snapshot(scope)}
 if name=='begin_project':return {'ok':True,'data':start(scope,args,state,run)}
 if name=='add_project_input':return {'ok':True,'data':add_input(scope,args)}
 if name=='relate_project_input':return {'ok':True,'data':relate_input(scope,args)}
 if name=='index_project_technical_source':return {'ok':True,'data':index_technical_source(scope,args)}
 if name=='index_project_technical_pages':return {'ok':True,'data':index_technical_source(scope,args)}
 if name=='save_project_checkpoint':return {'ok':True,'data':checkpoint(scope,args)}
 if name=='capture_project_file':return {'ok':True,'data':import_file(scope,args)}
 if name=='stage_project_file':return {'ok':True,'data':export_file(scope,args)}
 if name=='audit_project_delivery':return {'ok':True,'data':audit(scope)}
 if name=='start_project_auto_continue':return {'ok':True,'data':request_continuation(scope)}
 if name=='get_project_auto_continue_status':return {'ok':True,'data':continuation_status(scope)}
 b=binding(scope)
 if not b:raise ValueError('No active project: call resume_project and begin_project. Legacy project access refused.')
 from worker import dispatch
 return dispatch(state,b['run'],{**request,'arguments':args,'scoped':True})

def runtime(req):
 with sqlite3.connect(DB) as c:
  r=c.execute('select tenant_id,user_id,agent_id,deleted_at,agent_config from sessions where id=?',(req['session_id'],)).fetchone()
 if not r or r[3] or (r[2] or json.loads(r[4] or '{}').get('agent_id'))!=AGENT or str(r[0])!=str(req['tenant_id']):raise ValueError('Session scope mismatch')
 scope={'tenant_id':r[0],'session_id':req['session_id'],'user_id':r[1]}
 raw='\n'.join(str(scope[k]) for k in ('tenant_id','session_id','user_id'))
 scope['signature']=hmac.new(os.environ['DETA_PROJECT_SCOPE_KEY'].encode(),raw.encode(),hashlib.sha256).hexdigest()
 verify_scope(scope)
 if req['operation']=='resume':return {'ok':True,'data':snapshot(scope)}
 if req['operation']=='audit':return {'ok':True,'data':audit(scope)}
 raise ValueError('Unknown runtime operation')
