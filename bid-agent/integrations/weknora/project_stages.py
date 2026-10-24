"""Project-neutral ten-stage receipts and deterministic quality gates."""
from __future__ import annotations

import hashlib
import json
import subprocess
import time
from pathlib import Path

REQUIRED={
 1:['input_manifest','stage_validation'],
 2:['source_extraction','technical_source_index','tender_analysis','requirement_atomization','atomization_review','stage_validation'],
 3:['skeleton_bundle','slot_manifest','skeleton_lock','structure_comparison','render_manifest','stage_validation'],
 4:['project_facts','qualification_ledger','filled_docx','evidence_manifest','stage_validation'],
 5:['scoring_ledger','commitments_ledger','commercial_docx','commercial_manifest','stage_validation'],
 6:['configuration_lock','evidence_queue','technical_evidence_ledger','stage_validation'],
 7:['eight_column_docx','eight_column_xlsx','traceability','eight_column_manifest','stage_validation'],
 8:['content_quality_review','cross_document_review','stage_validation'],
 9:['final_bid_docx','eight_column_docx','eight_column_xlsx','internal_render_pdf','output_manifest','stage_validation'],
 10:['page_review','delivery_decision','stage_validation']}


def sha(path):
    with Path(path).open('rb') as stream:return hashlib.file_digest(stream,'sha256').hexdigest()


def load(path):return json.loads(Path(path).read_text())


def project(root):return load(Path(root)/'project.json')


def bound_json(path,root,versions=None):
    data=load(path);meta=project(root)
    if versions and data.get('version') not in versions:raise ValueError('Unsupported artifact version: '+str(data.get('version')))
    if data.get('project_id') not in (None,meta['project_id']):raise ValueError('Artifact belongs to another project')
    source=data.get('source_sha256',data.get('tender_source_sha256'))
    if source not in (None,meta['source_sha256']):raise ValueError('Artifact belongs to another tender source')
    return data


def validate_input_manifest(path,root):
    data=bound_json(path,root,{'source-lock-v2'});items=data.get('inputs',[])
    if not items or data.get('input_count')!=len(items) or data.get('all_readable') is not True:raise ValueError('Input manifest is incomplete')
    from project_gateway import safe
    main=[]
    for item in items:
        f=safe(root,item.get('project_path',''))
        if not f.is_file() or f.stat().st_size!=item.get('bytes') or sha(f)!=item.get('sha256') or item.get('readable') is not True:raise ValueError('Locked input missing, changed, or unreadable')
        if item.get('role')=='招标主文件':main.append(item)
        if 'priority' not in item or not isinstance(item.get('relation'),dict):raise ValueError('Input role priority/relation missing')
    if len(main)!=1 or main[0]['sha256']!=project(root)['source_sha256']:raise ValueError('Exactly one current tender main file is required')
    return {'inputs':len(items),'scanned_pages':sum(int(x.get('scanned_page_count') or 0) for x in items)}

def validate_source_extraction(path,root):
    data=bound_json(path,root,{'project-source-extraction-v1'});manifest=load(Path(root)/'analysis/input-manifest.json')
    expected={x['sha256']:x for x in manifest.get('inputs',[])};documents=data.get('documents',[])
    if {x.get('sha256') for x in documents}!=set(expected):raise ValueError('Source extraction does not cover every locked input')
    total=0
    for document in documents:
        item=expected[document['sha256']]
        if document.get('path')!=item.get('project_path') or document.get('role')!=item.get('role'):raise ValueError('Source extraction input binding changed')
        if Path(document['path']).suffix.lower()=='.pdf':
            pages=document.get('pages',[]);count=int(item.get('page_count') or 0)
            if document.get('status')!='extracted' or [x.get('physical_page') for x in pages]!=list(range(1,count+1)):raise ValueError('Source extraction page coverage is incomplete')
            if any(hashlib.sha256(str(x.get('text','')).encode()).hexdigest()!=x.get('text_sha256') for x in pages):raise ValueError('Extracted page text hash changed')
            total+=len(pages)
    if not total:raise ValueError('Source extraction contains no physical pages')
    return {'documents':len(documents),'physical_pages':total}

def validate_tender_analysis(path,root,source_path):
    data=bound_json(path,root,{'tender-analysis-v1'});source=Path(root)/data.get('source_extraction_path','')
    submitted=Path(source_path)
    if not source.is_file() or sha(source)!=data.get('source_extraction_sha256') or not submitted.is_file():raise ValueError('Tender analysis source extraction is stale')
    # Re-extraction is idempotent but receives a new path and created_at.  Do
    # not invalidate a complete analysis when the extracted documents/pages
    # are byte-for-byte equivalent after removing that volatile timestamp.
    if source.resolve()!=submitted.resolve():
        embedded=load(source);latest=load(submitted)
        embedded.pop('created_at',None);latest.pop('created_at',None)
        if embedded!=latest:raise ValueError('Tender analysis source extraction is stale')
    extraction=load(source);page_map={(d['sha256'],p['physical_page']):p['text'] for d in extraction.get('documents',[]) for p in d.get('pages',[])}
    allowed={'qualification','rejection','scoring','business','technical','pricing','format','signature'};items=data.get('items',[]);ids=[]
    if not items:raise ValueError('Tender analysis is empty')
    for item in items:
        if item.get('category') not in allowed or not item.get('requirement_id') or item['requirement_id'] in ids:raise ValueError('Tender analysis category/ID is invalid')
        ids.append(item['requirement_id']);quote=item.get('source_quote','')
        if not quote or quote not in page_map.get((item.get('source_sha256'),item.get('physical_page')),''):raise ValueError('Tender analysis contains a non-verbatim or unbound quote')
    reviews=data.get('page_reviews',[]);keys=[(x.get('source_sha256'),x.get('physical_page')) for x in reviews]
    if keys!=list(page_map) or data.get('reviewed_page_count')!=len(page_map):raise ValueError('Tender analysis page review coverage is incomplete or out of source order')
    item_pairs={(x['source_sha256'],x['physical_page'],x['category']) for x in items}
    for review in reviews:
        categories=review.get('categories') or [];key=(review.get('source_sha256'),review.get('physical_page'))
        if review.get('disposition')=='requirements_extracted':
            if not categories or any((key[0],key[1],category) not in item_pairs for category in categories):raise ValueError('Tender page review category lacks a same-page item')
        elif review.get('disposition')=='no_actionable_requirements':
            if categories or not str(review.get('reason') or '').strip():raise ValueError('No-actionable tender page lacks a reason')
        else:raise ValueError('Invalid tender page review disposition')
    counts={x:sum(i['category']==x for i in items) for x in sorted(allowed)}
    if data.get('category_counts')!=counts:raise ValueError('Tender analysis category counts are stale')
    technical_pages=sorted({x['physical_page'] for x in items if x['category']=='technical' and x.get('source_sha256')==project(root)['source_sha256']})
    return {'items':len(items),'reviewed_pages':len(reviews),'category_counts':counts,'amendment_effects':len(data.get('amendment_effects',[])),'technical_pages':technical_pages}


def validate_source_index(path,root):
    data=bound_json(path,root,{'technical-source-index-v1'});pages=data.get('pages',[])
    selected=data.get('physical_pages')
    if selected is None:
        start,end=data.get('physical_page_range',[0,0]);selected=list(range(int(start),int(end)+1))
    if not selected or selected!=sorted(set(selected)) or [x.get('physical_page') for x in pages]!=selected:raise ValueError('Technical source page selection is empty, duplicated, or out of order')
    ids=[]
    for p in pages:
        for table in p.get('tables',[]):
            for row in table.get('rows',[]):
                if hashlib.sha256(str(row.get('text','')).encode()).hexdigest()!=row.get('text_sha256'):raise ValueError('Source table row hash changed')
                ids.append(row.get('id'))
        for block in p.get('text_blocks',[]):
            if hashlib.sha256(str(block.get('text','')).encode()).hexdigest()!=block.get('text_sha256'):raise ValueError('Source block hash changed')
            ids.append(block.get('id'))
    if not ids or len(ids)!=len(set(ids)):raise ValueError('Technical source IDs are empty or duplicated')
    return data


def validate_atomization(path,root):
    data=bound_json(path,root,{'deta-requirement-atomization-v2'});index=Path(root)/data.get('source_index_path','')
    if not index.is_file() or sha(index)!=data.get('source_index_sha256'):raise ValueError('Atomization source index changed')
    source=validate_source_index(index,root)
    from bid_agent.requirement_atomization import validate_atomization as strict_validate
    strict_validate(data,source)
    if data.get('validation',{}).get('responses_or_conclusions_generated') is not False:raise ValueError('Atomization crossed into response generation')
    return {'rows':len(data.get('rows',[])),'requirements':sum(x.get('kind')=='requirement' for x in data.get('rows',[]))}


def validate_atomization_review(path,atomization_path,root):
    from openpyxl import load_workbook
    atom=bound_json(atomization_path,root,{'deta-requirement-atomization-v2'});book=load_workbook(path,read_only=True,data_only=False)
    try:
        if '条款拆分与对象归属' not in book.sheetnames:raise ValueError('Atomization review sheet is missing')
        sheet=book['条款拆分与对象归属'];headers=['原文顺序','层级编号','条款类型','适用对象','招标要求原文','来源页码及位置']
        iterator=sheet.iter_rows(values_only=True)
        actual_headers=list(next(iterator,()))
        if actual_headers!=headers or sheet.max_column!=6:raise ValueError('Atomization review must have the exact six columns')
        rows=atom.get('rows',[])
        if sheet.max_row!=len(rows)+1:raise ValueError('Atomization review row count differs from locked atomization')
        actual_rows=list(iterator)
        if len(actual_rows)!=len(rows):raise ValueError('Atomization review row count differs from locked atomization')
        for pos,(row,actual) in enumerate(zip(rows,actual_rows),2):
            values=list(actual[:5])
            expected=[pos-1,row.get('hierarchy_number') or '',row.get('clause_type'),row.get('applicable_object'),row.get('source_text')]
            if [values[0],str(values[1] or ''),*values[2:]]!=[expected[0],str(expected[1]),*expected[2:]]:raise ValueError('Atomization review differs from locked source row '+str(pos))
    finally:book.close()
    return {'rows':len(atom.get('rows',[])),'columns':6}


def validate_skeleton(path,root):
    data=bound_json(path,root,{'dynamic-skeleton-lock-v1'})
    bound={}
    for key in ('bundle','slot_manifest','structure_comparison','render_manifest'):
        item=Path(root)/data.get(key+'_path','')
        if not item.is_file() or sha(item)!=data.get(key+'_sha256'):raise ValueError('Skeleton lock is stale for '+key)
        bound[key]=load(item)
    slots=Path(root)/data['slot_manifest_path'];slot_data=bound['slot_manifest'];items=slot_data.get('slots',[]);ids=[x.get('slot_id') for x in items]
    if len(ids)!=len(set(ids)):raise ValueError('Skeleton slots are duplicated')
    for item in items:
        if not item.get('slot_type') or not isinstance(item.get('allowed_stages'),list):raise ValueError('Slot type/stage policy missing')
    bundle=bound['bundle'];doc=Path(root)/bundle.get('skeleton_docx_path','')
    if not doc.is_file() or sha(doc)!=bundle.get('skeleton_docx_sha256'):raise ValueError('Skeleton DOCX changed after lock')
    render=bound['render_manifest'];pdf=Path(root)/render.get('pdf_path','')
    if not pdf.is_file() or sha(pdf)!=render.get('pdf_sha256') or int(render.get('page_count') or 0)<1:raise ValueError('Skeleton render changed after lock')
    comparison=bound['structure_comparison'];coverage=comparison.get('coverage',{})
    if coverage.get('rate')!=1 or coverage.get('expected')!=coverage.get('mapped'):raise ValueError('Skeleton does not account for every tender-analysis item')
    return {'slots':len(items),'sections':len(bundle.get('sections',[])),'fixed_tables':len(bundle.get('fixed_tables',[])),'render_pages':render['page_count']}


def validate_facts(path,root):
    data=bound_json(path,root,{'project-facts-v1'});facts=data.get('facts',[]);allowed={'source_verified','business_provided','derived_verified','conflict','missing','not_applicable'}
    if not facts or len({x.get('field_code') for x in facts})!=len(facts):raise ValueError('Project facts are empty or duplicated')
    for item in facts:
        if item.get('state') not in allowed:raise ValueError('Invalid project fact state')
        if item['state'] in {'source_verified','business_provided','derived_verified'} and item.get('value') in (None,''):raise ValueError('Usable project fact has no value')
    expected=[x['field_code'] for x in facts if x['state'] in {'source_verified','business_provided','derived_verified'} and x.get('allowed_slots')]
    if data.get('auto_fill_denominator')!=expected:raise ValueError('Known-field denominator is inaccurate')
    expected_slots=[{'field_code':x['field_code'],'slot_id':slot} for x in facts if x['state'] in {'source_verified','business_provided','derived_verified'} for slot in x.get('allowed_slots',[])]
    if data.get('auto_fill_expected_slots')!=expected_slots:raise ValueError('Known-field slot denominator is inaccurate')
    return {'facts':len(facts),'auto_fill_expected':len(expected),'auto_fill_slots':len(expected_slots)}


def validate_qualification(path,root):
    data=bound_json(path,root);rows=data.get('items',[]);allowed={'met','not_met','missing','conflict','dynamic_pending','not_applicable'}
    if not rows or len({x.get('requirement_id') for x in rows})!=len(rows):raise ValueError('Qualification ledger is empty or duplicated')
    for row in rows:
        if row.get('status') not in allowed or not row.get('source_quote'):raise ValueError('Qualification row is incomplete')
        if row['status']=='met' and (not row.get('evidence') or not all(x.get('physical_page') and x.get('source_sha256') and x.get('applicability')=='applicable' for x in row['evidence'])):raise ValueError('Qualification met without page-bound applicable evidence')
    return {'requirements':len(rows)}

def validate_stage4_manifest(path,root,facts_path,qualification_path,filled_docx_path):
    data=bound_json(path,root,{'stage4-evidence-manifest-v1'})
    bindings=[('project_facts',facts_path),('qualification_ledger',qualification_path),('filled_docx',filled_docx_path)]
    for key,rel in bindings:
        file=Path(root)/rel
        if data.get(key+'_path')!=rel or not file.is_file() or data.get(key+'_sha256')!=sha(file):raise ValueError('Stage 4 evidence manifest is stale for '+key)
    if data.get('fill_rate')!=1 or data.get('filled_slots')!=data.get('expected_fill_slots'):raise ValueError('Known trusted fields are not filled into every allowed slot')
    if len(data.get('qualification_evidence_placements',[]))<int(data.get('required_qualification_placements') or 0):raise ValueError('Qualification evidence placement coverage is incomplete')
    return {'filled_slots':data['filled_slots'],'qualification_evidence_placements':len(data.get('qualification_evidence_placements',[]))}


def validate_scoring(path,root):
    data=bound_json(path,root);rows=data.get('items',[])
    if not rows:raise ValueError('Scoring ledger is empty')
    for row in rows:
        kind=row.get('kind')
        if kind not in {'objective','subjective','price','formula'}:raise ValueError('Invalid scoring kind')
        if kind=='subjective' and row.get('awarded_points') not in (None,0):raise ValueError('Subjective points cannot be self-awarded')
        if kind=='price' and row.get('price_approved') is not True and row.get('calculated_points') not in (None,0):raise ValueError('Price points calculated before price approval')
        if kind=='objective' and row.get('claimed_points',0) and not row.get('evidence'):raise ValueError('Objective score claim has no evidence')
    return {'items':len(rows)}

def validate_commitments(path,root):
    data=bound_json(path,root,{'commitments-ledger-v1'});rows=data.get('items',[])
    if not rows or len({x.get('commitment_id') for x in rows})!=len(rows):raise ValueError('Commitments ledger is empty or duplicated')
    if any(not x.get('source_quote') or not x.get('commencement') or not x.get('scope') for x in rows):raise ValueError('Commitment lost its trigger or scope')
    return {'items':len(rows),'business_approved':sum(x.get('business_approved') is True for x in rows)}


def validate_stage5_manifest(path,root,scoring_path,commitments_path,docx_path):
    data=bound_json(path,root,{'stage5-commercial-manifest-v1'})
    for key,rel in [('scoring_ledger',scoring_path),('commitments_ledger',commitments_path),('commercial_docx',docx_path)]:
        file=Path(root)/rel
        if data.get(key+'_path')!=rel or not file.is_file() or data.get(key+'_sha256')!=sha(file):raise ValueError('Stage 5 commercial manifest is stale for '+key)
    if data.get('coverage')!=1 or data.get('placed_items')!=data.get('expected_placements') or len(data.get('placements',[]))!=data.get('placed_items'):raise ValueError('Stage 5 commercial placement coverage is incomplete')
    return {'placements':data['placed_items'],'coverage':1}


def validate_configuration_lock(path,root):
    data=bound_json(path,root,{'configuration-lock-v1'});selection=Path(root)/data.get('selection_path','')
    if not selection.is_file() or sha(selection)!=data.get('selection_sha256'):raise ValueError('Configuration selection source changed')
    for item in data.get('procurement_items',[]):
        if item.get('selected_model') and not item.get('evidence'):raise ValueError('Selected model has no locked evidence set')
        for evidence in item.get('evidence',[]):
            f=Path(root)/evidence['path']
            if not f.is_file() or sha(f)!=evidence['sha256']:raise ValueError('Configuration evidence changed')
    return {'items':len(data.get('procurement_items',[])),'lock_version':data.get('lock_version')}

def validate_evidence_queue(path,root,lock_path):
    data=bound_json(path,root,{'evidence-task-queue-v1'});lock=bound_json(lock_path,root,{'configuration-lock-v1'})
    if data.get('configuration_lock_version')!=lock.get('lock_version'):raise ValueError('Evidence queue uses a stale configuration lock')
    tasks=data.get('tasks',[]);ids=[x.get('requirement_id') for x in tasks]
    if not tasks or len(ids)!=len(set(ids)) or data.get('task_count')!=len(tasks):raise ValueError('Evidence queue is empty, duplicated, or miscounted')
    batches=data.get('batches',[]);visited=[x for b in batches for x in b.get('task_ids',[])]
    if any(len(b.get('task_ids',[]))>20 for b in batches) or visited!=[x.get('task_id') for x in tasks]:raise ValueError('Evidence batches do not cover tasks exactly in order')
    return {'tasks':len(tasks),'batches':len(batches)}

def validate_evidence_ledger(path,root,queue_path):
    data=bound_json(path,root,{'technical-evidence-ledger-v1'});queue=bound_json(queue_path,root,{'evidence-task-queue-v1'})
    if data.get('evidence_queue_path')!=str(Path(queue_path).relative_to(root)) or data.get('evidence_queue_sha256')!=sha(queue_path):raise ValueError('Technical evidence ledger queue binding is stale')
    for batch in data.get('batch_results',[]):
        file=Path(root)/batch.get('path','')
        if not file.is_file() or sha(file)!=batch.get('sha256'):raise ValueError('Technical evidence batch changed after ledger creation')
    expected=[x['requirement_id'] for x in queue['tasks']];rows=data.get('items',[]);actual=[x.get('requirement_id') for x in rows]
    if actual!=expected:raise ValueError('Technical evidence ledger must visit every queued requirement exactly once in order')
    allowed={'正偏离','无偏离','负偏离','不可判定','无适用证据','证据冲突'}
    if any(x.get('machine_conclusion') not in allowed for x in rows):raise ValueError('Invalid machine evidence conclusion')
    if any(x.get('machine_conclusion') in {'正偏离','无偏离','负偏离'} and not x.get('comparison') for x in rows):raise ValueError('Deviation conclusion lacks comparison details')
    if data.get('manual_columns_written') is not False or any(x.get('manual_response_written') is not False or x.get('manual_deviation_written') is not False for x in rows):raise ValueError('Technical evidence ledger crossed into business response/deviation columns')
    return {'items':len(rows),'issues':sum(x.get('machine_conclusion') in {'负偏离','不可判定','无适用证据','证据冲突'} for x in rows)}


def validate_eight(docx_path,xlsx_path,root):
    from docx import Document
    from openpyxl import load_workbook
    headers=['车型','子系统模块','部件模块','功能模块','标准车型答案（历史投标响应）','本次招标要求','投标人响应','偏离说明']
    doc=Document(docx_path)
    if len(doc.tables)!=1 or [x.text for x in doc.tables[0].rows[0].cells]!=headers:raise ValueError('Eight-column DOCX header differs')
    for row_number,row in enumerate(doc.tables[0].rows[1:],2):
        if row.cells[6].text or row.cells[7].text:raise ValueError('DOCX manual response/deviation columns are not blank at row '+str(row_number))
    book=load_workbook(xlsx_path,read_only=False,data_only=False);sheet=book['八列内部核对表']
    try:
        if [sheet.cell(1,c).value for c in range(1,9)]!=headers:raise ValueError('Eight-column XLSX header differs')
        requirements=[]
        for row in range(2,sheet.max_row+1):
            if sheet.cell(row,7).value not in (None,'') or sheet.cell(row,8).value not in (None,''):raise ValueError('Manual response/deviation columns are not blank')
            if sheet.cell(row,6).value not in (None,'') and any(sheet.cell(row,c).value not in (None,'') for c in range(1,6)):
                requirements.append(sheet.cell(row,6).value)
    finally:book.close()
    if not requirements:raise ValueError('Eight-column table contains no tender requirements')
    return {'requirements':len(requirements)}


def validate_eight_bundle(manifest_path,traceability_path,docx_path,xlsx_path,root):
    manifest=bound_json(manifest_path,root,{'deta-eight-column-internal-review-v3'})
    trace=bound_json(traceability_path,root,{'deta-eight-column-row-traceability-v1'})
    if manifest.get('status')!='internal_unreviewed' or manifest.get('formal_bid_output') is not False or manifest.get('business_review_performed') is not False:raise ValueError('Eight-column manifest crossed the business-review boundary')
    rel_docx=str(Path(docx_path).relative_to(root));rel_xlsx=str(Path(xlsx_path).relative_to(root));rel_trace=str(Path(traceability_path).relative_to(root))
    if manifest.get('outputs')!={'docx':rel_docx,'xlsx':rel_xlsx} or manifest.get('sha256')!={'docx':sha(docx_path),'xlsx':sha(xlsx_path)}:raise ValueError('Eight-column manifest output binding is stale')
    if manifest.get('traceability_path')!=rel_trace or manifest.get('traceability_sha256')!=sha(traceability_path):raise ValueError('Eight-column traceability binding is stale')
    if trace.get('outputs')!={'docx':rel_docx,'xlsx':rel_xlsx} or trace.get('status')!='machine_traceability_complete':raise ValueError('Eight-column traceability output binding is invalid')
    bindings=trace.get('bindings',[]);ids=[x.get('requirement_id') for x in bindings]
    if not bindings or len(ids)!=len(set(ids)) or len(bindings)!=manifest.get('summary',{}).get('requirement_rows'):raise ValueError('Eight-column traceability does not bind each requirement exactly once')
    if any(x.get('manual_columns')!={'bidder_response_column':7,'deviation_column':8,'both_blank':True} for x in bindings):raise ValueError('Eight-column traceability crossed manual columns')
    references=[(manifest.get('source_index_path'),manifest.get('source_index_sha256')),(manifest.get('atomization_path'),manifest.get('atomization_sha256')),(manifest.get('history_match_path'),manifest.get('history_match_sha256')),(manifest.get('spec_path'),manifest.get('spec_sha256'))]
    references += [(x.get('path'),x.get('sha256')) for x in trace.get('evidence_batches',[])]
    for rel,digest in references:
        file=Path(root)/str(rel or '')
        if not rel or not file.is_file() or sha(file)!=digest:raise ValueError('Eight-column source binding is missing or stale')
    return {'requirements':len(bindings),'history_found':trace.get('summary',{}).get('history_found',0),'manual_columns_blank':True}


def validate_content_review(path,root):
    data=bound_json(path,root,{'content-quality-review-v2'});gates=data.get('gates',{})
    required={'known_fields','evidence','object_parameter','cross_document','manual_column_boundary'}
    if set(gates)!=required or any(gates[x].get('status') not in {'pass','issues_disclosed'} for x in required):raise ValueError('Content quality gates are incomplete')
    if gates['object_parameter'].get('method') not in {'deterministic','hybrid_rule_model_flags_only'}:raise ValueError('Semantic model cannot independently grant a pass')
    bindings=data.get('source_bindings',{})
    required_bindings={'stage4_manifest','technical_evidence_ledger','eight_column_docx','eight_column_xlsx','cross_document_review'}
    if set(bindings)!=required_bindings:raise ValueError('Content quality review source bindings are incomplete')
    for value in bindings.values():
        file=Path(root)/value.get('path','')
        if not file.is_file() or sha(file)!=value.get('sha256'):raise ValueError('Content quality review source changed after review')
    return {'issues':sum(len(x.get('issues',[])) for x in gates.values())}

def validate_cross_document(path,root):
    data=bound_json(path,root,{'cross-document-review-v1'});checks=data.get('checks',[])
    if not checks:raise ValueError('Cross-document review is empty')
    if any(x.get('status') not in {'consistent','conflict_disclosed','missing_disclosed','not_applicable'} for x in checks):raise ValueError('Cross-document review contains an unsupported conclusion')
    for item in checks:
        if item['status']=='not_applicable':
            if not item.get('reason'):raise ValueError('Not-applicable cross-document check has no reason')
            continue
        left=item.get('left') or {};right=item.get('right') or {};lv=left.get('value');rv=right.get('value')
        expected='missing_disclosed' if lv in (None,'') or rv in (None,'') else ('consistent' if str(lv).strip()==str(rv).strip() else 'conflict_disclosed')
        if item['status']!=expected:raise ValueError('Cross-document status does not match the recorded values')
    return {'checks':len(checks),'conflicts':sum(x.get('status')=='conflict_disclosed' for x in checks)}

def validate_output_manifest(path,root,artifacts):
    data=bound_json(path,root,{'stage9-output-manifest-v1'});expected=[]
    for kind in ('final_bid_docx','eight_column_docx','eight_column_xlsx'):
        rel=artifacts[kind];expected.append({'kind':kind,'path':rel,'sha256':sha(Path(root)/rel)})
    internal=Path(root)/artifacts['internal_render_pdf']
    if data.get('website_outputs')!=expected or data.get('internal_pdf_path')!=artifacts['internal_render_pdf'] or data.get('internal_pdf_sha256')!=sha(internal) or data.get('submission_ready') is not False or data.get('business_reviewed') is not False:
        raise ValueError('Output manifest does not describe the exact three website files and internal PDF')
    return {'website_outputs':3}


def validate_outputs(artifacts,root):
    allowed={Path(artifacts[x]).name for x in ('final_bid_docx','eight_column_docx','eight_column_xlsx')}
    output=Path(root)/'outputs';actual={x.name for x in output.iterdir() if x.is_file()} if output.exists() else set()
    if actual!=allowed or any(x.is_dir() for x in output.iterdir()):raise ValueError('Website outputs must contain exactly the three approved files')
    if any(x.lower().endswith('.pdf') for x in actual):raise ValueError('PDF may not be a website output')
    validate_eight(Path(root)/artifacts['eight_column_docx'],Path(root)/artifacts['eight_column_xlsx'],root)
    return {'outputs':3}


def validate_page_review(path,root):
    data=bound_json(path,root,{'stage10-page-review-v1'});pdf=Path(root)/data.get('pdf_path','')
    if not pdf.is_file() or sha(pdf)!=data.get('pdf_sha256'):raise ValueError('Page review PDF changed')
    pages=data.get('pages',[]);count=data.get('page_count')
    if [x.get('page') for x in pages]!=list(range(1,int(count)+1)):raise ValueError('Visual page coverage is incomplete')
    if any(not x.get('image_sha256') or x.get('status') not in {'pass','issue_resolved','issue_disclosed'} for x in pages):raise ValueError('Page observation missing or invalid')
    for item in pages:
        image=Path(root)/item.get('image_path','')
        if not image.is_file() or sha(image)!=item.get('image_sha256') or not str(item.get('observation') or '').strip():raise ValueError('Visual page observation is stale or missing')
    return {'pages':count}


def validate_stage_validation(path,root,stage,artifacts):
    data=bound_json(path,root);expected={k:sha(Path(root)/v) for k,v in artifacts.items() if k!='stage_validation'}
    if data.get('stage')!=stage or data.get('status')!='passed' or data.get('artifact_hashes')!=expected:raise ValueError('Stage validation is stale or not passed')
    checks=data.get('checks',[])
    if not checks or any(x.get('status')!='pass' for x in checks):raise ValueError('Stage validation checks are incomplete')
    if data.get('awarded_quality_state') not in (None,'in_progress','agent_scope_complete'):raise ValueError('Stage validation attempted to award a prohibited quality state')
    return data


def record(root,stage,artifacts,open_items):
    from project_gateway import safe,atomic
    if type(stage) is not int or stage not in REQUIRED:raise ValueError('Stage must be 1..10')
    if stage>1:
        prior=status(root)[:stage-1]
        invalid=[row['stage'] for row in prior if row.get('status')!='passed']
        if invalid:raise ValueError('Earlier stages must be current and passed before recording this stage: '+','.join(map(str,invalid)))
    missing=[x for x in REQUIRED[stage] if x not in artifacts]
    if missing:raise ValueError('Missing stage artifact kinds: '+', '.join(missing))
    if set(artifacts)-set(REQUIRED[stage]):raise ValueError('Unknown stage artifact kinds: '+', '.join(sorted(set(artifacts)-set(REQUIRED[stage]))))
    public={9:{'final_bid_docx','eight_column_docx','eight_column_xlsx'}}
    receipt=[]
    for kind,rel in artifacts.items():
        expected='outputs/' if kind in public.get(stage,set()) else 'analysis/'
        if not rel.startswith(expected):raise ValueError(kind+' must be stored under '+expected)
        f=safe(root,rel)
        if not f.is_file() or not f.stat().st_size:raise ValueError('Missing/empty artifact: '+rel)
        receipt.append({'kind':kind,'path':rel,'sha256':sha(f),'bytes':f.stat().st_size})
    results={}
    if stage==1:results['input']=validate_input_manifest(safe(root,artifacts['input_manifest']),root)
    if stage==2:
        results['source_extraction']=validate_source_extraction(safe(root,artifacts['source_extraction']),root)
        results['tender_analysis']=validate_tender_analysis(safe(root,artifacts['tender_analysis']),root,safe(root,artifacts['source_extraction']))
        source_data=validate_source_index(safe(root,artifacts['technical_source_index']),root)
        indexed_pages={x['physical_page'] for x in source_data['pages']};required_pages=set(results['tender_analysis']['technical_pages'])
        if not required_pages or indexed_pages!=required_pages:raise ValueError('Technical source index pages must exactly equal the main-tender pages classified as technical')
        results['source']={'physical_page_range':source_data['physical_page_range'],'pages':len(source_data['pages']),'stats':source_data.get('stats',{})}
        results['atomization']=validate_atomization(safe(root,artifacts['requirement_atomization']),root)
        results['atomization_review']=validate_atomization_review(safe(root,artifacts['atomization_review']),safe(root,artifacts['requirement_atomization']),root)
    if stage==3:results['skeleton']=validate_skeleton(safe(root,artifacts['skeleton_lock']),root)
    if stage==4:
        results['facts']=validate_facts(safe(root,artifacts['project_facts']),root)
        results['qualification']=validate_qualification(safe(root,artifacts['qualification_ledger']),root)
        results['filled_evidence']=validate_stage4_manifest(safe(root,artifacts['evidence_manifest']),root,artifacts['project_facts'],artifacts['qualification_ledger'],artifacts['filled_docx'])
    if stage==5:
        results['scoring']=validate_scoring(safe(root,artifacts['scoring_ledger']),root)
        results['commitments']=validate_commitments(safe(root,artifacts['commitments_ledger']),root)
        results['commercial']=validate_stage5_manifest(safe(root,artifacts['commercial_manifest']),root,artifacts['scoring_ledger'],artifacts['commitments_ledger'],artifacts['commercial_docx'])
    if stage==6:
        results['configuration']=validate_configuration_lock(safe(root,artifacts['configuration_lock']),root)
        results['evidence_queue']=validate_evidence_queue(safe(root,artifacts['evidence_queue']),root,safe(root,artifacts['configuration_lock']))
        results['evidence_ledger']=validate_evidence_ledger(safe(root,artifacts['technical_evidence_ledger']),root,safe(root,artifacts['evidence_queue']))
    if stage==7:
        results['eight_column']=validate_eight(safe(root,artifacts['eight_column_docx']),safe(root,artifacts['eight_column_xlsx']),root)
        results['eight_column_bundle']=validate_eight_bundle(safe(root,artifacts['eight_column_manifest']),safe(root,artifacts['traceability']),safe(root,artifacts['eight_column_docx']),safe(root,artifacts['eight_column_xlsx']),root)
    if stage==8:
        results['content']=validate_content_review(safe(root,artifacts['content_quality_review']),root)
        results['cross_document']=validate_cross_document(safe(root,artifacts['cross_document_review']),root)
    if stage==9:
        results['outputs']=validate_outputs(artifacts,root)
        results['manifest']=validate_output_manifest(safe(root,artifacts['output_manifest']),root,artifacts)
    if stage==10:
        results['pages']=validate_page_review(safe(root,artifacts['page_review']),root)
        decision=bound_json(safe(root,artifacts['delivery_decision']),root,{'delivery-decision-v1'})
        if decision.get('status')!='agent_scope_complete' or decision.get('submission_ready') is not False or decision.get('business_reviewed') is not False:raise ValueError('Final decision must stop at agent_scope_complete')
        review=safe(root,artifacts['page_review'])
        if decision.get('page_review_path')!=artifacts['page_review'] or decision.get('page_review_sha256')!=sha(review):raise ValueError('Final decision page review binding is stale')
    validate_stage_validation(safe(root,artifacts['stage_validation']),root,stage,artifacts)
    folder=Path(root)/'stage_receipts';folder.mkdir(exist_ok=True)
    target=folder/f'{stage:02}.json'
    if target.exists():atomic(folder/f'{stage:02}-{int(time.time()*1000)}.history.json',target.read_bytes())
    from workflow_policy import load_policy
    p=load_policy();state='agent_scope_complete' if stage==10 else 'in_progress'
    payload={'stage':stage,'name':p['stages'][stage-1],'status':'passed','quality_state':state,'semantic_verified':False,
             'artifacts':receipt,'validation_results':results,'open_items':open_items,'created':time.time(),
             'policy_version':p['version'],'policy_sha256':p['policy_sha256']}
    atomic(target,json.dumps(payload,ensure_ascii=False,indent=2).encode())
    for later in folder.glob('??.json'):
        if int(later.stem)>stage:
            old=load(later);old['status']='stale';old['reason']='Upstream stage '+str(stage)+' changed';atomic(later,json.dumps(old,ensure_ascii=False,indent=2).encode())
    return payload


def status(root):
    from project_gateway import safe
    from workflow_policy import load_policy
    p=load_policy();result=[]
    for stage in range(1,11):
        file=Path(root)/'stage_receipts'/f'{stage:02}.json'
        if not file.exists():result.append({'stage':stage,'name':p['stages'][stage-1],'status':'pending','required_artifacts':REQUIRED[stage]});continue
        row=load(file);row['name']=p['stages'][stage-1]
        for item in row.get('artifacts',[]):
            artifact=safe(root,item['path'])
            if not artifact.is_file() or sha(artifact)!=item['sha256']:row['status']='stale';row['reason']='Artifact changed or missing';break
        if row.get('policy_sha256')!=p['policy_sha256']:row['status']='stale';row['reason']='Authoritative workflow policy changed'
        result.append(row)
    return result
