"""Source-bound batch matching of verified historical standard answers."""
from __future__ import annotations

import hashlib
import json
import re
import tempfile
import unicodedata
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .library_tools import invoke

MATCH_VERSION = "deta-history-standard-match-v2"

ALIASES = {
    "参考型号": ["防爆锂离子蓄电池无轨胶轮车（5人座指挥车）型号*"],
    "乘坐人数（人）": ["额定乘员", "额定乘员*", "驾驶室准载人数*"],
    "驱动方式": ["驱动型式", "驱动型式*"],
    "最高行驶速度（km/h）": ["最高速度", "最高速度*", "最高车速*"],
    "最小转弯半径（mm）": ["最小转弯半径", "最小转弯通道圆外圆半径", "最大通道圆外半径"],
    "电机额定功率（kW）": ["电机额定功率", "驱动电机额定总功率", "驱动电机总额定功率", "总额定功率"],
    "外形参考尺寸（mm）": ["整车外形尺寸", "整车尺寸", "外形尺寸*"],
    "电源系统总能量（kWh）": ["电源系统总能量", "电池总能量*", "额定能量*"],
    "蓄电池形式": ["蓄电池型式", "电池型式", "电池类型"],
    "载重（kg）": ["额定载质量", "额定载质量*", "额定载重*"],
    "续航里程（km）": ["续航里程"],
    "轮胎数": ["轮胎数量", "轮胎数*"],
    "电池型式": ["蓄电池型式", "电池型式", "电池类型"],
    "电压/容量(V/Ah)": ["单体电池电压容量*", "单体蓄电池*"],
    "单体蓄电池个数及连接方式": ["单体蓄电池个数*", "电池串并联*"],
    "总电压(V)": ["总电压"],
    "电机型式": ["驱动电机类型", "驱动电机型式", "电机型式"],
    "电机控制器功率(kW)": ["电机控制器功率", "控制器额定功率*"],
    "转向器型式": ["转向型式", "转向器型式"],
    "驻车制动型式": ["驻车制动", "驻车制动型式"],
    "应急制动型式": ["应急制动", "应急制动型式"],
    "辅助制动系型式": ["辅助制动", "辅助制动系型式"],
    "灭火装置型式": ["灭火装置型式", "灭火器型式*"],
    "灭火装置数量及位置": ["灭火装置数量及位置", "灭火器数量及位置*"],
}


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _load(path: Path) -> dict:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict): raise ValueError("JSON root must be an object")
    return value


def _norm(value: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(value).upper().replace("（", "(").replace("）", ")"))


def _answer_equivalence_key(value: str) -> str:
    """Compare historical answers without changing the retained source text.

    Historical forms often carry a project-specific leading clause number.  A
    different leading number alone does not create a different product answer;
    any other wording or value difference remains a conflict.
    """
    text = unicodedata.normalize("NFKC", str(value or "")).strip()
    text = re.sub(r"^\s*(?:\d+(?:\.\d+)+|\d+[、．])\s*", "", text)
    return re.sub(r"\s+", "", text)


def _model_for_object(value: str):
    key = _norm(value)
    if "WLR5" in key:
        return {"requested_model": "WLR-5", "query_model": "WLR-5C(E)",
                "scope": "approved_model_family", "rule_id": "WLR5-PLATFORM-001"}
    if "WC5JA" in key:
        return {"requested_model": "WC5J(A)", "query_model": "WC5J(A)",
                "scope": "exact_model", "rule_id": "PUNCTUATION-NORMALIZATION"}
    if "GZC2500" in key:
        return {"requested_model": "GZC2/500", "query_model": "GZC2/500",
                "scope": "exact_model", "rule_id": None}
    return None


def _parameter(row: dict, elements: dict[str, dict]) -> str:
    ids = row["source_ids"]
    if len(ids) == 1 and elements[ids[0]]["kind"] == "table_row":
        cells = elements[ids[0]]["cells"]
        if len(cells) > 1 and str(cells[1] or "").strip(): return str(cells[1]).replace("\n", "").strip()
    text = row["source_text"]
    text = re.sub(r"^\s*(?:[（(]?\d+(?:\.\d+)*[）)、.．]?|[一二三四五六七八九十]+、)\s*", "", text)
    return re.split(r"[：:；;。]", text, maxsplit=1)[0].strip()[:80]


def _elements(index: dict) -> dict[str, dict]:
    out = {}
    for page in index.get("pages", []):
        for block in page.get("text_blocks", []): out[block["id"]] = {"kind": "block", "text": block.get("text", "")}
        for table in page.get("tables", []):
            for row in table.get("rows", []): out[row["id"]] = {"kind": "table_row", "cells": row.get("cells", [])}
    return out


def _batch(queries: list[dict]) -> dict:
    with tempfile.NamedTemporaryFile("w+b") as output:
        invoke({"operation": "model_standard_batch", "library": "history", "queries": queries}, stdout=output)
        output.flush(); output.seek(0); data = output.read()
    if len(data) > 20_000_000: raise ValueError("Historical batch result exceeds safe bound")
    return json.loads(data)


def _classify(query: dict, raw: dict) -> dict:
    candidates = []
    for item in raw.get("results", []):
        # The one approved WLR-5 platform mapping is resolved before querying.
        # Results must then be the exact resolved model.  Expanding the result
        # a second time leaks different configurations such as WLR5B.
        if item.get("match_scope") != "exact_model": continue
        if _norm(item.get("matched_model", "")) != _norm(query.get("query_model", "")): continue
        if item.get("match_type") not in {"exact_parameter", "explicit_alias", "compound_alias"}: continue
        if not item.get("id") or not item.get("source_triplet") or not str(item.get("standard_answer", "")).strip(): continue
        candidates.append(item)
    answers = {}
    for item in candidates: answers.setdefault(_answer_equivalence_key(item["standard_answer"]), []).append(item)
    base = {key: query.get(key) for key in ("requirement_id", "requested_model", "query_model", "match_scope", "family_rule_id", "parameter", "aliases")}
    if not answers:
        return {**base, "status": "not_found", "standard_answer": "", "selected_record_id": None,
                "source_triplet": None, "candidates": [], "reason": "no_applicable_same_model_history"}
    if len(answers) > 1:
        return {**base, "status": "conflict", "standard_answer": "", "selected_record_id": None,
                "source_triplet": None, "candidates": candidates,
                "reason": "multiple_different_authority_answers_require_business_selection"}
    _, records = next(iter(answers.items())); selected = records[0]
    return {**base, "status": "found", "standard_answer": selected["standard_answer"].strip(), "selected_record_id": selected["id"],
            "source_triplet": selected["source_triplet"], "source_ref": selected.get("source_ref"),
            "match_type": selected["match_type"], "candidates": records,
            "reason": "unique_authority_answer" if len(records) == 1 else "equivalent_authority_answers_leading_number_or_whitespace_only"}


def install_history_standard_tools(registry):
    from .tools import STRING, schema

    def build(atomization_path, output_stem):
        if "history" not in registry.store.get(registry.run_id)["scope"].get("libraries", []):
            raise ValueError("Historical authority library is not authorized for this task")
        source = registry.path(atomization_path); atom = _load(source)
        if atom.get("version") != "deta-requirement-atomization-v2" or atom.get("downstream_status") != "approved_for_history_matching":
            raise ValueError("A user-confirmed six-column parsing artifact is required")
        index_path = registry.path(atom.get("source_index_path", "")); index = _load(index_path)
        if atom.get("source_index_sha256") != _sha(index_path): raise ValueError("Parsing artifact is stale")
        elements = _elements(index); queries = []; by_id = {}
        for order, row in enumerate(atom.get("rows", []), 1):
            if row.get("kind") != "requirement": continue
            requirement_id = row["requirement_id"]; model = _model_for_object(row["applicable_object"]); parameter = _parameter(row, elements)
            base = {"requirement_id": requirement_id, "source_order": order, "hierarchy_number": row.get("hierarchy_number", ""),
                    "applicable_object": row["applicable_object"], "source_text": row["source_text"], "parameter": parameter}
            if not model:
                by_id[requirement_id] = {**base, "requested_model": None, "query_model": None, "match_scope": None,
                    "family_rule_id": None, "aliases": [], "status": "not_applicable", "standard_answer": "",
                    "selected_record_id": None, "source_triplet": None, "candidates": [], "reason": "requirement_has_no_single_exact_model"}
                continue
            query = {"requirement_id": requirement_id, "model": model["query_model"], "parameter": parameter,
                     "aliases": ALIASES.get(parameter, []), "limit": 10,
                     "requested_model": model["requested_model"], "query_model": model["query_model"],
                     "match_scope": model["scope"], "family_rule_id": model["rule_id"]}
            queries.append(query); by_id[requirement_id] = base
        raw = _batch([{key: q[key] for key in ("requirement_id", "model", "parameter", "aliases", "limit")} for q in queries]) if queries else {"results": [], "index_sha256": None}
        returned = {x["requirement_id"]: x for x in raw.get("results", [])}
        for query in queries:
            if query["requirement_id"] not in returned: raise ValueError("Historical batch omitted a query")
            by_id[query["requirement_id"]] = {**by_id[query["requirement_id"]], **_classify(query, returned[query["requirement_id"]])}
        items = [by_id[row["requirement_id"]] for row in atom["rows"] if row.get("kind") == "requirement"]
        result = {"version": MATCH_VERSION, "atomization_path": atomization_path, "atomization_sha256": _sha(source),
                  "history_index_sha256": raw.get("index_sha256"), "matching_policy": "approved object-to-model mapping followed by exact returned model; parameter-title prefix matching; leading clause-number/whitespace-only duplicates are equivalent; substantive conflicts remain blank",
                  "items": items, "counts": {status: sum(x["status"] == status for x in items) for status in ("found", "not_found", "conflict", "not_applicable")},
                  "rules": {"historical_answer_is_current_evidence": False, "response_generated": False, "deviation_generated": False,
                            "not_found_answer_blank": True, "conflict_answer_blank": True}}
        if not re.fullmatch(r"[\w\u4e00-\u9fff-]{1,80}", output_stem): raise ValueError("Invalid output_stem")
        target = registry.root / "analysis/step2" / (output_stem + ".json")
        if target.exists(): raise ValueError("Output exists; use a new output_stem")
        target.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        return {"path": str(target.relative_to(registry.root)), "sha256": _sha(target), "counts": result["counts"],
                "next_action": "人工处理conflict；found可进入八列表历史标准答案列，not_found/not_applicable保持空白"}

    def render(match_path, output_stem):
        data = _load(registry.path(match_path))
        if data.get("version") != MATCH_VERSION: raise ValueError("Invalid history match artifact")
        headers = ["原文顺序", "层级编号", "适用对象", "参数名", "招标要求原文", "匹配状态", "历史标准答案", "权威来源"]
        wb = Workbook(); ws = wb.active; ws.title = "历史标准答案匹配"
        thin = Side(style="thin", color="B7B7B7"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for col, value in enumerate(headers, 1):
            c=ws.cell(1,col,value);c.font=Font(name="宋体",size=9,bold=True);c.fill=PatternFill("solid",fgColor="D9EAF7");c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True);c.border=border
        for item in data["items"]:
            values=[item["source_order"],item["hierarchy_number"],item["applicable_object"],item["parameter"],item["source_text"],item["status"],item["standard_answer"],item.get("source_triplet") or ""]
            row=ws.max_row+1
            for col,value in enumerate(values,1):
                c=ws.cell(row,col,value);c.font=Font(name="宋体",size=8);c.alignment=Alignment(vertical="top",wrap_text=True);c.border=border
            if item["status"]=="conflict": ws.cell(row,6).fill=PatternFill("solid",fgColor="F4CCCC")
        for col,width in enumerate([10,16,24,26,70,16,70,48],1):ws.column_dimensions[get_column_letter(col)].width=width
        ws.freeze_panes="E2";ws.auto_filter.ref=f"A1:H{ws.max_row}"
        target=registry.root/"analysis/step2"/(output_stem+".xlsx")
        if target.exists():raise ValueError("Output exists; use a new output_stem")
        wb.save(target);check=load_workbook(target,read_only=True);sheet=check["历史标准答案匹配"]
        if sheet.max_row!=len(data["items"])+1 or sheet.max_column!=8:raise ValueError("Rendered match workbook is incomplete")
        check.close();return {"path":str(target.relative_to(registry.root)),"sha256":_sha(target),"counts":data["counts"]}

    registry.add("build_historical_standard_answer_matches",
                 "读取人工确认的六列条款解析基线，按适用对象和参数名批量查询历史响应权威库。唯一答案直接保留原文和来源；多答案冲突及未找到均保持空白。不查检验报告、不生成响应或偏离。",
                 schema({"atomization_path":STRING,"output_stem":STRING}),build,True)
    registry.add("render_historical_standard_answer_review",
                 "将历史标准答案匹配结果渲染为内部XLSX审查表，不生成投标人响应或偏离。",
                 schema({"match_path":STRING,"output_stem":STRING}),render,True)
