"""Editable chapter assembly and traceable full-page evidence insertion."""
import hashlib
import tempfile
import zipfile
from pathlib import Path


def digest(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def publish(source,target):
    with Path(source).open('rb') as inp, Path(target).open('xb') as out:
        while chunk:=inp.read(1024*1024): out.write(chunk)
    return {'sha256':digest(target),'bytes':Path(target).stat().st_size}


def assemble_docx(master, sources, target, header_policy):
    from docx import Document
    from docxcompose.composer import Composer
    if header_policy!='master':
        raise ValueError('Only explicit master header/footer policy supported; original headers require a different operation')
    if not sources: raise ValueError('At least one chapter required')
    paths=[master,*sources]
    hashes=[digest(p) for p in paths]
    docs=[Document(p) for p in paths]
    body_text=lambda d: ''.join(n.text or '' for n in d.element.body.iter() if n.tag.endswith('}t'))
    originals=[body_text(d) for d in docs]
    expected_images=set()
    for path in paths:
        with zipfile.ZipFile(path) as archive:
            expected_images.update(hashlib.sha256(archive.read(n)).hexdigest() for n in archive.namelist() if n.startswith('word/media/'))
    composer=Composer(docs[0],preserve_styles=True)
    manifest=[{'source_sha256':hashes[0],'role':'master','order':0}]
    for index,doc in enumerate(docs[1:],1):
        composer.doc.add_page_break()
        composer.append(doc)
        manifest.append({'source_sha256':hashes[index],'role':'chapter','order':index})
    with tempfile.TemporaryDirectory() as tmp:
        file=Path(tmp)/'assembled.docx';composer.save(file)
        check=Document(file);actual=body_text(check)
        # Sequential exact containment; cannot silently lose source body text.
        position=0
        for original in originals:
            found=actual.find(original,position)
            if found<0: raise ValueError('Source body text not retained in order')
            position=found+len(original)
        with zipfile.ZipFile(file) as archive:
            actual_images={hashlib.sha256(archive.read(n)).hexdigest() for n in archive.namelist() if n.startswith('word/media/')}
        if not expected_images.issubset(actual_images): raise ValueError('Source images lost during assembly')
        if hashes!=[digest(p) for p in paths]: raise ValueError('Source changed during assembly')
        return {**publish(file,target),'sources':manifest,'header_policy':'master','body_text_retained_in_order':True,'source_media_hashes_retained':True,
                'limitations':['Header/footer inherited from master by explicit policy','Images, numbering, sections and page layout require rendered verification']}


def append_pdf_evidence(source, source_sha256, pdf, pages, target, pdf_sha256):
    import pymupdf
    from docx import Document
    from docx.enum.section import WD_SECTION_START
    from docx.shared import Inches
    if digest(source)!=source_sha256: raise ValueError('Source DOCX version changed')
    if not pages or len(pages)>100 or len(set(pages))!=len(pages): raise ValueError('Provide 1–100 distinct page numbers')
    pdf_hash=digest(pdf)
    if pdf_hash != pdf_sha256: raise ValueError(f"Evidence PDF version/source mismatch; no output created. Supplied pdf_sha256={pdf_sha256}; actual SHA256 of selected PDF={pdf_hash}. Check the selected PDF path against the task/input manifest. If this is the intended file and manifest agrees, retry with its exact SHA256; otherwise select the intended PDF. No separate evidence-source-record file is required.")
    doc=Document(source);manifest=[]
    with tempfile.TemporaryDirectory() as tmp:
        with pymupdf.open(pdf) as original:
            for page_number in pages:
                if type(page_number) is not int or not 1<=page_number<=len(original): raise ValueError('Invalid PDF page number')
                page=original[page_number-1]
                # Full source page rendered; no crop, redaction, or stamp synthesis.
                pix=page.get_pixmap(matrix=pymupdf.Matrix(2,2),alpha=False)
                image=Path(tmp)/f'page-{page_number}.png';pix.save(image)
                section=doc.add_section(WD_SECTION_START.NEW_PAGE)
                section.page_width=Inches(page.rect.width/72)
                section.page_height=Inches(page.rect.height/72)
                section.top_margin=section.bottom_margin=Inches(.2)
                section.left_margin=section.right_margin=Inches(.2)
                # Empty inherited headers still reserve space in LibreOffice.
                # Reset their distance so full-page evidence stays on the page.
                section.header_distance=section.footer_distance=Inches(0)
                section.different_first_page_header_footer=False
                for header in (section.header,section.footer,section.first_page_header,section.first_page_footer,section.even_page_header,section.even_page_footer):
                    header.is_linked_to_previous=False
                    for paragraph in header.paragraphs: paragraph.clear()
                width=min(float(section.page_width-section.left_margin-section.right_margin)/914400,
                          (float(section.page_height-section.top_margin-section.bottom_margin)/914400-.15)*pix.width/pix.height)
                paragraph=doc.add_paragraph();paragraph.paragraph_format.space_after=0
                paragraph.paragraph_format.space_before=0;paragraph.paragraph_format.line_spacing=1
                run=paragraph.add_run();run.add_picture(str(image),width=Inches(width))
                manifest.append({'pdf_sha256':pdf_hash,'page':page_number,'image_sha256':digest(image),'pixels':[pix.width,pix.height]})
        file=Path(tmp)/'with-evidence.docx';doc.save(file)
        with zipfile.ZipFile(file) as archive:
            images={hashlib.sha256(archive.read(n)).hexdigest() for n in archive.namelist() if n.startswith('word/media/')}
        if any(item['image_sha256'] not in images for item in manifest): raise ValueError('Evidence image missing from output package')
        if digest(source)!=source_sha256 or digest(pdf)!=pdf_hash: raise ValueError('Source changed during evidence insertion')
        return {**publish(file,target),'source_sha256':source_sha256,'evidence_pages':manifest,
                'limitations':['Evidence pages are full-page images, not editable source text','Source signatures and markings are preserved only as historical source content','Must render and visually inspect all inserted pages']}
