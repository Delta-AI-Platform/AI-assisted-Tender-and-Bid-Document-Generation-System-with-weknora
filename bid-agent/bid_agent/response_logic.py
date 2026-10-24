"""Evidence-aware consistency diagnostics. Never upgrades model claims to approval."""
import hashlib,json,re,unicodedata
from pathlib import Path
from docx import Document
from pypdf import PdfReader

def norm(s):return re.sub(r'\s+','',unicodedata.normalize('NFKC',s))
def text_of(path):
 path=Path(path)
 if path.suffix.lower()=='.pdf':return '\n'.join(p.extract_text() or '' for p in PdfReader(path).pages)
 if path.suffix.lower()=='.docx':
  d=Document(path);return '\n'.join([p.text for p in d.paragraphs]+[c.text for t in d.tables for r in t.rows for c in r.cells])
 return path.read_text()

def check_logic(tender,docx):
 source=text_of(tender);body=text_of(docx);a,b=norm(source),norm(body);findings=[]
 def add(code,detail):findings.append({'code':code,'detail':detail,'status':'needs_review'})
 d=Document(docx)
 unsupported=sum(1 for t in d.tables for r in t.rows if r.cells and r.cells[-1].text.strip() in {'满足','完全满足','响应','完全响应'})
 if unsupported:add('bare_response',f'{unsupported}条正式详细响应被填为满足/响应；当前规则要求正式填写区留空，实际参数与拟应答应进入内部对标XLSX')
 if ('无偏离' in b or '完全响应' in b) and unsupported:add('unsupported_deviation','总括无偏离声明与未核证的概括响应并存')
 if '随附' in b and not d.inline_shapes:add('attachment_claim','正文声称随附证明，但未检测到内嵌证明图片；须检查实际附件，不得把清单当原件')
 # Compare explicit service timing in this tender; flag absence, not automatic rejection.
 for hours in re.findall(r'(\d+)小时内(?:作出)?响应',a):
  if not re.search(re.escape(hours)+r'小时内(?:作出)?响应',b):add('response_deadline','招标明确'+hours+'小时内响应，正文未找到对应承诺')
 if '尾气净化装置' in a and '尾气净化装置' not in b:add('missing_technical_section','招标含尾气净化装置要求，正文无对应专项响应')
 if '供货通知单' in a and '合同签订后' in b and ('日历天内' in b or '供货进度' in b):add('delivery_anchor','供货通知单起算与合同签订起算并存，需核对统一')
 if '参考型号' in a:
  references=set(re.findall(r'\b(?:WLR|WC|GZC)[A-Za-z0-9()/（）-]+',source))
  used=[x for x in references if norm(x) in b]
  if used:add('reference_model','出现招标参考型号'+','.join(sorted(used))+ '；须核实是否确为本企业拟投型号，不能直接当事实')
 if '备品备件明细' in a and not any(len(t.rows)>25 for t in d.tables):add('spares_detail','招标有备件明细，未见展开的长明细表；检查物料总数量和分项报价附表')
 if '非递交' not in b:add('draft_label','未核证草稿正文无非递交标识')
 return {'findings':findings,'quality_status':'not_verified','scope':'文本逻辑诊断；不替代原件适用性、法律判断或逐页视觉审核'}

def validate_matrix(registry,path):
 source=registry.path(path);data=json.loads(source.read_text());rows=data.get('requirements')
 if not isinstance(rows,list) or not rows:raise ValueError('Non-empty requirements array required')
 ids=set();errors=[];pending=[]
 for i,r in enumerate(rows):
  ident=r.get('id',str(i));label=str(ident)
  if ident in ids:errors.append(label+': duplicate requirement id')
  ids.add(ident)
  for key in ['requirement','requirement_source','response','response_kind','evidence','deviation']:
   if key not in r:errors.append(label+': missing '+key)
  kind=r.get('response_kind')
  if kind not in ['verified_fact','proposed_commitment','pending']:errors.append(label+': invalid response_kind')
  def anchor(ref):
   p=registry.path(ref['path']);digest=hashlib.sha256(p.read_bytes()).hexdigest()
   if digest!=ref['sha256']:raise ValueError('source hash changed')
   quote=norm(ref.get('quote',''))
   if len(quote)<4 or quote not in norm(text_of(p)):raise ValueError('source quotation not found')
  try:anchor(r['requirement_source'])
  except Exception as e:errors.append(label+': requirement source '+str(e))
  evidence=r.get('evidence') or []
  if kind=='verified_fact' and not evidence:errors.append(label+': fact has no evidence')
  for ref in evidence:
   try:anchor(ref)
   except Exception as e:errors.append(label+': evidence '+str(e))
  if kind!='verified_fact' and r.get('deviation') in ['无偏离','正偏离','none','positive']:
   if not (data.get('internal_review_only') is True and r.get('technical_confirmation')=='pending' and r.get('report_actual') is not None and r.get('proposed_configuration') is not None):errors.append(label+': proposed deviation needs internal-only label, separate actual/proposed configuration and pending technical confirmation')
  # Quote anchoring alone cannot prove applicability. Preserve human/model review distinction.
  pending.append(label+': subject/model/validity and requirement applicability must be reviewed')
 return {'rows':len(rows),'errors':errors,'pending_reviews':pending,'quality_status':'not_verified','source_sha256':hashlib.sha256(source.read_bytes()).hexdigest()}

def install_logic_tools(registry):
 from .tools import schema,STRING
 registry.add('check_response_logic','核对本项目招标与DOCX中承诺、参考型号、响应时间、附件声明和技术漏项；不能授予合格。',schema({'tender':STRING,'docx':STRING}),lambda tender,docx:check_logic(registry.path(tender),registry.path(docx)))
 registry.add('check_requirement_matrix','验证条款矩阵来源哈希、引文锚定及事实/拟承诺/待核与偏离的一致性。JSON requirements数组，每条id,requirement,requirement_source(path,sha256,quote),response,response_kind(verified_fact/proposed_commitment/pending),evidence(同样来源对象数组),deviation。引文存在不代表适用性通过。',schema({'path':STRING}),lambda path:validate_matrix(registry,path))
