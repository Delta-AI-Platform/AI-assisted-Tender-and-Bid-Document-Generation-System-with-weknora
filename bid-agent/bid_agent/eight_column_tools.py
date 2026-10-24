"""Deterministic renderer and validator for the internal eight-column review table."""
from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADERS = ["车型", "子系统模块", "部件模块", "功能模块", "标准车型答案（历史投标响应）", "本次招标要求", "投标人响应", "偏离说明"]
SPEC_VERSION = "deta-eight-column-internal-review-v3"
TRACE_VERSION = "deta-eight-column-row-traceability-v1"


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _load(path: Path) -> dict:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError("JSON root must be an object")
    return value


def _source_elements(index: dict) -> dict[str, dict]:
    result = {}
    for page in index.get("pages", []):
        for block in page.get("text_blocks", []):
            result[block["id"]] = {"kind": "block", "values": [str(block.get("text", ""))]}
        for table in page.get("tables", []):
            for row in table.get("rows", []):
                values = [str(x) for x in row.get("cells", []) if x not in (None, "")]
                result[row["id"]] = {"kind": "table_row", "values": values + ([" ".join(values)] if values else [])}
    return result


def _history_records(history: dict) -> dict[str, dict]:
    records = {}
    for item in history.get("items", []):
        if item.get("status") != "found":
            continue
        for result in item.get("results", []):
            rid = str(result.get("id", ""))
            if rid:
                records[rid] = {"requirement_id": item.get("requirement_id"), "query_model": item.get("model"), **result}
        selected = str(item.get("selected_record_id") or "")
        if selected:
            candidate = next((x for x in item.get("candidates", []) if str(x.get("id")) == selected), {})
            records[selected] = {
                "requirement_id": item.get("requirement_id"),
                "query_model": item.get("query_model") or item.get("requested_model"),
                "id": selected,
                "matched_model": candidate.get("matched_model") or item.get("query_model"),
                "match_scope": item.get("match_scope") or candidate.get("match_scope"),
                "match_type": item.get("match_type") or candidate.get("match_type"),
                "standard_answer": item.get("standard_answer"),
                "source_triplet": item.get("source_triplet") or candidate.get("source_triplet"),
                "source_ref": item.get("source_ref") or candidate.get("source_ref"),
            }
    return records


def _history_items(history: dict) -> dict[str, dict]:
    result = {}
    for item in history.get("items", []):
        rid = item.get("requirement_id")
        if not isinstance(rid, str) or not rid or rid in result:
            raise ValueError("history_match_path must contain unique non-empty requirement_id values")
        status = item.get("status")
        if status not in {"found", "not_found", "not_applicable", "conflict"}:
            raise ValueError(f"history item {rid}: unsupported status")
        reason = str(item.get("reason") or "").strip()
        if status != "found" and not reason:
            raise ValueError(f"history item {rid}: explicit reason required for {status}")
        result[rid] = item
    return result


def _evidence_items(batch_documents: list[tuple[str, dict]], requirement_ids: set[str]) -> dict[str, dict]:
    bound = {}
    for path, batch in batch_documents:
        if batch.get("all_items_visited") is not True:
            raise ValueError(f"evidence batch {path}: all_items_visited must be true")
        if batch.get("manual_response_or_deviation_cells_written") is not False:
            raise ValueError(f"evidence batch {path}: manual columns must remain untouched")
        results = batch.get("results")
        if not isinstance(results, list) or not results:
            raise ValueError(f"evidence batch {path}: non-empty results required")
        for item in results:
            rid = item.get("task_id")
            if not isinstance(rid, str) or not rid or rid in bound:
                raise ValueError(f"evidence batch {path}: duplicate or missing task_id")
            if rid not in requirement_ids:
                raise ValueError(f"evidence batch {path}: unknown requirement_id {rid}")
            item_status = item.get("item_status")
            if item_status == "item_error":
                reason = str(item.get("error") or "").strip()
                if not reason: raise ValueError(f"evidence item {rid}: item_error requires error text")
                binding = {"status": "item_error", "reason": reason, "source": None, "mathematical_relation": None}
            elif item_status == "processed":
                result = item.get("result") or {}
                if result.get("requirement_id") not in (None, rid):
                    raise ValueError(f"evidence item {rid}: result requirement_id mismatch")
                status = result.get("evaluation_status")
                if status not in {"calculated", "no_applicable_evidence", "blocked_by_applicability", "blocked_by_tender_anchor"}:
                    raise ValueError(f"evidence item {rid}: unsupported evaluation_status")
                applicability = result.get("applicability") or {}
                source = None
                if applicability.get("source_id"):
                    source = {k: applicability.get(k) for k in ("library", "source_id", "source_sha256", "physical_page", "applicability_status")}
                    original = applicability.get("original_page_verification") or {}
                    source["original_page_verification_status"] = original.get("verification_status")
                    source["page_text_sha256"] = original.get("page_text_sha256")
                    source["approved_mappings_used"] = applicability.get("approved_mappings_used", [])
                if status == "calculated":
                    if (not source or source.get("library") not in {"enterprise_facts", "inspection_reports"}
                            or not isinstance(source.get("source_id"), str)
                            or not re.fullmatch(r"[0-9a-f]{64}", str(source.get("source_sha256") or ""))
                            or not isinstance(source.get("physical_page"), int) or source["physical_page"] < 1
                            or source.get("original_page_verification_status") != "verified_original_page"
                            or source.get("applicability_status") != "applicable"):
                        raise ValueError(f"evidence item {rid}: calculated relation requires a verified applicable original page binding")
                reason = str(result.get("gap_reason") or "").strip()
                if status == "no_applicable_evidence" and not reason:
                    raise ValueError(f"evidence item {rid}: no_applicable_evidence requires gap_reason")
                deviation = result.get("deviation_conclusion")
                if deviation is not None and (deviation not in {"无偏离", "负偏离"}
                                               or result.get("deviation_conclusion_allowed") is not True
                                               or result.get("deviation_decision_status") != "concluded"
                                               or not isinstance(result.get("locked_configuration"), dict)):
                    raise ValueError(f"evidence item {rid}: deviation conclusion lacks a valid locked-configuration decision")
                binding = {"status": status, "reason": reason or None, "source": source,
                           "mathematical_relation": result.get("mathematical_relation"),
                           "deviation_conclusion": deviation,
                           "deviation_decision_status": result.get("deviation_decision_status"),
                           "locked_configuration": result.get("locked_configuration")}
            else:
                raise ValueError(f"evidence item {rid}: item_status must be processed or item_error")
            bound[rid] = {"batch_path": path, "position": item.get("position"), **binding}
    missing = sorted(requirement_ids - set(bound))
    if missing:
        raise ValueError("evidence batches do not cover every requirement_id: " + ", ".join(missing[:10]))
    return bound


def _validate_and_expand(spec: dict, index: dict, history: dict | None) -> list[dict]:
    if spec.get("version") != SPEC_VERSION:
        raise ValueError(f"spec.version must be {SPEC_VERSION}")
    rows = spec.get("rows")
    if not isinstance(rows, list) or not rows:
        raise ValueError("spec.rows must be a non-empty array")
    elements = _source_elements(index)
    history_records = _history_records(history or {})
    expanded = []
    seen_requirement_ids = set()
    for position, raw in enumerate(rows, 1):
        if not isinstance(raw, dict) or raw.get("kind") not in {"heading", "requirement"}:
            raise ValueError(f"row {position}: kind must be heading or requirement")
        if raw["kind"] == "heading":
            text = raw.get("text")
            if not isinstance(text, str) or not text.strip():
                raise ValueError(f"row {position}: heading text required")
            expanded.append({"kind": "heading", "text": text, "source_ids": raw.get("source_ids", [])})
            continue
        classification = raw.get("classification")
        if not isinstance(classification, list) or len(classification) != 4 or not all(isinstance(x, str) and x.strip() for x in classification):
            raise ValueError(f"row {position}: four non-empty hierarchy labels required")
        requirement_id = raw.get("requirement_id")
        if not isinstance(requirement_id, str) or not requirement_id or requirement_id in seen_requirement_ids:
            raise ValueError(f"row {position}: unique requirement_id required")
        seen_requirement_ids.add(requirement_id)
        sources = raw.get("sources")
        if not isinstance(sources, list) or not sources:
            raise ValueError(f"row {position}: at least one exact source quote required")
        quotes = []
        source_ids = []
        for source in sources:
            if not isinstance(source, dict) or not isinstance(source.get("id"), str) or not isinstance(source.get("quote"), str) or not source["quote"]:
                raise ValueError(f"row {position}: invalid source quote")
            element = elements.get(source["id"])
            if element is None or not any(source["quote"] in value for value in element["values"]):
                raise ValueError(f"row {position}: quote is not present in locked source element {source['id']}")
            quotes.append(source["quote"]); source_ids.append(source["id"])
        joiner = raw.get("source_joiner", "")
        if not isinstance(joiner, str):
            raise ValueError(f"row {position}: source_joiner must be text")
        source_text = joiner.join(quotes)
        removed = raw.get("removed_source_prefix", "")
        if removed:
            if not isinstance(removed, str) or not source_text.startswith(removed):
                raise ValueError(f"row {position}: removed_source_prefix does not match source")
            source_text = source_text[len(removed):]
        number = raw.get("display_number", "")
        if not isinstance(number, str):
            raise ValueError(f"row {position}: display_number must be text")
        expected = ((number + " ") if number else "") + source_text
        if raw.get("tender_requirement") != expected:
            raise ValueError(f"row {position}: tender requirement differs from exact source quotes or numbering")
        record_id = raw.get("history_record_id")
        answer = ""
        history_source = None
        if record_id:
            record = history_records.get(str(record_id))
            if record is None:
                raise ValueError(f"row {position}: history record is not a found authority result")
            if record.get("requirement_id") != requirement_id:
                raise ValueError(f"row {position}: history result belongs to another requirement")
            if record.get("match_scope") not in {"exact_model", "approved_model_family"}:
                raise ValueError(f"row {position}: history result is not exact-model or approved-family")
            answer = record.get("standard_answer")
            if not isinstance(answer, str) or not answer.strip():
                raise ValueError(f"row {position}: history result has no standard answer")
            history_source = {k: record.get(k) for k in ("id", "matched_model", "match_scope", "match_type", "source_triplet", "source_ref")}
        if raw.get("bidder_response") not in (None, "") or raw.get("deviation") not in (None, ""):
            raise ValueError(f"row {position}: bidder response and deviation must remain blank before manual review")
        expanded.append({"kind": "requirement", "requirement_id": requirement_id,
                         "classification": classification, "standard_answer": answer,
                         "tender_requirement": expected, "source_ids": source_ids,
                         "history_source": history_source})
    return expanded


def _validate_atomization_binding(spec: dict, atomization: dict, expanded: list[dict], index_path: Path) -> None:
    if atomization.get("version") != "deta-requirement-atomization-v2" or atomization.get("schema_role") != "requirement_parsing_and_object_attribution":
        raise ValueError("atomization_path is not an approved requirement atomization artifact")
    if atomization.get("source_index_sha256") != _sha(index_path):
        raise ValueError("atomization artifact is not bound to the current locked source index")
    if atomization.get("downstream_status") != "approved_for_eight_column":
        raise ValueError("atomization schema/object attribution has not been approved for eight-column use")
    atomic_rows = atomization.get("rows")
    if not isinstance(atomic_rows, list) or len(atomic_rows) != len(expanded):
        raise ValueError("eight-column rows differ from the approved atomization row count")
    for position, (atomic, row) in enumerate(zip(atomic_rows, expanded), 1):
        if atomic.get("kind") != row.get("kind"):
            raise ValueError(f"row {position}: row kind differs from approved atomization")
        if row["kind"] == "heading":
            if atomic.get("text", atomic.get("source_text")) != row.get("text") or atomic.get("source_ids") != row.get("source_ids"):
                raise ValueError(f"row {position}: heading differs from approved atomization")
            continue
        checks = ("requirement_id", "source_ids")
        if any(atomic.get(key) != row.get(key) for key in checks):
            raise ValueError(f"row {position}: requirement or object attribution differs from approved atomization")


def _build_traceability(rows: list[dict], history: dict, evidence: dict[str, dict]) -> dict:
    history_by_requirement = _history_items(history)
    requirement_ids = {row["requirement_id"] for row in rows if row["kind"] == "requirement"}
    missing_history = sorted(requirement_ids - set(history_by_requirement))
    extra_history = sorted(set(history_by_requirement) - requirement_ids)
    if missing_history or extra_history:
        raise ValueError(f"history_match_path requirement_id coverage mismatch; missing={missing_history[:10]}, extra={extra_history[:10]}")
    bindings = []
    for output_position, row in enumerate(rows, 1):
        if row["kind"] != "requirement":
            continue
        rid = row["requirement_id"]
        history_item = history_by_requirement[rid]
        history_status = history_item["status"]
        history_binding = {
            "status": history_status,
            "record_id": row.get("history_source", {}).get("id") if row.get("history_source") else None,
            "source_triplet": row.get("history_source", {}).get("source_triplet") if row.get("history_source") else None,
            "source_ref": row.get("history_source", {}).get("source_ref") if row.get("history_source") else None,
            "reason": history_item.get("reason"),
        }
        if history_status == "found" and not history_binding["record_id"]:
            raise ValueError(f"row {output_position}: found history was not selected into column 5")
        if history_status != "found" and row.get("standard_answer"):
            raise ValueError(f"row {output_position}: non-found history cannot populate column 5")
        bindings.append({
            "output_position": output_position,
            "docx_table_row": output_position + 1,
            "xlsx_row": output_position + 1,
            "requirement_id": rid,
            "source_ids": row["source_ids"],
            "tender_requirement_sha256": hashlib.sha256(row["tender_requirement"].encode("utf-8")).hexdigest(),
            "history": history_binding,
            "evidence": evidence[rid],
            "manual_columns": {"bidder_response_column": 7, "deviation_column": 8, "both_blank": True},
        })
    return {"version": TRACE_VERSION, "status": "machine_traceability_complete", "bindings": bindings,
            "summary": {"requirement_rows": len(bindings), "history_found": sum(x["history"]["status"] == "found" for x in bindings),
                        "evidence_calculated": sum(x["evidence"]["status"] == "calculated" for x in bindings),
                        "evidence_gaps": sum(x["evidence"]["status"] == "no_applicable_evidence" for x in bindings),
                        "evidence_blocked": sum(str(x["evidence"]["status"]).startswith("blocked_") for x in bindings),
                        "evidence_item_errors": sum(x["evidence"]["status"] == "item_error" for x in bindings),
                        "deviation_conclusions": sum(bool(x["evidence"].get("deviation_conclusion")) for x in bindings),
                        "negative_deviations": sum(x["evidence"].get("deviation_conclusion") == "负偏离" for x in bindings)},
            "rules": {"one_binding_per_requirement": True, "output_rows_stable": True,
                      "manual_response_and_deviation_untouched": True, "traceability_is_internal_only": True}}


def _merge_runs(rows: list[dict]) -> list[tuple[int, int, int, str]]:
    runs = []
    for col in range(4):
        start = None; path = None; value = None
        for pos, row in enumerate(rows, 1):
            current = tuple(row["classification"][:col + 1]) if row["kind"] == "requirement" else None
            if current and current == path:
                continue
            if start is not None and pos - 1 > start:
                runs.append((col, start, pos - 1, value))
            if current:
                start, path, value = pos, current, current[-1]
            else:
                start = path = value = None
        if start is not None and len(rows) > start:
            runs.append((col, start, len(rows), value))
    return runs


def _render(rows: list[dict], docx_path: Path, xlsx_path: Path):
    doc = Document(); sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width, sec.page_height = Cm(42), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.2); sec.left_margin = sec.right_margin = Cm(1.2)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("完整招标要求八列内部核对表（未核对）"); run.bold = True; run.font.name = "宋体"; run.font.size = Pt(15)
    note = doc.add_paragraph("说明：本表仅供业务人工核对。第5列只引用权威库中同车型或已批准同车型族的历史原文；无适用历史时保持空白。第7、8列尚未填写，不得作为递交稿或直接装入投标文件。")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in note.runs: r.font.name = "宋体"; r.font.size = Pt(9); r.font.color.rgb = RGBColor(192, 0, 0)
    table = doc.add_table(rows=1, cols=8); table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (cell, text) in enumerate(zip(table.rows[0].cells, HEADERS)):
        cell.text = text; cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in cell.paragraphs[0].runs: r.bold = True; r.font.name = "宋体"; r.font.size = Pt(8)
        shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "D9EAF7"); cell._tc.get_or_add_tcPr().append(shd)
    for row in rows:
        cells = table.add_row().cells
        if row["kind"] == "heading":
            # Headings are part of the verbatim tender-requirement stream, so
            # keep them in column 6.  Never merge through columns 7/8: Word
            # exposes a horizontally merged title as content in those cells,
            # which violates the manual-column boundary even when it looks
            # visually like a single heading row.
            cells[5].text = row["text"]
            for cell in cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "EDEDED"); cell._tc.get_or_add_tcPr().append(shd)
            for r in cells[5].paragraphs[0].runs: r.bold = True; r.font.name = "宋体"; r.font.size = Pt(9)
        else:
            values = row["classification"] + [row["standard_answer"], row["tender_requirement"], "", ""]
            for cell, value in zip(cells, values):
                cell.text = value; cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for para in cell.paragraphs:
                    for r in para.runs: r.font.name = "宋体"; r.font.size = Pt(8)
    merges = _merge_runs(rows)
    for col, start, end, value in merges:
        cell = table.cell(start, col).merge(table.cell(end, col)); cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in cell.paragraphs[0].runs: r.font.name = "宋体"; r.font.size = Pt(8)
    repeat = OxmlElement("w:tblHeader"); repeat.set(qn("w:val"), "true"); table.rows[0]._tr.get_or_add_trPr().append(repeat)
    doc.save(docx_path)

    wb = Workbook(); ws = wb.active; ws.title = "八列内部核对表"
    thin = Side(style="thin", color="000000"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col, value in enumerate(HEADERS, 1):
        cell = ws.cell(1, col, value); cell.font = Font(name="宋体", size=9, bold=True); cell.fill = PatternFill("solid", fgColor="D9EAF7"); cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = border
    for r_idx, row in enumerate(rows, 2):
        if row["kind"] == "heading":
            cell = ws.cell(r_idx, 6, row["text"]); cell.font = Font(name="宋体", size=9, bold=True)
            for col in range(1, 9):
                ws.cell(r_idx, col).border = border; ws.cell(r_idx, col).fill = PatternFill("solid", fgColor="EDEDED")
        else:
            values = row["classification"] + [row["standard_answer"], row["tender_requirement"], "", ""]
            for col, value in enumerate(values, 1):
                cell = ws.cell(r_idx, col, value); cell.font = Font(name="宋体", size=8); cell.alignment = Alignment(vertical="center", wrap_text=True); cell.border = border
        ws.row_dimensions[r_idx].height = 28
    ranges = []
    for col, start, end, value in merges:
        min_row, max_row, excel_col = start + 1, end + 1, col + 1
        ws.merge_cells(start_row=min_row, start_column=excel_col, end_row=max_row, end_column=excel_col)
        ws.cell(min_row, excel_col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ranges.append(f"{get_column_letter(excel_col)}{min_row}:{get_column_letter(excel_col)}{max_row}")
    for col, width in enumerate([20, 18, 20, 24, 48, 70, 48, 22], 1): ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "E2"; ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.print_title_rows = "1:1"
    wb.save(xlsx_path)
    return ranges


def _validate_outputs(docx_path: Path, xlsx_path: Path, rows: list[dict]) -> dict:
    doc = Document(docx_path)
    if len(doc.tables) != 1 or [c.text for c in doc.tables[0].rows[0].cells] != HEADERS:
        raise ValueError("DOCX does not have the required eight-column table")
    doc_rows = doc.tables[0].rows[1:]
    if len(doc_rows) != len(rows):
        raise ValueError("DOCX row count differs from validated spec")
    for position, (actual, expected) in enumerate(zip(doc_rows, rows), 2):
        if actual.cells[6].text or actual.cells[7].text:
            raise ValueError(f"DOCX bidder response or deviation is not blank at row {position}")
        expected_text = expected["text"] if expected["kind"] == "heading" else expected["tender_requirement"]
        if actual.cells[5].text != expected_text:
            raise ValueError(f"DOCX tender requirement differs from validated spec at row {position}")
    wb = load_workbook(xlsx_path, data_only=False); ws = wb["八列内部核对表"]
    if [ws.cell(1, col).value for col in range(1, 9)] != HEADERS:
        raise ValueError("XLSX headers differ from the required eight columns")
    if ws.max_row != len(rows) + 1:
        raise ValueError("XLSX row count differs from validated spec")
    for position, expected in enumerate(rows, 2):
        if ws.cell(position, 7).value not in (None, "") or ws.cell(position, 8).value not in (None, ""):
            raise ValueError(f"XLSX bidder response or deviation is not blank at row {position}")
        expected_text = expected["text"] if expected["kind"] == "heading" else expected["tender_requirement"]
        if ws.cell(position, 6).value != expected_text:
            raise ValueError(f"XLSX tender requirement order/content differs from validated spec at row {position}")
    return {"eight_columns": "passed", "source_order_and_text": "passed", "hierarchy_complete": "passed", "bidder_response_and_deviation_blank": "passed", "docx_readable": "passed", "xlsx_readable": "passed"}


def install_eight_column_tools(registry):
    from .tools import STRING, schema

    def get_schema():
        return {
            "version": SPEC_VERSION,
            "purpose": "八列内部核对表；业务人工核对前第7/8列必须为空",
            "top_level": {
                "version": SPEC_VERSION,
                "atomization_path": "analysis/step2/通过build_requirement_atomization生成的JSON",
                "source_index_path": "analysis/step2/technical-source-index.json",
                "history_match_path": "analysis/step7/history-match-log.json（必填；每个requirement_id必须有明确状态）",
                "evidence_batch_paths": ["analysis/step7/evidence-batch-001.json（一个或多个；合计必须覆盖所有requirement_id）"],
                "rows": "按招标文件物理顺序排列的 heading/requirement 数组"
            },
            "heading_row": {"kind": "heading", "text": "逐字标题", "source_ids": ["来源元素ID"]},
            "requirement_row": {
                "kind": "requirement", "requirement_id": "项目内稳定且唯一的ID",
                "classification": ["车型", "子系统模块", "部件模块", "功能模块"],
                "sources": [{"id": "来源元素ID", "quote": "来源元素中的逐字片段"}],
                "source_joiner": "多个片段的原文连接符，默认空字符串",
                "removed_source_prefix": "仅用于剥离原局部编号；必须与拼接原文开头完全一致",
                "display_number": "按实际章节层级补全的编号；不含多余层级",
                "tender_requirement": "display_number + 空格 + 剥离原编号后的逐字原文",
                "history_record_id": "history_match_path中status=found且属于本requirement_id的同车型记录ID；无则省略",
                "bidder_response": "", "deviation": ""
            },
            "hard_rules": [
                "source引文必须真实存在于锁定来源索引", "第1至4列均非空",
                "rows必须与已通过的原子化及对象归属产物完全一致",
                "历史记录只允许exact_model或approved_model_family；每条必须有found/not_found/not_applicable/conflict状态", "第7/8列必须为空",
                "证据批次必须all_items_visited=true且合计逐条覆盖全部requirement_id",
                "顺序按rows原序，工具不按车型或模块重排", "只生成DOCX和XLSX；manifest为内部追溯文件"
            ]
        }

    def build_internal(spec_path, output_stem):
        spec_file = registry.path(spec_path)
        if spec_file.suffix.lower() != ".json" or not str(Path(spec_path)).startswith("analysis/"):
            raise ValueError("spec_path must be an analysis/*.json project file")
        if not re.fullmatch(r"[\w\u4e00-\u9fff-]{1,80}", output_stem):
            raise ValueError("output_stem must be a short file name without directories")
        spec = _load(spec_file)
        index_path = registry.path(spec.get("source_index_path", ""))
        if index_path.suffix.lower() != ".json": raise ValueError("source_index_path must be JSON")
        atomization_path_value = spec.get("atomization_path")
        atomization_path = registry.path(atomization_path_value) if atomization_path_value else None
        if atomization_path is None or atomization_path.suffix.lower() != ".json":
            raise ValueError("atomization_path from build_requirement_atomization is required")
        history_path_value = spec.get("history_match_path")
        history_path = registry.path(history_path_value) if history_path_value else None
        if history_path is None or history_path.suffix.lower() != ".json":
            raise ValueError("history_match_path with one status per requirement is required")
        evidence_path_values = spec.get("evidence_batch_paths")
        if not isinstance(evidence_path_values, list) or not evidence_path_values:
            raise ValueError("evidence_batch_paths must be a non-empty array")
        evidence_paths = [registry.path(value) for value in evidence_path_values]
        if any(path.suffix.lower() != ".json" for path in evidence_paths):
            raise ValueError("every evidence batch path must be JSON")
        history_data = _load(history_path)
        rows = _validate_and_expand(spec, _load(index_path), history_data)
        _validate_atomization_binding(spec, _load(atomization_path), rows, index_path)
        requirement_ids = {row["requirement_id"] for row in rows if row["kind"] == "requirement"}
        evidence = _evidence_items([(value, _load(path)) for value, path in zip(evidence_path_values, evidence_paths)], requirement_ids)
        traceability = _build_traceability(rows, history_data, evidence)
        target_dir = registry.root / "analysis" / "step7"; target_dir.mkdir(parents=True, exist_ok=True)
        docx_path = target_dir / (output_stem + ".docx"); xlsx_path = target_dir / (output_stem + ".xlsx"); manifest_path = target_dir / (output_stem + ".manifest.json"); traceability_path = target_dir / (output_stem + ".traceability.json")
        for target in (docx_path, xlsx_path, manifest_path, traceability_path):
            if target.exists(): raise ValueError("Output exists; use a new output_stem")
        merge_ranges = _render(rows, docx_path, xlsx_path)
        validation = _validate_outputs(docx_path, xlsx_path, rows)
        traceability.update({"spec_path": spec_path, "spec_sha256": _sha(spec_file),
                             "history_match_path": history_path_value, "history_match_sha256": _sha(history_path),
                             "evidence_batches": [{"path": value, "sha256": _sha(path)} for value, path in zip(evidence_path_values, evidence_paths)],
                             "outputs": {"docx": str(docx_path.relative_to(registry.root)), "xlsx": str(xlsx_path.relative_to(registry.root))}})
        traceability_path.write_text(json.dumps(traceability, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        manifest = {"version": SPEC_VERSION, "status": "internal_unreviewed", "formal_bid_output": False,
                    "business_review_performed": False, "source_index_path": spec["source_index_path"], "source_index_sha256": _sha(index_path),
                    "atomization_path": atomization_path_value, "atomization_sha256": _sha(atomization_path),
                    "history_match_path": history_path_value, "history_match_sha256": _sha(history_path) if history_path else None,
                    "evidence_batch_paths": evidence_path_values,
                    "traceability_path": str(traceability_path.relative_to(registry.root)), "traceability_sha256": _sha(traceability_path),
                    "spec_path": spec_path, "spec_sha256": _sha(spec_file), "columns": HEADERS,
                    "summary": {"rows": len(rows), "requirement_rows": sum(r["kind"] == "requirement" for r in rows), "heading_rows": sum(r["kind"] == "heading" for r in rows), "historical_answers_filled": sum(bool(r.get("standard_answer")) for r in rows), "hierarchy_merge_ranges": len(merge_ranges)},
                    "outputs": {"docx": str(docx_path.relative_to(registry.root)), "xlsx": str(xlsx_path.relative_to(registry.root))},
                    "sha256": {"docx": _sha(docx_path), "xlsx": _sha(xlsx_path)}, "validation": validation,
                    "rules": {"only_docx_xlsx_deliverables": True, "pdf_not_generated": True, "bidder_response_blank": True, "deviation_blank": True, "unreviewed_may_enter_bid": False,
                              "one_traceability_binding_per_requirement": True}}
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        return {**manifest, "manifest": str(manifest_path.relative_to(registry.root)), "manifest_sha256": _sha(manifest_path)}

    registry.add("get_eight_column_internal_review_schema",
                 "返回build_eight_column_internal_review所需spec JSON的完整字段契约和硬规则。制作八列内部核对表前先调用本工具，不得猜测字段。",
                 schema({}), get_schema)
    registry.add("build_eight_column_internal_review",
                 "依据已通过的条款原子化产物、锁定招标来源索引、逐条历史匹配状态和覆盖全部条款的证据批次，生成未经过业务核对的八列内部核对表DOCX与XLSX，并输出逐行追溯JSON。追溯文件稳定绑定表格行号、requirement_id、历史记录、原件物理页或缺项原因。工具逐字校验第6列来源、验证第5列权威记录、保持原顺序、合并连续同层级单元格，并强制第7/8列为空；不得据此生成正式三列表或装入标书。先调用build_requirement_atomization，再读取bid-model-response-matrix并调用get_eight_column_internal_review_schema。",
                 schema({"spec_path": STRING, "output_stem": STRING}), build_internal, True)
