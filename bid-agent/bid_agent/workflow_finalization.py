"""Build the final page review and Agent-scope decision from real visual events."""
from __future__ import annotations

import hashlib
import json
import re
import shutil
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

from .tools import STRING, schema


def _sha(path):
    with Path(path).open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def _new(folder, stem, payload):
    folder.mkdir(parents=True, exist_ok=True);version=1
    while (folder/f'{stem}-v{version}.json').exists():version+=1
    path=folder/f'{stem}-v{version}.json';path.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8');return path


def install_finalization_tools(registry):
    page_schema={'type':'object','additionalProperties':False,'properties':{
        'page':{'type':'integer','minimum':1},'image_path':STRING,
        'status':{'type':'string','enum':['pass','issue_resolved','issue_disclosed']},
        'issues':{'type':'array','items':STRING},
    },'required':['page','image_path','status','issues']}

    def build(pdf_path,pages):
        project=json.loads((registry.root/'project.json').read_text());pdf=registry.path(pdf_path)
        if pdf.suffix.lower()!='.pdf' or not pdf.is_file():raise ValueError('Internal final render PDF is required')
        from pypdf import PdfReader
        count=len(PdfReader(str(pdf)).pages);digest=_sha(pdf)
        if [x['page'] for x in pages]!=list(range(1,count+1)):raise ValueError('Page review must cover every final PDF page exactly once in order')
        events=registry.store.context(registry.run_id)['events']
        artifacts=[x['payload'] for x in events if x['kind']=='artifact_created' and x['payload'].get('operation')=='page_image']
        observations=[x['payload'] for x in events if x['kind']=='visual_observation']
        normalized=[]
        for item in pages:
            image=registry.path(item['image_path']);image_hash=_sha(image) if image.is_file() else None
            render=next((x for x in reversed(artifacts) if x.get('output')==item['image_path'] and x.get('page')==item['page'] and x.get('source_sha256')==digest and x.get('sha256')==image_hash),None)
            if render is None:raise ValueError('Page image is not a hash-bound render of the stated final PDF page')
            visual=next((x for x in reversed(observations) if x.get('image')==item['image_path'] and x.get('image_sha256')==image_hash and str(x.get('observation') or '').strip()),None)
            if visual is None:raise ValueError('Page has no actual vision-model observation bound to its image hash')
            if item['status']=='pass' and item['issues']:raise ValueError('A passed page cannot retain disclosed issues')
            if item['status']!='pass' and not item['issues']:raise ValueError('A non-pass page requires issue details')
            normalized.append({'page':item['page'],'image_path':item['image_path'],'image_sha256':image_hash,'status':item['status'],'issues':item['issues'],'vision_model':visual.get('model'),'observation':visual.get('observation')})
        receipts=registry.root/'stage_receipts';missing=[]
        for stage in range(1,10):
            path=receipts/f'{stage:02}.json'
            if not path.is_file() or json.loads(path.read_text()).get('status')!='passed':missing.append(stage)
        if missing:raise ValueError('Stages 1-9 must be passed before final page review: '+','.join(map(str,missing)))
        folder=registry.path('analysis/step10')
        review=_new(folder,'page-review',{'version':'stage10-page-review-v1','project_id':project['project_id'],'source_sha256':project['source_sha256'],'pdf_path':pdf_path,'pdf_sha256':digest,'page_count':count,'pages':normalized,'actual_vision_observation_count':len(normalized),'created_at':time.time()})
        open_issues=sum(len(x['issues']) for x in normalized if x['status']=='issue_disclosed')
        decision=_new(folder,'delivery-decision',{'version':'delivery-decision-v1','project_id':project['project_id'],'source_sha256':project['source_sha256'],'status':'agent_scope_complete','business_reviewed':False,'submission_ready':False,'page_review_path':str(review.relative_to(registry.root)),'page_review_sha256':_sha(review),'open_visual_issues':open_issues,'label':'Agent范围完成、待业务核对、不可递交','created_at':time.time()})
        return {'page_review':str(review.relative_to(registry.root)),'delivery_decision':str(decision.relative_to(registry.root)),'pages':count,'open_visual_issues':open_issues,'quality_state':'agent_scope_complete'}

    def complete_visual_review(pdf_path,image_folder,max_workers):
        """Resume a full visual review without asking the Agent to make 100s of calls.

        Page images are still required to have been produced by ``page_image``.
        Every missing observation is obtained through the exact public
        ``inspect_page_image`` handler, so the normal image hash and visual
        event gates remain authoritative.  Existing hash-matched observations
        are reused, which makes retries safe after a native session restart.

        The vision model is deliberately observation-only.  This composite
        therefore discloses each observation for business review rather than
        allowing free-form model prose to award a visual pass.
        """
        from pypdf import PdfReader
        pdf=registry.path(pdf_path)
        if pdf.suffix.lower()!='.pdf' or not pdf.is_file():
            raise ValueError('Internal final render PDF is required')
        count=len(PdfReader(str(pdf)).pages);pdf_digest=_sha(pdf)
        folder=registry.path(image_folder);folder.mkdir(parents=True,exist_ok=True)
        workers=max(1,min(int(max_workers),6))
        events=registry.store.context(registry.run_id)['events']
        artifacts=[x['payload'] for x in events if x['kind']=='artifact_created'
                   and x['payload'].get('operation')=='page_image']
        # A new stage-9 PDF invalidates every old page render.  Never overwrite
        # or trust those images: switch to a hash-scoped folder and render the
        # current PDF there.  This also makes retries deterministic.
        def render_bound(relative,number):
            image=registry.path(relative)
            digest=_sha(image) if image.is_file() else None
            return image.is_file() and any(x.get('output')==relative and x.get('page')==number
                                           and x.get('source_sha256')==pdf_digest and x.get('sha256')==digest
                                           for x in artifacts)
        default_paths=[str((Path(image_folder)/f'page-{n:03d}.png').as_posix()) for n in range(1,count+1)]
        if any(registry.path(path).exists() and not render_bound(path,n)
               for n,path in enumerate(default_paths,1)):
            image_folder=str((Path(image_folder)/('render-'+pdf_digest[:12])).as_posix())
            folder=registry.path(image_folder);folder.mkdir(parents=True,exist_ok=True)
        pages=[]
        for number in range(1,count+1):
            relative=str((Path(image_folder)/f'page-{number:03d}.png').as_posix())
            image=registry.path(relative)
            if not render_bound(relative,number):
                if image.exists():raise ValueError(f'Stale page image occupies current hash-scoped path: {number}')
                registry.tools['page_image'].handler(source=pdf_path,page=number,target=relative)
            pages.append((number,relative,_sha(image)))
        events=registry.store.context(registry.run_id)['events']
        observations={(x['payload'].get('image'),x['payload'].get('image_sha256')):x['payload']
                      for x in events if x['kind']=='visual_observation'
                      and str(x['payload'].get('observation') or '').strip()}
        missing=[(number,path,digest) for number,path,digest in pages if (path,digest) not in observations]
        focus=('逐页终审：检查空白、裁切、重叠、模糊、方向、跨页断表、表头、页码、'
               '证明材料可读性、内部槽位或过程文字泄露、已知字段遗漏；只报告本页可见事实。')
        failures=[]
        if missing:
            with ThreadPoolExecutor(max_workers=workers,thread_name_prefix='page-vision') as pool:
                futures={pool.submit(registry.inspect_page_image,path,focus):(number,path,digest)
                         for number,path,digest in missing}
                for future in as_completed(futures):
                    number,path,digest=futures[future]
                    try:observations[(path,digest)]=future.result()
                    except Exception as exc:failures.append(f'page {number}: {type(exc).__name__}: {exc}')
        if failures:
            raise RuntimeError('Visual review did not cover every page: '+'; '.join(failures[:10]))
        review_pages=[]
        for number,path,digest in pages:
            observation=str(observations[(path,digest)]['observation']).strip()
            one_line=re.sub(r'\s+',' ',observation)
            review_pages.append({'page':number,'image_path':path,'status':'issue_disclosed',
                                 'issues':['视觉观察待业务复核：'+one_line[:500]]})
        result=build(pdf_path,review_pages)
        return {**result,'observations_reused':count-len(missing),'observations_created':len(missing),
                'max_workers':workers,'visual_authority':'observation_only_no_pass_award'}

    def publish(final_bid_candidate,eight_column_docx_candidate,eight_column_xlsx_candidate,internal_render_pdf_path,
                final_bid_filename,eight_column_docx_filename,eight_column_xlsx_filename):
        project=json.loads((registry.root/'project.json').read_text())
        candidates={'final_bid_docx':registry.path(final_bid_candidate),'eight_column_docx':registry.path(eight_column_docx_candidate),'eight_column_xlsx':registry.path(eight_column_xlsx_candidate)}
        pdf=registry.path(internal_render_pdf_path)
        expected_names={'final_bid_docx':final_bid_filename,'eight_column_docx':eight_column_docx_filename,'eight_column_xlsx':eight_column_xlsx_filename}
        for key,name in expected_names.items():
            if Path(name).name!=name or '/' in name or '\\' in name or not re.fullmatch(r'[^\x00-\x1f]{1,180}',name):raise ValueError('Website output name must be a safe basename')
        if Path(final_bid_filename).suffix.lower()!='.docx' or Path(eight_column_docx_filename).suffix.lower()!='.docx' or Path(eight_column_xlsx_filename).suffix.lower()!='.xlsx':raise ValueError('Website outputs must be DOCX, DOCX and XLSX')
        if any(not x.is_file() for x in candidates.values()) or not pdf.is_file() or pdf.suffix.lower()!='.pdf':raise ValueError('Three real candidates and an internal render PDF are required')
        output_root=(registry.root/'outputs').resolve()
        if any(x.resolve().is_relative_to(output_root) for x in candidates.values()):raise ValueError('Publication candidates must come from the internal analysis area')
        from docx import Document
        final_doc=Document(candidates['final_bid_docx']);final_text='\n'.join([p.text for p in final_doc.paragraphs]+[c.text for t in final_doc.tables for r in t.rows for c in r.cells])
        if 'Agent范围完成、待业务核对、不可递交' not in final_text:raise ValueError('Final bid candidate lacks the required non-submittable Agent-scope label')
        headers=['车型','子系统模块','部件模块','功能模块','标准车型答案（历史投标响应）','本次招标要求','投标人响应','偏离说明']
        review_doc=Document(candidates['eight_column_docx'])
        if len(review_doc.tables)!=1 or [x.text for x in review_doc.tables[0].rows[0].cells]!=headers:raise ValueError('Eight-column DOCX structure is invalid')
        # Heading rows are intentionally merged across all eight columns.  In
        # python-docx every grid position of such a row aliases the same XML
        # cell, so reading columns 7/8 repeats the heading text.  Manual-column
        # emptiness applies to real eight-cell requirement rows only.
        for row in review_doc.tables[0].rows[1:]:
            cells=row.cells
            if cells[6]._tc is cells[0]._tc and cells[7]._tc is cells[0]._tc:continue
            if cells[6].text or cells[7].text:raise ValueError('Eight-column DOCX manual columns are not blank')
        from openpyxl import load_workbook
        book=load_workbook(candidates['eight_column_xlsx'],read_only=False,data_only=False)
        try:
            if '八列内部核对表' not in book.sheetnames:raise ValueError('Eight-column XLSX sheet is missing')
            sheet=book['八列内部核对表']
            if [sheet.cell(1,c).value for c in range(1,9)]!=headers:raise ValueError('Eight-column XLSX structure is invalid')
            if any(sheet.cell(r,7).value not in (None,'') or sheet.cell(r,8).value not in (None,'') for r in range(2,sheet.max_row+1)):raise ValueError('Eight-column XLSX manual columns are not blank')
        finally:book.close()
        outputs=registry.path('outputs');outputs.mkdir(parents=True,exist_ok=True)
        existing=list(outputs.iterdir())
        if existing:
            archive=registry.path('analysis/archives/step9')/str(int(time.time()*1000));archive.mkdir(parents=True,exist_ok=False)
            for old in existing:shutil.move(str(old),archive/old.name)
        published={}
        for key,source in candidates.items():
            target=outputs/expected_names[key];temp=outputs/(target.name+'.tmp');shutil.copy2(source,temp);temp.replace(target);published[key]=str(target.relative_to(registry.root))
        website=[{'kind':key,'path':published[key],'sha256':_sha(registry.root/published[key])} for key in ('final_bid_docx','eight_column_docx','eight_column_xlsx')]
        folder=registry.path('analysis/step9');manifest=_new(folder,'output-manifest',{'version':'stage9-output-manifest-v1','project_id':project['project_id'],'source_sha256':project['source_sha256'],'website_outputs':website,'internal_pdf_path':internal_render_pdf_path,'internal_pdf_sha256':_sha(pdf),'submission_ready':False,'business_reviewed':False,'quality_state':'agent_scope_complete','created_at':time.time()})
        return {**published,'internal_render_pdf':internal_render_pdf_path,'output_manifest':str(manifest.relative_to(registry.root)),'website_output_count':3,'submission_ready':False}

    registry.add('build_final_page_review_and_decision','在最终DOCX已内部渲染、每一页均通过page_image生成真实页图并调用inspect_page_image获得视觉观察后调用。服务器核对PDF哈希、页图来源、视觉事件和1—9步回执，生成第10步page_review与delivery_decision；最高只到agent_scope_complete，始终不可递交。',schema({'pdf_path':STRING,'pages':{'type':'array','items':page_schema}}),build,True)
    registry.add('a_complete_stage10_visual_review','可恢复地完成最终PDF逐页视觉终审。复用已存在且哈希一致的视觉观察，并以受限并发通过inspect_page_image补齐缺页；服务器随后核对全部页图、视觉事件及1—9步回执，生成page_review与delivery_decision。视觉模型只提供观察，不能自行授予通过。',schema({'pdf_path':STRING,'image_folder':STRING,'max_workers':{'type':'integer','minimum':1,'maximum':6}}),complete_visual_review,True)
    registry.add('publish_final_website_outputs','第9步组装、八列表和内部PDF完成后调用。服务器验证完整标书非递交标识、八列表固定结构及第7/8列全空，把旧网站文件移入analysis内部归档，原子发布且只保留完整标书DOCX、八列表DOCX、八列表XLSX，并生成output_manifest。',schema({'final_bid_candidate':STRING,'eight_column_docx_candidate':STRING,'eight_column_xlsx_candidate':STRING,'internal_render_pdf_path':STRING,'final_bid_filename':STRING,'eight_column_docx_filename':STRING,'eight_column_xlsx_filename':STRING}),publish,True)
