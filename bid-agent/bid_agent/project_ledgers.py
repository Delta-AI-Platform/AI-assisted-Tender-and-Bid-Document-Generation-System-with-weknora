"""Deterministic project fact and selected-configuration ledgers."""
from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from .tools import STRING, schema

FACT_STATES = {'source_verified','business_provided','derived_verified','conflict','missing','not_applicable'}
AUTO_FILL_STATES = {'source_verified','business_provided','derived_verified'}


def _sha(path: Path) -> str:
    with path.open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def _atomic_new(path: Path, payload: dict):
    if path.exists():
        raise FileExistsError('Ledger version already exists; choose a new versioned target')
    path.parent.mkdir(parents=True, exist_ok=True)
    raw=json.dumps(payload,ensure_ascii=False,indent=2).encode()
    temp=path.with_name(path.name+'.tmp')
    temp.write_bytes(raw);temp.replace(path)
    return {'path':path,'sha256':hashlib.sha256(raw).hexdigest(),'bytes':len(raw)}


def _project(registry):
    path=registry.root/'project.json'
    if not path.is_file():raise ValueError('project.json is required')
    return json.loads(path.read_text())


def _verify_quote_on_source(path: Path, page, quote: str):
    suffix=path.suffix.lower()
    if suffix=='.pdf':
        if type(page) is not int or page<1:raise ValueError('PDF source requires a physical_page')
        import pymupdf
        with pymupdf.open(path) as document:
            if page>len(document):raise ValueError('physical_page exceeds source PDF')
            text=document[page-1].get_text('text')
        if quote not in text:raise ValueError('Source quote is not verbatim on the stated PDF physical page')
    elif suffix in {'.txt','.md','.json','.csv'}:
        if quote not in path.read_text(errors='replace'):raise ValueError('Source quote is not present in the stated source file')
    elif suffix=='.docx':
        from docx import Document
        document=Document(path);text='\n'.join([p.text for p in document.paragraphs]+[c.text for t in document.tables for r in t.rows for c in r.cells])
        if quote not in text:raise ValueError('Source quote is not present in the stated DOCX')
    elif suffix=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=False)
        try:text='\n'.join(str(c.value) for sheet in book.worksheets for row in sheet.iter_rows() for c in row if c.value is not None)
        finally:book.close()
        if quote not in text:raise ValueError('Source quote is not present in the stated XLSX')


def _source_ref(registry, ref, require_quote=True):
    path=registry.path(ref.get('path','')); expected=ref.get('sha256','')
    if not path.is_file() or _sha(path)!=expected:raise ValueError('Fact source is missing or changed: '+str(ref.get('path')))
    page=ref.get('physical_page');location=str(ref.get('location') or '').strip();quote=str(ref.get('quote') or '').strip()
    if page is not None and (type(page) is not int or page<1):raise ValueError('physical_page must be a positive integer')
    if require_quote and not quote:raise ValueError('A short source quote is required')
    if not location:raise ValueError('Source table/cell/paragraph location is required')
    if quote:_verify_quote_on_source(path,page,quote)
    return {'path':ref['path'],'sha256':expected,'physical_page':page,'location':location,'quote':quote}


def build_project_facts(registry, facts, target):
    project=_project(registry);seen=set();normal=[]
    if not isinstance(facts,list) or not facts:raise ValueError('Non-empty facts array required')
    for item in facts:
        key=str(item.get('field_code') or '').strip();state=item.get('state')
        if not key or key in seen:raise ValueError('Fact field_code must be non-empty and unique')
        seen.add(key)
        if state not in FACT_STATES:raise ValueError('Invalid fact state: '+str(state))
        value=item.get('value');slots=item.get('allowed_slots')
        if state in AUTO_FILL_STATES and value in ('',None):raise ValueError(key+': verified/provided/derived fact requires a value')
        if not isinstance(slots,list):raise ValueError(key+': allowed_slots must be an array')
        refs=item.get('sources',[])
        if state=='source_verified' and not refs:raise ValueError(key+': source_verified requires a source')
        checked=[_source_ref(registry,x) for x in refs]
        normal.append({'field_code':key,'label':str(item.get('label') or key),'value':value,'state':state,
                       'sources':checked,'allowed_slots':slots,'reason':str(item.get('reason') or '')})
    denominator=[x['field_code'] for x in normal if x['state'] in AUTO_FILL_STATES and x['allowed_slots']]
    expected_slots=[{'field_code':x['field_code'],'slot_id':slot} for x in normal if x['state'] in AUTO_FILL_STATES for slot in x['allowed_slots']]
    payload={'version':'project-facts-v1','project_id':project['project_id'],'source_sha256':project['source_sha256'],
             'facts':normal,'auto_fill_denominator':denominator,'auto_fill_expected_count':len(denominator),
             'auto_fill_expected_slots':expected_slots,'auto_fill_expected_slot_count':len(expected_slots),
             'status_counts':{state:sum(x['state']==state for x in normal) for state in sorted(FACT_STATES)}}
    dest=registry.path(target)
    if not str(Path(target)).startswith('analysis/') or dest.suffix!='.json':raise ValueError('Facts target must be a versioned analysis JSON')
    result=_atomic_new(dest,payload)
    return {**{k:(str(v.relative_to(registry.root)) if k=='path' else v) for k,v in result.items()},'summary':payload['status_counts']}


def create_configuration_lock(registry, selection_path, target):
    project=_project(registry);selection_file=registry.path(selection_path)
    selection=json.loads(selection_file.read_text());items=selection.get('procurement_items')
    if not isinstance(items,list) or not items:raise ValueError('Selection requires procurement_items')
    if selection.get('project_id')!=project['project_id'] or selection.get('source_sha256')!=project['source_sha256']:
        raise ValueError('Selection is not bound to the active project and tender source')
    seen=set();locked=[];all_hashes=[]
    for item in items:
        pid=str(item.get('procurement_id') or '').strip()
        if not pid or pid in seen:raise ValueError('procurement_id must be non-empty and unique')
        seen.add(pid);status=str(item.get('status') or '')
        model=str(item.get('selected_model') or '').strip()
        parent=str(item.get('parent_procurement_id') or '').strip()
        if model and parent:raise ValueError(pid+': select a model or inherit a parent, not both')
        if not model and not parent and not status.startswith(('gap','blocked')):raise ValueError(pid+': unlocked item requires an explicit gap/blocked status')
        evidence=[]
        for rel in item.get('evidence_paths',[]):
            source=registry.path(rel)
            if not source.is_file():raise ValueError(pid+': evidence file missing: '+rel)
            digest=_sha(source);all_hashes.append(digest);evidence.append({'path':rel,'sha256':digest})
        if model and not evidence:raise ValueError(pid+': selected model requires a hash-bound evidence set')
        locked.append({'procurement_id':pid,'name':str(item.get('name') or ''),'selected_model':model or None,
                       'submodel':str(item.get('submodel') or ''),'report_number':str(item.get('report_number') or ''),
                       'report_version':str(item.get('report_version') or ''),'parent_procurement_id':parent or None,
                       'status':status,'applicable_requirement_ids':item.get('applicable_requirement_ids',[]),
                       'evidence':evidence,'gap':str(item.get('gap') or '')})
    unknown={x['parent_procurement_id'] for x in locked if x['parent_procurement_id']} - seen
    if unknown:raise ValueError('Unknown parent procurement item: '+','.join(sorted(unknown)))
    selection_sha=_sha(selection_file)
    version=hashlib.sha256((project['project_id']+'\n'+project['source_sha256']+'\n'+selection_sha+'\n'+'\n'.join(sorted(all_hashes))).encode()).hexdigest()
    payload={'version':'configuration-lock-v1','lock_version':version,'project_id':project['project_id'],
             'tender_source_sha256':project['source_sha256'],'selection_path':selection_path,
             'selection_sha256':selection_sha,'procurement_items':locked,'immutable':True}
    dest=registry.path(target)
    if not re.fullmatch(r'analysis/step6/configuration-lock-[0-9A-Za-z._-]+\.json',target):
        raise ValueError('Use a versioned analysis/step6/configuration-lock-*.json target')
    result=_atomic_new(dest,payload)
    return {'path':str(result['path'].relative_to(registry.root)),'sha256':result['sha256'],'bytes':result['bytes'],'lock_version':version,'items':len(locked)}


def load_configuration_lock(registry, path, version):
    source=registry.path(path);data=json.loads(source.read_text());project=_project(registry)
    if data.get('version')!='configuration-lock-v1' or data.get('lock_version')!=version:raise ValueError('Configuration lock version mismatch')
    if data.get('project_id')!=project['project_id'] or data.get('tender_source_sha256')!=project['source_sha256']:raise ValueError('Configuration lock belongs to another or stale project')
    selection=registry.path(data.get('selection_path',''))
    if not selection.is_file() or _sha(selection)!=data.get('selection_sha256'):raise ValueError('Configuration selection source changed')
    for item in data.get('procurement_items',[]):
        for evidence in item.get('evidence',[]):
            f=registry.path(evidence['path'])
            if not f.is_file() or _sha(f)!=evidence['sha256']:raise ValueError('Configuration evidence changed: '+evidence['path'])
    return data


def install_project_ledger_tools(registry):
    source={'type':'object','properties':{'path':STRING,'sha256':STRING,'physical_page':{'type':['integer','null']},'location':STRING,'quote':STRING},'required':['path','sha256','physical_page','location','quote'],'additionalProperties':False}
    fact={'type':'object','properties':{'field_code':STRING,'label':STRING,'value':{},'state':{'type':'string','enum':sorted(FACT_STATES)},'sources':{'type':'array','items':source},'allowed_slots':{'type':'array','items':STRING},'reason':STRING},'required':['field_code','label','value','state','sources','allowed_slots','reason'],'additionalProperties':False}
    registry.add('build_project_facts','建立真实来源绑定的项目事实账本。只有source_verified、business_provided、derived_verified且有允许槽位的字段进入自动填写分母；冲突和缺失只记录并继续。',schema({'facts':{'type':'array','minItems':1,'items':fact},'target':STRING}),lambda facts,target:build_project_facts(registry,facts,target),True)
    registry.add('create_configuration_lock','读取当前项目report-selection并生成不可变、真实文件哈希绑定的选型锁。选型项必须有成套证据，备件只能引用父采购项；缺车型须明确记录gap并继续。',schema({'selection_path':STRING,'target':STRING}),lambda selection_path,target:create_configuration_lock(registry,selection_path,target),True)
    registry.add('read_configuration_lock','重新核验选型源和全部证据哈希后读取选型锁；任何源文件变化都会拒绝旧锁。',schema({'path':STRING,'lock_version':STRING}),lambda path,lock_version:load_configuration_lock(registry,path,lock_version))
