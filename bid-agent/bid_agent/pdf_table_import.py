"""Experimental editable PDF-table import; source page review still required."""
import contextlib
import io
import tempfile
import collections
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pymupdf


def import_table(source,page,table_index,target,include_surrounding_text=False,source_sha256=None):
    try:
        from .document_ops import digest,publish
    except ImportError:
        from document_ops import digest,publish
    before=digest(source)
    if source_sha256 != before:raise ValueError('Source version binding required')
    if type(page)is not int or page<1 or type(table_index)is not int or table_index<0:raise ValueError('Invalid page or table index')
    if type(include_surrounding_text)is not bool:raise ValueError('Explicit boolean required')
    with pymupdf.open(source) as pdf:
        if page>len(pdf):raise ValueError('Page outside source')
        with contextlib.redirect_stdout(io.StringIO()):tables=pdf[page-1].find_tables().tables
        if table_index>=len(tables):raise ValueError('Table outside source')
        page_text=pdf[page-1].get_text()
        t=tables[table_index];values=t.extract()
        page_width,page_height=pdf[page-1].rect.width,pdf[page-1].rect.height
        table_bbox=t.bbox
        before_lines=[];after_lines=[];excluded=[]
        if include_surrounding_text:
            if len(tables)!=1:raise ValueError('Surrounding text import requires a single table page')
            if pdf[page-1].get_images(full=True):raise ValueError('Page contains images; cannot silently omit them')
            for block in pdf[page-1].get_text('dict')['blocks']:
                for line in block.get('lines',[]):
                    text=''.join(span['text'] for span in line['spans'])
                    x0,y0,x1,y1=line['bbox']
                    if y1<table_bbox[1]:before_lines.append((y0,x0,text))
                    elif y0>table_bbox[3]:
                        if y0>page_height*.9 and text.strip().strip('—- ').isdigit():excluded.append(text)
                        else:after_lines.append((y0,x0,text))
                    elif not (table_bbox[0]<=((x0+x1)/2)<=table_bbox[2]):
                        raise ValueError('Side text requires manual layout')
        xs=sorted(set(round(v,2) for box in t.cells for v in (box[0],box[2])))
        ys=sorted(set(round(v,2) for box in t.cells for v in (box[1],box[3])))
        if len(xs)!=t.col_count+1 or len(ys)!=t.row_count+1:raise ValueError('Inconsistent grid requires manual review')
        spans=[span for block in pdf[page-1].get_text('dict')['blocks'] for line in block.get('lines',[]) for span in line['spans']]
        cell_sizes={}
        cells={}
        for i,row in enumerate(t.rows):
            for j,box in enumerate(row.cells):
                if box is None:continue
                key=tuple(round(v,2) for v in box)
                if key in cells and cells[key]!=values[i][j]:raise ValueError('Conflicting merged cell text')
                cells[key]=values[i][j] or ''
                sizes=collections.Counter()
                for span in spans:
                    sx0,sy0,sx1,sy1=span['bbox']
                    if key[0] <= (sx0+sx1)/2 <= key[2] and key[1] <= (sy0+sy1)/2 <= key[3]:
                        sizes[round(span['size']*2)/2]+=len(span['text'])
                cell_sizes[key]=sizes.most_common(1)[0][0] if sizes else 10.5
    doc=Document();s=doc.sections[0];s.page_width=Pt(page_width);s.page_height=Pt(page_height)
    s.top_margin=s.bottom_margin=Cm(2.5);s.left_margin=Pt(table_bbox[0]);s.right_margin=Pt(page_width-table_bbox[2])
    style=doc.styles['Normal'];style.font.name='Noto Serif CJK SC';style.font.size=Pt(10.5);style.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
    def add_positioned_lines(lines):
        groups=[]
        for y,x,text in sorted(lines):
            if groups and abs(y-groups[-1][0][0])<=2:groups[-1].append((y,x,text))
            else:groups.append([(y,x,text)])
        for group in groups:
            group.sort(key=lambda item:item[1]);p=doc.add_paragraph();origin=group[0][1]
            p.paragraph_format.left_indent=Pt(max(0,origin-table_bbox[0]))
            p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
            for j,(_,x,text) in enumerate(group):
                if j:
                    p.paragraph_format.tab_stops.add_tab_stop(Pt(x-origin));p.add_run('\t')
                p.add_run(text)
    add_positioned_lines(before_lines)
    table=doc.add_table(rows=t.row_count,cols=t.col_count);table.style='Table Grid';table.autofit=False
    available=float(s.page_width-s.left_margin-s.right_margin)
    for i,col in enumerate(table.columns):col.width=int(available*(xs[i+1]-xs[i])/(xs[-1]-xs[0]))
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    for i,row in enumerate(table.rows):
        row.height=Pt(ys[i+1]-ys[i]);row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
    mapping=[];coverage=collections.Counter()
    for box,text in cells.items():
        x0,y0,x1,y1=box;r0,r1=ys.index(y0),ys.index(y1);c0,c1=xs.index(x0),xs.index(x1)
        for rr in range(r0,r1):
            for cc in range(c0,c1):coverage[rr,cc]+=1
        cell=table.cell(r0,c0)
        if r1-r0>1 or c1-c0>1:cell=cell.merge(table.cell(r1-1,c1-1))
        cell.width=int(available*(box[2]-box[0])/(xs[-1]-xs[0]))
        cell.text=text
        margins=OxmlElement('w:tcMar')
        for side,width in [('top',0),('bottom',0),('left',20),('right',20)]:
            edge=OxmlElement('w:'+side);edge.set(qn('w:w'),str(width));edge.set(qn('w:type'),'dxa');margins.append(edge)
        cell._tc.get_or_add_tcPr().append(margins)
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.size=Pt(cell_sizes[box]);run.font.name='Noto Serif CJK SC'
                run._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
                sz=OxmlElement('w:szCs');sz.set(qn('w:val'),str(int(cell_sizes[box]*2)));run._r.get_or_add_rPr().append(sz)
            paragraph.paragraph_format.space_after=Pt(0)
            paragraph.paragraph_format.space_before=Pt(0)
            paragraph.paragraph_format.line_spacing=1
        mapping.append({'row':r0,'column':c0,'row_span':r1-r0,'column_span':c1-c0,'text':text,'pdf_bbox':box,'source_font_size_pt_rounded':cell_sizes[box]})
    if len(coverage)!=t.row_count*t.col_count or any(n!=1 for n in coverage.values()):raise ValueError('Table cells overlap or leave unclassified holes')
    add_positioned_lines(after_lines)
    if include_surrounding_text:
        for footer in excluded:page_text=page_text.replace(footer,'')
        word_text=''.join(doc.element.xpath('.//w:t/text()'))
        compact=lambda text:collections.Counter(re.sub(r'\s+','',text))
        if compact(page_text)!=compact(word_text):raise ValueError('Reconstruction omits or adds source characters')
    with tempfile.TemporaryDirectory() as tmp:
        out=Path(tmp)/'table.docx';doc.save(out);check=Document(out).tables[0]
        for item in mapping:
            if check.cell(item['row'],item['column']).text!=item['text']:raise ValueError('Saved cell mismatch')
        if digest(source)!=before:raise ValueError('Source changed')
        return {**publish(out,target),'source_sha256':before,'page':page,'table_index':table_index,'cells':mapping,
                'surrounding_text_included':include_surrounding_text,'excluded_source_footer':excluded,'scope':'Reconstructed editable table/page; font fallback Noto Serif CJK SC; source table detection and layout require visual review'}
