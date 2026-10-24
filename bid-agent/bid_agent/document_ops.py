"""Versioned document operations executed in the dedicated document environment."""
import hashlib
import json
import os
from pathlib import Path
import subprocess
import tempfile


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def publish(source, target):
    # Exclusive creation: a concurrent writer cannot cause an overwrite.
    with Path(source).open('rb') as inp, Path(target).open('xb') as out:
        while chunk := inp.read(1024 * 1024):
            out.write(chunk)
        out.flush()
        os.fsync(out.fileno())
    return {'sha256': digest(target), 'bytes': Path(target).stat().st_size}


def create_docx(target, blocks):
    from docx import Document
    from docx.shared import Pt
    from docx.oxml.ns import qn
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Noto Sans CJK SC'
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    if not blocks or len(blocks) > 10000:
        raise ValueError('Provide 1–10000 blocks')
    for index, block in enumerate(blocks):
        if not isinstance(block, dict) or block.get('type') not in {'paragraph','heading','page_break','table'}:
            raise ValueError(f'Block {index}: use type=paragraph, heading, page_break or table. Heading example: {{"type":"heading","level":1,"text":"Title"}}; block_type and heading1 are invalid.')
        if block['type'] in {'paragraph','heading'} and not isinstance(block.get('text'),str):
            raise ValueError(f'Block {index}: explicit string text required')
    for block in blocks:
        kind = block['type']
        if kind == 'paragraph':
            doc.add_paragraph(block['text'])
        elif kind == 'heading':
            level = block.get('level', 1)
            if level not in (1, 2, 3): raise ValueError('Heading level must be 1–3')
            doc.add_heading(block['text'], level=level)
        elif kind == 'page_break':
            doc.add_page_break()
        elif kind == 'table':
            rows = block['rows']
            if not rows or not rows[0] or any(len(r) != len(rows[0]) for r in rows):
                raise ValueError('Table must be rectangular and nonempty')
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = 'Table Grid'
            for i, row in enumerate(rows):
                for j, value in enumerate(row):
                    if not isinstance(value, str): raise ValueError('Table cells must be explicit strings, including blanks')
                    table.cell(i, j).text = value
        else:
            raise ValueError('Unknown document block')
    with tempfile.TemporaryDirectory() as tmp:
        file = Path(tmp) / 'document.docx'
        doc.save(file)
        return publish(file, target)


def edit_docx(source, source_sha256, target, edits):
    from docx import Document
    from docx.oxml.ns import qn
    if digest(source) != source_sha256: raise ValueError('Source version changed')
    doc = Document(source)
    if not isinstance(edits, list) or not edits:
        raise ValueError('edits requires at least one {block: nonnegative integer, before: exact original text, after: new text}')
    for edit in edits:
        if (not isinstance(edit, dict) or set(edit) != {'block', 'before', 'after'}
                or type(edit['block']) is not int or edit['block'] < 0
                or not isinstance(edit['before'], str) or not isinstance(edit['after'], str)):
            raise ValueError('Each edit requires exactly block (nonnegative integer), before (original text), after (new text); old_text/new_text are invalid. For empty BID_SLOT use fill_empty_slots.')
        if edit['block'] >= len(doc.element.body):
            raise ValueError('block is outside current document; use the body block index returned by read_document')
        block = doc.element.body[edit['block']]
        if block.tag != qn('w:p'): raise ValueError('Only paragraph blocks supported by this edit operation')
        nodes = list(block.iter(qn('w:t')))
        before = ''.join(n.text or '' for n in nodes)
        if before != edit['before']: raise ValueError('Expected paragraph text does not match')
        # Keep existing OOXML structure/runs. Whole-paragraph replacements are
        # allowed only for a single text run, so mixed formatting isn't flattened.
        if len(nodes) != 1: raise ValueError('Mixed-run paragraph requires a more precise editing tool')
        nodes[0].text = edit['after']
        nodes[0].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    if digest(source) != source_sha256: raise ValueError('Source changed during editing')
    with tempfile.TemporaryDirectory() as tmp:
        file = Path(tmp) / 'document.docx'; doc.save(file)
        return {**publish(file, target), 'source_sha256': source_sha256}


def render_pdf(source, target):
    before = digest(source)
    with tempfile.TemporaryDirectory(prefix='bid-render-') as tmp:
        tmp = Path(tmp)
        profile = (tmp / 'profile').as_uri()
        result = subprocess.run(['/usr/bin/libreoffice', '-env:UserInstallation=' + profile,
                                 '--headless', '--convert-to', 'pdf', '--outdir', str(tmp), str(source)],
                                capture_output=True, text=True, timeout=120)
        pdf = tmp / (Path(source).stem + '.pdf')
        if result.returncode or not pdf.is_file(): raise ValueError('LibreOffice conversion failed: ' + result.stderr[:500])
        if digest(source) != before: raise ValueError('Source changed during rendering')
        from pypdf import PdfReader
        pages = len(PdfReader(pdf).pages)
        return {**publish(pdf, target), 'source_sha256': before, 'pages': pages, 'visually_checked': False}


def assemble_pdf(sources, target):
    from pypdf import PdfReader, PdfWriter
    writer = PdfWriter(); manifest = []
    for source in sources:
        before = digest(source); reader = PdfReader(source)
        start = len(writer.pages) + 1
        for page in reader.pages: writer.add_page(page)
        manifest.append({'source_sha256': before, 'page_start': start, 'page_count': len(reader.pages)})
        if digest(source) != before: raise ValueError('Source changed during assembly')
    with tempfile.TemporaryDirectory() as tmp:
        file = Path(tmp) / 'assembled.pdf'
        with file.open('wb') as out: writer.write(out)
        return {**publish(file, target), 'sources': manifest, 'pages': len(writer.pages)}


def page_image(source, page, target):
    import pymupdf as fitz
    before = digest(source)
    with fitz.open(source) as doc:
        if page < 1 or page > len(doc): raise ValueError('Page outside document')
        pix = doc[page - 1].get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
        with tempfile.TemporaryDirectory() as tmp:
            file = Path(tmp) / 'page.png'; pix.save(file)
            if digest(source) != before: raise ValueError('Source changed during page rendering')
            return {**publish(file, target), 'source_sha256': before, 'page': page,
                    'width': pix.width, 'height': pix.height, 'visually_checked': False}


try:
    from .section_layout import format_docx
except ImportError:
    from section_layout import format_docx


OPERATIONS = {f.__name__: f for f in (create_docx, edit_docx, format_docx, render_pdf, assemble_pdf, page_image)}
try:
    from .quotation import write_reconciliation, restore_quotation
except ImportError:
    from quotation import write_reconciliation, restore_quotation
OPERATIONS.update(write_reconciliation=write_reconciliation, restore_quotation=restore_quotation)
try:
    from .assembly import assemble_docx, append_pdf_evidence
except ImportError:
    from assembly import assemble_docx, append_pdf_evidence
OPERATIONS.update(assemble_docx=assemble_docx, append_pdf_evidence=append_pdf_evidence)
try:
    from .table_edit import edit_table_cells
except ImportError:
    from table_edit import edit_table_cells
OPERATIONS['edit_table_cells'] = edit_table_cells
try:
    from .source_capture import capture_requirements
except ImportError:
    from source_capture import capture_requirements
OPERATIONS['capture_requirements'] = capture_requirements
try:
    from .pdf_table_import import import_table
except ImportError:
    from pdf_table_import import import_table
OPERATIONS['import_pdf_table'] = import_table
try:
    from .technical_workbook import create_technical_workbook
except ImportError:
    from technical_workbook import create_technical_workbook
OPERATIONS['create_technical_workbook'] = create_technical_workbook
try:
    from .tag_docx_slots import tag_docx_slots
    from .slot_bridge import fill_empty_slots
except ImportError:
    from tag_docx_slots import tag_docx_slots
    from slot_bridge import fill_empty_slots
OPERATIONS.update(tag_docx_slots=tag_docx_slots, fill_empty_slots=fill_empty_slots)
try:
    from .insert_docx import insert_docx
except ImportError:
    from insert_docx import insert_docx
OPERATIONS['insert_docx'] = insert_docx
try:
    from .repeat_blocks import repeat_docx_blocks
except ImportError:
    from repeat_blocks import repeat_docx_blocks
OPERATIONS['repeat_docx_blocks'] = repeat_docx_blocks
try:
    from .docx_styles import set_docx_styles
except ImportError:
    from docx_styles import set_docx_styles
OPERATIONS['set_docx_styles'] = set_docx_styles


try:
    from .tag_inline_slots import tag_inline_slots
except ImportError:
    from tag_inline_slots import tag_inline_slots
OPERATIONS['tag_inline_slots'] = tag_inline_slots
try:
    from .partition_sections import partition_sections
except ImportError:
    from partition_sections import partition_sections
OPERATIONS['partition_sections'] = partition_sections
try:
    from .replace_table_layouts import replace_table_layouts
except ImportError:
    from replace_table_layouts import replace_table_layouts
OPERATIONS['replace_table_layouts'] = replace_table_layouts

try:
    from .xlsx_spec_tables import xlsx_spec_tables
except ImportError:
    from xlsx_spec_tables import xlsx_spec_tables
OPERATIONS['xlsx_spec_tables'] = xlsx_spec_tables

try:
    from .expand_template_table import expand_template_table
except ImportError:
    from expand_template_table import expand_template_table
OPERATIONS['expand_template_table'] = expand_template_table

try:
    from .table_records import expand_table_from_records, edit_table_cells_from_records
except ImportError:
    from table_records import expand_table_from_records, edit_table_cells_from_records
OPERATIONS['expand_table_from_records'] = expand_table_from_records
OPERATIONS['edit_table_cells_from_records'] = edit_table_cells_from_records

try:
    from .table_pagination import set_table_pagination
except ImportError:
    from table_pagination import set_table_pagination
OPERATIONS['set_table_pagination'] = set_table_pagination

try:
    from .paragraph_pagination import set_paragraph_pagination
except ImportError:
    from paragraph_pagination import set_paragraph_pagination
OPERATIONS['set_paragraph_pagination'] = set_paragraph_pagination

try:
    from .technical_annotations import annotate_technical_workbook
except ImportError:
    from technical_annotations import annotate_technical_workbook
OPERATIONS['annotate_technical_workbook'] = annotate_technical_workbook

if __name__ == '__main__':
    try:
        request = json.load(__import__('sys').stdin)
        result = OPERATIONS[request['operation']](**request['args'])
        print(json.dumps({'ok': True, 'data': result}, ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({'ok': False, 'error': f'{type(exc).__name__}: {exc}'}))
        raise SystemExit(1)

try:
    from .empty_sections import remove_redundant_empty_section
except ImportError:
    from empty_sections import remove_redundant_empty_section
OPERATIONS['remove_redundant_empty_section'] = remove_redundant_empty_section
