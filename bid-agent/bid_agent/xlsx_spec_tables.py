"""Explicit, version-bound spreadsheet values to editable grouped DOCX tables."""
from pathlib import Path
import tempfile,math
from openpyxl import load_workbook
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def xlsx_spec_tables(source,source_sha256,target,sheet,header_row,start_row,end_row,group_column,columns,title,page_width_mm,page_height_mm,margin_mm):
    try:
        from .document_ops import digest,publish
    except ImportError:
        from document_ops import digest,publish
    if digest(source)!=source_sha256:raise ValueError('Source version changed')
    wb=load_workbook(source,data_only=False)
    if sheet not in wb.sheetnames:raise ValueError('Explicit sheet not found')
    ws=wb[sheet]
    for value in (header_row,start_row,end_row,group_column):
        if type(value) is not int:raise ValueError('Integer source coordinates required')
    if not 1<=header_row<start_row<=end_row<=ws.max_row or end_row-start_row>=1000:raise ValueError('Invalid source rows or more than 1000 rows')
    if not 1<=group_column<=ws.max_column:raise ValueError('Invalid grouping column')
    if not isinstance(columns,list) or not 1<=len(columns)<=20:raise ValueError('Select 1-20 columns')
    seen=set()
    for c in columns:
        if not isinstance(c,dict) or set(c)!={'column','width_mm'}:raise ValueError('Column and explicit width required')
        n=c['column'];width=c['width_mm']
        if type(n) is not int or not 1<=n<=ws.max_column or n in seen:raise ValueError('Unique valid columns required')
        seen.add(n)
        if type(width) not in (int,float) or not 0<width<=500:raise ValueError('Invalid column width')
    for v in (page_width_mm,page_height_mm,margin_mm):
        if type(v) not in (int,float) or not math.isfinite(v):raise ValueError('Finite page geometry required')
    if not 100<=page_width_mm<=500 or not 100<=page_height_mm<=500 or not 5<=margin_mm<min(page_width_mm,page_height_mm)/2:raise ValueError('Invalid page geometry')
    if sum(c['width_mm'] for c in columns)>page_width_mm-2*margin_mm+.01:raise ValueError('Selected columns exceed page content width')
    if not isinstance(title,str) or not title.strip():raise ValueError('Explicit title required')
    def value(row,col):
        cell=ws.cell(row,col)
        if cell.data_type in {'f','e'}:raise ValueError('Formula or error value requires separate review')
        return '' if cell.value is None else str(cell.value)
    headers=[value(header_row,c['column']) for c in columns]
    records=[{'excel_row':n,'group':value(n,group_column),'values':[value(n,c['column']) for c in columns]} for n in range(start_row,end_row+1)]
    doc=Document();sec=doc.sections[0];sec.page_width=Mm(page_width_mm);sec.page_height=Mm(page_height_mm);sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Mm(margin_mm)
    for name,size in [('Normal',9),('Heading 1',16),('Heading 2',12)]:
        st=doc.styles[name];st.font.name='Noto Serif CJK SC';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
    doc.add_heading(title,1);previous=None;table=None;mapping=[]
    for record in records:
        if table is None or record['group']!=previous:
            if table is not None:doc.add_page_break()
            doc.add_heading(record['group'],2);table=doc.add_table(rows=1,cols=len(columns));table.style='Table Grid';table.autofit=False
            for i,c in enumerate(columns):table.columns[i].width=Mm(c['width_mm']);table.cell(0,i).width=Mm(c['width_mm']);table.cell(0,i).text=headers[i]
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'));previous=record['group']
        cells=table.add_row().cells
        for i,text in enumerate(record['values']):cells[i].text=text;cells[i].width=Mm(columns[i]['width_mm'])
        mapping.append({**record,'table_index':len(doc.tables)-1,'row_index':len(table.rows)-1})
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
    with tempfile.TemporaryDirectory() as tmp:
        out=Path(tmp)/'tables.docx';doc.save(out);check=Document(out)
        for row in mapping:
            if [c.text for c in check.tables[row['table_index']].rows[row['row_index']].cells]!=row['values']:raise ValueError('Saved values differ')
        if digest(source)!=source_sha256:raise ValueError('Source changed')
        return {**publish(out,target),'source_sha256':source_sha256,'sheet':sheet,'columns':columns,'mapping':mapping,'quality_status':'not_verified','scope':'Selected source cell values only, no factual certification; formulas/errors rejected. Contiguous grouping preserves original row order. Source displayed number formatting, images and drawings are not reproduced. Noto fallback requires rendered review.'}
