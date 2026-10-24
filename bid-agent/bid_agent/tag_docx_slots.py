"""Prototype: tag explicitly selected empty paragraphs/cells for skill filling."""
import re
import tempfile
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from zipfile import ZipFile


def tag_docx_slots(source, source_sha256, target, slots):
    try:
        from .document_ops import digest, publish
    except ImportError:
        from document_ops import digest, publish
    if digest(source) != source_sha256:
        raise ValueError('Source version changed')
    if not isinstance(slots, list) or not slots:
        raise ValueError('Explicit slots required')
    doc = Document(source)
    known = set(doc.element.xpath('//w:sdtPr/w:tag/@w:val'))
    selected = set()
    for slot in slots:
        ident = slot.get('slot_id', '')
        if not isinstance(ident, str) or not re.fullmatch(r'[A-Za-z0-9_-]{1,100}', ident):
            raise ValueError('Stable ASCII slot identifier required')
        tag = 'BID_SLOT:'+ident
        if tag in known:
            raise ValueError('Duplicate slot identifier')
        known.add(tag)
        for key in ('paragraph',) if slot.get('kind') == 'paragraph' else ('table','row','column'):
            if type(slot.get(key)) is not int or slot[key] < 0:
                raise ValueError('Explicit nonnegative integer coordinates required')
        if slot.get('kind') == 'paragraph':
            paragraph = doc.paragraphs[slot['paragraph']]
        elif slot.get('kind') == 'cell':
            table = doc.tables[slot['table']]
            if slot['row'] >= len(table.rows) or slot['column'] >= len(table.columns):
                raise ValueError('Cell outside table')
            cell = table.cell(slot['row'], slot['column'])
            if len(cell.paragraphs) != 1 or cell.tables:
                raise ValueError('Complex cell cannot be tagged automatically')
            paragraph = cell.paragraphs[0]
        else:
            raise ValueError('Unknown slot kind')
        node = paragraph._p
        if node in selected:
            raise ValueError('Merged-cell alias or duplicate position')
        selected.add(node)
        if node.xpath('.//w:t|.//w:drawing|.//w:pict|.//w:fldChar|.//w:instrText|.//w:sdt|.//w:br|.//w:tab'):
            raise ValueError('Only truly empty positions can be tagged')
        sdt = OxmlElement('w:sdt')
        props = OxmlElement('w:sdtPr')
        marker = OxmlElement('w:tag');marker.set(qn('w:val'), tag)
        props.append(marker);sdt.append(props)
        content = OxmlElement('w:sdtContent')
        for child in list(node):
            if child.tag != qn('w:pPr'):
                node.remove(child);content.append(child)
        if not len(content):
            content.append(OxmlElement('w:r'))
        sdt.append(content);node.append(sdt)
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp)/'tagged.docx';doc.save(out)
        with ZipFile(source) as a, ZipFile(out) as b:
            if set(a.namelist()) != set(b.namelist()):
                raise ValueError('Document package parts changed')
            for name in a.namelist():
                if name != 'word/document.xml' and a.read(name) != b.read(name):
                    raise ValueError('Unrelated document part changed')
        if digest(source) != source_sha256:
            raise ValueError('Source changed')
        return {**publish(out,target), 'source_sha256':source_sha256, 'slots':slots,
                'quality_status':'prototype_not_frozen'}
