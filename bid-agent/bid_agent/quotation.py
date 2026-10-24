"""Read-only quotation reconciliation with cell/page provenance."""
import contextlib
import io
from decimal import Decimal, InvalidOperation
from pathlib import Path
import json
import hashlib

def digest(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def compact(value):
    return ''.join(str(value or '').split())


def reconcile_quotation(workbook, pdf, start_page, end_page):
    import pymupdf
    from openpyxl import load_workbook
    before = {'workbook': digest(workbook), 'pdf': digest(pdf)}
    book = load_workbook(workbook, data_only=False)
    if len(book.worksheets) != 1: raise ValueError('Specify a single-sheet quotation workbook')
    sheet = book.active
    if sheet.max_column != 27: raise ValueError('Expected 27-column quotation schema')
    expected_headers = {1:'序号',4:'设备型号',7:'部件配件名称',8:'图号/规格型号',11:'不含税单价（元）',12:'含税单价（元）'}
    for column, title in expected_headers.items():
        if compact(sheet.cell(4,column).value) != title:
            raise ValueError('Unexpected quotation header; explicit schema mapping required')
    records = []; problems = []
    with pymupdf.open(pdf) as doc:
        if not 1 <= start_page <= end_page <= len(doc) or end_page-start_page>49:
            raise ValueError('Invalid page range; maximum 50 pages')
        for page_no in range(start_page,end_page+1):
            with contextlib.redirect_stdout(io.StringIO()):
                tables = doc[page_no-1].find_tables()
            for table in tables.tables:
                if table.col_count != 27:
                    problems.append({'page':page_no,'error':'Unexpected table width','columns':table.col_count});continue
                for table_row, cells in enumerate(table.extract()):
                    raw = [v or '' for v in cells]
                    first = compact(raw[0])
                    if first in ('序号',''):  # continuation rows retain their physical sources
                        if not first and any(raw):
                            if records:
                                records[-1]['cells'] = [a+'\n'+b if b else a for a,b in zip(records[-1]['cells'],raw)]
                                records[-1]['sources'].append({'page':page_no,'table_row':table_row,'raw':raw})
                            else: problems.append({'page':page_no,'error':'Orphan continuation row'})
                        continue
                    if not first.isdigit():
                        problems.append({'page':page_no,'error':'Unrecognized sequence','raw':first});continue
                    records.append({'sequence':int(first),'cells':raw,'sources':[{'page':page_no,'table_row':table_row,'raw':raw}]})
    by_id = {}
    for record in records: by_id.setdefault(record['sequence'],[]).append(record)
    rows=[];expected_sequence=[]
    for excel_row in range(5,sheet.max_row+1):
        original=[sheet.cell(excel_row,c).value for c in range(1,28)]
        seq=original[0];expected_sequence.append(seq); candidates=by_id.get(seq,[])
        item={'excel_row':excel_row,'sequence':seq,'status':'unmatched','problems':[]}
        if len(candidates)!=1:
            item['problems'].append('Missing or duplicate PDF sequence')
        else:
            record=candidates[0];values=[compact(v) for v in record['cells']]
            # Strong identity: order + exact model + exact name + exact drawing number.
            for index in (3,6,7):
                if compact(original[index]) != values[index]:
                    item['problems'].append({'column':index+1,'expected':original[index],'observed':values[index]})
            prices=[]
            for index in (10,11):
                try:
                    value=Decimal(values[index])
                    if not value.is_finite() or value<0: raise InvalidOperation
                    prices.append(str(value))
                except InvalidOperation:
                    item['problems'].append({'column':index+1,'error':'Invalid or missing price','raw':record['cells'][index]})
            item['pdf_sources']=record['sources']
            item['identity']={str(i+1):values[i] for i in (3,6,7)}
            if not item['problems']:
                item['status']='matched';item['untaxed_price']=prices[0];item['taxed_price']=prices[1]
        rows.append(item)
    if [r['sequence'] for r in records] != expected_sequence:
        problems.append({'error':'PDF and workbook sequence/order differ'})
    book.close()
    if before != {'workbook':digest(workbook),'pdf':digest(pdf)}: raise ValueError('Input changed during reconciliation')
    return {'inputs':before,'rows':rows,'problems':problems,'matched':sum(r['status']=='matched' for r in rows),
            'total':len(rows),'scope':'identity and extracted price matching; not commercial price approval',
            'normalization':'Whitespace from PDF wrapping removed for comparisons; raw cells retained.'}


def write_reconciliation(workbook, pdf, start_page, end_page, target):
    report = reconcile_quotation(workbook, pdf, start_page, end_page)
    with Path(target).open('x', encoding='utf-8') as out:
        json.dump(report, out, ensure_ascii=False, indent=2)
    return {'sha256': digest(target), 'matched': report['matched'], 'total': report['total'],
            'problems': report['problems'], 'input_hashes': report['inputs'],
            'unmatched': [{'sequence': r['sequence'], 'problems': r['problems']} for r in report['rows'] if r['status'] != 'matched']}


def restore_quotation(workbook, pdf, start_page, end_page, target):
    from openpyxl import load_workbook
    import tempfile
    report = reconcile_quotation(workbook, pdf, start_page, end_page)
    if report['problems'] or report['matched'] != report['total']:
        raise ValueError('Incomplete or ambiguous matching: no restored workbook created')
    book = load_workbook(workbook); sheet = book.active
    original = {(c.row,c.column): c.value for row in sheet for c in row}
    for row in report['rows']:
        for column, key in ((11,'untaxed_price'),(12,'taxed_price')):
            cell = sheet.cell(row['excel_row'],column)
            if cell.value is not None:
                raise ValueError('Price cell already populated: will not overwrite')
            cell.value = float(Decimal(row[key])); cell.number_format = '0.00'
    with tempfile.TemporaryDirectory() as tmp:
        file = Path(tmp)/'restored.xlsx'; book.save(file);book.close()
        check = load_workbook(file); actual=check.active
        for (row,column),value in original.items():
            if row >= 5 and column in (11,12): continue
            if actual.cell(row,column).value != value:
                raise ValueError('Non-price cell changed during workbook serialization')
        for row in report['rows']:
            for column,key in ((11,'untaxed_price'),(12,'taxed_price')):
                if Decimal(str(actual.cell(row['excel_row'],column).value)) != Decimal(row[key]):
                    raise ValueError('Price changed during serialization')
        check.close()
        if report['inputs'] != {'workbook':digest(workbook),'pdf':digest(pdf)}:
            raise ValueError('Input changed before publication')
        with file.open('rb') as src, Path(target).open('xb') as out:
            while chunk := src.read(1024*1024): out.write(chunk)
    return {'sha256':digest(target),'matched':report['matched'],'total':report['total'],
            'changed_columns':['K','L'],'input_hashes':report['inputs'],
            'purpose':'Historical gold-price restoration for validation; not a newly approved quotation'}
