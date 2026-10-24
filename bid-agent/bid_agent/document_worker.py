"""Runs with the dedicated document environment; no arbitrary code or shell input."""
import hashlib
import json
import sys
from pathlib import Path


def read_document(path, start_page=None, end_page=None):
    path = Path(path)
    if path.stat().st_size > 50_000_000:
        raise ValueError('Document exceeds 50 MB limit')
    suffix = path.suffix.lower()
    if suffix == '.docx':
        from docx import Document
        from docx.oxml.ns import qn
        doc = Document(path)
        blocks = []
        for index, element in enumerate(doc.element.body):
            if element.tag == qn('w:p'):
                blocks.append({'block': index, 'type': 'paragraph',
                               'text': ''.join(t.text or '' for t in element.iter(qn('w:t')))})
            elif element.tag == qn('w:tbl'):
                rows = []
                for row in element.findall(qn('w:tr')):
                    rows.append(['\n'.join(''.join(t.text or '' for t in p.iter(qn('w:t'))) for p in cell.findall(qn('w:p')))
                                 for cell in row.findall(qn('w:tc'))])
                blocks.append({'block': index, 'type': 'table', 'rows': rows})
        result = {'blocks': blocks, 'limitations': ['Main body XML extraction; headers, text boxes, nested tables and visual layout require separate inspection.']}
    elif suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        start = 1 if start_page is None else start_page
        end = len(reader.pages) if end_page is None else end_page
        if not 1 <= start <= end <= len(reader.pages): raise ValueError('Invalid page range')
        if start_page is not None and end-start > 19: raise ValueError('Read at most 20 pages per call')
        result = {'total_pages': len(reader.pages), 'pages': [{'page': n+1, 'text': reader.pages[n].extract_text() or ''} for n in range(start-1, end)],
                  'limitations': ['Text extraction only; scanned pages require OCR and visual inspection.']}
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False)
        result = {'sheets': {s.title: [[str(v) if v is not None else None for v in row] for row in s.iter_rows(values_only=True)] for s in book},
                  'limitations': ['Cell values and formulas only; appearance is not verified.']}
        book.close()
    else:
        raise ValueError('Supported: DOCX PDF XLSX')
    result['sha256'] = hashlib.sha256(path.read_bytes()).hexdigest()
    return result


if __name__ == '__main__':
    try:
        print(json.dumps({'ok': True, 'data': read_document(sys.argv[1], *[int(x) for x in sys.argv[2:]])}, ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({'ok': False, 'error': f'{type(exc).__name__}: {exc}'}))
        sys.exit(1)
