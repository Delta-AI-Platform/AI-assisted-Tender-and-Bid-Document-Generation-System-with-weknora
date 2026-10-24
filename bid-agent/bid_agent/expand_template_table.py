"""Prefreeze expansion of a flat, blank template table; no business defaults."""
from pathlib import Path
from copy import deepcopy
import hashlib, tempfile, zipfile
from lxml import etree
from docx import Document
from docx.oxml.ns import qn

def digest(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def expand_template_table(source,source_sha256,target,table_index,header_rows,prototype_row,sequence_column,records):
    if digest(source)!=source_sha256:raise ValueError('Source changed')
    if Path(target).exists():raise ValueError('Target exists')
    d=Document(source)
    if not 0<=table_index<len(d.tables):raise ValueError('Invalid table')
    t=d.tables[table_index];n=len(t.columns)
    if not 1<=header_rows<=prototype_row<len(t.rows):raise ValueError('Invalid prototype/header')
    if not 0<=sequence_column<n or not 1<=len(records)<=1000:raise ValueError('Invalid sequence/records')
    for row in records:
        if not isinstance(row,list) or len(row)!=n or any(not isinstance(v,str) for v in row):raise ValueError('Explicit string for every cell required')
    for row in t.rows[header_rows:]:
        if len(row._tr.tc_lst)!=n:raise ValueError('Merged/irregular template body unsupported')
        for j,cell in enumerate(row.cells):
            if any(x.tag in {qn('w:gridSpan'),qn('w:vMerge'),qn('w:drawing'),qn('w:sdt'),qn('w:fldChar'),qn('w:hyperlink')} for x in cell._tc.iter()):raise ValueError('Complex cell unsupported')
            if j!=sequence_column and cell.text.strip():raise ValueError('Refuse replacing populated body')
            if j==sequence_column and cell.text.strip() and not (cell.text.strip().isdigit() or cell.text.strip() in {'...','…'}):raise ValueError('Sequence body contains fixed text')
    proto=deepcopy(t.rows[prototype_row]._tr)
    for row in list(t.rows)[header_rows:]:t._tbl.remove(row._tr)
    for values in records:
        tr=deepcopy(proto);t._tbl.append(tr)
        for cell,value in zip(t.rows[-1].cells,values):
            p=cell.paragraphs[0];rpr=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
            p.clear()
            for extra in list(cell.paragraphs)[1:]:cell._tc.remove(extra._p)
            run=p.add_run(value)
            if rpr is not None:run._r.insert(0,rpr)
    xml=etree.tostring(d.element,encoding='UTF-8',xml_declaration=True,standalone=True)
    with tempfile.TemporaryDirectory() as tmp:
        f=Path(tmp)/'expanded.docx'
        with zipfile.ZipFile(source) as a,zipfile.ZipFile(f,'w') as b:
            for info in a.infolist():b.writestr(info,xml if info.filename=='word/document.xml' else a.read(info.filename))
        check=Document(f).tables[table_index]
        if [[c.text for c in row.cells] for row in check.rows[header_rows:]]!=records:raise ValueError('Readback failed')
        if digest(source)!=source_sha256:raise ValueError('Source changed')
        with f.open('rb') as src,Path(target).open('xb') as out:out.write(src.read())
    return {'sha256':digest(target),'records':len(records),'columns':n,'quality_status':'not_verified','purpose':'prefreeze template expansion; business values require source verification'}
