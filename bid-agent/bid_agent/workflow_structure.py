"""Server-built dynamic skeleton manifests bound to real project artifacts."""
from __future__ import annotations

import hashlib
import json
import os
import re
import subprocess
import tempfile
import time
from pathlib import Path

from .tools import STRING, schema


def _sha(path):
    with Path(path).open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def _write_new(folder, stem, payload):
    folder.mkdir(parents=True, exist_ok=True)
    sequence = 1
    while (folder / f'{stem}-v{sequence}.json').exists():
        sequence += 1
    path = folder / f'{stem}-v{sequence}.json'
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
    return path


def _source_pdf(registry, project):
    """Return the hash-locked primary PDF, never a similarly named attachment."""
    for path in sorted(registry.path('input').glob('*.pdf')):
        if _sha(path) == project['source_sha256']:
            return path
    raise ValueError('Hash-locked tender PDF is missing')


def _project_metadata(registry, project):
    """Extract only stable cover facts used to declare the skeleton slots."""
    import pymupdf
    source = _source_pdf(registry, project)
    pages=[]
    with pymupdf.open(source) as document:
        for index in range(min(20, len(document))):
            pages.append((index+1, document[index].get_text('text')))
    number=''
    for _, text in pages:
        match=re.search(r'(?:项目编号|招标编号|采购项目编号|采购编号)\s*[：:]\s*([^\s，,；;）)】]+)',text)
        if match:
            number=match.group(1).strip();break
    return {'project_title':project.get('title') or '', 'project_number':number, 'bidder_name':''}


def install_structure_tools(registry):
    section_schema = {
        'type': 'object', 'additionalProperties': False,
        'properties': {
            'section_id': STRING, 'order': {'type': 'integer', 'minimum': 1},
            'title': STRING, 'source_requirement_ids': {'type': 'array', 'items': STRING},
        }, 'required': ['section_id', 'order', 'title', 'source_requirement_ids'],
    }
    slot_schema = {
        'type': 'object', 'additionalProperties': False,
        'properties': {
            'slot_id': STRING,
            'slot_type': {'type': 'string', 'enum': ['text', 'table', 'evidence', 'price', 'signature', 'date', 'manual_response']},
            'data_type': STRING,
            'allowed_stages': {'type': 'array', 'items': {'type': 'integer', 'minimum': 4, 'maximum': 9}},
            'source_requirement_ids': {'type': 'array', 'items': STRING},
        }, 'required': ['slot_id', 'slot_type', 'data_type', 'allowed_stages', 'source_requirement_ids'],
    }
    table_schema = {
        'type': 'object', 'additionalProperties': False,
        'properties': {
            'table_id': STRING, 'title': STRING, 'docx_table_index': {'type': 'integer', 'minimum': 0},
            'source_requirement_ids': {'type': 'array', 'items': STRING},
        }, 'required': ['table_id', 'title', 'docx_table_index', 'source_requirement_ids'],
    }
    disposition_schema = {
        'type': 'object', 'additionalProperties': False,
        'properties': {
            'requirement_id': STRING,
            'disposition': {'type': 'string', 'enum': ['section', 'slot', 'fixed_table', 'internal_control_only']},
            'target_id': STRING, 'reason': STRING,
        }, 'required': ['requirement_id', 'disposition', 'target_id', 'reason'],
    }

    def build(tender_analysis_path, skeleton_docx_path, internal_render_pdf_path,
              sections, slots, fixed_tables, requirement_dispositions):
        project_path = registry.path('project.json'); project = json.loads(project_path.read_text())
        tender_path = registry.path(tender_analysis_path); tender = json.loads(tender_path.read_text())
        if tender.get('project_id') != project['project_id'] or tender.get('source_sha256') != project['source_sha256']:
            raise ValueError('Tender analysis belongs to another project/source')
        docx = registry.path(skeleton_docx_path); pdf = registry.path(internal_render_pdf_path)
        if docx.suffix.lower() != '.docx' or pdf.suffix.lower() != '.pdf' or not docx.is_file() or not pdf.is_file():
            raise ValueError('A real skeleton DOCX and its internal render PDF are required')
        from docx import Document
        from pypdf import PdfReader
        document = Document(docx)
        text = '\n'.join([p.text for p in document.paragraphs] + [c.text for t in document.tables for r in t.rows for c in r.cells])
        section_ids = [x['section_id'] for x in sections]
        if not sections or len(section_ids) != len(set(section_ids)) or [x['order'] for x in sections] != list(range(1, len(sections)+1)):
            raise ValueError('Skeleton sections must have unique IDs and continuous source order')
        if any(not x['title'].strip() or x['title'] not in text for x in sections):
            raise ValueError('Every declared skeleton section title must exist in the real DOCX')
        slot_ids = [x['slot_id'] for x in slots]
        if len(slot_ids) != len(set(slot_ids)) or any(not x['allowed_stages'] for x in slots):
            raise ValueError('Skeleton slots are duplicated or have no allowed stage')
        table_ids = [x['table_id'] for x in fixed_tables]
        if len(table_ids) != len(set(table_ids)) or any(x['docx_table_index'] >= len(document.tables) for x in fixed_tables):
            raise ValueError('Fixed table declaration is duplicated or outside the real DOCX')
        requirement_ids = [x['requirement_id'] for x in tender.get('items', [])]
        dispositions = [x['requirement_id'] for x in requirement_dispositions]
        if dispositions != requirement_ids or len(dispositions) != len(set(dispositions)):
            raise ValueError('Requirement dispositions must cover every tender-analysis item exactly once in source order')
        targets = {'section': set(section_ids), 'slot': set(slot_ids), 'fixed_table': set(table_ids)}
        for item in requirement_dispositions:
            if item['disposition'] == 'internal_control_only':
                if item['target_id']:
                    raise ValueError('Internal-only disposition must have an empty target_id')
            elif item['target_id'] not in targets[item['disposition']]:
                raise ValueError('Requirement disposition points to an unknown structure target')
        known = set(requirement_ids)
        for item in sections + slots + fixed_tables:
            if any(x not in known for x in item['source_requirement_ids']):
                raise ValueError('Structure item references an unknown tender requirement')
        page_count = len(PdfReader(str(pdf)).pages)
        if page_count < 1:
            raise ValueError('Internal render PDF has no pages')
        folder = registry.path('analysis/step3')
        bundle = _write_new(folder, 'skeleton-bundle', {
            'version': 'dynamic-skeleton-bundle-v1', 'project_id': project['project_id'], 'source_sha256': project['source_sha256'],
            'skeleton_docx_path': skeleton_docx_path, 'skeleton_docx_sha256': _sha(docx),
            'sections': sections, 'fixed_tables': fixed_tables,
        })
        slot_manifest = _write_new(folder, 'slot-manifest', {
            'version': 'dynamic-slot-manifest-v1', 'project_id': project['project_id'], 'source_sha256': project['source_sha256'], 'slots': slots,
        })
        comparison = _write_new(folder, 'structure-comparison', {
            'version': 'structure-comparison-v1', 'project_id': project['project_id'], 'source_sha256': project['source_sha256'],
            'tender_analysis_path': tender_analysis_path, 'tender_analysis_sha256': _sha(tender_path),
            'requirement_dispositions': requirement_dispositions, 'coverage': {'expected': len(requirement_ids), 'mapped': len(dispositions), 'rate': 1},
        })
        render = _write_new(folder, 'render-manifest', {
            'version': 'skeleton-render-manifest-v1', 'project_id': project['project_id'], 'source_sha256': project['source_sha256'],
            'pdf_path': internal_render_pdf_path, 'pdf_sha256': _sha(pdf), 'page_count': page_count,
            'section_geometry': [{'section_index': i, 'width_mm': round(s.page_width.mm, 3), 'height_mm': round(s.page_height.mm, 3)} for i, s in enumerate(document.sections)],
        })
        lock = _write_new(folder, 'skeleton-lock', {
            'version': 'dynamic-skeleton-lock-v1', 'project_id': project['project_id'], 'source_sha256': project['source_sha256'],
            'skeleton_id': hashlib.sha256((_sha(docx) + _sha(slot_manifest) + _sha(comparison)).encode()).hexdigest(),
            'bundle_path': str(bundle.relative_to(registry.root)), 'bundle_sha256': _sha(bundle),
            'slot_manifest_path': str(slot_manifest.relative_to(registry.root)), 'slot_manifest_sha256': _sha(slot_manifest),
            'structure_comparison_path': str(comparison.relative_to(registry.root)), 'structure_comparison_sha256': _sha(comparison),
            'render_manifest_path': str(render.relative_to(registry.root)), 'render_manifest_sha256': _sha(render),
            'created_at': time.time(),
        })
        return {
            'skeleton_bundle': str(bundle.relative_to(registry.root)), 'slot_manifest': str(slot_manifest.relative_to(registry.root)),
            'skeleton_lock': str(lock.relative_to(registry.root)), 'structure_comparison': str(comparison.relative_to(registry.root)),
            'render_manifest': str(render.relative_to(registry.root)), 'sections': len(sections), 'slots': len(slots),
            'fixed_tables': len(fixed_tables), 'requirements_mapped': len(dispositions), 'render_pages': page_count,
        }

    registry.add(
        'build_dynamic_skeleton_artifacts',
        '在真实可编辑骨架DOCX及其内部渲染PDF完成后调用。按当前tender_analysis核对章节、固定表、可填槽位及每条要求的结构去向，服务器一次生成并哈希锁定第3步五项JSON产物；禁止模型自行声称骨架冻结。',
        schema({
            'tender_analysis_path': STRING, 'skeleton_docx_path': STRING, 'internal_render_pdf_path': STRING,
            'sections': {'type': 'array', 'items': section_schema}, 'slots': {'type': 'array', 'items': slot_schema},
            'fixed_tables': {'type': 'array', 'items': table_schema},
            'requirement_dispositions': {'type': 'array', 'items': disposition_schema},
        }), build, True,
    )

    def build_project_skeleton_draft(tender_analysis_path):
        """Build a compact editable skeleton without streaming the tender back through the model."""
        project = json.loads(registry.path('project.json').read_text())
        tender_path = registry.path(tender_analysis_path); tender = json.loads(tender_path.read_text())
        if tender.get('project_id') != project['project_id'] or tender.get('source_sha256') != project['source_sha256']:
            raise ValueError('Tender analysis belongs to another project/source')
        items = tender.get('items') or []
        if not items:
            raise ValueError('Tender analysis is empty')
        labels = {'qualification':'资格与符合性要求','rejection':'否决与禁止事项','scoring':'评审与评分要求',
                  'business':'商务要求','technical':'技术要求','pricing':'报价要求','format':'格式要求','signature':'签章要求'}
        categories=[]
        for item in items:
            if item['category'] not in categories: categories.append(item['category'])
        sections=[]; slots=[]; dispositions=[]
        for order, category in enumerate(categories,1):
            category_items=[x for x in items if x['category']==category]
            section_id='section_'+category
            sections.append({'section_id':section_id,'order':order,'title':labels[category],
                             'source_requirement_ids':[x['requirement_id'] for x in category_items]})
            for item in category_items:
                rid=item['requirement_id'];slot_id='slot_'+hashlib.sha256(rid.encode()).hexdigest()[:16]
                slot_type={'pricing':'price','signature':'signature','technical':'manual_response'}.get(category,'text')
                slots.append({'slot_id':slot_id,'slot_type':slot_type,'data_type':'string',
                              'allowed_stages':[4,5,6,7,8,9], 'source_requirement_ids':[rid]})
        for item in items:
            dispositions.append({'requirement_id':item['requirement_id'],'disposition':'section','target_id':'section_'+item['category'],
                                 'reason':'按招标来源顺序归入对应要求章节；填写槽位保持空白'})
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        folder=registry.path('analysis/step3');folder.mkdir(parents=True,exist_ok=True)
        docx=folder/'dynamic-skeleton.docx';pdf=folder/'dynamic-skeleton.pdf'
        document=Document();normal=document.styles['Normal'];normal.font.name='SimSun';normal._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
        title=document.add_heading(project.get('title') or '投标文件',0);title.alignment=WD_ALIGN_PARAGRAPH.CENTER
        marker=document.add_paragraph('Agent范围草稿骨架；待业务核对；不可递交');marker.alignment=WD_ALIGN_PARAGRAPH.CENTER
        metadata=_project_metadata(registry,project)
        metadata_table=document.add_table(rows=3,cols=2);metadata_table.style='Table Grid'
        metadata_rows=[('项目名称','slot_project_title','project_title'),('项目编号','slot_project_number','project_number'),('投标人','slot_bidder_name','bidder_name')]
        for row,(label,slot_id,key) in zip(metadata_table.rows,metadata_rows):
            row.cells[0].text=label;row.cells[1].text=metadata[key]
            slots.insert(0,{'slot_id':slot_id,'slot_type':'text','data_type':'string','allowed_stages':[4,9],'source_requirement_ids':[]})
        fixed_tables=[{'table_id':'table_project_identity','title':'项目基本信息','docx_table_index':0,'source_requirement_ids':[]}]
        for section in sections:
            document.add_heading(section['title'],level=1)
            for item in [x for x in items if x['category']==section['section_id'][8:]]:
                document.add_heading(item['requirement_id'],level=2)
                document.add_paragraph(item['source_quote'])
                document.add_paragraph('来源：物理页%s，%s'%(item['physical_page'],item.get('source_location') or ''))
                slot_id=next(x['slot_id'] for x in slots if x['source_requirement_ids']==[item['requirement_id']])
                table=document.add_table(rows=2,cols=2);table.style='Table Grid'
                table.cell(0,0).text='填写项';table.cell(0,1).text='内容'
                table.cell(1,0).text=slot_id;table.cell(1,1).text=''
        document.save(docx)
        with tempfile.TemporaryDirectory(prefix='deta-lo-') as profile:
            command=['/usr/bin/libreoffice','--headless',f'-env:UserInstallation=file://{profile}','--convert-to','pdf','--outdir',str(folder),str(docx)]
            completed=subprocess.run(command,capture_output=True,text=True,timeout=300,env={**os.environ,'HOME':'/home/ubuntu'})
            if completed.returncode or not pdf.is_file():
                raise RuntimeError('LibreOffice skeleton render failed: '+(completed.stderr or completed.stdout)[-500:])
        result=build(tender_analysis_path,str(docx.relative_to(registry.root)),str(pdf.relative_to(registry.root)),
                     sections,slots,fixed_tables,dispositions)
        result.update({'skeleton_docx':str(docx.relative_to(registry.root)),'internal_render_pdf':str(pdf.relative_to(registry.root))})
        return result

    draft_description='服务器直接读取当前项目tender_analysis，按真实来源顺序生成可编辑DOCX空白骨架、内部PDF及五项哈希锁定产物。用于长文件，避免Agent逐页重读原件造成上下文膨胀；所有响应槽保持空白。'
    draft_schema=schema({'tender_analysis_path': STRING})
    registry.add('a_build_project_skeleton_draft',draft_description,draft_schema,build_project_skeleton_draft,True)
    registry.add('build_project_skeleton_draft',draft_description,draft_schema,build_project_skeleton_draft,True)
