"""Explicit annotations on generated technical review sheets; requirements immutable."""
import json,hashlib,tempfile
from pathlib import Path
from zipfile import ZipFile
from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

def annotate_technical_workbook(source,source_sha256,target,records_file,records_sha256):
    try:
        from .document_ops import digest,publish
    except ImportError:
        from document_ops import digest,publish
    if digest(source)!=source_sha256:raise ValueError('Workbook version changed')
    raw=Path(records_file).read_bytes()
    if hashlib.sha256(raw).hexdigest()!=records_sha256:raise ValueError('Annotation records version changed')
    def unique(pairs):
        result={}
        for k,v in pairs:
            if k in result:raise ValueError('Duplicate JSON key')
            result[k]=v
        return result
    records=json.loads(raw,object_pairs_hook=unique)
    edits=records.get('edits') if isinstance(records,dict) else records
    if not isinstance(edits,list) or not 1<=len(edits)<=1000:raise ValueError('Provide 1–1000 explicit annotations')
    with ZipFile(source) as z:
        if any(any(part in n for part in ('drawings/','vbaProject','embeddings/','externalLinks/','pivotTables/')) for n in z.namelist()):raise ValueError('Unsupported workbook package features')
    wb=load_workbook(source);sheet=wb['技术对标'];index=wb['证据索引']
    if [sheet.cell(2,c).value for c in range(1,7)]!=['子系统模块','部件模块','功能模块','本次招标要求','投标人响应','偏离说明']:raise ValueError('Unexpected review template')
    ids={}
    for row in index.iter_rows(min_row=2,values_only=True):
        ident,number=row[:2]
        if not isinstance(ident,str) or ident in ids or type(number) is not int or not 3<=number<=sheet.max_row:raise ValueError('Invalid requirement index')
        quote=sheet.cell(number,4).value
        if not isinstance(quote,str) or hashlib.sha256(quote.encode()).hexdigest()!=row[4]:raise ValueError('Requirement quote fingerprint mismatch')
        ids[ident]=number
    if len(set(ids.values()))!=len(ids):raise ValueError('Aliased requirement rows')
    selected=set()
    for e in edits:
        if not isinstance(e,dict) or set(e)!={'requirement_id','column','before','after'}:raise ValueError('Explicit ID,column,before,after required')
        number=ids.get(e['requirement_id']);col=e['column']
        if number is None or type(col) is not int or col not in (1,2,3,5,6):raise ValueError('Unknown requirement or protected column')
        if (number,col) in selected:raise ValueError('Duplicate annotation cell')
        selected.add((number,col));cell=sheet.cell(number,col)
        if cell.value!=e['before']:raise ValueError('Annotation before value mismatch')
        if not isinstance(e['after'],str) or not e['after'] or len(e['after'])>32767 or ILLEGAL_CHARACTERS_RE.search(e['after']):raise ValueError('Invalid annotation text')
    before={(s.title,c.coordinate):(c.value,c.data_type,c.style_id) for s in wb for row in s for c in row}
    for e in edits:
        cell=sheet.cell(ids[e['requirement_id']],e['column']);cell.value=e['after'];cell.data_type='s'
    with tempfile.TemporaryDirectory() as tmp:
        f=Path(tmp)/'annotated.xlsx';wb.save(f);check=load_workbook(f,data_only=False)
        allowed={('技术对标',sheet.cell(r,c).coordinate) for r,c in selected}
        for key,value in before.items():
            cell=check[key[0]][key[1]]
            if key not in allowed and (cell.value,cell.data_type,cell.style_id)!=value:raise ValueError('Protected workbook cell changed')
        for e in edits:
            cell=check['技术对标'].cell(ids[e['requirement_id']],e['column'])
            if cell.value!=e['after'] or cell.data_type!='s':raise ValueError('Annotation readback mismatch')
        if digest(source)!=source_sha256:raise ValueError('Source changed during annotation')
        return {**publish(f,target),'source_sha256':source_sha256,'records_sha256':records_sha256,'annotated_cells':len(edits),'requirements_preserved':True,'quality_status':'not_verified','limitations':['Explicit annotations only; semantic correctness and coverage not approved','Evidence status remains unverified; formula-like annotations stored as text']}
