"""Create an internal review workbook from version-bound captured requirements."""
import hashlib
import json
import tempfile
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter


def create_technical_workbook(source, source_sha256, target):
    try:
        from .document_ops import digest, publish
    except ImportError:
        from document_ops import digest, publish
    if digest(source) != source_sha256:
        raise ValueError('Source version changed')
    data = json.loads(Path(source).read_text())
    rows = data.get('requirements')
    if not isinstance(rows, list) or not rows:
        raise ValueError('Nonempty captured requirements required')
    seen = set()
    for row in rows:
        ident, quote = row.get('requirement_id'), row.get('quote')
        if not isinstance(ident, str) or not ident or len(ident) > 32767 or ILLEGAL_CHARACTERS_RE.search(ident) or ident in seen:
            raise ValueError('Unique stable requirement IDs required')
        seen.add(ident)
        if not isinstance(quote, str) or not quote.strip():
            raise ValueError('Original quote required')
        if len(quote) > 32767 or ILLEGAL_CHARACTERS_RE.search(quote):
            raise ValueError('Quote cannot be stored losslessly in one Excel cell')
        pages = row.get('source_physical_pages')
        if not isinstance(pages, list) or not pages or any(type(p) is not int or p < 1 for p in pages):
            raise ValueError('Source page references required')
    wb = Workbook()
    sheet = wb.active
    sheet.title = '技术对标'
    sheet.append(['内部技术审查稿：仅转换所提供要求清单，未证明完整覆盖；未选定车型报告，不生成标准答案列；响应和偏离待证据核实。'])
    sheet.merge_cells('A1:F1')
    sheet['A1'].alignment = Alignment(wrap_text=True, vertical='top')
    sheet.row_dimensions[1].height = 45
    sheet.append(['子系统模块', '部件模块', '功能模块', '本次招标要求', '投标人响应', '偏离说明'])
    evidence = wb.create_sheet('证据索引')
    evidence.append(['要求ID', '对标行', '原物理页', '要求清单SHA256', '原文SHA256', '证据状态'])
    checks = wb.create_sheet('核对记录')
    checks.append(['项目', '记录'])
    checks.append(['范围', '要求顺序与输入清单一致；不是完整标书或已确认符合性结论。'])
    checks.append(['长文本显示', '单元格保留全文；Excel显示高度有上限，须展开编辑栏检查，不以打印预览证明全文可见。'])
    checks.append(['车型与报告', '尚未选择或绑定报告；不得从历史车型填入参数。'])
    checks.append(['模块分类', '未作语义匹配，前三列保持空白。'])
    for number, row in enumerate(rows, 3):
        cell = sheet.cell(number, 4, row['quote'])
        cell.data_type = 's'  # Text beginning with '=' must never become a formula.
        for c in sheet[number]:
            c.alignment = Alignment(wrap_text=True, vertical='top')
        sheet.row_dimensions[number].height = min(300, max(45, (len(row['quote']) // 45 + 1) * 15))
        evidence.append([row['requirement_id'], number, ','.join(map(str, row['source_physical_pages'])),
                         source_sha256, hashlib.sha256(row['quote'].encode()).hexdigest(), '待核实'])
        evidence.cell(number-1, 1).data_type = 's'
    for s in wb:
        s.freeze_panes = 'D3' if s == sheet else 'A2'
        for cell in s[2 if s == sheet else 1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill('solid', fgColor='DCEAF7')
        for col in range(1, s.max_column+1):
            s.column_dimensions[get_column_letter(col)].width = 24
    sheet.column_dimensions['D'].width = 85
    sheet.column_dimensions['E'].width = 60
    sheet.column_dimensions['F'].width = 30
    sheet.auto_filter.ref = f'A2:F{sheet.max_row}'
    checks.column_dimensions['B'].width = 110
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp)/'review.xlsx'
        wb.save(out)
        saved = load_workbook(out, data_only=False)
        for number, row in enumerate(rows, 3):
            cell = saved['技术对标'].cell(number, 4)
            if cell.value != row['quote'] or cell.data_type != 's':
                raise ValueError('Saved requirement changed')
            if any(saved['技术对标'].cell(number, col).value is not None for col in (1, 2, 3, 5, 6)):
                raise ValueError('Unexpected inferred response or classification')
        if digest(source) != source_sha256:
            raise ValueError('Source changed during conversion')
        return {**publish(out, target), 'source_sha256': source_sha256, 'requirements': len(rows),
                'scope': 'Internal unfilled workbook; exact supplied quote preservation only',
                'quality_status': 'not_verified'}
