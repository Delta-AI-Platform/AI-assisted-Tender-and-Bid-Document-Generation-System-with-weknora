"""Guarded table edits. No inferred values and no merged-cell aliases."""
import hashlib
import tempfile
from pathlib import Path
from zipfile import ZipFile
from copy import deepcopy


def edit_table_cells(source, source_sha256, target, edits):
    from docx import Document
    from docx.oxml.ns import qn
    from lxml import etree
    try:
        from .document_ops import digest, publish
    except ImportError:
        from document_ops import digest, publish
    if digest(source) != source_sha256:
        raise ValueError('Source version changed')
    if not edits or len(edits)>1000:
        raise ValueError('Provide 1–1000 explicitly bound cell edits')
    doc=Document(source)
    before=deepcopy(doc.element)
    selected=[]; seen=set()
    for edit in edits:
        indices=[edit[k] for k in ('table','row','column')]
        if any(type(i) is not int or i<0 for i in indices):
            raise ValueError('Physical indices must be nonnegative integers')
        table,row,column=indices
        cell=doc.tables[table].rows[row].cells[column]
        if cell._tc in seen: raise ValueError('Duplicate or merged-cell alias')
        seen.add(cell._tc)
        if not isinstance(edit['before'],str) or not isinstance(edit['after'],str):
            raise ValueError('Before/after must be explicit strings')
        if cell.text != edit['before']: raise ValueError('Cell text binding mismatch')
        if any(c.tag not in {qn('w:tcPr'),qn('w:p')} for c in cell._tc):
            raise ValueError('Nested table or complex cell cannot be flattened')
        if len(cell.paragraphs)!=1:
            raise ValueError('Multi-paragraph cell needs precise paragraph editing')
        p=cell.paragraphs[0]
        if any(c.tag not in {qn('w:pPr'),qn('w:r')} for c in p._p):
            raise ValueError('Complex cell paragraph cannot be flattened')
        for run in p.runs:
            for child in run._r:
                if child.tag in {qn('w:rPr'),qn('w:t')}:
                    continue
                # Plain text line breaks are represented in cell.text and are
                # therefore covered by the exact before/after binding. Page,
                # column and clearing breaks carry layout semantics: reject.
                if (child.tag == qn('w:br') and not len(child)
                        and set(child.attrib).issubset({qn('w:type')})
                        and child.get(qn('w:type'), 'textWrapping') == 'textWrapping'):
                    continue
                raise ValueError('Cell contains fields, drawings or other protected content')
        if len(p.runs)>1: raise ValueError('Mixed-run cell requires precise editing')
        selected.append((cell,p,edit))
    for cell,p,edit in selected:
        if p.runs: p.runs[0].text=edit['after']
        else: p.add_run(edit['after'])
    # Exclude only explicitly authorized paragraph contents from structural
    # comparison. All other XML, including adjacent text and table properties,
    # must remain identical.
    after=deepcopy(doc.element)
    tree=doc.element.getroottree()
    for cell,p,edit in selected:
        path=tree.getpath(p._p)
        for root in (before,after):
            node=root.getroottree().xpath(path,namespaces=doc.element.nsmap)[0]
            for child in list(node):
                if child.tag!=qn('w:pPr'): node.remove(child)
    if etree.tostring(before)!=etree.tostring(after):
        raise ValueError('Protected document structure changed')
    if digest(source)!=source_sha256: raise ValueError('Source changed during editing')
    with tempfile.TemporaryDirectory() as tmp:
        out=Path(tmp)/'edited.docx';doc.save(out)
        with ZipFile(source) as a,ZipFile(out) as b:
            protected={name:hashlib.sha256(a.read(name)).hexdigest() for name in a.namelist() if name!='word/document.xml'}
            changed=[name for name,sha in protected.items() if name not in b.namelist() or hashlib.sha256(b.read(name)).hexdigest()!=sha]
            if changed: raise ValueError('Protected package parts changed: '+', '.join(changed))
        verified=Document(out)
        for _,_,edit in selected:
            if verified.tables[edit['table']].cell(edit['row'],edit['column']).text!=edit['after']:
                raise ValueError('Saved cell value differs')
        return {**publish(out,target),'source_sha256':source_sha256,'edited_cells':len(edits),
                'scope':'Only bound cell text edited; protected XML and package parts unchanged; no factual or visual approval'}
