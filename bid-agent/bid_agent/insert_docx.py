"""Insert an editable source at an explicitly bound unique placeholder."""
from copy import deepcopy
from docx.oxml import OxmlElement
import hashlib
import tempfile
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn
from docxcompose.composer import Composer


def insert_docx(source, source_sha256, fragment, fragment_sha256, marker, target, header_policy, section_policy="master", placement="replace"):
    try:
        from .document_ops import digest, publish
    except ImportError:
        from document_ops import digest, publish
    if digest(source) != source_sha256 or digest(fragment) != fragment_sha256:
        raise ValueError('Source or fragment version changed')
    if header_policy != 'master':
        raise ValueError('Explicit master header/footer policy required')
    if not isinstance(marker, str) or not marker.strip():
        raise ValueError('Explicit unique placeholder text required')
    base, part = Document(source), Document(fragment)
    text = lambda element: ''.join(t.text or '' for t in element.iter(qn('w:t')))
    hits = [p for p in base.element.body if p.tag == qn('w:p') and text(p) == marker]
    if len(hits) != 1:
        raise ValueError('Marker must match exactly one top-level paragraph')
    anchor = hits[0]
    if anchor.xpath('.//w:drawing|.//w:pict|.//w:fldChar|.//w:instrText|.//w:sectPr|.//w:sdt'):
        raise ValueError('Placeholder contains protected non-text content')
    if placement not in {'replace','before','after'}:raise ValueError('Explicit valid placement required')
    index = base.element.body.index(anchor)
    left = ''.join(text(e) for e in list(base.element.body)[:index])
    right = ''.join(text(e) for e in list(base.element.body)[index+1:])
    expected = (left+text(part.element.body)+right if placement=='replace' else left+text(part.element.body)+text(anchor)+right if placement=='before' else left+text(anchor)+text(part.element.body)+right)
    if placement=='after':index+=1
    images = set()
    for path in (source, fragment):
        with ZipFile(path) as z:
            images.update(hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/'))
    if section_policy not in {'master','preserve_fragment'}:raise ValueError('Explicit valid section policy required')
    if section_policy=='master' and part.element.xpath('//w:drawing|//w:pict'):
        raise ValueError('Image fragments require preserve_fragment section policy')
    if section_policy=='preserve_fragment':
        # Preserve the current master section before the inserted material.
        active=None
        for block in list(base.element.body)[index:]:
            candidates=([block] if block.tag==qn('w:sectPr') else block.xpath('./w:pPr/w:sectPr'))
            if candidates:active=deepcopy(candidates[0]);break
        if active is None:raise ValueError('Master section not found')
        if index:
            previous=base.element.body[index-1]
            if previous.tag!=qn('w:p'):
                previous=OxmlElement('w:p');base.element.body.insert(index,previous);index+=1
            pr=previous.find(qn('w:pPr'))
            if pr is None:pr=OxmlElement('w:pPr');previous.insert(0,pr)
            if pr.find(qn('w:sectPr')) is None:pr.append(active)
        # docxcompose omits the fragment's terminal sectPr. Move a copy to
        # its last content paragraph so the last proof page keeps its geometry.
        blocks=list(part.element.body);last=blocks[-2]
        if last.tag!=qn('w:p'):
            last=OxmlElement('w:p');part.element.body.insert(len(blocks)-1,last)
        pr=last.find(qn('w:pPr'))
        if pr is None:pr=OxmlElement('w:pPr');last.insert(0,pr)
        if pr.find(qn('w:sectPr')) is not None:raise ValueError('Fragment tail already closes a section')
        pr.append(deepcopy(part.element.body.sectPr))
    composer = Composer(base, preserve_styles=True)
    if placement=='replace':base.element.body.remove(anchor)
    composer.insert(index, part)
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp)/'inserted.docx';composer.save(out)
        if text(Document(out).element.body) != expected:
            raise ValueError('Inserted document text differs or appears in wrong order')
        with ZipFile(out) as z:
            observed = {hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
        if not images.issubset(observed):
            raise ValueError('Source image missing after insertion')
        if digest(source) != source_sha256 or digest(fragment) != fragment_sha256:
            raise ValueError('Source changed during insertion')
        return {**publish(out,target),'source_sha256':source_sha256,'fragment_sha256':fragment_sha256,
                'anchor_text':marker,'placement':placement,'body_index':index,'header_policy':header_policy,'section_policy':section_policy,
                'body_text_in_exact_order':True,'source_media_retained':True,'quality_status':'not_verified',
                'limitations':['Requires rendered review of numbering, tables, section geometry and pagination',
                               'Anchor retained for before/after; removed only for replace; master headers apply']}
