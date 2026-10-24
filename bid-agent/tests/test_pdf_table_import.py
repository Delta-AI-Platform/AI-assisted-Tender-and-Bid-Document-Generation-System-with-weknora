import tempfile
import unittest
from pathlib import Path

import pymupdf
from docx import Document
from bid_agent.document_ops import digest
from bid_agent.pdf_table_import import import_table


class PdfTableImportTests(unittest.TestCase):
    def source(self, path, second=False):
        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_text((50, 45), 'Required form')
        for y in (70, 100, 130):
            page.draw_line((50, y), (350, y))
        for x in (50, 350):
            page.draw_line((x, 70), (x, 130))
        page.draw_line((200, 100), (200, 130))
        page.insert_text((60, 90), 'Merged heading')
        page.insert_text((60, 120), 'Fixed requirement')
        page.insert_text((210, 120), 'Response')
        page.insert_text((50, 160), 'Keep this mandatory note.')
        if second:
            for y in (200, 230, 260):
                page.draw_line((50, y), (350, y))
            for x in (50, 200, 350):
                page.draw_line((x, 200), (x, 260))
            page.insert_text((60, 220), 'Second form')
        pdf.save(path)
        pdf.close()

    def test_merged_heading_and_mandatory_note_survive(self):
        with tempfile.TemporaryDirectory() as tmp:
            src, out = Path(tmp)/'source.pdf', Path(tmp)/'form.docx'
            self.source(src)
            before = digest(src)
            import_table(src, 1, 0, out, True, before)
            doc = Document(out)
            table = doc.tables[0]
            self.assertEqual(table.cell(0, 0).text, 'Merged heading')
            self.assertIs(table.cell(0, 0)._tc, table.cell(0, 1)._tc)
            self.assertEqual(table.cell(1, 0).text, 'Fixed requirement')
            self.assertEqual(table.cell(1, 1).text, 'Response')
            self.assertIn('Keep this mandatory note.', [p.text for p in doc.paragraphs])
            self.assertEqual(digest(src), before)

    def test_ambiguous_multiple_forms_and_stale_source_do_not_publish(self):
        with tempfile.TemporaryDirectory() as tmp:
            src, out = Path(tmp)/'source.pdf', Path(tmp)/'form.docx'
            self.source(src, second=True)
            with self.assertRaisesRegex(ValueError, 'single table'):
                import_table(src, 1, 0, out, True, digest(src))
            self.assertFalse(out.exists())
            with self.assertRaisesRegex(ValueError, 'version binding'):
                import_table(src, 1, 0, out, False, 'stale')
            self.assertFalse(out.exists())
