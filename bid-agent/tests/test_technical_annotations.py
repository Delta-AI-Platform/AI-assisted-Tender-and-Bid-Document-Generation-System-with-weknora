import json,tempfile,unittest
from pathlib import Path
from openpyxl import load_workbook
from bid_agent.document_ops import digest
from bid_agent.technical_workbook import create_technical_workbook
from bid_agent.technical_annotations import annotate_technical_workbook
from bid_agent.store import Store
from bid_agent.tools import ProjectTools

class TechnicalAnnotationTests(unittest.TestCase):
 def test_registered_annotations_preserve_quote_and_never_evaluate_formula(self):
  with tempfile.TemporaryDirectory() as tmp:
   root=Path(tmp);src=root/'q.json';src.write_text(json.dumps({'requirements':[{'requirement_id':'R1','quote':'完整原文 ≥ 1000','source_physical_pages':[5]}]}));create_technical_workbook(src,digest(src),root/'a.xlsx');r=root/'edits.json';r.write_text(json.dumps({'edits':[{'requirement_id':'R1','column':5,'before':None,'after':'=HYPERLINK("not-executed")'}]}));store=Store(root/'db');run=store.create('test',{'project':str(root)});tools=ProjectTools(root,store,run)
   result=tools.tools['annotate_technical_workbook'].handler(source='a.xlsx',source_sha256=digest(root/'a.xlsx'),target='out.xlsx',records_file='edits.json',records_sha256=digest(r));w=load_workbook(root/'out.xlsx');self.assertEqual(w['技术对标']['D3'].value,'完整原文 ≥ 1000');self.assertEqual(w['技术对标']['E3'].data_type,'s');self.assertEqual(w['证据索引']['F2'].value,'待核实');self.assertTrue(result['requirements_preserved'])
 def test_protected_quote_unknown_id_stale_value_and_duplicate_rejected(self):
  with tempfile.TemporaryDirectory() as tmp:
   root=Path(tmp);src=root/'q.json';src.write_text(json.dumps({'requirements':[{'requirement_id':'R1','quote':'固定要求','source_physical_pages':[5]}]}));create_technical_workbook(src,digest(src),root/'a.xlsx');e={'requirement_id':'R1','column':5,'before':None,'after':'历史内容'};f=root/'r.json'
   for edits in [[{**e,'column':4}],[{**e,'requirement_id':'missing'}],[{**e,'before':'wrong'}],[e,e]]:
    f.write_text(json.dumps({'edits':edits}))
    with self.assertRaises(ValueError):annotate_technical_workbook(root/'a.xlsx',digest(root/'a.xlsx'),root/'bad.xlsx',f,digest(f))
    self.assertFalse((root/'bad.xlsx').exists())
   w=load_workbook(root/'a.xlsx');w['技术对标']['D3']='篡改要求';w.save(root/'a.xlsx');f.write_text(json.dumps({'edits':[e]}))
   with self.assertRaisesRegex(ValueError,'fingerprint'):annotate_technical_workbook(root/'a.xlsx',digest(root/'a.xlsx'),root/'bad.xlsx',f,digest(f))
