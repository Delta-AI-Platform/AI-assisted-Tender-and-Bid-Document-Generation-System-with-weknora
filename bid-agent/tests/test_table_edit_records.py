import json,tempfile,unittest
from pathlib import Path
from docx import Document
from bid_agent.document_ops import digest
from bid_agent.store import Store
from bid_agent.tools import ProjectTools
from bid_agent.table_records import edit_table_cells_from_records

class TableEditRecordsTests(unittest.TestCase):
 def test_registered_snapshot_exact_values_and_stale_rejection(self):
  with tempfile.TemporaryDirectory() as tmp:
   root=Path(tmp);d=Document();t=d.add_table(rows=2,cols=2);t.cell(0,0).text='固定';t.cell(1,0).text='原\n值';d.save(root/'s.docx')
   edits=[{'table':0,'row':1,'column':0,'before':'原\n值','after':'原值'},{'table':0,'row':1,'column':1,'before':'','after':'明确值'}]
   records=root/'records.json';records.write_text(json.dumps({'edits':edits},ensure_ascii=False));sha=digest(records)
   store=Store(root/'db');run=store.create('test',{'project':str(root)});tools=ProjectTools(root,store,run)
   r=tools.tools['edit_table_cells_from_records'].handler(source='s.docx',source_sha256=digest(root/'s.docx'),target='out.docx',records_file='records.json',records_sha256=sha)
   self.assertEqual(r['edited_cells'],2);self.assertEqual(r['records_file_sha256'],sha)
   saved=Document(root/'out.docx');self.assertEqual(saved.tables[0].cell(0,0).text,'固定');self.assertEqual(saved.tables[0].cell(1,0).text,'原值');self.assertEqual(saved.tables[0].cell(1,1).text,'明确值')
   records.write_text('{}')
   with self.assertRaisesRegex(ValueError,'version mismatch'):edit_table_cells_from_records(root/'s.docx',digest(root/'s.docx'),root/'bad.docx',records,sha)
   self.assertFalse((root/'bad.docx').exists())
 def test_duplicate_keys_stale_values_and_merged_alias_reject(self):
  with tempfile.TemporaryDirectory() as tmp:
   root=Path(tmp);d=Document();t=d.add_table(rows=1,cols=2);t.cell(0,0).merge(t.cell(0,1));d.save(root/'s.docx');f=root/'r.json'
   for content in ['{"edits":[],"edits":[]}',json.dumps([{'table':0,'row':0,'column':0,'before':'wrong','after':'x'}]),json.dumps([{'table':0,'row':0,'column':i,'before':'','after':'x'} for i in (0,1)])]:
    f.write_text(content)
    with self.assertRaises(ValueError):edit_table_cells_from_records(root/'s.docx',digest(root/'s.docx'),root/'bad.docx',f,digest(f))
    self.assertFalse((root/'bad.docx').exists())
