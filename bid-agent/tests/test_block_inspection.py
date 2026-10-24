import tempfile,unittest
from pathlib import Path
from docx import Document
from bid_agent.block_inspection import inspect_blocks
from bid_agent.document_ops import edit_docx,digest
from bid_agent.store import Store
from bid_agent.tools import ProjectTools

class BlockInspectionTests(unittest.TestCase):
 def test_filtered_indices_drive_guarded_edit_and_package_unchanged(self):
  with tempfile.TemporaryDirectory() as tmp:
   p=Path(tmp);d=Document();d.add_paragraph('重复');d.add_table(rows=1,cols=1).cell(0,0).text='表格';d.add_paragraph('重复');d.add_paragraph('唯一目标');d.save(p/'s.docx');before=digest(p/'s.docx')
   store=Store(p/'db');run=store.create('test',{'project':str(p)});tools=ProjectTools(p,store,run)
   result=tools.tools['inspect_docx_blocks'].handler(path='s.docx',offset=0,count=2,query='重复')
   self.assertEqual([b['block'] for b in result['blocks']],[0,2]);self.assertEqual(result['matching_blocks'],2)
   target=inspect_blocks(p/'s.docx',0,1,'唯一目标')['blocks'][0]
   edit_docx(p/'s.docx',result['sha256'],p/'out.docx',[{'block':target['block'],'before':target['text'],'after':'已准确定位'}])
   self.assertEqual(Document(p/'out.docx').paragraphs[-1].text,'已准确定位');self.assertEqual(digest(p/'s.docx'),before)
   with self.assertRaises(ValueError):tools.tools['inspect_docx_blocks'].handler(path='../outside.docx',offset=0,count=1,query='')
 def test_bounds_pagination_and_explicit_truncation(self):
  with tempfile.TemporaryDirectory() as tmp:
   p=Path(tmp)/'s.docx';d=Document();d.add_paragraph('x'*5000);d.add_paragraph('last');d.save(p)
   a=inspect_blocks(p,0,1,'');self.assertTrue(a['blocks'][0]['text_truncated']);self.assertEqual(len(a['blocks'][0]['text']),4000);self.assertEqual(a['next_offset'],1)
   self.assertEqual(inspect_blocks(p,1,1,'')['blocks'][0]['text'],'last')
   for args in [(-1,1,''),(0,41,''),(False,1,''),(0,1,'x'*301)]:
    with self.assertRaises(ValueError):inspect_blocks(p,*args)
