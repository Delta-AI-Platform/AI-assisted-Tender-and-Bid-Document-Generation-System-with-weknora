import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from lxml import etree
from bid_agent.section_layout import format_docx,NS
from bid_agent.document_ops import digest

class SectionLayoutTests(unittest.TestCase):
    def test_only_selected_geometry_changes(self):
        with tempfile.TemporaryDirectory() as tmp:
            r=Path(tmp);d=Document();d.add_paragraph('固定原文');d.add_section();d.add_paragraph('证明原页');d.save(r/'in.docx')
            args=dict(source=r/'in.docx',source_sha256=digest(r/'in.docx'),target=r/'out.docx',sections=[dict(section_index=0,width_mm=210,height_mm=297,top_mm=25,bottom_mm=25,left_mm=27,right_mm=27)])
            format_docx(**args)
            with ZipFile(r/'in.docx') as a,ZipFile(r/'out.docx') as b:
                for name in a.namelist():
                    if name!='word/document.xml':self.assertEqual(a.read(name),b.read(name))
                before=etree.fromstring(a.read('word/document.xml'));after=etree.fromstring(b.read('word/document.xml'))
                self.assertEqual(etree.tostring(before.xpath('//w:sectPr',namespaces=NS)[1]),etree.tostring(after.xpath('//w:sectPr',namespaces=NS)[1]))
                for tree in (before,after):
                    node=tree.xpath('//w:sectPr',namespaces=NS)[0]
                    for child in node.xpath('./w:pgSz|./w:pgMar',namespaces=NS):node.remove(child)
                self.assertEqual(etree.tostring(before),etree.tostring(after))
            with self.assertRaises(ValueError):format_docx(**{**args,'target':r/'bad.docx','sections':[]})
            self.assertFalse((r/'bad.docx').exists())
