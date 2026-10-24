import unittest,tempfile
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from bid_agent.document_ops import digest,render_pdf
from bid_agent.paragraph_pagination import set_paragraph_pagination
import pymupdf

class ParagraphPaginationTests(unittest.TestCase):
 def test_page_break_moves_second_form_and_preserves_text_tables_and_parts(self):
  with tempfile.TemporaryDirectory() as tmp:
   root=Path(tmp);d=Document();d.add_paragraph('First person');d.add_table(rows=1,cols=1).cell(0,0).text='First CV';d.add_paragraph('Second person');d.add_table(rows=1,cols=1).cell(0,0).text='Second CV';d.save(root/'s.docx')
   set_paragraph_pagination(root/'s.docx',digest(root/'s.docx'),root/'out.docx',[{'block':2,'before':'Second person','page_break_before':True}])
   out=Document(root/'out.docx');self.assertEqual([p.text for p in out.paragraphs],[p.text for p in d.paragraphs])
   self.assertEqual([etree.tostring(t._tbl) for t in out.tables],[etree.tostring(t._tbl) for t in d.tables])
   with ZipFile(root/'s.docx') as a,ZipFile(root/'out.docx') as b:
    for n in a.namelist():
     if n!='word/document.xml':self.assertEqual(a.read(n),b.read(n))
   render_pdf(root/'out.docx',root/'out.pdf');pdf=pymupdf.open(root/'out.pdf');self.assertEqual(len(pdf),2);self.assertIn('First CV',pdf[0].get_text());self.assertIn('Second CV',pdf[1].get_text())
 def test_wrong_text_table_and_section_rejected(self):
  with tempfile.TemporaryDirectory() as tmp:
   root=Path(tmp);d=Document();d.add_paragraph('First');d.add_table(rows=1,cols=1);d.add_section();d.save(root/'s.docx')
   for index,before in [(0,'wrong'),(1,''),(2,'')]:
    with self.assertRaises(ValueError):set_paragraph_pagination(root/'s.docx',digest(root/'s.docx'),root/'bad.docx',[{'block':index,'before':before,'page_break_before':True}])
    self.assertFalse((root/'bad.docx').exists())
