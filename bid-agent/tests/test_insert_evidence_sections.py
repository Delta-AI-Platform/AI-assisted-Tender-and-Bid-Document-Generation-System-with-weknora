import tempfile
import unittest
from pathlib import Path
from docx import Document
import pymupdf
from bid_agent.assembly import append_pdf_evidence
from bid_agent.document_ops import digest,render_pdf
from bid_agent.insert_docx import insert_docx

class EvidenceInsertionTests(unittest.TestCase):
    def test_last_image_keeps_fragment_geometry_and_following_master_restores(self):
        with tempfile.TemporaryDirectory() as tmp:
            r=Path(tmp);a=Document();a.add_paragraph('Before');a.add_paragraph('[[PROOF]]');a.add_paragraph('After');a.save(r/'master.docx');b=Document();b.add_paragraph('Proof');b.save(r/'heading.docx')
            pdf=pymupdf.open()
            for n in range(2):p=pdf.new_page(width=595.28,height=841.89);p.insert_text((40,40),str(n+1))
            pdf.save(r/'source.pdf');pdf.close()
            append_pdf_evidence(r/'heading.docx',digest(r/'heading.docx'),r/'source.pdf',[1,2],r/'fragment.docx',pdf_sha256=digest(r/'source.pdf'))
            args=(r/'master.docx',digest(r/'master.docx'),r/'fragment.docx',digest(r/'fragment.docx'),'[[PROOF]]',r/'out.docx','master')
            with self.assertRaisesRegex(ValueError,'preserve_fragment'):insert_docx(*args)
            self.assertFalse((r/'out.docx').exists())
            insert_docx(*args,section_policy='preserve_fragment');render_pdf(r/'out.docx',r/'out.pdf');d=pymupdf.open(r/'out.pdf');images=[]
            for page in d:
                for image in page.get_image_info():self.assertTrue(page.rect.contains(pymupdf.Rect(image['bbox'])));images.append(image)
            self.assertEqual(len(images),2);self.assertIn('After',d[-1].get_text());self.assertAlmostEqual(d[-1].rect.width,612,delta=1)
