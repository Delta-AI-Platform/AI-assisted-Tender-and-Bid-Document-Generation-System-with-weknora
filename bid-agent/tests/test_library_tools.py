import json,tempfile,unittest,importlib.util
from pathlib import Path
from unittest.mock import patch
from bid_agent.store import Store
from bid_agent.tools import ProjectTools
from bid_agent.document_ops import digest

class LibraryToolsTests(unittest.TestCase):
 def test_scope_binding_copy_and_version_rejection(self):
  with tempfile.TemporaryDirectory() as tmp:
   root=Path(tmp);s=Store(root/'state.db');rid=s.create('test',{'project':str(root)});r=ProjectTools(root,s,rid)
   with self.assertRaisesRegex(ValueError,'not authorized'):r.tools['library_status'].handler(library='knowledge')
   with s.connect() as c:c.execute('UPDATE runs SET scope=? WHERE id=?',(json.dumps({'project':str(root),'libraries':['knowledge']}),rid))
   raw=b'original document';import hashlib;sha=hashlib.sha256(raw).hexdigest()
   def fake(a,stdout=None):
    if stdout is not None:stdout.write(raw);return
    return {'sha256':sha,'suffix':'.pdf','indexed_sha256':'old','index_hash_matches':False}
   with patch('bid_agent.library_tools.invoke',side_effect=fake):
    f=r.tools['import_library_source'].handler
    with self.assertRaisesRegex(ValueError,'hash'):f(library='knowledge',id='kbfile-x',expected_sha256='wrong',target='inputs/a.pdf')
    out=f(library='knowledge',id='kbfile-x',expected_sha256=sha,target='inputs/a.pdf');self.assertEqual(digest(root/'inputs/a.pdf'),sha);self.assertFalse(out['source']['index_hash_matches'])
    with self.assertRaises(FileExistsError):f(library='knowledge',id='kbfile-x',expected_sha256=sha,target='inputs/a.pdf')
    with self.assertRaises(ValueError):f(library='knowledge',id='kbfile-x',expected_sha256=sha,target='../escape.pdf')
    with self.assertRaisesRegex(ValueError,'not authorized'):r.tools['library_status'].handler(library='history')
 def test_fixed_reader_rejects_traversal_and_symlink(self):
  spec=importlib.util.spec_from_file_location('reader',Path(__file__).parents[1]/'integrations/libraries/reader.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
  with tempfile.TemporaryDirectory() as t:
   root=Path(t);(root/'library').mkdir();(root/'outside').write_text('secret');(root/'library/link').symlink_to(root/'outside');m.ROOT=root/'library'
   for path in ['../outside',str(root/'outside'),'link']:
    with self.assertRaises(ValueError):m.safe_source(path)
   (m.ROOT/'ok').write_text('public');self.assertEqual(m.safe_source('ok').read_text(),'public')
   selected=m.safe_source('ok');selected.unlink();selected.symlink_to(root/'outside')
   with self.assertRaises(OSError):m.open_confined(selected)

 def test_history_search_skips_unindexed_files_without_source_id(self):
  spec=importlib.util.spec_from_file_location('reader',Path(__file__).parents[1]/'integrations/libraries/reader.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
  records=[{'status':'unreadable'}, {'source_id':'haf-test','filename':'售后.docx','project_name':'项目'}]
  pairs={'pairs':[{'source_id':'haf-test','case_id':'case','block':{'block_id':'har-test','section_title':'售后','requirement':'服务','response':'历史内容'}}]}
  with patch.object(m,'catalog',return_value=records),patch.object(m,'load',return_value=pairs):
   result=m.handle({'operation':'search','library':'history','query':'售后','limit':3});self.assertEqual(result['matches'],1);self.assertEqual(result['results'][0]['source_id'],'haf-test')

 def test_three_library_boundary_and_hybrid_pipeline_receipt(self):
  spec=importlib.util.spec_from_file_location('reader',Path(__file__).parents[1]/'integrations/libraries/reader.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
  fact={'id':'fact','display_name':'营业执照.jpg','category_code':'CORP','active_module':True,'product_model':'','search_text':'深圳市德塔 营业执照','tokens':['营业执照'],'members':[{'id':'kbfile-f','sha256':'a'*64,'source_rel':'facts/license.jpg'}]}
  report={'id':'report','display_name':'WLR-5C(E)额定载质量检验报告.pdf','category_code':'TEST','active_module':True,'product_model':'WLR-5CE','search_text':'WLR-5C(E) 额定载质量 检验报告','tokens':['额定载质量'],'members':[{'id':'kbfile-r','sha256':'b'*64,'source_rel':'reports/wlr.pdf','extension':'.pdf'}]}
  with patch.object(m,'load',return_value={'documents':[fact,report]}):
   self.assertEqual([x['id'] for x in m.catalog('enterprise_facts')],['fact'])
   self.assertEqual([x['id'] for x in m.catalog('inspection_reports')],['report'])
  def fake_post(url,payload,timeout=60):
   if url.endswith('/embeddings'):
    return {'data':[{'index':i,'embedding':[1.0,0.0]} for i in range(len(payload['input']))]}
   return {'results':[{'index':0,'relevance_score':0.9}]}
  with patch.object(m,'catalog',return_value=[report]),patch.object(m,'post_json',side_effect=fake_post):
   result=m.evidence_search('inspection_reports','额定载质量','', 'WLR-5C(E)',5)
  self.assertEqual(result['results'][0]['id'],'report')
  self.assertEqual(result['results'][0]['original_page_verification'],'required_before_fact_use')
  self.assertEqual(result['pipeline']['rerank']['model'],'BAAI/bge-reranker-v2-m3-int8')

 def test_evidence_search_uses_indexed_body_and_exact_page_locator(self):
  spec=importlib.util.spec_from_file_location('reader',Path(__file__).parents[1]/'integrations/libraries/reader.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
  with tempfile.TemporaryDirectory() as t:
   index=Path(t)/'chunks.sqlite3';db=__import__('sqlite3').connect(index)
   db.execute('''CREATE TABLE chunks(library TEXT,doc_id TEXT,member_id TEXT,member_page INTEGER,native_knowledge_id TEXT,native_chunk_id TEXT,chunk_index INTEGER,source_sha256 TEXT,title TEXT,content TEXT,source_locators TEXT)''')
   loc=json.dumps([{'type':'pdf','mapping':'exact','page':15,'quote':'额定载质量为460kg','bbox':[0.1,0.2,0.3,0.4]}],ensure_ascii=False)
   db.execute('INSERT INTO chunks VALUES(?,?,?,?,?,?,?,?,?,?,?)',('inspection_reports','report','member',None,'knowledge','chunk',0,'b'*64,'普通检验报告.pdf','深圳市德塔工业智能电动汽车有限公司 额定载质量为460kg',loc));db.commit();db.close();m.EVIDENCE_TEXT_INDEX=index
   row={'id':'report','display_name':'普通检验报告.pdf','category_code':'TEST','active_module':True,'product_model':'WLR-5CE','search_text':'WLR-5C(E)','tokens':[],'members':[{'id':'kbfile-r','sha256':'b'*64,'source_rel':'reports/wlr.pdf','extension':'.pdf','page_no':None}]}
   def fake_post(url,payload,timeout=60):
    if url.endswith('/embeddings'):return {'data':[{'index':i,'embedding':[1.0,0.0]} for i in range(len(payload['input']))]}
    return {'results':[{'index':0,'relevance_score':0.9}]}
   with patch.object(m,'catalog',return_value=[row]),patch.object(m,'post_json',side_effect=fake_post):
    result=m.evidence_search('inspection_reports','额定载质量','深圳市德塔工业智能电动汽车有限公司','WLR-5C(E)',5)
   self.assertEqual(result['results'][0]['page_hints'][0]['physical_page_hint'],15)
   self.assertIsNone(result['results'][0]['members'][0]['page_no'])
   self.assertTrue(result['results'][0]['indexed_text_available'])

 def test_original_page_verification_requires_hash_and_parameter_value_alignment(self):
  spec=importlib.util.spec_from_file_location('reader',Path(__file__).parents[1]/'integrations/libraries/reader.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
  with tempfile.TemporaryDirectory() as t:
   root=Path(t);m.ROOT=root;original=root/'report.pdf';original.write_bytes(b'bound original')
   import hashlib;digest=hashlib.sha256(original.read_bytes()).hexdigest();member={'id':'kbfile-r','sha256':digest,'source_rel':'report.pdf','extension':'.pdf','page_no':None}
   with patch.object(m,'catalog',return_value=[{'active_module':True,'members':[member]}]),patch.object(m,'extract_original_page',return_value=('深圳市德塔工业智能电动汽车有限公司 WLR-5C(E) 额定载质量(kg)：460','test_page_read',50,60)):
    passed=m.verify_evidence_page('inspection_reports','kbfile-r',digest,15,{'subject':'深圳市德塔工业智能电动汽车有限公司','model':'WLR-5C(E)','parameter':'额定载质量','value':'460','quote':''})
    self.assertTrue(passed['fact_use_allowed']);self.assertEqual(passed['verification_status'],'verified_original_page')
    failed=m.verify_evidence_page('inspection_reports','kbfile-r',digest,15,{'subject':'','model':'WLR-5C(E)','parameter':'额定载质量','value':'999','quote':''})
    self.assertFalse(failed['fact_use_allowed']);self.assertFalse(failed['checks']['value']['matched'])
    with self.assertRaisesRegex(ValueError,'SHA256'):m.verify_evidence_page('inspection_reports','kbfile-r','0'*64,15,{'subject':'','model':'','parameter':'额定载质量','value':'460','quote':''})

 def test_applicability_rejects_different_parameter_before_comparison(self):
  spec=importlib.util.spec_from_file_location('reader',Path(__file__).parents[1]/'integrations/libraries/reader.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
  verified={'source_sha256':'a'*64,'verification_status':'verified_original_page','fact_use_allowed':True,'checks':{},'page_text_sha256':'b'*64,'excerpt':'原页'}
  base_req={'requirement_id':'R1','source_quote':'额定载质量≥400kg','object_type':'vehicle','vehicle_model':'WLR-5C(E)','submodel':'','component':'整车','component_model':'','parameter':'额定载质量','unit':'kg','test_condition':''}
  base_ev={'subject':'深圳市德塔工业智能电动汽车有限公司','object_type':'vehicle','vehicle_model':'WLR-5C(E)','submodel':'','component':'整车','component_model':'','document_number':'BF-2024-0047','test_condition':'','parameter':'额定载质量','unit':'kg','value':'460','quote':''}
  with patch.object(m,'verify_evidence_page',return_value=verified):
   passed=m.validate_evidence_applicability('inspection_reports','kbfile-r','a'*64,15,base_req,base_ev)
   self.assertTrue(passed['comparison_allowed']);self.assertFalse(passed['compliance_or_deviation_conclusion_allowed'])
   wrong={**base_ev,'parameter':'变速机构参数'}
   failed=m.validate_evidence_applicability('inspection_reports','kbfile-r','a'*64,15,base_req,wrong)
   self.assertFalse(failed['comparison_allowed']);self.assertIn('parameter:mismatch',failed['reasons'])
   unanchored={**base_req,'source_quote':'自动挡车型'}
   failed=m.validate_evidence_applicability('inspection_reports','kbfile-r','a'*64,15,unanchored,base_ev)
   self.assertIn('requirement_parameter_anchor:mismatch',failed['reasons'])

 def test_approved_library_mappings_allow_ai_to_continue_without_pause(self):
  spec=importlib.util.spec_from_file_location('reader',Path(__file__).parents[1]/'integrations/libraries/reader.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
  with tempfile.TemporaryDirectory() as t:
   mapping=Path(t)/'mappings.json';mapping.write_text(json.dumps({'model_groups':[{'id':'WLR5-PLATFORM-001','values':['WLR-5','WLR-5C(E)']}],'parameter_groups':[{'id':'LOAD-001','values':['载重','额定载质量']}]}));m.APPROVED_MAPPINGS=mapping;m.approved_mappings.cache_clear()
   self.assertEqual(m.applicability_check('vehicle_model','WLR-5','WLR-5C(E)')['status'],'pass_approved_mapping')
   self.assertEqual(m.applicability_check('parameter','载重','额定载质量')['mapping_id'],'LOAD-001')

 def test_structured_value_comparison_units_ranges_and_enums(self):
  spec=importlib.util.spec_from_file_location('reader',Path(__file__).parents[1]/'integrations/libraries/reader.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
  self.assertEqual(m.compare_structured_value({'operator':'ge','required_value':'500','required_max':'','unit':'kg','enum_values':[]},'0.6','t')['relation'],'true')
  self.assertEqual(m.compare_structured_value({'operator':'between','required_value':'200','required_max':'750','unit':'V','enum_values':[]},'660','V')['relation'],'true')
  self.assertEqual(m.compare_structured_value({'operator':'covers','required_value':'200','required_max':'750','unit':'V','enum_values':[]},'150','V','800')['relation'],'true')
  self.assertEqual(m.compare_structured_value({'operator':'enum_eq','required_value':'','required_max':'','unit':'','enum_values':['自动挡']},'机械式、手动')['relation'],'false')
  with self.assertRaisesRegex(ValueError,'different dimensions'):m.compare_structured_value({'operator':'ge','required_value':'500','required_max':'','unit':'kg','enum_values':[]},'600','V')

 def test_value_evaluation_stops_when_applicability_fails(self):
  spec=importlib.util.spec_from_file_location('reader',Path(__file__).parents[1]/'integrations/libraries/reader.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
  with patch.object(m,'validate_evidence_applicability',return_value={'comparison_allowed':False,'reasons':['parameter:mismatch']}):
   x=m.evaluate_evidence_value_relation('inspection_reports','id','a'*64,1,{'requirement_id':'R'}, {'value':'460','unit':'kg'}, {'operator':'ge','required_value':'500','required_max':'','unit':'kg','enum_values':[]})
  self.assertEqual(x['evaluation_status'],'blocked_by_applicability');self.assertIsNone(x['mathematical_relation'])

 def test_value_evaluation_rejects_operator_relabeling(self):
  spec=importlib.util.spec_from_file_location('reader',Path(__file__).parents[1]/'integrations/libraries/reader.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
  applicable={'comparison_allowed':True}
  requirement={'requirement_id':'R','source_quote':'载重≥500kg'};evidence={'value':'600','value_max':'','unit':'kg'}
  with patch.object(m,'validate_evidence_applicability',return_value=applicable):
   x=m.evaluate_evidence_value_relation('inspection_reports','id','a'*64,1,requirement,evidence,{'operator':'le','required_value':'500','required_max':'','unit':'kg','enum_values':[]})
  self.assertEqual(x['evaluation_status'],'blocked_by_tender_anchor');self.assertFalse(x['tender_anchor_checks']['operator'])

 def test_locked_configuration_false_relation_concludes_negative_deviation(self):
  spec=importlib.util.spec_from_file_location('reader',Path(__file__).parents[1]/'integrations/libraries/reader.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
  import hashlib
  requirement={'requirement_id':'AR0049','source_quote':'载重（kg） ≥500'};evidence={'vehicle_model':'WLR-5C(E)'}
  locked={'selection_status':'locked_for_project','selected_vehicle_model':'WLR-5C(E)','selection_source':'项目车型锁定表','source_index_sha256':'a'*64,'requirement_text_sha256':hashlib.sha256(requirement['source_quote'].encode()).hexdigest()}
  evaluated={'requirement_id':'AR0049','evaluation_status':'calculated','mathematical_relation':{'relation':'false'},'applicability':{'comparison_allowed':True}}
  with patch.object(m,'evaluate_evidence_value_relation',return_value=evaluated):
   x=m.conclude_locked_configuration_deviation('inspection_reports','id','b'*64,22,requirement,evidence,{},locked)
  self.assertEqual(x['deviation_conclusion'],'负偏离');self.assertTrue(x['deviation_conclusion_allowed'])
  with self.assertRaisesRegex(ValueError,'exactly match'):
   m.conclude_locked_configuration_deviation('inspection_reports','id','b'*64,22,requirement,{'vehicle_model':'其他车型'},{},locked)

 def test_batch_evaluation_visits_every_item_and_never_fail_fast(self):
  spec=importlib.util.spec_from_file_location('reader',Path(__file__).parents[1]/'integrations/libraries/reader.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
  def fake(*args,**kwargs):
   requirement=args[4]
   if requirement['requirement_id']=='BAD':raise ValueError('bad row')
   return {'requirement_id':requirement['requirement_id'],'evaluation_status':'calculated','mathematical_relation':{'relation':'true'},'execution_action':'continue_ai_execution'}
  items=[{'task_id':'a','mode':'evaluate','requirement':{'requirement_id':'OK'}},{'task_id':'b','mode':'evaluate','requirement':{'requirement_id':'BAD'}},{'task_id':'c','mode':'gap','gap_reason':'无原件','requirement':{'requirement_id':'GAP'}}]
  with patch.object(m,'evaluate_evidence_value_relation',side_effect=fake):result=m.batch_evaluate_evidence_relations('inspection_reports',items)
  self.assertTrue(result['all_items_visited']);self.assertEqual(result['summary']['processed'],2);self.assertEqual(result['summary']['item_errors'],1)
  self.assertEqual(result['summary']['no_applicable_evidence'],1);self.assertEqual(result['execution_action'],'continue_ai_execution')

 def test_batch_automatically_concludes_when_configuration_is_locked(self):
  spec=importlib.util.spec_from_file_location('reader',Path(__file__).parents[1]/'integrations/libraries/reader.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
  item={'task_id':'AR0049','mode':'evaluate','requirement':{'requirement_id':'AR0049'},'locked_configuration':{'selection_status':'locked_for_project'}}
  concluded={'requirement_id':'AR0049','evaluation_status':'calculated','mathematical_relation':{'relation':'false'},'deviation_decision_status':'concluded','deviation_conclusion':'负偏离','deviation_conclusion_allowed':True}
  with patch.object(m,'conclude_locked_configuration_deviation',return_value=concluded) as conclude,patch.object(m,'evaluate_evidence_value_relation') as evaluate:
   result=m.batch_evaluate_evidence_relations('inspection_reports',[item])
  conclude.assert_called_once();evaluate.assert_not_called();self.assertEqual(result['results'][0]['result']['deviation_conclusion'],'负偏离')

 def test_original_page_ocr_is_reused_within_batch_process(self):
  spec=importlib.util.spec_from_file_location('reader',Path(__file__).parents[1]/'integrations/libraries/reader.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
  m.cached_original_page.cache_clear()
  with patch.object(m,'extract_original_page',return_value=('原页文字','ocr',1,1)) as extract:
   first=m.cached_original_page('/tmp/same.pdf',22,None,'a'*64)
   second=m.cached_original_page('/tmp/same.pdf',22,None,'a'*64)
  self.assertEqual(first,second);extract.assert_called_once()

 def test_readability_opens_document_and_rejects_corruption_and_stale_hash(self):
  from docx import Document
  spec=importlib.util.spec_from_file_location('reader',Path(__file__).parents[1]/'integrations/libraries/reader.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
  with tempfile.TemporaryDirectory() as t:
   root=Path(t);m.ROOT=root/'originals';m.AUTH=root/'authority';m.ROOT.mkdir();m.AUTH.mkdir();(m.ROOT/'business').mkdir();(m.AUTH/'authority.meta.json').write_text(json.dumps({'source_roots':{'business':'business'}}))
   d=Document();d.add_paragraph('Actual readable text');d.save(m.ROOT/'business/good.docx');(m.ROOT/'business/bad.docx').write_bytes(b'not a Word document')
   rows=[{'source_id':'good','source_path':'good.docx','domain':'business','suffix':'.docx','sha256':digest(m.ROOT/'business/good.docx')},{'source_id':'bad','source_path':'bad.docx','domain':'business','suffix':'.docx','sha256':digest(m.ROOT/'business/bad.docx')},{'source_id':'stale','source_path':'good.docx','domain':'business','suffix':'.docx','sha256':'0'*64}]
   with patch.object(m,'catalog',return_value=rows):
    results=m.handle({'operation':'verify_readability','library':'history','offset':0,'count':5})['results'];self.assertTrue(results[0]['readable']);self.assertGreater(results[0]['text_characters'],0);self.assertFalse(results[1]['readable']);self.assertFalse(results[2]['readable']);self.assertIn('hash mismatch',results[2]['error'])

 def test_legacy_doc_with_embedded_zip_uses_verified_conversion(self):
  from docx import Document
  from zipfile import ZipFile
  spec=importlib.util.spec_from_file_location('reader',Path(__file__).parents[1]/'integrations/libraries/reader.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
  with tempfile.TemporaryDirectory() as t:
   root=Path(t);m.ROOT=root/'originals';m.AUTH=root/'authority';m.ROOT.mkdir();m.AUTH.mkdir();(m.ROOT/'business').mkdir();cache=root/'cache';cache.mkdir();m.CONVERSION_ROOTS=(cache,)
   (m.AUTH/'authority.meta.json').write_text(json.dumps({'source_roots':{'business':'business'}}));src=m.ROOT/'business/legacy.doc';src.write_bytes(b'legacy-binary-prefix')
   with ZipFile(src,'a') as z:z.writestr('theme/theme.xml','<theme/>')
   sha=digest(src);d=Document();d.add_paragraph('Converted content');d.save(cache/(sha+'.docx'))
   row={'source_id':'legacy','source_path':'legacy.doc','domain':'business','suffix':'.doc','sha256':sha}
   with patch.object(m,'catalog',return_value=[row]):
    x=m.verify_history(0,1)['results'][0];self.assertTrue(x['readable']);self.assertEqual(x['read_mode'],'hash_bound_docx_conversion');self.assertEqual(x['sha256'],sha)
