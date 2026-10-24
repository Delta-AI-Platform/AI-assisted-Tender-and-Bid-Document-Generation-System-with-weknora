import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from bid_agent.document_ops import digest
from bid_agent.store import Store
from bid_agent.tools import ProjectTools


class SlotBridgeTests(unittest.TestCase):
    def setup_project(self, root):
        doc = Document();doc.add_paragraph('固定要求：不得改变')
        table = doc.add_table(rows=1, cols=2);table.cell(0, 0).text = '供应商名称'
        doc.save(root/'source.docx')
        store = Store(root/'state.sqlite');run = store.create('test', {'project':str(root)})
        return ProjectTools(root, store, run)

    def test_registered_slot_fill_preserves_fixed_parts(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp);tools = self.setup_project(root)
            tools.tools['tag_docx_slots'].handler(source='source.docx', source_sha256=digest(root/'source.docx'),
                target='tagged.docx', slots=[{'slot_id':'name','kind':'cell','table':0,'row':0,'column':1}])
            result = tools.tools['fill_empty_slots'].handler(source='tagged.docx', source_sha256=digest(root/'tagged.docx'),
                target='filled.docx', values={'name':'测试名称\n第二行'})
            self.assertEqual(result['updated_slots'], ['name'])
            self.assertEqual(result['quality_status'], 'not_verified')
            from bid_agent.document_worker import read_document
            self.assertIn('测试名称', str(read_document(root/'filled.docx')))
            self.assertEqual(Document(root/'filled.docx').tables[0].cell(0,0).text, '供应商名称')
            with ZipFile(root/'tagged.docx') as a, ZipFile(root/'filled.docx') as b:
                for name in a.namelist():
                    if name != 'word/document.xml':self.assertEqual(a.read(name),b.read(name))
            for source, values in [('tagged.docx',{'missing':'x'}),('filled.docx',{'name':'覆盖'}),
                                   ('tagged.docx',{'name':'末尾换行\n'})]:
                with self.assertRaises(ValueError):
                    tools.tools['fill_empty_slots'].handler(source=source, source_sha256=digest(root/source),
                        target='bad.docx', values=values)
                self.assertFalse((root/'bad.docx').exists())

    def test_fixed_cell_and_merged_alias_are_not_tagged(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp);tools = self.setup_project(root)
            with self.assertRaises(ValueError):
                tools.tools['tag_docx_slots'].handler(source='source.docx',source_sha256=digest(root/'source.docx'),
                    target='bad.docx',slots=[{'slot_id':'fixed','kind':'cell','table':0,'row':0,'column':0}])
            self.assertFalse((root/'bad.docx').exists())
            doc=Document();table=doc.add_table(rows=1,cols=2);table.cell(0,0).merge(table.cell(0,1));doc.save(root/'merged.docx')
            with self.assertRaises(ValueError):
                tools.tools['tag_docx_slots'].handler(source='merged.docx',source_sha256=digest(root/'merged.docx'),
                    target='bad.docx',slots=[{'slot_id':str(i),'kind':'cell','table':0,'row':0,'column':i} for i in range(2)])
            self.assertFalse((root/'bad.docx').exists())
