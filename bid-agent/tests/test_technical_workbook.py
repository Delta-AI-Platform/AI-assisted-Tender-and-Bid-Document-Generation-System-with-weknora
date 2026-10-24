import json
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from bid_agent.document_ops import digest
from bid_agent.technical_workbook import create_technical_workbook
from bid_agent.store import Store
from bid_agent.tools import ProjectTools


class TechnicalWorkbookTests(unittest.TestCase):
    def source(self, root, quotes):
        src = root/'source.json'
        src.write_text(json.dumps({'requirements': [
            {'requirement_id': '=ID'+str(i), 'quote': text, 'source_physical_pages': [i+1]}
            for i, text in enumerate(quotes)]}, ensure_ascii=False))
        return src

    def test_registered_worker_preserves_text_and_never_infers_response(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            quotes = ['=SUM(A1:A2) 是原文而不是公式', '（2）最后一次交货后清点完毕起计算质保期。\n完整限定语']
            src = self.source(root, quotes)
            store = Store(root/'state.sqlite')
            run = store.create('test', {'project': str(root)})
            tools = ProjectTools(root, store, run)
            receipt = tools.tools['create_technical_workbook'].handler(
                source='source.json', source_sha256=digest(src), target='review.xlsx')
            self.assertEqual(receipt['quality_status'], 'not_verified')
            workbook = load_workbook(root/'review.xlsx', data_only=False)
            for row, quote in enumerate(quotes, 3):
                self.assertEqual(workbook['技术对标'].cell(row, 4).value, quote)
                self.assertEqual(workbook['技术对标'].cell(row, 4).data_type, 's')
                self.assertEqual(workbook['证据索引'].cell(row-1, 1).data_type, 's')
                for column in (1, 2, 3, 5, 6):
                    self.assertIsNone(workbook['技术对标'].cell(row, column).value)

    def test_cell_capacity_and_invalid_characters_fail_without_partial_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for quote in ['x'*32768, '原文\x00不能静默删除']:
                src = self.source(root, [quote])
                with self.assertRaises(ValueError):
                    create_technical_workbook(src, digest(src), root/'bad.xlsx')
                self.assertFalse((root/'bad.xlsx').exists())

    def test_changed_source_cannot_be_converted_under_old_binding(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            src = self.source(root, ['原有要求'])
            sha = digest(src)
            self.source(root, ['已变更要求'])
            with self.assertRaisesRegex(ValueError, 'version changed'):
                create_technical_workbook(src, sha, root/'bad.xlsx')
            self.assertFalse((root/'bad.xlsx').exists())
