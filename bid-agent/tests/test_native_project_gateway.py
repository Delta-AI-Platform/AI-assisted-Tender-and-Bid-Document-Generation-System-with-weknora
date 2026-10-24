import hashlib,hmac,json,os,sqlite3,sys,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from docx import Document
from pypdf import PdfWriter
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'integrations/weknora'))
import project_gateway as g
import project_stages as ps
from bid_agent.store import Store
from bid_agent.response_logic import check_logic,validate_matrix

class GatewayTest(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();p=Path(self.tmp.name);self.p=p
  self.patch=patch.multiple(g,DB=p/'native.db',STATE=p/'bindings.db',ROOT=p/'projects');self.patch.start()
  self.env=patch.dict(os.environ,{'DETA_PROJECT_SCOPE_KEY':'test-key'});self.env.start()
  with sqlite3.connect(g.DB) as c:
   c.execute('create table sessions(id,tenant_id,user_id,agent_id,deleted_at)')
   c.execute('insert into sessions values(?,?,?,?,null)',('session',1,'owner',g.AGENT))
  self.s={'tenant_id':1,'session_id':'session','user_id':'owner'}
  self.s['signature']=hmac.new(b'test-key',b'1\nsession\nowner',hashlib.sha256).hexdigest()
  self.store=Store(p/'runs.db');old=p/'legacy';old.mkdir()
  self.run=self.store.create('legacy',{'project':str(old),'skills':{},'libraries':['knowledge']})
 def pdf(self,label):
  path=self.p/(label+'.pdf');writer=PdfWriter();writer.add_blank_page(width=595+(ord(label[0])%10),height=842)
  with path.open('wb') as out:writer.write(out)
  return path.read_bytes()
 def text_pdf(self,text):
  import pymupdf
  path=self.p/'text.pdf';doc=pymupdf.open();page=doc.new_page();page.insert_text((72,72),text);doc.save(path);doc.close();return path.read_bytes()
 def tearDown(self):self.env.stop();self.patch.stop();self.tmp.cleanup()
 def test_long_tender_payload_recovery_tool_is_in_catalog_first_page(self):
  names=[item['name'] for item in g.catalog()[:50]]
  self.assertIn('analyze_tender_payload',names)
  self.assertLess(names.index('analyze_tender_payload'),5)
 def test_scope_rejects_forged_tenant_and_user(self):
  self.assertEqual(g.verify_scope(self.s),self.s)
  for key,val in [('tenant_id',2),('user_id','other'),('session_id','other')]:
   with self.assertRaises(ValueError):g.verify_scope({**self.s,key:val})
 def test_no_legacy_fallback(self):
  with self.assertRaisesRegex(ValueError,'No active project'):
   g.route(self.store.path,self.run,{'method':'call','name':'list_files','arguments':{'__deta_scope':self.s}})
 def test_switch_and_repeat_are_isolated(self):
  blobs={'a':self.pdf('a'),'b':self.pdf('b')}
  with patch.object(g,'sandbox_read',side_effect=lambda s,p:blobs[Path(p).stem]):
   def begin(name):return g.start(self.s,{'source_path':'/workspace/input/'+name+'.pdf','sha256':hashlib.sha256(blobs[name]).hexdigest(),'project_title':name},self.store.path,self.run)
   a=begin('a');again=begin('a');b=begin('b')
  self.assertEqual(a['project_id'],again['project_id']);self.assertNotEqual(a['project_id'],b['project_id'])
  self.assertFalse(b['historical_replay']);self.assertNotEqual(g.binding(self.s)['run'],self.run)
  result=g.route(self.store.path,self.run,{'method':'call','name':'list_files','arguments':{'__deta_scope':self.s}})
  self.assertTrue(result['ok']);self.assertIn('input/b.pdf',result['data']['files']);self.assertNotIn('input/a.pdf',result['data']['files'])
 def test_path_escape_rejected(self):
  for p in ['../secret','/etc/passwd','outputs/../../secret']:
   with self.assertRaises(ValueError):g.safe(self.p,p)
 def test_draft_not_approved_by_checkpoint(self):
  blob=self.pdf('a')
  with patch.object(g,'sandbox_read',return_value=blob):
   g.start(self.s,{'source_path':'/workspace/input/a.pdf','sha256':hashlib.sha256(blob).hexdigest(),'project_title':'a'},self.store.path,self.run)
  g.checkpoint(self.s,{'phase':'passed','next_action':'done','open_items':[],'blocked':False})
  with patch.object(g,'container',side_effect=ValueError('not running')):result=g.audit(self.s)
  self.assertFalse(result['submission_ready']);self.assertEqual(result['status'],'draft_not_verified')
 def test_server_builds_stage_validation_and_stage_one_records(self):
  blob=self.pdf('a')
  with patch.object(g,'sandbox_read',return_value=blob):
   project=g.start(self.s,{'source_path':'/workspace/input/a.pdf','sha256':hashlib.sha256(blob).hexdigest(),'project_title':'a'},self.store.path,self.run)
  validation=g.create_stage_validation(self.s,{'stage':1,'artifacts':{'input_manifest':'analysis/input-manifest.json'}})
  self.assertEqual(validation['stage'],1)
  self.assertTrue(validation['path'].startswith('analysis/step1/stage-validation-v'))
  artifacts={'input_manifest':'analysis/input-manifest.json','stage_validation':validation['path']}
  result=g.route(self.store.path,self.run,{'method':'call','name':'record_project_stage','arguments':{'stage':1,'artifacts':artifacts,'open_items':[],'__deta_scope':self.s}})
  self.assertTrue(result['ok']);self.assertEqual(result['data']['status'],'passed')
  receipt=json.loads((g.ROOT/project['project_id']/'stage_receipts/01.json').read_text())
  self.assertEqual(receipt['validation_results']['input']['inputs'],1)
 def test_stage_validation_requires_exact_real_artifacts(self):
  blob=self.pdf('a')
  with patch.object(g,'sandbox_read',return_value=blob):
   g.start(self.s,{'source_path':'/workspace/input/a.pdf','sha256':hashlib.sha256(blob).hexdigest(),'project_title':'a'},self.store.path,self.run)
  with self.assertRaisesRegex(ValueError,'exact non-validation artifact set'):
   g.create_stage_validation(self.s,{'stage':1,'artifacts':{}})
 def test_resume_selects_first_stale_stage(self):
  blob=self.pdf('a')
  with patch.object(g,'sandbox_read',return_value=blob):
   project=g.start(self.s,{'source_path':'/workspace/input/a.pdf','sha256':hashlib.sha256(blob).hexdigest(),'project_title':'a'},self.store.path,self.run)
  receipts=g.ROOT/project['project_id']/'stage_receipts';receipts.mkdir()
  (receipts/'01.json').write_text(json.dumps({'stage':1,'status':'passed','policy_sha256':'obsolete','artifacts':[]}))
  resumed=g.snapshot(self.s)
  self.assertEqual((resumed['next_stage'],resumed['next_stage_status']),(1,'stale'))
  self.assertEqual(resumed['required_artifacts'],ps.REQUIRED[1])
 def test_record_rejects_when_an_earlier_stage_is_stale(self):
  root=self.p/'stale-project';(root/'stage_receipts').mkdir(parents=True)
  (root/'project.json').write_text(json.dumps({'project_id':'P','source_sha256':'0'*64}))
  (root/'stage_receipts/01.json').write_text(json.dumps({'stage':1,'status':'passed','policy_sha256':'obsolete','artifacts':[]}))
  with self.assertRaisesRegex(ValueError,'Earlier stages must be current and passed'):
   ps.record(root,2,{},[])
 def test_tender_analysis_requires_verbatim_page_quote(self):
  blob=self.text_pdf('qualification requirement ABC')
  with patch.object(g,'sandbox_read',return_value=blob):
   project=g.start(self.s,{'source_path':'/workspace/input/text.pdf','sha256':hashlib.sha256(blob).hexdigest(),'project_title':'text'},self.store.path,self.run)
  extracted=g.extract_project_source(self.s)
  item={'requirement_id':'Q1','category':'qualification','source_quote':'qualification requirement ABC','source_sha256':hashlib.sha256(blob).hexdigest(),'physical_page':1,'source_location':'page text','scope':'bidder'}
  reviews=[{'source_sha256':hashlib.sha256(blob).hexdigest(),'physical_page':1,'disposition':'requirements_extracted','categories':['qualification'],'reason':''}]
  result=g.build_tender_analysis(self.s,{'source_extraction_path':extracted['path'],'items':[item],'page_reviews':reviews,'amendment_effects':[]})
  self.assertEqual(result['items'],1)
  source=g.ROOT/project['project_id']/extracted['path'];equivalent=source.with_name('source-extraction-equivalent.json')
  copy=json.loads(source.read_text());copy['created_at']=copy['created_at']+1;equivalent.write_text(json.dumps(copy))
  validated=ps.validate_tender_analysis(g.ROOT/project['project_id']/result['path'],g.ROOT/project['project_id'],equivalent)
  self.assertEqual(validated['items'],1)
  with self.assertRaisesRegex(ValueError,'not verbatim'):
   g.build_tender_analysis(self.s,{'source_extraction_path':extracted['path'],'items':[{**item,'requirement_id':'Q2','source_quote':'invented'}],'page_reviews':reviews,'amendment_effects':[]})
 def test_tender_analysis_defaults_single_source_and_expands_no_action_ranges(self):
  import pymupdf
  path=self.p/'two-pages.pdf';doc=pymupdf.open()
  for text in ('qualification requirement ABC','cover information only'):
   page=doc.new_page();page.insert_text((72,72),text)
  doc.save(path);doc.close();blob=path.read_bytes()
  with patch.object(g,'sandbox_read',return_value=blob):
   g.start(self.s,{'source_path':'/workspace/input/two-pages.pdf','sha256':hashlib.sha256(blob).hexdigest(),'project_title':'two'},self.store.path,self.run)
  extracted=g.extract_project_source(self.s)
  indexed=g.index_technical_source(self.s,{'physical_pages':[1,2]})
  index_payload=json.loads((g.ROOT/g.binding(self.s)['project']/indexed['path']).read_text())
  self.assertEqual(index_payload['physical_pages'],[1,2])
  item={'requirement_id':'Q1','category':'qualification','source_quote':'qualification requirement ABC','physical_page':1,'source_location':'page text','scope':'bidder'}
  reviews=[
   {'physical_page':1,'disposition':'requirements_extracted','categories':['qualification'],'reason':''},
   {'physical_page':2,'end_physical_page':2,'disposition':'no_actionable_requirements','categories':[],'reason':'封面说明，无可执行要求'},
  ]
  batch=g.save_tender_analysis_batch(self.s,{'source_extraction_path':extracted['path'],'source_sha256':hashlib.sha256(blob).hexdigest(),'start_page':1,'end_page':2,'items':[item],'page_reviews':reviews})
  finalized=g.finalize_tender_analysis_batches(self.s,{'source_extraction_path':extracted['path'],'batch_paths':[batch['path']],'amendment_effects':[]})
  self.assertEqual(finalized['items'],1)
  result=g.build_tender_analysis(self.s,{'source_extraction_path':extracted['path'],'items':[item],'page_reviews':reviews,'amendment_effects':[]})
  payload=json.loads((g.ROOT/g.binding(self.s)['project']/result['path']).read_text())
  self.assertEqual(payload['reviewed_page_count'],2)
  self.assertTrue(all(x['source_sha256']==hashlib.sha256(blob).hexdigest() for x in payload['page_reviews']))
 def test_import_tender_analysis_payload_validates_and_finalizes(self):
  import pymupdf
  path=self.p/'payload-source.pdf';doc=pymupdf.open()
  for text in ('qualification requirement ABC','cover information only'):
   page=doc.new_page();page.insert_text((72,72),text)
  doc.save(path);doc.close();blob=path.read_bytes();source_hash=hashlib.sha256(blob).hexdigest()
  with patch.object(g,'sandbox_read',return_value=blob):
   project=g.start(self.s,{'source_path':'/workspace/input/payload-source.pdf','sha256':source_hash,'project_title':'payload'},self.store.path,self.run)
  extracted=g.extract_project_source(self.s)
  payload={'source_extraction_path':extracted['path'],'items':[{'requirement_id':'Q1','category':'qualification','source_quote':'qualification requirement ABC','physical_page':1,'source_location':'page text','scope':'bidder'}],'page_reviews':[{'physical_page':1,'disposition':'requirements_extracted','categories':['qualification'],'reason':''},{'physical_page':2,'end_physical_page':2,'disposition':'no_actionable_requirements','categories':[],'reason':'封面说明，无可执行要求'}],'amendment_effects':[]}
  raw=json.dumps(payload,ensure_ascii=False).encode();sandbox_path='/workspace/output/'+project['project_id']+'/analysis/step2/tender-analysis-payload.json'
  with patch.object(g,'sandbox_read',return_value=raw):
   result=g.import_tender_analysis_payload(self.s,{'source_extraction_path':extracted['path'],'payload_path':sandbox_path,'payload_sha256':hashlib.sha256(raw).hexdigest()})
  self.assertEqual(result['batch_count'],1);self.assertEqual(result['items'],1)
  final=json.loads((g.ROOT/project['project_id']/result['path']).read_text())
  self.assertEqual(final['reviewed_page_count'],2)
  self.assertEqual(final['items'][0]['source_sha256'],source_hash)
 def test_import_persistent_pending_tender_analysis_payload(self):
  blob=self.text_pdf('qualification requirement ABC');source_hash=hashlib.sha256(blob).hexdigest()
  with patch.object(g,'sandbox_read',return_value=blob):
   project=g.start(self.s,{'source_path':'/workspace/input/text.pdf','sha256':source_hash,'project_title':'persistent payload'},self.store.path,self.run)
  extracted=g.extract_project_source(self.s)
  payload={'source_extraction_path':extracted['path'],'items':[{'requirement_id':'Q1','category':'qualification','source_quote':'qualification requirement ABC','physical_page':1,'source_location':'page text','scope':'bidder'}],'page_reviews':[{'physical_page':1,'disposition':'requirements_extracted','categories':['qualification'],'reason':''}],'amendment_effects':[]}
  raw=json.dumps(payload,ensure_ascii=False).encode();folder=g.ROOT/project['project_id']/'analysis/step2/tender-analysis-payloads';folder.mkdir(parents=True);pending=folder/'pending.json';pending.write_bytes(raw)
  resumed=g.snapshot(self.s);self.assertEqual(resumed['pending_tender_analysis_payloads'][0]['path'],'analysis/step2/tender-analysis-payloads/pending.json')
  result=g.import_tender_analysis_payload(self.s,{'source_extraction_path':extracted['path'],'payload_path':'analysis/step2/tender-analysis-payloads/pending.json','payload_sha256':hashlib.sha256(raw).hexdigest()})
  self.assertEqual(result['items'],1);self.assertFalse(pending.exists())
  self.assertEqual(len(list(folder.glob('imported-*.json'))),1)
 def test_tender_analysis_batch_rejects_more_than_twenty_pages(self):
  blob=self.pdf('a')
  with patch.object(g,'sandbox_read',return_value=blob):
   g.start(self.s,{'source_path':'/workspace/input/a.pdf','sha256':hashlib.sha256(blob).hexdigest(),'project_title':'a'},self.store.path,self.run)
  extracted=g.extract_project_source(self.s)
  with self.assertRaisesRegex(ValueError,'cannot exceed 20'):
   g.save_tender_analysis_batch(self.s,{'source_extraction_path':extracted['path'],'source_sha256':hashlib.sha256(blob).hexdigest(),'start_page':1,'end_page':21,'items':[],'page_reviews':[]})
 def test_docx_logic_conflicts(self):
  tender=self.p/'tender.txt';tender.write_text('2小时内响应；尾气净化装置；具体供货日期以供货通知单为准。')
  d=Document();d.add_paragraph('完全响应，无偏离；证明随附。合同签订后60日历天内供货。');t=d.add_table(rows=1,cols=2);t.cell(0,0).text='技术参数';t.cell(0,1).text='满足';p=self.p/'bid.docx';d.save(p)
  codes={r['code'] for r in check_logic(tender,p)['findings']}
  self.assertTrue({'bare_response','unsupported_deviation','attachment_claim','response_deadline','missing_technical_section','delivery_anchor'}.issubset(codes))
 def test_matrix_cannot_certify_pending(self):
  tender=self.p/'source.txt';tender.write_text('设备交付要求为60日历天')
  data={'requirements':[{'id':'Q1','requirement':'期限','requirement_source':{'path':'source.txt','sha256':g.sha(tender),'quote':'60日历天'},'response':'满足','response_kind':'pending','evidence':[],'deviation':'无偏离'}]}
  (self.p/'matrix.json').write_text(json.dumps(data))
  class R:
   def path(_,p):return g.safe(self.p,p)
  result=validate_matrix(R(),'matrix.json');self.assertTrue(result['errors']);self.assertEqual(result['quality_status'],'not_verified')

if __name__=='__main__':unittest.main()
