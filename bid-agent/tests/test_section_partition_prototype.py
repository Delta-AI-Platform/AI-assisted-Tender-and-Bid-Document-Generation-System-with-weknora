import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from bid_agent.partition_sections import partition_sections
from bid_agent.document_ops import digest

class SectionPartitionPrototypeTests(unittest.TestCase):
    def test_table_boundary_partition_retains_table_and_unrelated_parts(self):
        with tempfile.TemporaryDirectory() as tmp:
            r=Path(tmp);d=Document();d.add_paragraph('宽表标题');t=d.add_table(rows=1,cols=2);t.cell(0,0).text='原文';t.cell(0,1).text='';d.add_paragraph('正文');d.save(r/'in.docx')
            geometry=dict(width_mm=210,height_mm=297,top_mm=25,bottom_mm=25,left_mm=12,right_mm=12)
            sections=[dict(end_body_index=1,**geometry),dict(end_body_index=2,**{**geometry,'left_mm':27,'right_mm':27})]
            partition_sections(r/'in.docx',digest(r/'in.docx'),r/'out.docx',sections)
            out=Document(r/'out.docx');self.assertEqual(len(out.sections),2);self.assertEqual(out.tables[0]._tbl.xml,d.tables[0]._tbl.xml)
            self.assertEqual([p.text for p in out.paragraphs if p.text],['宽表标题','正文'])
            with ZipFile(r/'in.docx') as a,ZipFile(r/'out.docx') as b:
                for name in a.namelist():
                    if name!='word/document.xml':self.assertEqual(a.read(name),b.read(name))
            with self.assertRaises(ValueError):partition_sections(r/'in.docx',digest(r/'in.docx'),r/'bad.docx',sections[:1])
            self.assertFalse((r/'bad.docx').exists())
            with self.assertRaises(ValueError):partition_sections(r/'out.docx',digest(r/'out.docx'),r/'bad.docx',sections)
