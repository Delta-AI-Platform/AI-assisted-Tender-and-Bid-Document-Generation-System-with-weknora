import json,tempfile,unittest
from pathlib import Path
from docx import Document
from bid_agent.expand_template_table import digest
from bid_agent.table_records import expand_table_from_records
class TableRecordsTests(unittest.TestCase):
 def test_bound_values_and_mismatch(self):
  with tempfile.TemporaryDirectory() as tmp:
   p=Path(tmp);d=Document();t=d.add_table(rows=2,cols=2);t.cell(0,0).text='ID';t.cell(0,1).text='Price';t.cell(1,0).text='1';d.save(p/'a.docx')
   records=[['1','001.2300'],['2','']];f=p/'records.json';f.write_text(json.dumps({'records':records}))
   args=(p/'a.docx',digest(p/'a.docx'),p/'out.docx',0,1,1,0,f)
   with self.assertRaisesRegex(ValueError,'version mismatch'):expand_table_from_records(*args,'wrong')
   self.assertFalse((p/'out.docx').exists())
   result=expand_table_from_records(*args,digest(f));out=Document(p/'out.docx')
   self.assertEqual([[c.text for c in r.cells] for r in out.tables[0].rows[1:]],records)
   self.assertEqual(result['records_file_sha256'],digest(f))
