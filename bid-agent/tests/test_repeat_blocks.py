import unittest,tempfile
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.oxml import OxmlElement
from bid_agent.document_ops import digest
from bid_agent.repeat_blocks import repeat_docx_blocks
from bid_agent.store import Store
from bid_agent.tools import ProjectTools

class RepeatBlocksTests(unittest.TestCase):
    def test_registered_repeat_keeps_merges_other_parts_and_tail(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);d=Document();d.add_paragraph('前文');d.add_paragraph('简历');t=d.add_table(rows=2,cols=2);t.cell(0,0).merge(t.cell(0,1));t.cell(0,0).text='固定表头';d.add_paragraph('后文');d.save(root/'s.docx')
            store=Store(root/'db');rid=store.create('test',{'project':str(root)});tools=ProjectTools(root,store,rid)
            tools.tools['repeat_docx_blocks'].handler(source='s.docx',source_sha256=digest(root/'s.docx'),target='out.docx',start_block=1,end_block=2,copies=3,expected_texts=['简历','固定表头'])
            out=Document(root/'out.docx');self.assertEqual([p.text for p in out.paragraphs],['前文','简历','简历','简历','后文']);self.assertEqual(len(out.tables),3)
            for table in out.tables:self.assertIs(table.cell(0,0)._tc,table.cell(0,1)._tc)
            with ZipFile(root/'s.docx') as a,ZipFile(root/'out.docx') as b:
                for n in a.namelist():
                    if n!='word/document.xml':self.assertEqual(a.read(n),b.read(n))
    def test_rejects_changed_bindings_and_protected_elements(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            for tag in ('sectPr','bookmarkStart','drawing','fldChar','numPr','sdt'):
                d=Document();p=d.add_paragraph('模板');p._p.append(OxmlElement('w:'+tag));d.save(root/'s.docx')
                with self.assertRaisesRegex(ValueError,'protected'):
                    repeat_docx_blocks(root/'s.docx',digest(root/'s.docx'),root/'bad.docx',0,0,2,['模板'])
                self.assertFalse((root/'bad.docx').exists())
            d=Document();d.add_paragraph('模板');d.save(root/'s.docx')
            for sha,expected in [('stale',['模板']),(digest(root/'s.docx'),['错误'])]:
                with self.assertRaises(ValueError):repeat_docx_blocks(root/'s.docx',sha,root/'bad.docx',0,0,2,expected)
            self.assertFalse((root/'bad.docx').exists())
