import hashlib,json,tempfile,unittest
from pathlib import Path

from bid_agent.project_ledgers import build_project_facts,create_configuration_lock
from bid_agent.workflow_ledgers import _next_stage2_composite_version,_next_stage7_composite_version,build_qualification,build_scoring,build_commitments,build_stage5_commercial_manifest,build_evidence_queue,build_stage6_gap_baseline,build_quality_review,build_deterministic_quality_review,build_fill_evidence_manifest,build_technical_evidence_ledger,build_cross_document_review


class Registry:
 def __init__(self,root):self.root=Path(root)
 def path(self,relative):
  value=self.root/relative
  if Path(relative).is_absolute() or '..' in Path(relative).parts or not value.resolve().is_relative_to(self.root.resolve()):raise ValueError('bad path')
  return value

def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()

class WorkflowLedgerTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);self.r=Registry(self.root)
  (self.root/'input').mkdir();(self.root/'analysis/step2').mkdir(parents=True)
  self.source=self.root/'input/tender.txt';self.source.write_text('载重≥500kg\n投标人须具有营业执照\n')
  (self.root/'project.json').write_text(json.dumps({'project_id':'P1','source_sha256':sha(self.source)}))
 def tearDown(self):self.tmp.cleanup()
 def ref(self,quote='投标人须具有营业执照'):
  return {'path':'input/tender.txt','sha256':sha(self.source),'physical_page':1,'location':'第1段','quote':quote}
 def test_fact_denominator_excludes_missing(self):
  result=build_project_facts(self.r,[
   {'field_code':'name','label':'项目名','value':'项目','state':'source_verified','sources':[self.ref('载重≥500kg')],'allowed_slots':['cover'],'reason':''},
   {'field_code':'price','label':'报价','value':None,'state':'missing','sources':[],'allowed_slots':['price'],'reason':'未批准'}],
   'analysis/step4/project-facts-v1.json')
  data=json.loads((self.root/result['path']).read_text());self.assertEqual(data['auto_fill_denominator'],['name'])
  self.assertEqual(data['auto_fill_expected_slots'],[{'field_code':'name','slot_id':'cover'}])
 def test_stage2_composite_skips_any_partial_version(self):
  step2=self.root/'analysis/step2'
  (step2/'atomization-review-v1.xlsx').write_bytes(b'old review')
  (step2/'requirement-atomization-server-v2.json').write_text('{}')
  (step2/'requirement-atomization-server-v3-server-spec.json').write_text('{}')
  self.assertEqual(_next_stage2_composite_version(self.r),4)
 def test_stage7_composite_skips_any_partial_version(self):
  step7=self.root/'analysis/step7';step7.mkdir(parents=True)
  (step7/'eight-column-internal-review-v1.xlsx').write_bytes(b'old workbook')
  (step7/'eight-column-internal-review-v2.manifest.json').write_text('{}')
  (step7/'eight-column-internal-review-v3.traceability.json').write_text('{}')
  self.assertEqual(_next_stage7_composite_version(self.r),4)
 def test_stage4_manifest_requires_every_trusted_slot(self):
  from docx import Document
  facts=build_project_facts(self.r,[{'field_code':'name','label':'项目名','value':'项目','state':'source_verified','sources':[self.ref('载重≥500kg')],'allowed_slots':['cover'],'reason':''}],'analysis/step4/project-facts.json')
  qualification=build_qualification(self.r,[{'requirement_id':'Q1','tender_source':self.ref(),'status':'missing','subject':'公司','evidence':[],'reason':'无原件'}],'analysis/step4/qualification.json')
  doc=Document();doc.add_paragraph('项目');doc.save(self.root/'analysis/step4/filled.docx')
  result=build_fill_evidence_manifest(self.r,facts['path'],qualification['path'],'analysis/step4/filled.docx',[{'field_code':'name','slot_id':'cover','value':'项目'}],[],'analysis/step4/evidence-manifest.json')
  self.assertEqual(result['filled_slots'],1)
 def test_qualification_met_requires_page_evidence(self):
  with self.assertRaisesRegex(ValueError,'page-bound'):
   build_qualification(self.r,[{'requirement_id':'Q1','tender_source':self.ref(),'status':'met','subject':'公司','evidence':[],'reason':''}],'analysis/step4/qualification-v1.json')
 def test_source_quote_and_certificate_fields_are_enforced(self):
  bad=self.ref('不存在的引文')
  with self.assertRaisesRegex(ValueError,'not present'):
   build_project_facts(self.r,[{'field_code':'x','label':'x','value':'x','state':'source_verified','sources':[bad],'allowed_slots':['x'],'reason':''}],'analysis/step4/bad-facts.json')
  evidence={**self.ref(),'applicability':'applicable','subject':'德塔公司','evidence_type':'certificate','certificate_number':'','validity':'','issuer':''}
  with self.assertRaisesRegex(ValueError,'requires number'):
   build_qualification(self.r,[{'requirement_id':'Q2','tender_source':self.ref(),'status':'met','subject':'德塔公司','evidence':[evidence],'reason':''}],'analysis/step4/bad-qualification.json')
 def test_scoring_and_commitment_business_boundaries(self):
  with self.assertRaisesRegex(ValueError,'Subjective'):
   build_scoring(self.r,[{'scoring_id':'S1','kind':'subjective','tender_source':self.ref(),'awarded_points':1}],'analysis/step5/scoring-bad.json')
  result=build_commitments(self.r,[{'commitment_id':'C1','tender_source':self.ref(),'commencement':'收到通知后','scope':'本项目','time_limit':'2小时','business_approved':False}],'analysis/step5/commitments-v1.json')
  self.assertEqual(result['business_approved'],0)
 def test_stage5_manifest_binds_supported_content_and_approved_commitment(self):
  from docx import Document
  scoring=build_scoring(self.r,[{'scoring_id':'S1','kind':'objective','tender_source':self.ref(),'claimed_points':0}],'analysis/step5/scoring.json')
  commitments=build_commitments(self.r,[{'commitment_id':'C1','tender_source':self.ref(),'commencement':'收到通知后','scope':'本项目','time_limit':'2小时','business_approved':True}],'analysis/step5/commitments.json')
  doc=Document();doc.add_paragraph('投标人须具有营业执照');doc.add_paragraph('收到通知后2小时，本项目适用。');doc.save(self.root/'analysis/step5/commercial.docx')
  result=build_stage5_commercial_manifest(self.r,scoring['path'],commitments['path'],'analysis/step5/commercial.docx',[
   {'ledger_kind':'scoring','item_id':'S1','docx_location':'第1段','quote':'投标人须具有营业执照'},
   {'ledger_kind':'commitment','item_id':'C1','docx_location':'第2段','quote':'收到通知后2小时，本项目适用。'}], 'analysis/step5/commercial-manifest.json')
  self.assertEqual(result['coverage'],1)
 def test_configuration_lock_and_queue_are_file_bound_and_batched(self):
  evidence=self.root/'input/report.txt';evidence.write_text('车型A报告')
  selection={'project_id':'P1','source_sha256':sha(self.source),'procurement_items':[{'procurement_id':'G1','name':'车辆','status':'locked','selected_model':'车型A','evidence_paths':['input/report.txt'],'applicable_requirement_ids':['R1']}]}
  selection_path=self.root/'analysis/step2/selection.json';selection_path.write_text(json.dumps(selection))
  lock=create_configuration_lock(self.r,'analysis/step2/selection.json','analysis/step6/configuration-lock-v1.json')
  atom={'project_id':'P1','source_sha256':sha(self.source),'rows':[{'kind':'requirement','requirement_id':f'R{i}','source_text':f'要求{i}'} for i in range(1,42)]}
  atom_path=self.root/'analysis/step2/atom.json';atom_path.write_text(json.dumps(atom))
  queue=build_evidence_queue(self.r,'analysis/step2/atom.json',lock['path'],lock['lock_version'],'analysis/step6/evidence-queue-v1.json')
  self.assertEqual((queue['tasks'],queue['batches']),(41,3))
  queue_data=json.loads((self.root/queue['path']).read_text())
  self.assertEqual((queue_data['project_id'],queue_data['source_sha256']),('P1',sha(self.source)))
 def test_stage6_gap_baseline_visits_every_requirement_without_guessing_model(self):
  atom={'project_id':'P1','source_sha256':sha(self.source),'rows':[{'kind':'requirement','requirement_id':f'R{i}','source_text':f'要求{i}'} for i in range(1,23)]}
  atom_path=self.root/'analysis/step2/atom-gap.json';atom_path.write_text(json.dumps(atom))
  result=build_stage6_gap_baseline(self.r,'analysis/step2/atom-gap.json')
  self.assertEqual(result['requirements_visited'],22);self.assertTrue(result['selection_gap'])
  lock=json.loads((self.root/result['configuration_lock']).read_text())
  self.assertIsNone(lock['procurement_items'][0]['selected_model'])
  ledger=json.loads((self.root/result['technical_evidence_ledger']).read_text())
  self.assertEqual(ledger['conclusion_counts']['无适用证据'],22)
  self.assertFalse(ledger['manual_columns_written'])
 def test_semantic_review_cannot_award_pass(self):
  gates={name:{'status':'pass','method':'deterministic','issues':[],'evidence':[]} for name in ['known_fields','evidence','object_parameter','cross_document','manual_column_boundary']}
  gates['object_parameter']['method']='model_only'
  with self.assertRaisesRegex(ValueError,'cannot independently'):
   build_quality_review(self.r,gates,[],'analysis/step8/content-v1.json')
 def test_stage8_quality_statuses_are_recomputed_from_bound_artifacts(self):
  from docx import Document
  from openpyxl import Workbook
  step8=self.root/'analysis/step8';step8.mkdir(parents=True,exist_ok=True)
  stage4={'version':'stage4-evidence-manifest-v1','project_id':'P1','source_sha256':sha(self.source),'fill_rate':1,'filled_slots':1,'expected_fill_slots':1}
  technical={'version':'technical-evidence-ledger-v1','project_id':'P1','source_sha256':sha(self.source),'items':[{'requirement_id':'R1','machine_conclusion':'无适用证据','comparison':None}]}
  cross={'version':'cross-document-review-v1','project_id':'P1','source_sha256':sha(self.source),'checks':[{'check_id':'C1','status':'consistent'}]}
  for name,data in [('stage4.json',stage4),('technical.json',technical),('cross.json',cross)]:(step8/name).write_text(json.dumps(data))
  headers=['车型','子系统模块','部件模块','功能模块','标准车型答案（历史投标响应）','本次招标要求','投标人响应','偏离说明']
  doc=Document();table=doc.add_table(rows=2,cols=8)
  for i,v in enumerate(headers):table.cell(0,i).text=v
  table.cell(1,5).text='载重≥500kg';doc.save(step8/'review.docx')
  book=Workbook();sheet=book.active;sheet.title='八列内部核对表';sheet.append(headers);sheet.append(['','','','','','载重≥500kg','','']);book.save(step8/'review.xlsx')
  result=build_deterministic_quality_review(self.r,'analysis/step8/stage4.json','analysis/step8/technical.json','analysis/step8/review.docx','analysis/step8/review.xlsx','analysis/step8/cross.json',[],'analysis/step8/content-v2.json')
  data=json.loads((self.root/result['path']).read_text());self.assertEqual(data['gates']['evidence']['status'],'issues_disclosed');self.assertEqual(data['gates']['manual_column_boundary']['status'],'pass')
 def test_technical_evidence_ledger_preserves_locked_negative_deviation(self):
  queue={'version':'evidence-task-queue-v1','project_id':'P1','source_sha256':sha(self.source),'tasks':[{'task_id':'R1','requirement_id':'R1'}]}
  q=self.root/'analysis/step6/queue.json';q.parent.mkdir(parents=True);q.write_text(json.dumps(queue))
  batch={'version':'evidence-batch-result-v1','project_id':'P1','source_sha256':sha(self.source),'request_task_ids':['R1'],'all_items_visited':True,'manual_response_or_deviation_cells_written':False,
         'results':[{'task_id':'R1','item_status':'processed','result':{'evaluation_status':'calculated','deviation_conclusion':'负偏离','deviation_conclusion_allowed':True,'decision_basis':'locked','mathematical_relation':{'relation':'false','details':{'observed':'460','required':'500'}}}}]}
  b=self.root/'analysis/step6/evidence-batches/b1.json';b.parent.mkdir(parents=True);b.write_text(json.dumps(batch))
  result=build_technical_evidence_ledger(self.r,'analysis/step6/queue.json',['analysis/step6/evidence-batches/b1.json'],'analysis/step6/ledger.json')
  data=json.loads((self.root/result['path']).read_text());self.assertEqual(data['items'][0]['machine_conclusion'],'负偏离');self.assertFalse(data['manual_columns_written'])
 def test_cross_document_status_is_computed_from_real_values(self):
  left=self.root/'analysis/left.txt';right=self.root/'analysis/right.txt';left.parent.mkdir(exist_ok=True);left.write_text('项目编号 A');right.write_text('项目编号 B')
  def side(path,value):return {'path':str(path.relative_to(self.root)),'sha256':sha(path),'physical_page':None,'location':'正文','quote':'项目编号 '+value,'value':value}
  result=build_cross_document_review(self.r,[{'check_id':'C1','subject':'项目编号','left':side(left,'A'),'right':side(right,'B'),'not_applicable_reason':''}],'analysis/step8/cross.json')
  data=json.loads((self.root/result['path']).read_text());self.assertEqual(data['checks'][0]['status'],'conflict_disclosed')

if __name__=='__main__':unittest.main()
