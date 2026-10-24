import hashlib
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from pypdf import PdfWriter

from bid_agent.store import Store
from bid_agent.tools import ProjectTools


class WorkflowStructureTests(unittest.TestCase):
    def test_server_builds_compact_project_skeleton_from_analysis(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)/'project';root.mkdir();inputs=root/'input';inputs.mkdir();source=inputs/'tender.pdf'
            writer=PdfWriter();writer.add_blank_page(width=595,height=842)
            with source.open('wb') as stream:writer.write(stream)
            digest=hashlib.sha256(source.read_bytes()).hexdigest()
            (root/'project.json').write_text(json.dumps({'project_id':'P2','source_sha256':digest,'title':'真实项目'}))
            tender=root/'analysis/step2/tender.json';tender.parent.mkdir(parents=True)
            tender.write_text(json.dumps({'version':'tender-analysis-v1','project_id':'P2','source_sha256':digest,'items':[
                {'requirement_id':'Q1','category':'qualification','source_quote':'具备资格','physical_page':3,'source_location':'正文'},
                {'requirement_id':'T1','category':'technical','source_quote':'载重不小于500kg','physical_page':8,'source_location':'技术表'},
                {'requirement_id':'Q2','category':'qualification','source_quote':'资格补充','physical_page':9,'source_location':'正文'},
            ]},ensure_ascii=False))
            store=Store(Path(tmp)/'state.sqlite');run=store.create('structure',{'project':str(root)});tools=ProjectTools(root,store,run)
            result=tools.tools['build_project_skeleton_draft'].handler('analysis/step2/tender.json')
            self.assertEqual(result['requirements_mapped'],3);self.assertEqual(result['sections'],2)
            self.assertTrue((root/result['skeleton_docx']).is_file());self.assertTrue((root/result['internal_render_pdf']).is_file())
            document=Document(root/result['skeleton_docx'])
            self.assertIn('载重不小于500kg','\n'.join(p.text for p in document.paragraphs))

    def test_server_builds_and_binds_all_skeleton_manifests(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / 'project'; root.mkdir()
            source = root / 'tender.txt'; source.write_text('资格要求')
            source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
            (root / 'project.json').write_text(json.dumps({'project_id': 'P1', 'source_sha256': source_hash}))
            tender = root / 'analysis/step2/tender.json'; tender.parent.mkdir(parents=True)
            tender.write_text(json.dumps({'version': 'tender-analysis-v1', 'project_id': 'P1', 'source_sha256': source_hash,
                                          'items': [{'requirement_id': 'Q1'}]}))
            docx = root / 'analysis/step3/skeleton.docx'; docx.parent.mkdir(parents=True)
            document = Document(); document.add_heading('资格审查', level=1); document.add_table(rows=1, cols=1); document.save(docx)
            pdf = root / 'analysis/step3/skeleton.pdf'; writer = PdfWriter(); writer.add_blank_page(width=595, height=842)
            with pdf.open('wb') as stream: writer.write(stream)
            store = Store(Path(tmp) / 'state.sqlite'); run = store.create('structure', {'project': str(root)})
            tools = ProjectTools(root, store, run)
            result = tools.tools['build_dynamic_skeleton_artifacts'].handler(
                'analysis/step2/tender.json', 'analysis/step3/skeleton.docx', 'analysis/step3/skeleton.pdf',
                [{'section_id': 'S1', 'order': 1, 'title': '资格审查', 'source_requirement_ids': ['Q1']}],
                [{'slot_id': 'SL1', 'slot_type': 'evidence', 'data_type': 'document', 'allowed_stages': [4], 'source_requirement_ids': ['Q1']}],
                [{'table_id': 'T1', 'title': '资格表', 'docx_table_index': 0, 'source_requirement_ids': ['Q1']}],
                [{'requirement_id': 'Q1', 'disposition': 'section', 'target_id': 'S1', 'reason': '正文要求'}],
            )
            self.assertEqual(result['requirements_mapped'], 1)
            for key in ('skeleton_bundle', 'slot_manifest', 'skeleton_lock', 'structure_comparison', 'render_manifest'):
                self.assertTrue((root / result[key]).is_file())
            import sys
            sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'integrations/weknora'))
            from project_stages import validate_skeleton
            self.assertEqual(validate_skeleton(root / result['skeleton_lock'], root)['render_pages'], 1)


if __name__ == '__main__': unittest.main()
