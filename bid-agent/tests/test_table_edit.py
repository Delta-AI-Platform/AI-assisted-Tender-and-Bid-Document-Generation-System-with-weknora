import tempfile
import unittest
from pathlib import Path
from docx import Document
from bid_agent.document_ops import digest
from bid_agent.store import Store
from bid_agent.tools import ProjectTools

class TableEditTests(unittest.TestCase):
    def test_plain_line_break_repair_preserves_format_and_rejects_layout_breaks(self):
        from bid_agent.table_edit import edit_table_cells
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            for mode in ('plain','explicit','page','column','clear','field'):
                doc=Document();t=doc.add_table(rows=1,cols=2)
                t.cell(0,1).text='相邻固定文字'
                run=t.cell(0,0).paragraphs[0].add_run('优势：甲\n乙');run.bold=True
                br=run._r.find(qn('w:br'))
                if mode in ('page','column','explicit'):
                    br.set(qn('w:type'),'textWrapping' if mode=='explicit' else mode)
                elif mode=='clear':br.set(qn('w:clear'),'all')
                elif mode=='field':run._r.append(OxmlElement('w:fldChar'))
                source=root/(mode+'.docx');target=root/(mode+'-out.docx');doc.save(source)
                original=digest(source)
                edits=[{'table':0,'row':0,'column':0,'before':t.cell(0,0).text,'after':'优势：甲乙'}]
                if mode in ('plain','explicit'):
                    edit_table_cells(source,original,target,edits)
                    saved=Document(target)
                    self.assertEqual(saved.tables[0].cell(0,0).text,'优势：甲乙')
                    self.assertTrue(saved.tables[0].cell(0,0).paragraphs[0].runs[0].bold)
                    self.assertEqual(saved.tables[0].cell(0,1).text,'相邻固定文字')
                else:
                    with self.assertRaisesRegex(ValueError,'protected'):
                        edit_table_cells(source,original,target,edits)
                    self.assertFalse(target.exists())
                self.assertEqual(digest(source),original)

    def test_tool_preserves_fixed_cells_and_rejects_stale_binding(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); doc=Document();doc.add_paragraph('固定说明')
            table=doc.add_table(rows=2,cols=2);table.cell(0,0).text='要求';table.cell(0,1).text='响应';table.cell(1,0).text='全部原文'
            doc.save(root/'source.docx'); sha=digest(root/'source.docx')
            store=Store(root/'state.sqlite');run=store.create('edit',{'project':str(root)});tools=ProjectTools(root,store,run)
            handler=tools.tools['edit_table_cells'].handler
            edit={'table':0,'row':1,'column':1,'before':'','after':'有依据的值'}
            result=handler(source='source.docx',source_sha256=sha,target='v2.docx',edits=[edit])
            out=Document(root/'v2.docx')
            self.assertEqual(out.tables[0].cell(1,0).text,'全部原文')
            self.assertEqual(out.tables[0].cell(1,1).text,'有依据的值')
            self.assertEqual(digest(root/'source.docx'),sha)
            self.assertEqual(result['quality_status'],'not_verified')
            with self.assertRaises(ValueError):handler(source='source.docx',source_sha256=sha,target='bad.docx',edits=[{**edit,'before':'wrong'}])
            self.assertFalse((root/'bad.docx').exists())
    def test_merged_alias_and_mixed_runs_are_rejected(self):
        from bid_agent.table_edit import edit_table_cells
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);doc=Document();t=doc.add_table(rows=1,cols=2);t.cell(0,0).merge(t.cell(0,1));doc.save(root/'s.docx')
            edits=[{'table':0,'row':0,'column':i,'before':'','after':'x'} for i in range(2)]
            with self.assertRaises(ValueError):edit_table_cells(root/'s.docx',digest(root/'s.docx'),root/'out.docx',edits)
            doc=Document();t=doc.add_table(rows=1,cols=1);t.cell(0,0).paragraphs[0].add_run('a');t.cell(0,0).paragraphs[0].add_run('b').bold=True;doc.save(root/'s.docx')
            with self.assertRaises(ValueError):edit_table_cells(root/'s.docx',digest(root/'s.docx'),root/'out.docx',[{**edits[0],'before':'ab'}])
            self.assertFalse((root/'out.docx').exists())
