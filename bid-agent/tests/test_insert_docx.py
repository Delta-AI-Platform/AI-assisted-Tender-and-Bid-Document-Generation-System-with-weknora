import tempfile
import unittest
from pathlib import Path
from docx import Document
from bid_agent.document_ops import digest
from bid_agent.store import Store
from bid_agent.tools import ProjectTools


class InsertDocxTests(unittest.TestCase):
    def test_registered_insertion_preserves_position_and_merges(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);base=Document();base.add_paragraph('前文');base.add_paragraph('[[TABLE]]');base.add_paragraph('后文');base.save(root/'base.docx')
            fragment=Document();table=fragment.add_table(rows=2,cols=2);table.cell(0,0).merge(table.cell(0,1));table.cell(0,0).text='合并表头';table.cell(1,0).text='不可删减要求';fragment.save(root/'fragment.docx')
            store=Store(root/'state.sqlite');run=store.create('test',{'project':str(root)});tools=ProjectTools(root,store,run)
            args=dict(source='base.docx',source_sha256=digest(root/'base.docx'),fragment='fragment.docx',fragment_sha256=digest(root/'fragment.docx'),marker='[[TABLE]]',target='out.docx',header_policy='master')
            result=tools.tools['insert_docx'].handler(**args)
            out=Document(root/'out.docx')
            self.assertEqual([e.tag.split('}')[-1] for e in out.element.body],['p','tbl','p','sectPr'])
            self.assertEqual([p.text for p in out.paragraphs],['前文','后文'])
            self.assertIs(out.tables[0].cell(0,0)._tc,out.tables[0].cell(0,1)._tc)
            self.assertEqual(out.tables[0].cell(1,0).text,'不可删减要求')
            self.assertEqual(result['quality_status'],'not_verified')
            self.assertIn('fragment.docx',result['input_paths'])
            with self.assertRaises(ValueError):
                tools.tools['insert_docx'].handler(**{**args,'target':'bad.docx','fragment':'../fragment.docx'})
            self.assertFalse((root/'bad.docx').exists())

    def test_duplicate_marker_and_changed_fragment_do_not_publish(self):
        from bid_agent.insert_docx import insert_docx
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);base=Document();base.add_paragraph('[[TABLE]]');base.add_paragraph('[[TABLE]]');base.save(root/'base.docx')
            fragment=Document();fragment.add_paragraph('内容');fragment.save(root/'fragment.docx')
            with self.assertRaisesRegex(ValueError,'exactly one'):
                insert_docx(root/'base.docx',digest(root/'base.docx'),root/'fragment.docx',digest(root/'fragment.docx'),'[[TABLE]]',root/'out.docx','master')
            with self.assertRaisesRegex(ValueError,'version changed'):
                insert_docx(root/'base.docx',digest(root/'base.docx'),root/'fragment.docx','old','[[TABLE]]',root/'out.docx','master')
            self.assertFalse((root/'out.docx').exists())

    def test_before_and_after_retain_anchor_text(self):
        from bid_agent.insert_docx import insert_docx
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);base=Document();base.add_paragraph('前文');base.add_paragraph('原章标题');base.add_paragraph('后文');base.save(root/'base.docx');fragment=Document();fragment.add_paragraph('新增资料');fragment.save(root/'fragment.docx')
            for placement,expected in [('before',['前文','新增资料','原章标题','后文']),('after',['前文','原章标题','新增资料','后文'])]:
                out=root/(placement+'.docx')
                insert_docx(root/'base.docx',digest(root/'base.docx'),root/'fragment.docx',digest(root/'fragment.docx'),'原章标题',out,'master',placement=placement)
                self.assertEqual([p.text for p in Document(out).paragraphs],expected)
