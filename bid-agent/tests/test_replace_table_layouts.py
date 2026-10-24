import tempfile
import unittest
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from bid_agent.document_ops import digest
from bid_agent.store import Store
from bid_agent.tools import ProjectTools

class ReplaceTableLayoutTests(unittest.TestCase):
    def test_registered_tool_rejects_changed_words_and_path_escape(self):
        with tempfile.TemporaryDirectory() as tmp:
            r=Path(tmp);a=Document();a.add_table(rows=1,cols=1).cell(0,0).text='固定原文';a.save(r/'a.docx');b=Document();b.add_table(rows=1,cols=1).cell(0,0).text='篡改';b.save(r/'b.docx')
            store=Store(r/'state.sqlite');run=store.create('test',{'project':str(r)});tools=ProjectTools(r,store,run)
            args=dict(source='a.docx',source_sha256=digest(r/'a.docx'),target='out.docx',replacements=[dict(table_index=0,fragment='b.docx',fragment_sha256=digest(r/'b.docx'))])
            with self.assertRaisesRegex(ValueError,'Fixed text'):tools.tools['replace_table_layouts'].handler(**args)
            args['replacements'][0]['fragment']='../b.docx'
            with self.assertRaises(ValueError):tools.tools['replace_table_layouts'].handler(**args)
            self.assertFalse((r/'out.docx').exists())
            a.save(r/'b.docx');args['replacements'][0].update(fragment='b.docx',fragment_sha256=digest(r/'b.docx'))
            tools.tools['replace_table_layouts'].handler(**args)
            self.assertTrue((r/'out.docx').exists())
            a.tables[0].cell(0,0).paragraphs[0]._p.append(OxmlElement('w:sdt'));a.save(r/'protected.docx')
            with self.assertRaisesRegex(ValueError,'Protected'):
                tools.tools['replace_table_layouts'].handler(**{**args,'source':'protected.docx','source_sha256':digest(r/'protected.docx'),'target':'bad.docx'})
