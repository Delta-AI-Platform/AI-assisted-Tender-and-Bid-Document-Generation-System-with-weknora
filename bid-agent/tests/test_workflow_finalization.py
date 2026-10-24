import hashlib
import json
import tempfile
import unittest
from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import Workbook

from bid_agent.store import Store
from bid_agent.tools import ProjectTools


class WorkflowFinalizationTests(unittest.TestCase):
    def test_publish_replaces_outputs_with_exact_three_files(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)/'project';root.mkdir();(root/'analysis/step9').mkdir(parents=True);(root/'outputs').mkdir()
            (root/'project.json').write_text(json.dumps({'project_id':'P1','source_sha256':'a'*64}))
            final=Document();final.add_paragraph('Agent范围完成、待业务核对、不可递交');final.save(root/'analysis/step9/final.docx')
            headers=['车型','子系统模块','部件模块','功能模块','标准车型答案（历史投标响应）','本次招标要求','投标人响应','偏离说明']
            review=Document();table=review.add_table(rows=2,cols=8)
            for i,value in enumerate(headers):table.cell(0,i).text=value
            table.cell(1,5).text='招标原文';review.save(root/'analysis/step9/review.docx')
            book=Workbook();sheet=book.active;sheet.title='八列内部核对表';sheet.append(headers);sheet.append(['','','','','','招标原文','','']);book.save(root/'analysis/step9/review.xlsx')
            pdf=root/'analysis/step9/final.pdf';document=pymupdf.open();document.new_page();document.save(pdf);document.close()
            (root/'outputs/old.txt').write_text('old')
            (root/'outputs/old-folder').mkdir();(root/'outputs/old-folder/process.json').write_text('{}')
            store=Store(Path(tmp)/'state.sqlite');run=store.create('publish',{'project':str(root)});tools=ProjectTools(root,store,run)
            result=tools.tools['publish_final_website_outputs'].handler('analysis/step9/final.docx','analysis/step9/review.docx','analysis/step9/review.xlsx','analysis/step9/final.pdf','完整标书.docx','八列表.docx','八列表.xlsx')
            self.assertEqual({x.name for x in (root/'outputs').iterdir()},{'完整标书.docx','八列表.docx','八列表.xlsx'})
            self.assertTrue(list((root/'analysis/archives/step9').rglob('old.txt')));self.assertFalse(result['submission_ready'])
            self.assertTrue(list((root/'analysis/archives/step9').rglob('process.json')))

    def test_final_review_requires_real_page_render_and_visual_event(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)/'project';root.mkdir();(root/'analysis/step10/pages').mkdir(parents=True);(root/'stage_receipts').mkdir()
            (root/'project.json').write_text(json.dumps({'project_id':'P1','source_sha256':'a'*64}))
            for stage in range(1,10):(root/'stage_receipts'/f'{stage:02}.json').write_text(json.dumps({'status':'passed'}))
            pdf=root/'analysis/final.pdf';doc=pymupdf.open();doc.new_page().insert_text((72,72),'page');doc.save(pdf);doc.close()
            with pymupdf.open(pdf) as document:
                pix=document[0].get_pixmap(matrix=pymupdf.Matrix(1.5,1.5),alpha=False);image=root/'analysis/step10/pages/page-1.png';pix.save(image)
            def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()
            store=Store(Path(tmp)/'state.sqlite');run=store.create('final',{'project':str(root)});tools=ProjectTools(root,store,run)
            with store.connect() as conn:
                store.event(conn,run,'artifact_created',{'operation':'page_image','output':'analysis/step10/pages/page-1.png','page':1,'source_sha256':sha(pdf),'sha256':sha(image)})
                store.event(conn,run,'visual_observation',{'image':'analysis/step10/pages/page-1.png','image_sha256':sha(image),'model':'vision','observation':'文字清晰，无裁切。'})
            result=tools.tools['build_final_page_review_and_decision'].handler('analysis/final.pdf',[{'page':1,'image_path':'analysis/step10/pages/page-1.png','status':'pass','issues':[]}])
            self.assertEqual(result['quality_state'],'agent_scope_complete');self.assertEqual(result['pages'],1)
            decision=json.loads((root/result['delivery_decision']).read_text());self.assertFalse(decision['submission_ready'])

    def test_composite_visual_review_reuses_hash_bound_observation(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)/'project';root.mkdir();(root/'analysis/step10').mkdir(parents=True);(root/'stage_receipts').mkdir()
            (root/'project.json').write_text(json.dumps({'project_id':'P1','source_sha256':'a'*64}))
            for stage in range(1,10):(root/'stage_receipts'/f'{stage:02}.json').write_text(json.dumps({'status':'passed'}))
            pdf=root/'analysis/final.pdf';doc=pymupdf.open();doc.new_page().insert_text((72,72),'page');doc.save(pdf);doc.close()
            with pymupdf.open(pdf) as document:
                image=root/'analysis/step10/page-001.png';document[0].get_pixmap(alpha=False).save(image)
            def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()
            store=Store(Path(tmp)/'state.sqlite');run=store.create('final',{'project':str(root)});tools=ProjectTools(root,store,run)
            with store.connect() as conn:
                store.event(conn,run,'artifact_created',{'operation':'page_image','output':'analysis/step10/page-001.png','page':1,'source_sha256':sha(pdf),'sha256':sha(image)})
                store.event(conn,run,'visual_observation',{'image':'analysis/step10/page-001.png','image_sha256':sha(image),'model':'vision','observation':'页面文字清晰；业务状态仍需复核。'})
            result=tools.tools['a_complete_stage10_visual_review'].handler('analysis/final.pdf','analysis/step10',2)
            self.assertEqual(result['observations_reused'],1);self.assertEqual(result['observations_created'],0)
            review=json.loads((root/result['page_review']).read_text())
            self.assertEqual(review['pages'][0]['status'],'issue_disclosed')
            self.assertEqual(review['actual_vision_observation_count'],1)


if __name__=='__main__':unittest.main()
