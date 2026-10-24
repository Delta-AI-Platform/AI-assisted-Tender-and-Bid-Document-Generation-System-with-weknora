import subprocess
import unittest

class QuotationTests(unittest.TestCase):
    def test_restore_changes_only_prices_and_rejects_incomplete_matches(self):
        script = r'''
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from openpyxl import Workbook,load_workbook
from bid_agent.quotation import restore_quotation,digest
with TemporaryDirectory() as tmp:
    root=Path(tmp);source=root/'source.xlsx';pdf=root/'gold.pdf';pdf.write_bytes(b'test')
    book=Workbook();sheet=book.active
    sheet.cell(5,1,1);sheet.cell(5,4,'MODEL-A');sheet.cell(5,7,'PART');sheet.cell(5,8,'DRAWING');sheet.cell(5,27,'KEEP')
    book.save(source)
    report={'inputs':{'workbook':digest(source),'pdf':digest(pdf)},'matched':1,'total':1,'problems':[],
            'rows':[{'excel_row':5,'untaxed_price':'123.45','taxed_price':'139.50'}]}
    with patch('bid_agent.quotation.reconcile_quotation',return_value=report):
        restore_quotation(source,pdf,1,1,root/'result.xlsx')
    result=load_workbook(root/'result.xlsx').active
    assert result.cell(5,11).value==123.45 and result.cell(5,12).value==139.5
    assert result.cell(5,27).value=='KEEP' and load_workbook(source).active.cell(5,11).value is None
    with patch('bid_agent.quotation.reconcile_quotation',return_value={**report,'matched':0}):
        try: restore_quotation(source,pdf,1,1,root/'invalid.xlsx')
        except ValueError: pass
        else: raise AssertionError('Ambiguous match accepted')
    assert not (root/'invalid.xlsx').exists()
'''
        result=subprocess.run(['/home/ubuntu/.local/share/deta-bid-skill/venv/bin/python','-c',script],capture_output=True,text=True)
        self.assertEqual(result.returncode,0,result.stderr)
