import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from bid_agent.eight_column_tools import HEADERS, install_eight_column_tools


class Registry:
    def __init__(self, root):
        self.root = Path(root)
        self.tools = {}

    def path(self, relative):
        path = Path(relative)
        if path.is_absolute() or ".." in path.parts:
            raise ValueError("unsafe")
        return self.root / path

    def add(self, name, description, parameters, handler, side_effect=False):
        self.tools[name] = handler


class EightColumnToolsTest(unittest.TestCase):
    def fixture(self, root):
        analysis = root / "analysis" / "step2"
        analysis.mkdir(parents=True)
        index = {
            "pages": [{"text_blocks": [
                {"id": "B1", "text": "一、总则"},
                {"id": "B2", "text": "1.设备必须适用于井下环境"},
                {"id": "B3", "text": "2.环境温度：-20℃～40℃"},
            ], "tables": []}]
        }
        history = {"items": [{"requirement_id": "REQ-1", "model": "WLR-5", "status": "found", "reason": "unique_authority_answer", "results": [{
            "id": "H-1", "matched_model": "WLR-5", "match_scope": "exact_model", "match_type": "exact_parameter",
            "standard_answer": "1.设备适用于井下环境", "source_triplet": "项目｜WLR-5｜历史标书.docx", "source_ref": "p10"
        }]}, {"requirement_id": "REQ-2", "model": "WLR-5", "status": "not_found", "reason": "no_applicable_same_model_history", "results": []}]}
        evidence = {"all_items_visited": True, "manual_response_or_deviation_cells_written": False, "results": [
            {"task_id": "REQ-1", "position": 1, "item_status": "processed", "result": {"requirement_id": "REQ-1", "evaluation_status": "calculated", "mathematical_relation": {"relation": "true"}, "applicability": {"library": "inspection_reports", "source_id": "kbfile-1", "source_sha256": "a" * 64, "physical_page": 15, "applicability_status": "applicable", "original_page_verification": {"verification_status": "verified_original_page", "page_text_sha256": "b" * 64}, "approved_mappings_used": []}}},
            {"task_id": "REQ-2", "position": 2, "item_status": "processed", "result": {"requirement_id": "REQ-2", "evaluation_status": "no_applicable_evidence", "mathematical_relation": None, "gap_reason": "无适用原件。"}}
        ]}
        spec = {"version": "deta-eight-column-internal-review-v3",
                "atomization_path": "analysis/step2/requirement-atomization.json",
                "source_index_path": "analysis/step2/technical-source-index.json",
                "history_match_path": "analysis/step2/history-match-log.json",
                "evidence_batch_paths": ["analysis/step2/evidence-batch.json"],
                "rows": [
                    {"kind": "heading", "text": "一、总则", "source_ids": ["B1"]},
                    {"kind": "requirement", "requirement_id": "REQ-1",
                     "classification": ["全项目通用", "总则", "通用要求", "适用环境"],
                     "sources": [{"id": "B2", "quote": "1.设备必须适用于井下环境"}],
                     "removed_source_prefix": "1.", "display_number": "1.1",
                     "tender_requirement": "1.1 设备必须适用于井下环境", "history_record_id": "H-1",
                     "bidder_response": "", "deviation": ""},
                    {"kind": "requirement", "requirement_id": "REQ-2",
                     "classification": ["全项目通用", "总则", "通用要求", "环境温度"],
                     "sources": [{"id": "B3", "quote": "2.环境温度：-20℃～40℃"}],
                     "removed_source_prefix": "2.", "display_number": "1.2",
                     "tender_requirement": "1.2 环境温度：-20℃～40℃", "bidder_response": "", "deviation": ""}
                ]}
        (analysis / "technical-source-index.json").write_text(json.dumps(index, ensure_ascii=False))
        (analysis / "history-match-log.json").write_text(json.dumps(history, ensure_ascii=False))
        (analysis / "evidence-batch.json").write_text(json.dumps(evidence, ensure_ascii=False))
        atom_rows = [
            {"kind": "heading", "text": "一、总则", "source_ids": ["B1"]},
            {"kind": "requirement", "requirement_id": "REQ-1", "classification": ["全项目通用", "总则", "通用要求", "适用环境"], "source_ids": ["B2"], "tender_requirement": "1.1 设备必须适用于井下环境"},
            {"kind": "requirement", "requirement_id": "REQ-2", "classification": ["全项目通用", "总则", "通用要求", "环境温度"], "source_ids": ["B3"], "tender_requirement": "1.2 环境温度：-20℃～40℃"}
        ]
        digest = __import__('hashlib').sha256((analysis / "technical-source-index.json").read_bytes()).hexdigest()
        (analysis / "requirement-atomization.json").write_text(json.dumps({"version": "deta-requirement-atomization-v2", "schema_role": "requirement_parsing_and_object_attribution", "downstream_status": "approved_for_eight_column", "source_index_sha256": digest, "rows": atom_rows}, ensure_ascii=False))
        (analysis / "matrix-spec.json").write_text(json.dumps(spec, ensure_ascii=False))
        return spec

    def test_builds_docx_xlsx_and_keeps_manual_columns_blank(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); self.fixture(root); registry = Registry(root); install_eight_column_tools(registry)
            result = registry.tools["build_eight_column_internal_review"]("analysis/step2/matrix-spec.json", "八列内部核对表-v1")
            self.assertEqual(result["validation"]["source_order_and_text"], "passed")
            self.assertTrue(result["rules"]["pdf_not_generated"])
            doc = Document(root / result["outputs"]["docx"])
            self.assertEqual([c.text for c in doc.tables[0].rows[0].cells], HEADERS)
            self.assertEqual(doc.tables[0].rows[1].cells[5].text, "一、总则")
            self.assertEqual((doc.tables[0].rows[1].cells[6].text, doc.tables[0].rows[1].cells[7].text), ("", ""))
            wb = load_workbook(root / result["outputs"]["xlsx"]); ws = wb["八列内部核对表"]
            self.assertEqual(ws["F2"].value, "一、总则")
            self.assertIsNone(ws["G2"].value); self.assertIsNone(ws["H2"].value)
            self.assertEqual(ws["E3"].value, "1.设备适用于井下环境")
            self.assertEqual(ws["F3"].value, "1.1 设备必须适用于井下环境")
            self.assertIsNone(ws["G3"].value); self.assertIsNone(ws["H3"].value)
            self.assertFalse(list(root.rglob("*.pdf")))
            trace = json.loads((root / result["traceability_path"]).read_text())
            self.assertEqual(trace["summary"]["requirement_rows"], 2)
            self.assertEqual(trace["bindings"][0]["requirement_id"], "REQ-1")
            self.assertEqual(trace["bindings"][0]["xlsx_row"], 3)
            self.assertEqual(trace["bindings"][0]["history"]["record_id"], "H-1")
            self.assertEqual(trace["bindings"][1]["evidence"]["reason"], "无适用原件。")
            self.assertTrue(trace["bindings"][1]["manual_columns"]["both_blank"])

    def test_rejects_changed_tender_wording(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); spec = self.fixture(root)
            spec["rows"][1]["tender_requirement"] = "1.1 设备应当适用于井下环境"
            (root / "analysis/step2/matrix-spec.json").write_text(json.dumps(spec, ensure_ascii=False))
            registry = Registry(root); install_eight_column_tools(registry)
            with self.assertRaisesRegex(ValueError, "differs from exact source"):
                registry.tools["build_eight_column_internal_review"]("analysis/step2/matrix-spec.json", "错误稿")

    def test_rejects_evidence_batches_with_incomplete_requirement_coverage(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); self.fixture(root)
            path = root / "analysis/step2/evidence-batch.json"; data = json.loads(path.read_text()); data["results"].pop(); path.write_text(json.dumps(data))
            registry = Registry(root); install_eight_column_tools(registry)
            with self.assertRaisesRegex(ValueError, "do not cover every requirement_id"):
                registry.tools["build_eight_column_internal_review"]("analysis/step2/matrix-spec.json", "缺证据覆盖")

    def test_rejects_calculated_relation_without_verified_original_page(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); self.fixture(root)
            path = root / "analysis/step2/evidence-batch.json"; data = json.loads(path.read_text()); data["results"][0]["result"].pop("applicability"); path.write_text(json.dumps(data))
            registry = Registry(root); install_eight_column_tools(registry)
            with self.assertRaisesRegex(ValueError, "verified applicable original page binding"):
                registry.tools["build_eight_column_internal_review"]("analysis/step2/matrix-spec.json", "缺原件页")


if __name__ == "__main__":
    unittest.main()
