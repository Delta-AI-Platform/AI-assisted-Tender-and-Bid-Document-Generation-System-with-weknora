import tempfile,unittest,zipfile
from pathlib import Path
from docx import Document
from lxml import etree
from bid_agent.expand_template_table import expand_template_table,digest
class ExpandTableTests(unittest.TestCase):
 def test_preserve_and_reject_business_overwrite(self):
  with tempfile.TemporaryDirectory() as tmp:
   p=Path(tmp);d=Document();d.add_paragraph('Before');t=d.add_table(rows=3,cols=3)
   for c,v in zip(t.rows[0].cells,['ID','NAME','PRICE']):c.text=v
   t.cell(1,0).text='1';t.cell(2,0).text='...';d.add_paragraph('After');d.save(p/'a.docx')
   rows=[['1','A',''],['2','B',''],['3','C','']]
   expand_template_table(p/'a.docx',digest(p/'a.docx'),p/'b.docx',0,1,1,0,rows)
   b=Document(p/'b.docx');self.assertEqual(len(b.tables[0].rows),4)
   self.assertEqual(etree.tostring(t.rows[0]._tr),etree.tostring(b.tables[0].rows[0]._tr))
   with zipfile.ZipFile(p/'a.docx') as a,zipfile.ZipFile(p/'b.docx') as z:
    for n in a.namelist():
     if n!='word/document.xml':self.assertEqual(a.read(n),z.read(n))
   with self.assertRaisesRegex(ValueError,'populated'):
    expand_template_table(p/'b.docx',digest(p/'b.docx'),p/'c.docx',0,1,1,0,rows)
   self.assertFalse((p/'c.docx').exists())
 def test_merged_body_rejected(self):
  with tempfile.TemporaryDirectory() as tmp:
   p=Path(tmp);d=Document();t=d.add_table(rows=2,cols=3);t.cell(1,1).merge(t.cell(1,2));d.save(p/'a.docx')
   with self.assertRaisesRegex(ValueError,'Merged'):
    expand_template_table(p/'a.docx',digest(p/'a.docx'),p/'b.docx',0,1,1,0,[['1','','']])
