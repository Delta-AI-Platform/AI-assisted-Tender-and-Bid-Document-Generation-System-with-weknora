import tempfile,unittest
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml.ns import qn
from bid_agent.empty_sections import remove_redundant_empty_section
from bid_agent.document_ops import digest,render_pdf

class EmptySectionTests(unittest.TestCase):
 def make(self,p):
  d=Document();d.add_paragraph('Before');d.add_section(WD_SECTION_START.NEW_PAGE);d.add_section(WD_SECTION_START.NEW_PAGE);d.add_paragraph('After');d.save(p)
 def test_redundant_blank_page_removed_with_other_parts_preserved(self):
  with tempfile.TemporaryDirectory() as t:
   p=Path(t);self.make(p/'s.docx');d=Document(p/'s.docx');index=next(i for i,b in enumerate(d.element.body) if i>0 and b.xpath('./w:pPr/w:sectPr') and d.element.body[i-1].xpath('./w:pPr/w:sectPr'))
   before=render_pdf(p/'s.docx',p/'before.pdf');remove_redundant_empty_section(p/'s.docx',digest(p/'s.docx'),p/'out.docx',index);after=render_pdf(p/'out.docx',p/'after.pdf');self.assertEqual(before['pages']-after['pages'],1)
   with ZipFile(p/'s.docx') as a,ZipFile(p/'out.docx') as b:
    self.assertEqual(a.namelist(),b.namelist())
    for n in a.namelist():
     if n!='word/document.xml':self.assertEqual(a.read(n),b.read(n))
 def test_nonempty_or_different_geometry_rejected(self):
  with tempfile.TemporaryDirectory() as t:
   p=Path(t);self.make(p/'s.docx');d=Document(p/'s.docx');index=2;d.element.body[index].xpath('./w:pPr/w:sectPr/w:pgSz')[0].set(qn('w:w'),'14000');d.save(p/'different.docx')
   with self.assertRaisesRegex(ValueError,'differs'):remove_redundant_empty_section(p/'different.docx',digest(p/'different.docx'),p/'bad.docx',index)
   d=Document(p/'s.docx');d.paragraphs[index].add_run('Keep this');d.save(p/'text.docx')
   with self.assertRaisesRegex(ValueError,'text or protected'):remove_redundant_empty_section(p/'text.docx',digest(p/'text.docx'),p/'bad.docx',index)
   self.assertFalse((p/'bad.docx').exists())
