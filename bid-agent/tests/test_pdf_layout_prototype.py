import tempfile
import unittest
from pathlib import Path
import pymupdf
from docx import Document
from bid_agent.pdf_table_import import import_table
from bid_agent.document_ops import digest

class PdfLayoutPrototypeTests(unittest.TestCase):
    def test_separate_same_baseline_date_fragments_and_small_type_survive(self):
        with tempfile.TemporaryDirectory() as tmp:
            r=Path(tmp);pdf=pymupdf.open();p=pdf.new_page()
            for y in [70,100,130]:p.draw_line((50,y),(350,y))
            for x in [50,200,350]:p.draw_line((x,70),(x,130))
            for pos,text in [((60,90),'Requirement'),((210,90),'Response'),((60,120),'Preserve 123')]:p.insert_text(pos,text,fontsize=8)
            for x,text in [(180,'Year'),(250,'Month'),(320,'Day')]:p.insert_text((x,170),text,fontsize=12)
            pdf.save(r/'in.pdf');pdf.close()
            import_table(r/'in.pdf',1,0,r/'out.docx',True,digest(r/'in.pdf'))
            d=Document(r/'out.docx')
            self.assertIn('Year\tMonth\tDay',[p.text for p in d.paragraphs])
            self.assertEqual(d.tables[0].cell(0,0).text,'Requirement')
            self.assertEqual(d.tables[0].cell(0,0).paragraphs[0].runs[0].font.size.pt,8)
            self.assertEqual(d.tables[0].cell(1,1).text,'')
