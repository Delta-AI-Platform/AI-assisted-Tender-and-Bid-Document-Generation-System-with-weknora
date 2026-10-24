import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook
from docx import Document
from bid_agent.xlsx_spec_tables import xlsx_spec_tables
from bid_agent.document_ops import digest

class XlsxSpecTableTests(unittest.TestCase):
    def test_contiguous_groups_keep_original_order_and_reject_formula(self):
        with tempfile.TemporaryDirectory() as tmp:
            r=Path(tmp);w=Workbook();s=w.active;s.title='配件';s.append(['车型','原值']);s.append(['A','土1mm']);s.append(['B','原文']);s.append(['A',None]);w.save(r/'in.xlsx')
            args=dict(source=r/'in.xlsx',source_sha256=digest(r/'in.xlsx'),target=r/'out.docx',sheet='配件',header_row=1,start_row=2,end_row=4,group_column=1,columns=[{'column':2,'width_mm':100}],title='参数',page_width_mm=210,page_height_mm=297,margin_mm=25)
            result=xlsx_spec_tables(**args);d=Document(r/'out.docx')
            self.assertEqual([t.cell(1,0).text for t in d.tables],['土1mm','原文',''])
            self.assertEqual([x['excel_row'] for x in result['mapping']],[2,3,4])
            s['B2']='=1+1';w.save(r/'formula.xlsx')
            with self.assertRaisesRegex(ValueError,'Formula'):xlsx_spec_tables(**{**args,'source':r/'formula.xlsx','source_sha256':digest(r/'formula.xlsx'),'target':r/'bad.docx'})
            self.assertFalse((r/'bad.docx').exists())
            with self.assertRaisesRegex(ValueError,'exceed'):xlsx_spec_tables(**{**args,'target':r/'wide.docx','columns':[{'column':2,'width_mm':200}]})
            self.assertFalse((r/'wide.docx').exists())
