import subprocess
import unittest
from bid_agent.runtime import APP_PYTHON

class AssemblyTests(unittest.TestCase):
    def test_chapters_images_and_full_evidence_pages_survive(self):
        script=r'''
from tempfile import TemporaryDirectory
from pathlib import Path
from docx import Document
from docx.shared import Inches
import pymupdf
from pypdf import PdfReader
from bid_agent.assembly import assemble_docx,append_pdf_evidence,digest
from bid_agent.document_ops import render_pdf
with TemporaryDirectory() as tmp:
    root=Path(tmp)
    for name,color in [('one',(1,0,0)),('two',(0,0,1))]:
        pdf=pymupdf.open();page=pdf.new_page(width=60,height=60);page.draw_rect(page.rect,color=color,fill=color)
        page.get_pixmap().save(root/(name+'.png'));pdf.close()
    master=Document();master.add_paragraph('MASTER EXACT TEXT');master.sections[0].header.paragraphs[0].text='MASTER HEADER';master.add_picture(str(root/'one.png'),width=Inches(.3));master.save(root/'master.docx')
    chapter=Document();chapter.add_paragraph('CHAPTER EXACT TEXT');chapter.sections[0].header.paragraphs[0].text='OTHER HEADER';table=chapter.add_table(rows=1,cols=2);table.cell(0,0).text='REQ';table.cell(0,1).text='';chapter.add_picture(str(root/'two.png'),width=Inches(.3));chapter.save(root/'chapter.docx')
    result=assemble_docx(root/'master.docx',[root/'chapter.docx'],root/'merged.docx','master')
    assert result['body_text_retained_in_order'] and result['source_media_hashes_retained']
    assert Document(root/'merged.docx').tables[0].cell(0,1).text==''
    pdf=pymupdf.open()
    for i in range(2): pdf.new_page().insert_text((30,50),'EVIDENCE '+str(i+1))
    pdf.save(root/'evidence.pdf');pdf.close()
    inserted=append_pdf_evidence(root/'merged.docx',digest(root/'merged.docx'),root/'evidence.pdf',[1,2],root/'with-evidence.docx',pdf_sha256=digest(root/'evidence.pdf'))
    assert len(inserted['evidence_pages'])==2
    rendered=render_pdf(root/'with-evidence.docx',root/'final.pdf')
    assert rendered['pages']==4,rendered
    assert 'MASTER EXACT TEXT' in PdfReader(root/'final.pdf').pages[0].extract_text()
    assert 'CHAPTER EXACT TEXT' in PdfReader(root/'final.pdf').pages[1].extract_text()
    with pymupdf.open(root/'final.pdf') as preview:
        for page in list(preview)[-2:]:
            images=page.get_image_info()
            assert len(images)==1,images
            assert page.rect.contains(pymupdf.Rect(images[0]['bbox'])),(page.rect,images[0]['bbox'])
    assert digest(root/'master.docx')==result['sources'][0]['source_sha256']
'''
        result=subprocess.run([APP_PYTHON,'-c',script],capture_output=True,text=True,timeout=120)
        self.assertEqual(result.returncode,0,result.stderr+result.stdout)
